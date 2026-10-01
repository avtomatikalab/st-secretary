"""Дисциплины «по времени» (спелео, пешеходные): время, отсечки, баллы × 15/30 с, снятия, КВ — на выдуманных командах."""

from dataclasses import replace
from fractions import Fraction

import pytest

from st_secretary import psr_run as pr
from st_secretary import time_run as tr
from st_secretary.competition import Zachet
from st_secretary.disciplines import Status
from st_secretary.norms import achieved_norm
from st_secretary.psr_run import Member, TeamInput
from st_secretary.qualification import Qual
from st_secretary.reference import norm_edition


@pytest.fixture
def speleo(psr_card):
    return replace(psr_card, zachety=[Zachet("М/Ж", 2, "0840271811Я", team_size=4)])


def team(name, n, qual=Qual.II, chip=""):
    return TeamInput(f"{name}.xlsx", name, "г. N", str(n), [Member(f"{name} {i}", qual, qual.label, chip) for i in range(4)])


STAGES = [{"id": "s1", "tour": "", "name": "Колодец"}, {"id": "s2", "tour": "", "name": "Шкуродёр"}]


def test_times_penalties_removals_and_places(speleo):
    z = speleo.zachety[0]
    teams = [team("Кедр", 1), team("Сосна", 2), team("Ель", 3), team("Пихта", 4), team("Лиственница", 5),
             team("Берёза", 6), team("Осина", 7)]
    zdata = {"stages": STAGES, "expected": "25", "kv": "60", "teams": {
        "Кедр.xlsx": {"start": "10:00:00", "finish": "10:20:30", "cutoffs": "2:00", "points": {"s1": "0,3", "s2": "1"}},
        "Сосна.xlsx": {"start": "10:05:00", "finish": "10:24:00", "points": {"s1": "2"}},   # 19:00 + 2 × 15 = 19:30
        "Ель.xlsx": {"start": "10:10:00", "finish": "10:25:00", "points": {"s1": "с"}},     # быстрее всех, но со снятием
        "Пихта.xlsx": {"start": "10:15:00", "finish": "11:30:00"},                          # 75 мин > КВ 60 мин
        "Лиственница.xlsx": {"start": "23:50:00", "finish": "00:12:10"},                    # финиш после полуночи
        "Берёза.xlsx": {"start": "10:30", "finish": "10:52:00", "cutoffs": "abc"},
        "Осина.xlsx": {"start": "10:35:00"},                                                # финиша ещё нет
    }}
    run = tr.compute(speleo, z, zdata, teams)
    assert run.kind == "time" and run.seconds_per_point == 15  # расчётное время 25 мин → 15 с за балл
    rows = {r.inp.team: r for r in run.rows}
    k = rows["Кедр"]
    assert (k.distance_time, k.total) == (Fraction(18 * 60 + 30), Fraction(18 * 60 + 30) + Fraction(13, 10) * 15)
    assert tr.clock_text(k.total) == "18:49,5"
    assert rows["Лиственница"].distance_time == 22 * 60 + 10
    assert [(r.inp.team, r.place) for r in run.rows[:4]] == [("Кедр", 1), ("Сосна", 2), ("Лиственница", 3), ("Ель", 4)]
    assert rows["Пихта"].status is Status.OVER_TIME and rows["Пихта"].auto_status and rows["Пихта"].place is None
    assert rows["Осина"].place is None and rows["Берёза"].place is None
    texts = [i.text for i in run.issues]
    assert "«Берёза»: отсечки «abc» — не время (нужно мм:сс)" in texts
    assert any("«Пихта»: время на дистанции 1:15:00 больше КВ (60 мин)" in t for t in texts)
    assert any(t.startswith("нет времени старта или финиша — место не присуждается: ") and "Осина" in t for t in texts)
    assert rows["Кедр"].percent == 100 and rows["Сосна"].percent == Fraction(19 * 60 + 30) / k.total * 100
    assert rows["Ель"].removals == 1 and rows["Ель"].norm == ""  # со снятием норматив не присваивается
    norms = norm_edition(speleo.norms_edition)
    assert run.rank.value is not None
    want = achieved_norm(norms, 2, run.rank.value, Fraction(100), speleo.level).qual
    assert k.norm == (want.label if want else "")


def test_seconds_per_point_and_removal_order(speleo):
    z = speleo.zachety[0]
    teams = [team("А", 1), team("Б", 2), team("В", 3)]
    zdata = {"stages": STAGES, "spp": "30", "removed_order": "count", "teams": {
        "А.xlsx": {"start": "10:00:00", "finish": "10:30:00", "points": {"s1": "с", "s2": "с"}},
        "Б.xlsx": {"start": "10:00:00", "finish": "10:40:00", "points": {"s1": "с", "s2": "1"}},
        "В.xlsx": {"start": "10:00:00", "finish": "10:50:00", "points": {"s2": "2"}}}}
    run = tr.compute(speleo, z, zdata, teams)
    assert run.seconds_per_point == 30
    # без снятий — первые; среди снятых выше тот, у кого снятий меньше (так в Условиях)
    assert [(r.inp.team, r.place) for r in run.rows] == [("В", 1), ("Б", 2), ("А", 3)]
    assert run.rows[0].total == 50 * 60 + 60


@pytest.mark.parametrize("s, sec", [("10:05:23", 36323), ("10:05", 36300), ("9:05:23,4", Fraction(327234, 10)), ("", None)])
def test_parse_clock(s, sec):
    assert tr.parse_clock(s) == sec


@pytest.mark.parametrize("s, sec", [("5:30", 330), ("1:05:30", 3930), ("5", 300), ("", 0), ("2,5", 150)])
def test_parse_duration(s, sec):
    assert tr.parse_duration(s) == sec


def test_bad_values_and_texts():
    for bad in ("25:61:00", "10:5", "десять"):
        with pytest.raises(ValueError):
            tr.parse_clock(bad)
    with pytest.raises(ValueError):
        tr.parse_duration("5:75")
    assert tr.clock_text(Fraction(3930)) == "1:05:30" and tr.clock_text(Fraction(59)) == "0:59"
    assert tr.is_time_discipline(Zachet("М/Ж", 2, "0840131811Я")) and not tr.is_time_discipline(Zachet("М/Ж", 3, "0840161811Я"))


def test_speleo_norms_need_six_with_judges_points(speleo):
    """ЕВСК п. 25.4.2: в спелео результат содержит баллы судей — на муниципальных нужно не менее 6 участников."""
    z = speleo.zachety[0]
    teams = [team(n, i) for i, n in enumerate(["А", "Б", "В", "Г"], start=1)]
    zdata = {"stages": STAGES, "teams": {f"{n}.xlsx": {"start": "10:00:00", "finish": f"10:{20 + i}:00"}
                                         for i, n in enumerate(["А", "Б", "В", "Г"])}}
    run = tr.compute(speleo, z, zdata, teams)
    assert not run.norms_ok and any("нужно не менее 6" in i.text for i in run.issues)


def pedestrian(card, code="0840251811Я"):
    return replace(card, zachety=[Zachet("М/Ж", 2, code, team_size=4)])


def test_pedestrian_nopenalty_time_plus_penalty_time_and_removal_rules(psr_card):
    comp = pedestrian(psr_card)
    z = comp.zachety[0]
    assert tr.profile(z) == "pedestrian"
    teams = [team("А", 1), team("Б", 2), team("В", 3), team("Г", 4)]
    zdata = {"stages": STAGES, "system": "nopenalty", "kv": "60", "teams": {
        "А.xlsx": {"start": "10:00:00", "finish": "10:30:00", "pen_time": "3:00"},       # потеря снаряжения
        "Б.xlsx": {"start": "10:00:00", "finish": "10:31:00"},
        "В.xlsx": {"start": "10:00:00", "finish": "10:20:00", "points": {"s1": "с"}},   # снятие с этапа
        "Г.xlsx": {"start": "10:00:00", "finish": "10:25:00", "points": {"s2": "2"}}}}  # баллов в бесштрафовой нет
    run = tr.compute(comp, z, zdata, teams)
    rows = {r.inp.team: r for r in run.rows}
    assert (rows["А"].total, rows["А"].place) == (33 * 60, 2) and rows["Б"].place == 1
    assert rows["В"].status.value == "removed" and rows["В"].auto_status and rows["В"].place is None  # п. 6.2.8 а
    assert rows["Г"].place is None and "points" in rows["Г"].bad
    assert any("в бесштрафовой системе штрафных баллов нет" in i.text for i in run.issues)
    zdata["removal"] = "okv"  # п. 6.2.8 б — штрафное время, равное ОКВ, за каждое снятие
    del zdata["teams"]["Г.xlsx"]["points"]
    run = tr.compute(comp, z, zdata, teams)
    rows = {r.inp.team: r for r in run.rows}
    assert rows["В"].status.value == "finished" and rows["В"].total == 20 * 60 + 60 * 60 and rows["В"].place == 4
    zdata["system"] = "penalty"  # штрафная: баллы × 30 с по умолчанию (расчётное время не задано)
    zdata["teams"]["Г.xlsx"]["points"] = {"s2": "2"}
    run = tr.compute(comp, z, zdata, teams)
    assert {r.inp.team: r for r in run.rows}["Г"].total == 25 * 60 + 60 and run.system == "penalty"


def test_nordic_red_card_and_fixed_15_seconds(psr_card):
    comp = replace(psr_card, zachety=[Zachet("Ж", 2, "0840291811Л")])
    z = comp.zachety[0]
    people = [team(n, i) for i, n in enumerate(["А", "Б", "В"], start=1)]
    zdata = {"stages": [{"id": "k1", "tour": "", "name": "Точка 1"}], "spp": "30", "kv": "90", "teams": {
        "А.xlsx": {"start": "10:00:00", "finish": "10:40:00", "points": {"k1": "2"}},  # 2 × 15 с
        "Б.xlsx": {"start": "10:00:00", "finish": "10:35:00", "red": "к"},
        "В.xlsx": {"start": "10:00:00", "finish": "10:41:00", "pen_time": "0:30"}}}   # фальстарт
    run = tr.compute(comp, z, zdata, people)
    rows = {r.inp.team: r for r in run.rows}
    assert run.seconds_per_point == 15 and rows["А"].total == 40 * 60 + 30
    assert rows["Б"].status.value == "removed" and any("результат аннулирован" in i.text for i in run.issues)
    zdata["removal"] = "okv"  # красная карточка — штрафное время, равное ОКВ
    rows = {r.inp.team: r for r in tr.compute(comp, z, zdata, people).rows}
    assert rows["Б"].total == 35 * 60 + 90 * 60 and rows["Б"].place == 3


def test_mountain_points_for_time_technique_tactics(psr_card):
    comp = replace(psr_card, zachety=[Zachet("М/Ж", 3, "0840211811Я", team_size=4)])  # горная группа: 2 балла/мин
    z = comp.zachety[0]
    teams = [team("А", 1), team("Б", 2), team("В", 3), team("Г", 4)]
    zdata = {"stages": STAGES, "kv": "120", "teams": {
        "А.xlsx": {"start": "10:00:00", "finish": "11:00:00", "points": {"s1": "10", "s2": "3"}},
        "Б.xlsx": {"start": "10:00:00", "finish": "10:50:30", "no_tactics": "нет"},           # 50 % ОКВ = 60 мин → 120
        "В.xlsx": {"start": "10:00:00", "finish": "10:40:00", "declared": "50:00"},           # отклонение 20 % → 4 балла
        "Г.xlsx": {"start": "10:00:00", "finish": "10:30:00", "points": {"s1": "с"}}}}
    run = tr.compute(comp, z, zdata, teams)
    rows = {r.inp.team: r for r in run.rows}
    assert run.scoring == "points" and tr.profile(z) == "mountain"
    assert rows["А"].total == 60 * 2 + 13 and rows["В"].total == 40 * 2 + 4
    assert rows["Б"].total == 101 + 120 and pr.result_text(run, rows["Б"]) == "221,00"
    assert [r.inp.team for r in run.rows] == ["В", "А", "Б", "Г"]  # со снятием — после всех без снятий (п. 6.1.6)
    assert rows["В"].percent == 100  # методика % для баллов — из карточки
    pair = replace(comp, zachety=[Zachet("М/Ж", 3, "0840101811Я")])
    rows = {r.inp.team: r for r in tr.compute(pair, pair.zachety[0], zdata, teams).rows}
    assert rows["А"].total == 60 * 4 + 13  # связка: 4 балла/мин


def test_extra_result_part_topography_like_speleo_championship(speleo):
    """Правки, п. 32 (ЧК края 2021, группа спелео, 3 класс): к результату прибавляется «Топосъёмка» — 0:18:05 +
    0:10:39 = 0:28:44. Составляющая баллами — × эквивалент балла; не время — ошибка, место не присуждается."""
    z = speleo.zachety[0]
    zdata = {"stages": STAGES, "expected": "25", "adds": [{"id": "a1", "name": "Топосъёмка", "kind": "time"},
                                                          {"id": "a2", "name": "Описание", "kind": "points"}],
             "teams": {"Кедр.xlsx": {"start": "10:00:00", "finish": "10:18:05", "add-a1": "10:39"},
                       "Сосна.xlsx": {"start": "10:05:00", "finish": "10:20:00", "add-a1": "0:05:00", "add-a2": "2"},
                       "Ель.xlsx": {"start": "10:10:00", "finish": "10:20:00", "add-a1": "десять"}}}
    run = tr.compute(speleo, z, zdata, [team("Кедр", 1), team("Сосна", 2), team("Ель", 3)])
    rows = {r.inp.team: r for r in run.rows}
    assert [a["name"] for a in run.adds] == ["Топосъёмка", "Описание"]
    assert tr.clock_text(rows["Кедр"].total) == "28:44" and rows["Кедр"].extra["add-a1"] == 639
    assert rows["Сосна"].total == 15 * 60 + 5 * 60 + 2 * 15  # 15:00 + 5:00 + 2 балла × 15 с
    assert rows["Ель"].place is None and "add-a1" in rows["Ель"].bad
    assert "«Ель»: Топосъёмка «десять» — не время (мм:сс или ч:мм:сс)" in [i.text for i in run.issues]
    assert [r.inp.team for r in run.rows[:2]] == ["Сосна", "Кедр"]


def test_speleo_protocol_like_sekretar_st(speleo, tmp_path):
    """Правки, п. 33: протокол спелео — как у СЕКРЕТАРЬ_ST: № п/п, состав с разрядами, «п. N» — баллы по пунктам таблицы
    штрафов (только встречавшиеся, от судей), время прохождения, сумма баллов и штрафное время, результат, место,
    %, норматив; ранг — над таблицей; меньше 6 участников — «Разряды не присваиваются…» под таблицей."""
    from datetime import datetime

    from openpyxl import load_workbook

    from st_secretary.exporters.results_protocol import FEW, write_protocol

    z = speleo.zachety[0]
    zdata = {"stages": STAGES, "expected": "25", "teams": {
        "Кедр.xlsx": {"start": "10:00:00", "finish": "10:18:00", "points": {"s1": "0,3", "s2": "1"}},
        "Сосна.xlsx": {"start": "10:05:00", "finish": "10:25:00"}},
        "judge": {"s1": {"Кедр.xlsx": {"pens": [{"code": "7", "v": 0.3, "n": 1}]}},
                  "s2": {"Кедр.xlsx": {"pens": [{"code": "1", "v": 0.5, "n": 2}]}}}}
    run = tr.compute(speleo, z, zdata, [team("Кедр", 1), team("Сосна", 2)])
    assert run.rows[0].by_item == {"7": Fraction(3, 10), "1": Fraction(1)}
    path = write_protocol(speleo, run, "official", datetime(2025, 9, 20, 15, 0), tmp_path / "п.xlsx")
    rows = [[c for c in row if c not in (None, "")] for row in load_workbook(path)["Протокол"].iter_rows(values_only=True)]
    flat = [c for row in rows for c in row]
    assert "МУЖЧИНЫ/ЖЕНЩИНЫ. СМЕШАННЫЕ ГРУППЫ" in flat and "Квалификационный ранг дистанции:" in flat
    i = next(i for i, row in enumerate(rows) if "№ п/п" in row)  # две строки шапки: над пунктами и над итогом
    assert rows[i] == ["№ п/п", "Группа", "Состав группы", "Территория", "Прохождение дистанции", "Результат"]
    assert rows[i + 1] == ["п. 1", "п. 7", "Время прохождения дистанции", "Сумма штрафных баллов на этапах",
                        "Штрафное время на этапах", "Результат", "Место", "% от результата победителя",
                        "Выполненный норматив"]
    kedr = next(row for row in rows if row and row[0] == 1)
    assert kedr[1:3] == ["Кедр", "Кедр 0(II), Кедр 1(II), Кедр 2(II), Кедр 3(II)"]
    assert kedr[4:] == ["1", "0,3", "0:18:00", "1,3", "0:00:19,5", "0:18:19,5", 1, "100,00%", "-"]
    assert FEW in flat  # двое — меньше 6
