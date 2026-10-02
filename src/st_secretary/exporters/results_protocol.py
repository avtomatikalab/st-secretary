"""Протокол результатов ПСР (предварительный и официальный) — Excel, два листа.

«Протокол» — как итоговый протокол СЕКРЕТАРЬ_ST: шапка (проводящие организации, наименование, даты и место,
дисциплина, класс, код ВРВС, группа, квалификационный ранг), места, команды, составы с разрядами, баллы по
турам, результат, процент от победителя, выполненный разряд; подписи главного судьи и главного секретаря.
«По этапам» — баллы каждой команды по каждому этапу: для стенда информации и для протестов.
"""

from __future__ import annotations

import re
from datetime import datetime
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from st_secretary.competition import Competition
from st_secretary.norms import JUNIOR_UNTIL
from st_secretary.psr_run import ZachetRun, points_text, result_text
from st_secretary.results import group_words
from st_secretary.time_run import clock_text, percent_note

THIN = Side(style="thin", color="7F7F7F")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="center", wrap_text=True)
HEAD = PatternFill("solid", fgColor="DCE6F1")

PRELIMINARY, OFFICIAL = "preliminary", "official"


def _pct(x) -> str:
    return f"{float(x):.2f}".replace(".", ",") if x is not None else ""


def _by_stage(wb, comp: Competition, run: ZachetRun, kind: str, at: datetime) -> None:
    """Лист «По этапам» — баллы каждой команды по каждому этапу: для стенда информации и для протестов."""
    z = run.zachet
    timed = run.kind == "time"
    person = run.unit == "person"
    st = wb.create_sheet("По этапам")
    st.cell(1, 1, f"{comp.title} — баллы по этапам, зачёт {z.key}").font = Font(bold=True, size=12)
    what = "Предварительные результаты" if kind == PRELIMINARY else "Результаты" if comp.unofficial else "Официальные результаты"
    st.cell(2, 1, f"{what} на {at:%d.%m.%Y %H:%M}")
    heads = ["Место", "№", "Участник" if person else "Команда"] + [s.title for s in run.stages] + ["Итого"]
    for c, h in enumerate(heads, start=1):
        cell = st.cell(4, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True, size=9), HEAD, BOX
        cell.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True, text_rotation=90 if c > 3 else 0)
    for i, t in enumerate(run.rows, start=5):
        cells = [t.raw.get(s.id, "") if timed else  # «с» — снятие
                 f"{points_text(t.points.get(s.id))} {t.marks[s.id]}" if s.id in t.marks else points_text(t.points.get(s.id))
                 for s in run.stages]
        if run.show_codes:  # номера пунктов таблицы штрафов у баллов — по желанию (настройка зачёта)
            cells = [f"{c} ({t.codes[s.id]})" if s.id in t.codes and c != "" else c for c, s in zip(cells, run.stages,
                                                                                                    strict=True)]
        values = [t.place or "—", t.inp.number or "", t.inp.team] + cells + [result_text(run, t)]
        for c, v in enumerate(values, start=1):
            cell = st.cell(i, c, v)
            cell.border, cell.font = BOX, Font(size=9, bold=c in (1, len(values)))
            cell.alignment = CENTER if c != 3 else WRAP
    st.column_dimensions["C"].width = 22
    st.row_dimensions[4].height = 120
    for c in range(4, len(heads) + 1):
        marked = 4 <= c < 4 + len(run.stages) and any(run.stages[c - 4].id in t.marks for t in run.rows)
        coded = run.show_codes and 4 <= c < 4 + len(run.stages) and any(run.stages[c - 4].id in t.codes for t in run.rows)
        st.column_dimensions[get_column_letter(c)].width = 16 if coded else 13 if marked else 5
    if any(t.marks for t in run.rows):
        st.cell(5 + len(run.rows) + 1, 1, "«снята» — команда снята с этапа, «сверх КВ» — превышено КВ этапа: "
                                          "в клетке МШ этапа.").font = Font(size=9, italic=True)
    st.page_setup.orientation = "landscape"
    st.page_setup.fitToWidth, st.page_setup.fitToHeight = 1, 0
    st.sheet_properties.pageSetUpPr.fitToPage = True


def _published(kind: str, at: datetime, protests_until: datetime | None) -> str:
    """Пометка под таблицей: когда опубликован (и до какого часа протесты) или когда утверждён."""
    if kind == PRELIMINARY:
        return (f"Опубликован {at:%d.%m.%Y в %H:%M}. Протесты по результатам принимаются в течение 1 часа — "
                f"до {protests_until:%H:%M} (Правила, раздел 3, п. 8.17)." if protests_until else
                f"Опубликован {at:%d.%m.%Y в %H:%M}.")
    return f"Результаты утверждены {at:%d.%m.%Y в %H:%M}."


def _signatures(ws, comp: Competition, r: int) -> int:
    """Подписи главного судьи и главного секретаря; возвращает следующую свободную строку."""
    for role in ("Главный судья", "Главный секретарь"):
        o = comp.official(role)
        ws.cell(r, 1, f"{role} ________________ / {o.signature if o else ' ' * 30} /").font = Font(size=11)
        r += 2
    return r


def _hms(sec) -> str:
    """Время, как в протоколах СЕКРЕТАРЬ_ST: «0:18:05» (с часами; десятые — если есть)."""
    if sec is None:
        return ""
    whole = int(sec)
    out = f"{whole // 3600}:{whole % 3600 // 60:02d}:{whole % 60:02d}"
    return out + (f",{int((sec - whole) * 10)}" if sec != whole else "")


def _item_key(code: str):
    """Пункты по порядку: 7, 9, 11, 11.2 — числа по значению, остальное — после."""
    parts = re.findall(r"\d+", code)
    return (0, [int(x) for x in parts], code) if parts else (1, [], code)


FEW = ("Разряды не присваиваются, т.к. в соответствии с положением о ЕВСК (пункт 24) на дистанции принимало участие "
       "менее 6 участников (связок, групп)")


def _speleo(ws, comp: Competition, run: ZachetRun, kind: str, at: datetime, protests_until: datetime | None) -> None:
    """Протокол спелео — как у СЕКРЕТАРЬ_ST (Правки, п. 33; протоколы ЧК и ПК Красноярского края 2021): № п/п, участник
    (год, разряд, делегация) или состав с разрядами, территория; «Прохождение дистанции» — штрафные баллы по пунктам
    таблицы штрафов (только встречавшиеся); «Результат» — время прохождения, сумма баллов и штрафное время (если
    штрафы были), составляющие («Топосъёмка»), результат, место, % от победителя, норматив. Ранг — над таблицей,
    «разряды не присваиваются» — под ней. Что сделано иначе и почему — docs/decisions.md, решение 043."""
    z = run.zachet
    person, pair = run.unit == "person", run.unit == "pair"
    items = sorted({c for t in run.rows for c in t.by_item}, key=_item_key)
    pens = any(t.points and sum(t.points.values(), Fraction(0)) for t in run.rows)
    removals = any(t.removals for t in run.rows)
    norms = not comp.unofficial
    who = (["Участник", "Год", "Разряд", "Делегация", "Территория"] if person else
           ["Связка" if pair else "Группа", f"Состав {'связки' if pair else 'группы'}", "Территория"])
    result = (["Время прохождения дистанции"]
              + (["Сумма штрафных баллов на этапах", "Штрафное время на этапах"] if pens else [])
              + (["Снятий с этапов"] if removals else []) + [a["name"] for a in run.adds] + ["Результат", "Место"]
              + (["% от результата победителя", "Выполненный норматив"] if norms else []))
    head = ["№ п/п", *who, *[f"п. {c}" for c in items], *result]
    width = len(head)

    def line(r, text, bold=False, size=11, align="center"):
        c = ws.cell(r, 1, text)
        c.font, c.alignment = Font(bold=bold, size=size), Alignment(horizontal=align, wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)

    r = 1
    for org in comp.organizers:
        line(r, org, size=10)
        r += 1
    line(r, comp.title, True, 14)
    ws.cell(r + 1, 1, comp.dates_text).font = Font(italic=True, size=10)
    ws.cell(r + 1, width, comp.place).font = Font(italic=True, size=10)
    ws.cell(r + 1, width).alignment = Alignment(horizontal="right")
    line(r + 2, "Предварительный протокол соревнований" if kind == PRELIMINARY else "Протокол соревнований", True, 13)
    disc = (f"в дисциплине: \"{z.discipline_name}\" {z.distance_class} класса, код ВРВС {z.discipline_code}"
            if not z.is_custom else z.header_text)
    line(r + 3, disc, True, 12)
    line(r + 4, (f"Зачёт «{z.name}». " if z.name else "") + group_words(z.group, z.rank_format, long=True)
         .upper().replace(", ", ". "), True, 12)
    if not comp.unofficial:
        ws.cell(r + 5, 2, "Квалификационный ранг дистанции:").font = Font(size=10)
        ws.cell(r + 5, 4, run.rank.formatted() if run.rank and run.rank.value is not None else "не подсчитывался")
    r += 7
    # две строки шапки: «Прохождение дистанции» над пунктами, «Результат» над итогом
    first_item, first_res = 2 + len(who), 2 + len(who) + len(items)
    for c, h in enumerate(head, start=1):
        top = c < first_item  # участник, территория — на обе строки шапки
        cell = ws.cell(r if top else r + 1, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True, size=9), HEAD, BOX
        cell.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True,
                                   text_rotation=0 if top else 90)
        if top:
            ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c)
    for c0, c1, text in ((first_item, first_res - 1, "Прохождение дистанции"), (first_res, width, "Результат")):
        if c1 >= c0:
            cell = ws.cell(r, c0, text)
            cell.font, cell.fill, cell.border, cell.alignment = Font(bold=True, size=9), HEAD, BOX, CENTER
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
    ws.row_dimensions[r + 1].height = 110
    r += 2
    spp = run.seconds_per_point or 30
    for n, t in enumerate(run.rows, start=1):
        finished = t.place is not None
        if person:
            m = t.inp.members[0] if t.inp.members else None
            who_values = [t.inp.team, (m.birth or "")[:4] if m else "", (m.qual_label or "б/р") if m else "",
                          t.inp.club, t.inp.territory]
        else:
            who_values = [t.inp.club or t.inp.team, ", ".join(f"{m.fio}({m.qual_label or 'б/р'})" for m in t.inp.members),
                          t.inp.territory]
        pts = sum(t.points.values(), Fraction(0)) if t.points else Fraction(0)
        res = ([_hms(t.distance_time)]
               + ([points_text(pts), _hms(pts * spp) if pts else ""] if pens else [])
               + ([t.removals or ""] if removals else [])
               + [(_hms(t.extra[f"add-{a['id']}"]) if a["kind"] == "time" else points_text(t.extra[f"add-{a['id']}"]))
                  if f"add-{a['id']}" in t.extra else "" for a in run.adds]
               + [_hms(t.total) if finished else result_text(run, t) or "—", t.place or "—"]
               + ([f"{_pct(t.percent)}%" if t.percent is not None else "", _norm_cell(t.norm) or "-"] if norms else []))
        values = [n, *who_values, *[points_text(t.by_item.get(c)) if c in t.by_item else "" for c in items], *res]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            bold = head[c - 1] in ("Результат", "Место", "% от результата победителя")
            cell.border, cell.font = BOX, Font(size=10, bold=bold)
            cell.alignment = WRAP if 2 <= c < first_item and head[c - 1] not in ("Год", "Разряд") else CENTER
        r += 1
    r += 1
    notes = [_published(kind, at, protests_until)]
    if comp.unofficial:
        notes.append("Неофициальные соревнования: квалификационный ранг и разряды не определяются.")
    elif run.norms_why:
        notes.append(FEW if "менее" in run.norms_why else f"Разряды не присваиваются: {run.norms_why}.")
    if norms and percent_note(run.adds):
        notes.append(percent_note(run.adds))
    if norms and any(_junior(t.norm) for t in run.rows):
        notes.append(JUNIOR_NOTE)
    for text in notes:
        ws.cell(r, 1, text).font = Font(size=10, italic=True)
        r += 1
    r += 1
    r = _signatures(ws, comp, r)
    widths = ([5] + ([24, 6, 7, 24, 16] if person else [22, 46, 16]) + [5] * len(items)
              + [9] * (len(result) - (2 if norms else 0)) + ([9, 9] if norms else []))
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.page_setup.orientation = "landscape" if not person else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


JUNIOR_NOTE = f"* Юношеские спортивные разряды присваиваются спортсменам до {JUNIOR_UNTIL} лет."


def _junior(norm: str) -> bool:
    return any(x.strip().endswith("ю") for x in norm.split(","))


def _norm_cell(norm: str) -> str:
    """Норматив в клетке протокола: юношеские — со звёздочкой («2ю*»), как у СЕКРЕТАРЬ_ST (ПК края 2021)."""
    return ", ".join(x.strip() + ("*" if x.strip().endswith("ю") else "") for x in norm.split(",")) if norm else ""


def write_protocol(comp: Competition, run: ZachetRun, kind: str, at: datetime, path: str | Path,
                   protests_until: datetime | None = None) -> Path:
    """Протокол результатов (Excel): лист «Протокол» — у спелео как у СЕКРЕТАРЬ_ST (Правки, п. 33), у остальных —
    общий вид; лист «По этапам» — баллы по этапам для стенда и протестов."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Протокол"
    (_speleo if run.profile == "speleo" else _general)(ws, comp, run, kind, at, protests_until)
    _by_stage(wb, comp, run, kind, at)
    path = Path(path)
    wb.save(path)
    return path


def _general(ws, comp: Competition, run: ZachetRun, kind: str, at: datetime, protests_until: datetime | None) -> None:
    """Лист «Протокол» общего вида (ПСР, пешеходные, СХ, горные): шапка, места, составы, баллы по турам или время,
    результат, процент и разряд; пометки и подписи."""
    z = run.zachet
    timed = run.kind == "time"  # спелео, пешеходные: время на дистанции, штраф, снятия
    tours = [] if timed else run.tours
    middle = (["Время на дистанции", "Штраф, баллы", "Снятий"] + [a["name"] for a in run.adds]) if timed else tours
    show_class = any(r.actual_class for r in run.rows)
    show_marks = not timed and any(r.marks for r in run.rows)  # ПСР: сняты с этапов (МШ) — отметка в протоколе
    person = run.unit == "person"  # личная дисциплина: место у спортсмена
    who = ["Участник", "Команда", "Территория", "Разряд"] if person else ["Команда", "Территория", "Состав (разряд)"]
    norms = [] if comp.unofficial else ["% от победителя", "Выполнен разряд"]  # неофициальные — без разрядов
    head = (["Место", "№"] + who + middle
            + (["Снятия с этапов"] if show_marks else [])
            + ["Результат"] + norms + (["Факт. класс"] if show_class else []))
    width = len(head)

    def line(r, text, bold=False, size=11):
        c = ws.cell(r, 1, text)
        c.font, c.alignment = Font(bold=bold, size=size), Alignment(horizontal="center", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)

    r = 1
    for org in comp.organizers:
        line(r, org.upper(), size=10)
        r += 1
    line(r, comp.title, True, 13)
    line(r + 1, f"{comp.dates_text}, {comp.place}")
    title = "ПРЕДВАРИТЕЛЬНЫЙ ПРОТОКОЛ РЕЗУЛЬТАТОВ" if kind == PRELIMINARY else "ПРОТОКОЛ РЕЗУЛЬТАТОВ"
    line(r + 2, title, True, 13)
    line(r + 3, z.header_text)
    rank = run.rank.formatted() if run.rank else "не определялся"
    zname = f"Зачёт «{z.name}». " if z.name else ""
    line(r + 4, f"{zname}Группа: {group_words(z.group, z.rank_format, long=True)}"
         + ("" if comp.unofficial else f". Квалификационный ранг: {rank}"))
    r += 6
    for c, h in enumerate(head, start=1):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill, cell.border, cell.alignment = Font(bold=True, size=10), HEAD, BOX, CENTER
    r += 1
    for t in run.rows:
        members = ", ".join(f"{m.fio} ({m.qual_label or 'б/р'})" for m in t.inp.members)
        who_values = ([t.inp.team, t.inp.club, t.inp.territory, ", ".join(m.qual_label or "б/р" for m in t.inp.members)]
                      if person else [t.inp.team, t.inp.territory, members])
        result = result_text(run, t) or ("время не внесено" if timed else "баллы не внесены")
        adds = [(clock_text(t.extra.get(f"add-{a['id']}")) if a["kind"] == "time"
                 else points_text(t.extra.get(f"add-{a['id']}"))) if f"add-{a['id']}" in t.extra else ""
                for a in run.adds]  # дополнительные составляющие результата (Правки, п. 32)
        mid = ([clock_text(t.distance_time), points_text(sum(t.points.values())) if t.points else "", t.removals or "",
                *adds] if timed else [points_text(t.tours.get(x)) for x in tours])
        values = ([t.place or "—", t.inp.number or ""] + who_values + mid
                  + ([mark_text(run, t)] if show_marks else [])
                  + [result] + ([_pct(t.percent), _norm_cell(t.norm)] if norms else [])
                  + ([t.actual_class or ""] if show_class else []))
        res_col = 3 + len(who) + len(middle) + (1 if show_marks else 0)
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.border, cell.font = BOX, Font(size=10, bold=(c in (1, res_col)))
            cell.alignment = WRAP if 3 <= c < 3 + len(who) or (show_marks and c == res_col - 1) else CENTER
        r += 1
    r += 1
    notes = [_published(kind, at, protests_until)]
    if comp.unofficial:
        notes.append("Неофициальные соревнования: квалификационный ранг и разряды не определяются.")
    elif run.rank and run.rank.value is None and run.rank.reason:
        notes.append(f"Квалификационный ранг не определялся: {run.rank.reason}.")
    if norms and percent_note(run.adds):
        notes.append(percent_note(run.adds))
    if norms and any(_junior(t.norm) for t in run.rows):
        notes.append(JUNIOR_NOTE)
    for n in notes:
        ws.cell(r, 1, n).font = Font(size=10, italic=True)
        r += 1
    r += 1
    r = _signatures(ws, comp, r)
    for c, w in enumerate([7, 6] + ([28, 20, 16, 8] if person else [22, 16, 46]) + [10 if timed else 8] * len(middle)
                          + ([22] if show_marks else []) + [11] + ([10, 10] if norms else []) + ([8] if show_class else []),
                          start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def mark_text(run: ZachetRun, t) -> str:
    """ПСР: с каких этапов команда снята (в клетке МШ) — «Узлы; Бивак (сверх КВ)»."""
    return "; ".join(s.name + ("" if t.marks[s.id] == "снята" else f" ({t.marks[s.id]})")
                     for s in run.stages if s.id in t.marks)


def awards_rows(run: ZachetRun, source: str) -> dict:
    """Официальные результаты → данные страницы «Награждение» (как из протокола СЕКРЕТАРЬ_ST); если нормативы в
    составе разные (по возрасту, п. 46) — у каждого участника свой «norm»."""
    rows = []
    for t in run.rows:  # спортсмен и связка: «команда» — их команда, награждаются участники
        own = len(set(t.member_norms)) > 1
        members = [{"fio": m.fio, "qual": m.qual_label} | ({"norm": n} if own else {})
                   for m, n in zip(t.inp.members, t.member_norms or [""] * len(t.inp.members), strict=True)]
        rows.append({"team": t.inp.club or t.inp.team, "territory": t.inp.territory, "number": t.inp.number,
                     "place": t.place, "result": result_text(run, t), "norm": t.norm, "members": members})
    rank = run.rank.formatted() if run.rank and run.rank.value is not None else ""
    return {"source": source, "group_text": "", "rank": rank, "rows": rows}
