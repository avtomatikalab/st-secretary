"""Сверка документов — на выдуманных документах: каждая ошибка ловится, свои документы программы проходят чисто."""

from docx import Document
from openpyxl import Workbook
from test_contracts import CUSTOMER, PERSON, team_data

from st_secretary import staff as st
from st_secretary import verify as vf
from st_secretary.exporters import contracts as ct
from st_secretary.issues import ERROR, WARNING

PEOPLE = [vf.Known("Судьин Иван Петрович", "СС1К", "в карточке"),
          vf.Known("Секретарёва Анна Ивановна", "СС2К", "в карточке"),
          vf.Known("Лебедев Антон Игоревич", "", "в заявке")]


def docx(tmp_path, name, lines):
    d = Document()
    for line in lines:
        d.add_paragraph(line)
    p = tmp_path / name
    d.save(str(p))
    return p


def texts(issues, sev=None):
    return [i.text for i in issues if sev is None or i.severity == sev]


def test_single_document_checks(tmp_path, psr_card):
    p = docx(tmp_path, "Отчёт.docx", [
        "ОТЧЁТ главного судьи Чемпионата города N по спортивному туризму",
        "в период с «04» октября 2024 г. по «06» октября 2024 г.",
        "Спортивная дисциплина «дистанция - комбинированые», код ВРВС 0840271811Я; также 0840999999Я",
        "Первенство города N по спортивному туризму — следующий старт",
        "Главный судья ________ / И.П. Судьин, СС2К, г. Красноярск /",
        "Главный секретарь ________ / А.П. Секретарёва /",
        "Лебедев Антон Игорьевич — выполнил II разряд; Секретарёва АннаИвановна",
        "Сумма 4250 (четыре тысячи двести пятьдесят пять) рублей 00 копеек",
        "Дата выдачи паспорта 25.06.2011; 30 сентября 2025 г. — отчёт сдан",
    ])
    doc = vf.read_doc(p)
    issues = vf.verify([doc], psr_card, PEOPLE).files[0][1]
    t = texts(issues)
    assert "дата ««04» октября 2024» — не 2025 год" in t and "дата ««06» октября 2024» — не 2025 год" in t
    assert not any("25.06.2011" in x or "30 сентября" in x for x in t)  # паспорт и дата отчёта — не ошибки
    assert any(x.startswith("код ВРВС 0840271811Я — это «дистанция - спелео - группа»") for x in t)
    assert any(x.startswith("код ВРВС 0840999999Я — такого кода нет во ВРВС") for x in t)
    assert "«комбинированые» — похоже на опечатку: во ВРВС «комбинированная»" in t
    assert any("«Первенство города N по спортивному туризму» не как в карточке" in x for x in t)
    assert not any("Чемпионата города N" in x for x in t)  # то же название в родительном падеже — не ошибка
    assert "Судьин Иван Петрович: категория СС2К, а в карточке — СС1К" in t
    assert "«А.П. Секретарёва» — инициалы не сходятся: в карточке «Секретарёва Анна Ивановна»" in t
    assert "«Лебедев Антон Игорьевич» — в документе, а «Лебедев Антон Игоревич» (в заявке)" in texts(issues, ERROR)
    assert "«Секретарёва АннаИвановна» — имя и отчество слитно" in texts(issues, WARNING)
    assert any("прописью «четыре тысячи двести пятьдесят пять» — не та" in x for x in t)
    assert all(i.source == "Отчёт.docx" and i.before for i in issues)  # у каждого замечания — цитата
    assert len(issues) == 11


def tabel_xlsx(tmp_path):
    """Табель как в 2025 г.: шапка в две строки, отметки «р», числа — значения, а не формулы."""
    wb = Workbook()
    ws = wb.active
    ws.append(["ТАБЕЛЬ – НАРЯД"])
    ws.append(["№", "Фамилия, имя, отчество", "Должность", "Дни месяца", None, None, "ИТОГО", None, None])
    ws.append([None, None, None, "19 сентября", "20 сентября", "21 сентября", "кол-во дней", "оплата за (день)",
               "общая сумма (руб.)"])
    ws.append([1, "Судьин Иван Петрович", "гл.судья", "р", "р", "р", 3, 850, 2550])
    ws.append([2, "Работяга Семён Ильич", "рабочий", None, "р", "р", 3, 460, 1380])  # отмечено 2, к оплате 3
    ws.append([3, "Секретарёва Анна Ивановна", "гл. секретарь", None, "р", "р", 2, 700, 1500])  # 2 × 700 ≠ 1500
    ws.append(["Итого начислено", None, None, None, None, None, None, None, 5000])
    ws.append([None, "Гл. судья ______ / Судьин И.П."])
    p = tmp_path / "табель.xlsx"
    wb.save(p)
    return p


def test_tabel_contracts_and_acts(tmp_path, psr_card):
    tabel = vf.read_doc(tabel_xlsx(tmp_path))
    rows, issues = vf.parse_tabels(tabel)
    assert [(r.fio, r.marks, r.days) for r in rows] == [("Судьин Иван Петрович", 3, 3), ("Работяга Семён Ильич", 2, 3),
                                                        ("Секретарёва Анна Ивановна", 2, 2)]
    assert texts(issues) == ["табель, Работяга Семён Ильич: отмечено дней «р» — 2, а к оплате — 3",
                             "табель, Секретарёва Анна Ивановна: 2 дн. × 700 = 1400, а в сумме — 1500",
                             "табель: «Итого начислено» — 5000, а сумма по людям — 5430"]

    team = st.people(psr_card, team_data())
    judge = ct.contract_values(psr_card, team[0], CUSTOMER, PERSON)  # 4 дня × 850 = 3400, срок 19–22 сентября
    ct.write_contract(None, judge, tmp_path / "Судьин.docx")
    d = Document(str(tmp_path / "Судьин.docx"))
    for par in d.paragraphs:  # в договоре — срок соревнований, а в акте остался срок работы: как в 2025 г.
        if par.text.startswith("1.2. Срок"):
            par.runs[0].text = "1.2. Срок оказания услуг: 20–21 сентября 2025 г."
    d.save(str(tmp_path / "Судьин.docx"))
    contract = vf.read_doc(tmp_path / "Судьин.docx")
    info = vf.parse_contracts(contract)
    assert len(info) == 1 and info[0].fio == "Судьин Иван Петрович" and info[0].amount == 3400
    assert (info[0].act_days, info[0].act_price, info[0].act_sum) == (4, 850, 3400)
    rep = vf.verify([tabel, contract], psr_card, PEOPLE)
    c_issues = texts(rep.files[1][1])
    assert ("договор Судьин Иван Петрович: срок в договоре «20–21 сентября 2025 г», а в акте — «19–22 сентября 2025»"
            in c_issues)
    cross = texts(rep.cross)
    assert "Судьин Иван Петрович: в акте 4 дн., а в табеле — 3" in cross
    assert "Судьин Иван Петрович: сумма по договору 3400, в табеле — 2550" in cross
    assert ("есть в табеле, а договоров среди проверенных файлов нет — 2 чел.: Работяга Семён Ильич (1380 руб.), "
            "Секретарёва Анна Ивановна (1500 руб.)") in cross


def test_program_documents_are_consistent(tmp_path, psr_card):
    """Табель, договоры и акты программы между собой и с карточкой — без ошибок."""
    data = team_data()
    team = st.people(psr_card, data)
    s = st.settings(psr_card, data)
    ct.write_tabel(psr_card, team, s["days"], 30, CUSTOMER, {}, tmp_path / "Табель-наряд.xlsx")
    items = [(None, ct.contract_values(psr_card, p, CUSTOMER, PERSON if p.from_card else {})) for p in team if p.paid]
    ct.write_contracts(items, tmp_path / "Договоры и акты.docx")
    docs = [vf.read_doc(tmp_path / n) for n in ("Табель-наряд.xlsx", "Договоры и акты.docx")]
    people = PEOPLE + [vf.Known("Работяга Семён Ильич", "", "в табеле программы")]
    rep = vf.verify(docs, psr_card, people)
    assert rep.contracts == 3 and rep.tabel_rows == 3
    assert [i.text for i in rep.all_issues if i.severity == ERROR] == []


def test_unreadable_files(tmp_path, psr_card):
    (tmp_path / "старый.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    (tmp_path / "битый.docx").write_bytes(b"not a zip")
    rep = vf.verify([vf.read_doc(tmp_path / "старый.doc"), vf.read_doc(tmp_path / "битый.docx")], psr_card, PEOPLE)
    assert rep.files[0][0].error.startswith("старый формат") and rep.files[1][0].error.startswith("файл не открывается")
