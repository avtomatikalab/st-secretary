"""Договоры, акты и табель судейской бригады — на выдуманных людях и реквизитах."""

from datetime import date

import pytest
from docx import Document
from openpyxl import load_workbook

from st_secretary import staff as st
from st_secretary.exporters import contracts as ct


def valid_inn(base10: str) -> str:
    d = [int(c) for c in base10]
    d.append(sum(w * x for w, x in zip((7, 2, 4, 10, 3, 5, 9, 4, 6, 8), d)) % 11 % 10)
    d.append(sum(w * x for w, x in zip((3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8), d)) % 11 % 10)
    return "".join(map(str, d))


def valid_snils(base9: str) -> str:
    s = sum(int(c) * (9 - i) for i, c in enumerate(base9))
    check = s if s < 100 else 0 if s in (100, 101) else s % 101 % 100
    return f"{base9[:3]}-{base9[3:6]}-{base9[6:]} {check:02d}"


def valid_account(bik: str, base19: str) -> str:
    for k in range(10):
        acc = base19[:8] + str(k) + base19[8:]
        if st.account_ok(acc, bik):
            return acc
    raise AssertionError


BIK = "040000001"
PERSON = {"birth": "01.02.1990", "passport": "04 11 123456", "issued_by": "Отделом УФМС города N",
          "issued_on": "05.06.2010", "dept_code": "240-001", "address": "660000, г. N, ул. Лесная, д. 1, кв. 2",
          "inn": valid_inn("2400000001"), "snils": valid_snils("112233445"),
          "account": valid_account(BIK, "4081781000000000001"), "bank": "Банк города N", "bik": BIK,
          "phone": "+7 900 000-00-00", "judge_id": "0123"}
CUSTOMER = {"name": "Муниципальное автономное учреждение «Центр спорта города N»", "short": "МАУ «ЦС»",
            "head_post": "директор", "head_fio": "Начальников Пётр Сергеевич", "basis": "Устава",
            "requisites": "660000, г. N, ул. Главная, 1\nИНН 2400000000", "rates_basis": "норм расходов города N",
            "funding": "за счёт субсидии из бюджета города N"}


def team_data():
    return {"period": {"from": "2025-09-19", "to": "2025-09-22"},
            "rates": {"главный судья|1": 850, "главный секретарь|2": 700, "рабочий комендантской бригады|б/к": 460},
            "extra": [{"fio": "Работяга Семён Ильич", "role": "Рабочий комендантской бригады", "category": "б/к"},
                      {"fio": "Помощница Ольга Павловна", "role": "Волонтёр", "category": "б/к"}],
            "people": {"судьин иван петрович": {"days": ["2025-09-19", "2025-09-20", "2025-09-21", "2025-09-22"]},
                       "помощница ольга павловна": {"unpaid": True}}}


def test_people_days_rates_and_totals(psr_card):
    s = st.settings(psr_card, team_data())
    assert [d.day for d in s["days"]] == [19, 20, 21, 22]
    team = st.people(psr_card, team_data())
    assert [p.fio for p in team] == ["Судьин Иван Петрович", "Секретарёва Анна Ивановна", "Работяга Семён Ильич",
                                     "Помощница Ольга Павловна"]
    judge, secretary, worker, helper = team
    assert judge.amount == 4 * 850 and judge.span_text == "19–22 сентября 2025 г."
    assert [d.day for d in secretary.days] == [20, 21]  # по умолчанию — дни соревнований
    assert secretary.amount == 1400 and worker.amount == 920
    assert helper.unpaid and not helper.paid and helper.amount == 0
    assert judge.role_text == "главного судьи 1 категории"
    assert secretary.role_text == "главного секретаря 2 категории"
    assert worker.role_text == "рабочего комендантской бригады"
    t = st.totals(team, 30)
    assert t == {"people": 3, "days": 8, "sum": 5720, "accrual": 1716.0, "total": 7436.0}
    assert st.default_period(psr_card) == (date(2025, 9, 19), date(2025, 9, 22))


def test_checks(psr_card):
    data = team_data()
    data["rates"].pop("главный секретарь|2")
    data["people"]["работяга семен ильич"] = {"days": ["2025-09-18", "2025-09-20"]}
    team = st.people(psr_card, data)
    bad = dict(PERSON, inn="240000000199", snils="112-233-445 00", account=PERSON["account"][:-1] + "0")
    issues = st.check(team, {"судьин иван петрович": PERSON, "работяга семен ильич": bad}, {})
    texts = [i.text for i in issues]
    assert any("Секретарёва Анна Ивановна: не задана ставка" in t for t in texts)
    assert any("Работяга Семён Ильич: отмечены дни вне периода" in t and "18.09" in t for t in texts)
    assert any("Работяга Семён Ильич: инн — ИНН не сходится" in t for t in texts)
    assert any("снилс — СНИЛС не сходится" in t for t in texts)
    assert any("номер счёта — номер счёта не сходится с БИК" in t for t in texts)
    assert any("Секретарёва Анна Ивановна: для договора нет данных" in t for t in texts)
    assert not any(t.startswith("Судьин") for t in texts)  # всё заполнено и сходится
    assert any("не заполнены сведения о заказчике" in t for t in texts)
    assert not any("Помощница" in t for t in texts)  # без оплаты — договор не нужен


def test_personal_validation():
    assert st.inn_ok(PERSON["inn"]) and not st.inn_ok("123456789012")
    assert st.snils_ok(PERSON["snils"]) and not st.snils_ok("112-233-445 01")
    assert st.account_ok(PERSON["account"], BIK)
    assert st.personal_problems(dict(PERSON, passport="04 11 12345", issued_on="31.02.2010")) == [
        ("passport", "в серии и номере паспорта должно быть 10 цифр"),
        ("issued_on", "дата — в виде дд.мм.гггг, и такая дата должна существовать")]
    assert st.personal_problems({}) == [] and len(st.missing_personal({})) == len(st.FOR_CONTRACT)


def test_tabel(tmp_path, psr_card):
    team = st.people(psr_card, team_data())
    s = st.settings(psr_card, team_data())
    ws = load_workbook(ct.write_tabel(psr_card, team, s["days"], 30, CUSTOMER,
                                      {"судьин иван петрович": PERSON}, tmp_path / "t.xlsx")).active
    cells = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value not in (None, "")}
    assert cells["A1"] == "ТАБЕЛЬ – НАРЯД"
    assert cells["A3"] == ("Чемпионата города N по спортивному туризму в спортивной дисциплине "
                           "«дистанция - комбинированная»")
    assert "Директор МАУ «ЦС»" in cells.values() and "______________ П.С. Начальников" in cells.values()
    assert [cells[f"{c}7"] for c in "GHIJ"] == ["19.09", "20.09", "21.09", "22.09"]
    assert [cells.get(f"{c}8") for c in "ABCDEF"] == [1, "Судьин Иван Петрович", "Главный судья", "1990", "1", "0123"]
    assert [cells.get(f"{c}8") for c in "GHIJ"] == ["р", "р", "р", "р"]
    assert cells["K8"] == '=COUNTIF(G8:J8,"р")' and cells["L8"] == 850 and cells["M8"] == "=K8*L8"
    names = [cells.get(f"B{r}") for r in range(8, 11)]
    assert names == ["Судьин Иван Петрович", "Секретарёва Анна Ивановна", "Работяга Семён Ильич"]  # без «без оплаты»
    assert cells["B11"] == "Итого начислено" and cells["M11"] == "=SUM(M8:M10)"
    assert cells["B12"] == "Начисления на оплату 30 %" and cells["M12"] == "=ROUND(M11*30/100,2)"
    assert cells["M13"] == "=M11+M12"
    assert any(str(v).startswith("Главный судья ______________ / И.П. Судьин /") for v in cells.values())


def texts_of(doc) -> str:
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            out += [c.text for c in row.cells]
    return "\n".join(out)


def test_contract_default_template(tmp_path, psr_card):
    judge = st.people(psr_card, team_data())[0]
    values = ct.contract_values(psr_card, judge, CUSTOMER, PERSON)
    assert values["В лице"] == "директора Начальникова Петра Сергеевича"
    assert values["Сумма прописью"] == "3400 (три тысячи четыреста) рублей 00 копеек"
    unknown = ct.write_contract(None, values, tmp_path / "d.docx")
    assert unknown == set()
    t = texts_of(Document(str(tmp_path / "d.docx")))
    assert "{{" not in t
    assert ("Муниципальное автономное учреждение «Центр спорта города N», в лице директора Начальникова Петра "
            "Сергеевича, действующего на основании Устава") in t
    assert "в качестве главного судьи 1 категории на спортивных соревнованиях: Чемпионат города N" in t
    assert "1.2. Срок оказания услуг: 19–22 сентября 2025 г." in t
    assert "(850,00 руб. в день, дней — 4) и составляет 3400 (три тысячи четыреста) рублей 00 копеек" in t
    assert "Главный судья 1 категории — Чемпионат города N" in t and "3 400,00" in t
    assert "ИНН: " + PERSON["inn"] in t and "/ Судьин И.П. /" in t and "/ Начальников П.С. /" in t
    assert "г. Красноярск" in t  # место заключения — по территории организаторов


def test_blank_personal_and_custom_template(tmp_path, psr_card):
    worker = st.people(psr_card, team_data())[2]
    values = ct.contract_values(psr_card, worker, {}, {})
    assert values["Паспорт"] == ct.BLANK and values["Заказчик"] == ct.BLANK
    doc = Document()
    p = doc.add_paragraph("Исполнитель: ")
    p.add_run("{{Ф")  # Word часто разбивает текст поля на куски
    p.add_run("ИО}}, сумма {{Сумма}}; {{Нечто}}")
    tpl = tmp_path / "Шаблон договора.docx"
    doc.save(str(tpl))
    assert ct.template_fields(tpl) == {"ФИО", "Сумма", "Нечто"}
    unknown = ct.write_contract(tpl, values, tmp_path / "w.docx")
    assert unknown == {"Нечто"}
    assert Document(str(tmp_path / "w.docx")).paragraphs[0].text == \
        "Исполнитель: Работяга Семён Ильич, сумма 920,00; {{Нечто}}"


def test_all_contracts_one_file(tmp_path, psr_card):
    team = [p for p in st.people(psr_card, team_data()) if p.paid]
    many = [ct.contract_values(psr_card, p, CUSTOMER, {}) for p in team]
    ct.write_contracts(None, many, tmp_path / "all.docx")
    t = texts_of(Document(str(tmp_path / "all.docx")))
    assert t.count("ДОГОВОР № ______") == 3 and t.count("АКТ № ______") == 3
    for p in team:
        assert p.fio in t


@pytest.mark.parametrize("role, cat, key", [("Главный судья", "СС1К", "главный судья|1"),
                                            ("Рабочий", "", "рабочий|б/к"), ("Судья  этапа", "СС3К", "судья этапа|3")])
def test_rate_key(role, cat, key):
    assert st.rate_key(role, cat) == key
