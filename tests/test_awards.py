"""Итоги и документы награждения: протокол СЕКРЕТАРЬ_ST → места; дипломы, наклейки, список награждаемых."""

from datetime import date
from fractions import Fraction
from pathlib import Path
from types import SimpleNamespace

import pytest
from conftest import make_application
from docx import Document
from openpyxl import load_workbook

from st_secretary import results as res
from st_secretary.exporters import awards as aw
from st_secretary.importers.preapp_xlsx import read_preapplication
from st_secretary.importers.sekretar_xls import ResultRow
from st_secretary.preapp import process


def protocol():
    """Итоговый протокол, как его читает importers.sekretar_xls (люди выдуманы)."""
    rows = [
        ResultRow("3", "Кедр", "Лебедев Антон Игоревич(III), Зуева Мария Олеговна(III), Носов Глеб Андреевич(б/р)",
                  "Красноярск", {}, Fraction(-283), "1", None, "II"),
        ResultRow("9", "Сосна", "Орлов Павел Ильич(КМС), Белова Ирина Петровна(б/р), Пестов Юрий Андреевич(КМС)",
                  "Томск", {}, Fraction(-213), "2", None, "II"),
        ResultRow("6", "Ель", "Ким Олег Борисович(КМС), Лис Ева Павловна(I), Рыбин Иван Ильич(КМС)",
                  "Красноярск", {}, Fraction(-60), "3", None, ""),
        ResultRow("2", "Пихта", "Сом Яков Ильич(б/р), Ёж Инна Олеговна(б/р), Кот Лев Ильич(б/р)",
                  "Красноярск", {}, None, "", None, ""),
    ]
    head = ["ПРОВОДЯЩИЕ", "ЧЕМПИОНАТ Г. N ПО СПОРТИВНОМУ ТУРИЗМУ", "20-21 сентября 2025 г.", "окрестности г. N",
            'Протокол соревнований в дисциплине: "дистанция - комбинированная" 3 класса, код ВРВС 0840161811Я '
            "МУЖЧИНЫ/ЖЕНЩИНЫ. СМЕШАННЫЕ ГРУППЫ", "Квалификационный ранг дистанции: 71,7"]
    return SimpleNamespace(path=Path("Result_ПСР.xls"), header_lines=head, rows=rows, rank_text="71,7")


def zres(psr_card, preapps=None):
    return res.load({"zachety": {"М/Ж_3": res.from_protocol(protocol())}}, psr_card, preapps)


def test_protocol_import(tmp_path, psr_card):
    d = res.from_protocol(protocol())
    assert d["group_text"] == "МУЖЧИНЫ/ЖЕНЩИНЫ. СМЕШАННЫЕ ГРУППЫ" and d["rank"] == "71,7"
    assert d["rows"][0]["members"][0] == {"fio": "Лебедев Антон Игоревич", "qual": "III"}
    assert d["rows"][3]["place"] is None and d["rows"][1]["result"] == "-213"
    # дата рождения и пол — из заявки: в протоколе их нет
    a = make_application(tmp_path / "Кедр.xlsx", "Кедр", "Красноярск", "Лебедев Антон Игоревич", "89135550000", 3, [
        ["Кедр", "Красноярск", "Лебедев Антон Игоревич", "Лебедев Антон Игоревич", "02.02.1990", "III", "м", "М/Ж", 3]])
    [z] = zres(psr_card, process([read_preapplication(a)], psr_card))
    first = z.rows[0].members[0]
    assert first.birth == date(1990, 2, 2) and first.sex == "м" and first.first_last == "Антон Лебедев"
    assert [p.team for p in z.medalists] == ["Кедр", "Сосна", "Ель"] and z.rows[-1].place is None


def test_texts():
    assert res.rotations(["А", "Б", "В"]) == [["А", "Б", "В"], ["Б", "В", "А"], ["В", "А", "Б"]]
    comp = SimpleNamespace(title="Чемпионат г. Красноярска по спортивному туризму")
    assert res.title_in(comp) == "Чемпионате г. Красноярска по спортивному туризму"
    assert res.title_in(SimpleNamespace(title="Открытое первенство города")) == "Открытом первенстве города"
    # прилагательные в начале — по окончанию (Правки, п. 50)
    for title, of, in_ in [
            ("Учебный чемпионат г. Энска", "Учебного чемпионата г. Энска", "Учебном чемпионате г. Энска"),
            ("Краевой кубок", "Краевого кубка", "Краевом кубке"),
            ("Открытый Кубок Красноярского края", "Открытого Кубка Красноярского края",
             "Открытом Кубке Красноярского края"),
            ("Всероссийские соревнования", "Всероссийских соревнований", "Всероссийских соревнованиях"),
            ("Летнее первенство", "Летнего первенства", "Летнем первенстве"),
            ("Всероссийский фестиваль", "Всероссийского фестиваля", "Всероссийском фестивале"),
            ("Краевая спартакиада", "Краевой спартакиады", "Краевой спартакиаде"),
            ("Чемпионат", "Чемпионата", "Чемпионате"),
            ("XV слёт туристов", "XV слёт туристов", "XV слёт туристов")]:  # не знаем, как склонять, — как есть
        assert res.title_of(SimpleNamespace(title=title)) == of
        assert res.title_in(SimpleNamespace(title=title)) == in_
    assert res.date_text(date(2025, 9, 21)) == "21 сентября 2025 г."


def text_of(path):
    return [p.text for p in Document(path).paragraphs if p.text]


def test_diplomas_each_member_first(tmp_path, psr_card):
    out = aw.write_diplomas(zres(psr_card), psr_card, tmp_path / "d.docx")
    t = text_of(out)
    assert t.count("Награждаются") == 9  # 3 призёра × 3 участника — по диплому каждому
    firsts = [t[i + 1] for i, s in enumerate(t) if s == "Награждаются"][:3]
    assert firsts == ["Антон Лебедев, Мария Зуева, Глеб Носов", "Мария Зуева, Глеб Носов, Антон Лебедев",
                      "Глеб Носов, Антон Лебедев, Мария Зуева"]
    assert "за I место" in t and "за III место" in t and "за IV место" not in " ".join(t)
    assert "на Чемпионате города N по спортивному туризму" in t
    assert "в спортивной дисциплине «дистанция - комбинированная» 3 класса, смешанные группы" in t  # название — как в ВРВС
    assert "окрестности г. N, 21 сентября 2025 г." in t  # дата — из карточки, а не прошлогодняя


def test_stickers_and_awardees(tmp_path, psr_card):
    z = zres(psr_card)
    ws = load_workbook(aw.write_stickers(z, psr_card, tmp_path / "s.xlsx")).active
    stickers = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert len(stickers) == 9 + 2 and stickers[0].startswith("ЧЕМПИОНАТ ГОРОДА N")  # 9 медалей + 2 запасные
    t = text_of(aw.write_awardees(z, psr_card, tmp_path / "l.docx"))
    places = [s.split(" место")[0] for s in t if " место — " in s]
    assert places == ["III", "II", "I"]  # в порядке вызова
    assert any(s == "I место — команда «Кедр», Красноярск: Антон Лебедев, Мария Зуева, Глеб Носов" for s in t)


# ------------------------------------------------------------------ документы о судьях


def test_judging_certificates(tmp_path, psr_card):
    from st_secretary.exporters.judges import write_judging_certificates

    grades = {res.person_key("Судьин Иван Петрович"): "отлично"}
    t = text_of(write_judging_certificates(psr_card, grades, tmp_path / "j.docx"))
    assert t.count("СПРАВКА") == 2 and "ФЕДЕРАЦИЯ СПОРТИВНОГО ТУРИЗМА" in t
    assert ("Дана Судьину Ивану Петровичу в том, что он участвовал в судействе Чемпионата города N по "
            "спортивному туризму — дисциплина «дистанция - комбинированная», 20–21 сентября 2025 г., "
            "окрестности г. N, в качестве главного судьи.") in t
    assert any(s.startswith("Дана Секретарёвой Анне Ивановне в том, что она участвовала") for s in t)
    assert "Оценка судейства: «отлично»." in t and "Оценка судейства: «________________»." in t
    assert any("И.П. Судьин, СС1К, г. Красноярск" in s for s in t)


@pytest.mark.parametrize("role, expected", [
    ("Главный судья", "в качестве главного судьи"),
    ("Заместитель главного секретаря", "в качестве заместителя главного секретаря"),
    ("Начальник дистанции", "в качестве начальника дистанции"),
    ("Председатель комиссии по допуску", "в качестве председателя комиссии по допуску"),
    ("Старший судья этапа", "в качестве старшего судьи этапа"),
    ("Судья-хронометрист", "в качестве судьи-хронометриста"),
    ("Судья на старте", "в качестве судьи на старте"),
    ("Волонтёр", "в должности: волонтёр"),
    ("Главный", "в должности: главный"),
])
def test_role_phrase(role, expected):
    from st_secretary.exporters.judges import role_phrase

    assert role_phrase(role) == expected


def test_sk_and_subjects_certificates(tmp_path, psr_card):
    from st_secretary.exporters.judges import write_sk_certificate, write_subjects_certificate

    doc = Document(write_sk_certificate(psr_card, tmp_path / "sk.docx"))
    t = [p.text for p in doc.paragraphs if p.text]
    assert "о составе и квалификации судейской коллегии" in t
    assert "Всего судей 2 человека, в том числе из других регионов 0 человек." in t
    assert "Первая категория — 1 человек" in t and "Вторая категория — 1 человек" in t
    table = [[c.text for c in r.cells] for r in doc.tables[0].rows]
    assert table[1] == ["1", "Главный судья", "Судьин Иван Петрович", "СС1К", "г. Красноярск"]
    t = text_of(write_subjects_certificate(psr_card, ["Красноярск", "Томск", "Красноярск"], tmp_path / "s.docx"))
    assert "Количество субъектов РФ, принявших участие в соревнованиях: 2" in t and "2. Томск" in t


# ------------------------------------------------------------------ выписки на разряды и отчёт


def kedr_preapps(tmp_path, psr_card):
    a = make_application(tmp_path / "Кедр.xlsx", "Кедр", "Красноярск", "Лебедев Антон Игоревич", "89135550000", 3, [
        ["Кедр", "Красноярск", "Лебедев Антон Игоревич", "Лебедев Антон Игоревич", "02.02.1990", "III", "м", "М/Ж", 3],
        ["Кедр", "Красноярск", "Лебедев Антон Игоревич", "Зуева Мария Олеговна", "05.06.1996", "III", "ж", "М/Ж", 3],
        ["Кедр", "Красноярск", "Лебедев Антон Игоревич", "Носов Глеб Андреевич", "09.09.2010", "б/р", "м", "М/Ж", 3],
    ])
    return process([read_preapplication(a)], psr_card)


def test_extracts_only_norms_with_birth_dates(tmp_path, psr_card):
    from st_secretary.exporters.final import extract_count, write_extracts

    z = zres(psr_card, kedr_preapps(tmp_path, psr_card))
    assert extract_count(z) == 6  # Кедр и Сосна выполнили II — по 3 человека
    ws = load_workbook(write_extracts(z, psr_card, tmp_path / "e.xlsx")).active
    cells = [[c.value for c in row] for row in ws.iter_rows()]
    flat = [str(v) for row in cells for v in row if v is not None]
    assert "ВЫПИСКА ИЗ ПРОТОКОЛА СОРЕВНОВАНИЙ" in flat and any("код ВРВС 0840161811Я" in v for v in flat)
    assert any("квалификационный ранг соревнований: 71,7" in v for v in flat)
    rows = [r for r in cells if r[1] and r[5] == "II"]
    assert len(rows) == 6 and rows[0][:6] == ["1", "Лебедев Антон Игоревич", "02.02.1990", "III", "-283", "II"]
    assert next(r for r in rows if r[1] == "Орлов Павел Ильич")[2] == "нет в заявке"  # без даты — видно сразу
    assert not any(r[1] == "Ким Олег Борисович" for r in cells)  # 3 место без норматива — не в выписке
    assert any(str(v).startswith("Главный судья") for v in flat)


def test_chief_judge_report(tmp_path, psr_card):
    from st_secretary.exporters.final import write_report

    pre = kedr_preapps(tmp_path, psr_card)
    grades = {res.person_key("Судьин Иван Петрович"): "отлично"}
    doc = Document(write_report(psr_card, zres(psr_card, pre), pre.entries, ["Красноярск"], grades,
                                {"base": "Материальная база соответствовала требованиям."}, tmp_path / "r.docx"))
    t = [p.text for p in doc.paragraphs if p.text]
    assert "о проведении: Чемпионат города N по спортивному туризму" in t
    assert "в период с «20» сентября 2025 г. по «21» сентября 2025 г." in t
    assert ("2. Общее количество участников, допущенных до соревнований, — 3, из них мужчин — 2, "
            "женщин — 1.") in t
    assert any(s.startswith("4. Из общего числа участников по возрасту: до 16 лет — 1") for s in t)  # 2010 г. р.
    assert "I место — команда «Кедр» (Красноярск): Лебедев Антон, Зуева Мария, Носов Глеб" in t
    assert "Выполнили нормативы: II — 6 чел." in t
    assert "Протесты, жалобы в ГСК не подавались." in t  # текст по умолчанию
    assert "Материальная база соответствовала требованиям." in t
    table = [[c.text for c in r.cells] for r in doc.tables[0].rows]
    assert table[1] == ["1", "Судьин Иван Петрович", "СС1К", "г. Красноярск", "Главный судья", "отлично"]
