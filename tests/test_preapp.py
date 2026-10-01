"""Предзаявки: все дефекты, встреченные в реальных заявках, на выдуманных людях."""

from datetime import date

from conftest import make_application
from openpyxl import load_workbook

from st_secretary.cli import main
from st_secretary.exporters.preapp_xlsx import write_preapp_report
from st_secretary.importers.card_xlsx import write_card
from st_secretary.importers.preapp_xlsx import read_preapplication
from st_secretary.issues import ERROR, FIXED, WARNING
from st_secretary.preapp import process


def texts(result, severity=None, team=None):
    return " | ".join(i.text for i in result.issues
                      if (severity is None or i.severity == severity) and (team is None or i.team == team))


def build(tmp_path):
    a = make_application(tmp_path / "a.xlsx", "Лесовики", "Красноярск", "Иванов Пётр Сергеевич",
                         "89131234567, les@mail.ru", 3, [
                             ["Лесовики", "Красноярск", "Иванов Пётр Сергеевич", "Иванов ПётрСергеевич",
                              "17.10.1989", "II", "м", "м/ж", 3],
                             ["Лесов", "Краснорярск ", "Иванов Пётр Сергеевич", "Смирнова Анна Олеговна12.03.2001",
                              None, "КМС", "ж", "М/Ж", None],
                             ["Лесовики", "Красноярск", "Иванов Пётр Сергеевич", "Кузьмин Олег Игоревич",
                              "29.02.1995", "б/р", "ж", None, 3],
                         ])
    b = make_application(tmp_path / "b.xlsx", " ", "г. Томск", "Орлова Мария Ивановна", None, 4, [
        ["Сосна", "г.Томск", "Орлова Мария Ивановна", "Орлова МарияИвановна", "07/12/2000", "III", "ж", "М/Ж", 3],
        ["Сосна", "г.Томск", "Орлова Мария Ивановна", "Пестов Юрий Андреевич", "01.01.2009", "2ю", "м", "М/Ж", 3],
        ["Сосна", "г.Томск", "Орлова Мария Ивановна", "Смирнова Анна Олеговна", "12.03.2001", "?", "ж", "Ю/Д", 3],
    ])
    return [read_preapplication(a), read_preapplication(b)]


def test_reading_skips_sample_and_empty_rows(tmp_path):
    apps = build(tmp_path)
    assert [len(a.rows) for a in apps] == [3, 3]
    assert apps[0].team == "Лесовики" and apps[0].contacts.startswith("8913")


def test_cleaning_and_checks(tmp_path, psr_card):
    r = process(build(tmp_path), psr_card)
    first, second = r.teams
    # ФИО, даты, пол
    assert [e.name.full for e in first.entries] == ["Иванов Пётр Сергеевич", "Смирнова Анна Олеговна", "Кузьмин Олег Игоревич"]
    assert first.entries[1].birth == date(2001, 3, 12)
    assert "перенесена в колонку даты" in texts(r, FIXED)
    assert "такой даты не существует" in texts(r, ERROR)
    assert "отчество «Игоревич» — мужское" in texts(r, WARNING)
    # команда, территория, группа, класс
    assert {e.team for e in first.entries} == {"Лесовики"} and "«Лесов» приведено к «Лесовики»" in texts(r, FIXED)
    assert first.territory == "Красноярск"
    assert all(e.group == "М/Ж" and e.distance_class == 3 for e in first.entries)
    assert all(e.team_dist == "1" for e in first.entries)  # ПСР — только командная дистанция
    assert first.phone == "+7 913 123-45-67" and first.email == "les@mail.ru"
    # вторая заявка: пустая шапка, Томск, возраст, разряд, неизвестная группа
    assert second.team == "Сосна" and second.territory == "Томск"
    assert "не указано название команды" in texts(r, WARNING, "Сосна")
    assert "по решению ГСК" in texts(r, WARNING, "Сосна")  # 16 лет в 2025 г.
    assert "не распознан разряд" in texts(r, ERROR, "Сосна")
    assert "«Ю/Д» и класс «3» не совпадают" in texts(r, ERROR, "Сосна")
    assert "указано участников: 4, в таблице: 3" in texts(r, WARNING, "Сосна")
    assert "нет контактов" in texts(r, WARNING, "Сосна")
    assert "прочитана как день/месяц/год" in texts(r, FIXED, "Сосна")
    # один человек в двух командах
    assert "участник заявлен и в команде «Лесовики»" in texts(r, ERROR, "Сосна")


def test_every_error_and_warning_explains_why_and_what_to_do(tmp_path, psr_card):
    r = process(build(tmp_path), psr_card)
    for i in r.issues:
        if i.severity in (ERROR, WARNING):
            assert i.why and i.todo, i.text
    date_err = next(i for i in r.issues if "не существует" in i.text)
    assert "1995 год не високосный" in date_err.why and "представител" in date_err.todo
    assert date_err.row == 12 and date_err.person == "Кузьмин Олег Игоревич"  # строка файла, как видно в Excel
    gsk = next(i for i in r.issues if "по решению ГСК" in i.text)
    assert "16 лет" in gsk.why and "комиссию по допуску" in gsk.todo
    assert "исполняется 16 лет" in gsk.text


def test_written_application_reads_back_the_same(tmp_path, psr_card):
    """Заявка, записанная программой, читается и проверяется так же, как исходная."""
    from st_secretary.importers.preapp_xlsx import write_preapplication

    head = {"team": "Лесовики", "territory": "Красноярск", "representative": "Иванов Пётр Сергеевич",
            "contacts": "89131234567", "declared": "2"}
    rows = [{"fio": "Иванов Пётр Сергеевич", "birth": date(1989, 10, 17), "qual": "II", "sex": "м", "group": "М/Ж",
             "cls": "3", "team_dist": "1"},
            {"fio": "Смирнова Анна Олеговна", "birth": "29.02.1995", "qual": "КМС", "sex": "ж", "group": "М/Ж",
             "cls": "3", "team_dist": "1", "chip": "0123"}]
    p = write_preapplication(tmp_path / "w.xlsx", head, rows, "Исправлено в программе", ["б/р", "II", "КМС"])
    app = read_preapplication(p)
    assert (app.team, app.territory, app.declared_count) == ("Лесовики", "Красноярск", 2)
    assert [r.values["fio"] for r in app.rows] == ["Иванов Пётр Сергеевич", "Смирнова Анна Олеговна"]
    assert app.rows[1].values["chip"] == "0123"  # ведущий ноль не теряется
    r = process([app], psr_card)
    t = r.teams[0]
    assert t.entries[0].birth == date(1989, 10, 17) and t.entries[0].qual.label == "II"
    assert all(e.zachet and e.zachet.key == "М/Ж_3" and e.team == "Лесовики" for e in t.entries)
    assert "такой даты не существует" in texts(r, ERROR)  # ошибочная дата сохранена как есть и видна
    assert not texts(r, FIXED)  # программа пишет уже чистые данные — исправлять нечего


def test_host_territory_is_never_changed(tmp_path, psr_card):
    rows = [["Т", "Краснорярск", "Иванов Пётр Сергеевич", f"Иванов{i} Пётр Сергеевич", "01.01.1990", "б/р",
             "м" if i else "ж", "М/Ж", 3] for i in range(3)]
    typo = make_application(tmp_path / "t.xlsx", "Т", "Краснорярск", "Иванов Пётр Сергеевич", "89130000000", 3, rows)
    ok = make_application(tmp_path / "o.xlsx", "О", "Красноярск", "Петров Иван Ильич", "89130000001", 1,
                          [["О", "Красноярск", "Петров Иван Ильич", "Петров Иван Ильич", "01.01.1990", "б/р", "м",
                            "М/Ж", 3]])
    r = process([read_preapplication(typo), read_preapplication(ok)], psr_card)
    assert {t.territory for t in r.teams} == {"Красноярск"}
    assert "исправлена на «Краснорярск»" not in texts(r)


def test_team_composition(tmp_path, psr_card):
    rows = [["Мальчики", "Красноярск", "Сидоров Олег Петрович", n, "01.01.2008", "б/р", "м", "М/Ж", 3]
            for n in ("Сидоров Олег Петрович", "Попов Иван Ильич")]
    p = make_application(tmp_path / "m.xlsx", "Мальчики", "Красноярск", "Сидоров Олег Петрович", "89130000000", 2, rows)
    r = process([read_preapplication(p)], psr_card)
    t = texts(r, ERROR)
    assert "в команде 2 чел., по Положению — 3" in t
    assert "женщин в команде 0" in t
    assert "нет участника 18 лет и старше" in t


def test_broken_form_is_reported(tmp_path, psr_card):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "что-то другое"
    wb.save(tmp_path / "x.xlsx")
    r = process([read_preapplication(tmp_path / "x.xlsx")], psr_card)
    assert "не найдена таблица участников" in texts(r, ERROR)


def test_report_workbook(tmp_path, psr_card):
    r = process(build(tmp_path), psr_card)
    out = write_preapp_report(r, psr_card, tmp_path / "svodka.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Итог", "Заявка для СЕКРЕТАРЬ", "Проверка", "Команды"]
    z = wb["Заявка для СЕКРЕТАРЬ"]
    assert z["A1"].value == "№ п/п" and z["E1"].value == "Фамилия Имя"
    assert z["A2"].value.startswith("Лесовики")  # сразу под шапкой — строка команды, без пустой строки
    first = [c.value for c in z[3]]
    assert first[:10] == [1, "Лесовики", "Красноярск", "Иванов Пётр Сергеевич", "Иванов Пётр Сергеевич",
                          first[5], "II", "м", "М/Ж", 3]
    assert first[14] == 1 and first[15] == "М/Ж_3"  # номер группы — числом
    assert z.cell(3, 6).number_format == "DD.MM.YYYY"
    assert wb["Проверка"].max_row == len(r.issues) + 1


def test_cli_end_to_end(tmp_path, psr_card, capsys):
    folder = tmp_path / "zayavki"
    folder.mkdir()
    build(folder)
    card = write_card(tmp_path / "card.xlsx", psr_card)
    assert main(["preapp", str(folder), "--card", str(card), "--out", str(tmp_path / "s.xlsx")]) == 0
    out = capsys.readouterr().out
    assert "Заявок: 2, команд: 2, участников: 6" in out
    assert (tmp_path / "s.xlsx").exists()
    assert main(["card-check", str(card)]) == 0


ROWS_2 = [["Сталактит", "Энск", "Пещерин Олег Петрович", "Сводов Артём", "12.03.2009", "II", "м", "ЮН", None, None,
           None, None, None, 1],
          ["Сталактит", "Энск", "Пещерин Олег Петрович", "Гротова Вера", "05.07.2010", "III", "ж", "ЮН", None, None,
           None, None, None, 1]]
ROWS_3 = [["Сталактит", "Энск", "Пещерин Олег Петрович", "Колодцев Иван", "21.01.2006", "I", "м", "ЮН", 3, None,
           None, None, None, 1],
          ["Сталактит", "Энск", "Пещерин Олег Петрович", "Натёкова Лиза", "30.09.2007", "II", "ж", "ЮН", None, None,
           None, None, None, 1]]


def _check_two_sheets(app):
    assert not app.problems and not app.notes and app.sheets == ["2 КЛАСС", "3 КЛАСС"]
    assert [str(r.values["fio"]) for r in app.rows] == ["Сводов Артём", "Гротова Вера", "Колодцев Иван",
                                                        "Натёкова Лиза"]
    # класс не вписан — из названия листа; вписан — как вписан; строки второго листа — «лист 2, строка …»
    assert [str(r.values["cls"]).split(".")[0] for r in app.rows] == ["2", "2", "3", "3"]
    assert app.rows[0].row < 1000 <= app.rows[2].row and app.team == "Сталактит"


def test_application_on_several_sheets_reads_all_of_them(tmp_path):
    """Правки, п. 29: бланк делегации спелео — по листу на класс; читаются все листы с таблицей участников."""
    from conftest import make_class_sheets_application

    from st_secretary.importers.preapp_xlsx import row_place

    path = make_class_sheets_application(tmp_path / "Сталактит.xlsx", {"2 КЛАСС": ROWS_2, "3 КЛАСС": ROWS_3})
    _check_two_sheets(read_preapplication(path))
    assert row_place(1010) == "лист 2, строка 10" and row_place(10) == "строка 10"


def test_application_on_several_sheets_xls(tmp_path):
    """То же для .xls (как присылают на самом деле): выдуманная делегация, 2 + 2 участника."""
    import pytest

    pytest.importorskip("xlrd")
    from pathlib import Path

    _check_two_sheets(read_preapplication(Path(__file__).parent / "fixtures" / "class_sheets_application.xls"))


def test_sheet_without_table_is_reported_not_skipped_silently(tmp_path, psr_card):
    from conftest import make_class_sheets_application

    path = make_class_sheets_application(tmp_path / "Сталактит.xlsx", {"2 КЛАСС": ROWS_2})
    wb = load_workbook(path)
    ws = wb.create_sheet("Ещё участники")
    for r in (["Натёкова Лиза", "30.09.2007", "II"], ["Колодцев Иван", "21.01.2006", "I"]):
        ws.append(r)
    wb.save(path)
    app = read_preapplication(path)
    assert len(app.rows) == 2 and app.notes and "лист «Ещё участники» не прочитан" in app.notes[0][0]
    result = process([app], psr_card)
    assert any(i.severity == WARNING and "Ещё участники" in i.text for i in result.issues)
