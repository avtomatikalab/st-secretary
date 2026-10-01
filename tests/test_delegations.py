"""Делегации (Правки, п. 20; решение 040): документы человека проверяются один раз — во всех его командах и
соревнованиях фестиваля; делегация в комиссии (допуск врача, взнос одной строкой или по командам, папка сканов
по людям); свои колонки в заявке читаются без ошибок. Все данные — выдуманные."""

import html
import io
from dataclasses import replace

import pytest

pytest.importorskip("fastapi")
from conftest import make_application
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_web import base, kedr, sosna, team_form  # заготовки страниц и заявок

from st_secretary.web.app import create_app


@pytest.fixture
def opened():
    return []


@pytest.fixture
def client(tmp_path, opened):
    return TestClient(create_app(tmp_path / "данные", opener=opened.append, docs_dir=tmp_path / "документы",
                                 board_host="127.0.0.1"))


def kedr2(tmp_path) -> bytes:
    """Вторая команда той же делегации (та же территория и представитель)."""
    return make_application(tmp_path / "Кедр-2.xlsx", "Кедр-2", "Красноярск", "Лебедев Антон Игоревич",
                            "89135550000", 3, [
                                ["Кедр-2", "Красноярск", "Лебедев Антон Игоревич", "Орехов Пётр Ильич",
                                 "03.03.1991", "II", "м", "М/Ж", 3],
                                ["Кедр-2", "Красноярск", "Лебедев Антон Игоревич", "Вишнёва Ольга Петровна",
                                 "07.07.1995", "III", "ж", "М/Ж", 3],
                                ["Кедр-2", "Красноярск", "Лебедев Антон Игоревич", "Липин Игорь Сергеевич",
                                 "11.11.1992", "б/р", "м", "М/Ж", 3],
                            ]).read_bytes()


def test_person_documents_checked_once_in_all_competitions_of_festival(client, tmp_path, psr_card):
    store = client.app.state.store
    a = store.create(psr_card)
    b = store.create(replace(psr_card, title="Кубок города N по спортивному туризму"))
    a.add_preapp("Кедр.xlsx", kedr(tmp_path))
    b.add_preapp("Кедр.xlsx", kedr(tmp_path))
    form = team_form(client.get(base(b) + "/admission").text, "Кедр.xlsx")
    assert not any(k.startswith("p-0-d-") for k in form)  # не в фестивале — в другом соревновании не видно
    r = client.post("/festival/new", data={"title": "Выезд", "member": [a.id, b.id]}, follow_redirects=False)
    assert "festival_made" in r.headers["location"]

    # документы отметили в одном соревновании — в другом они уже отмечены (тот же человек: ФИО + дата рождения)
    client.post(base(a) + "/admission/team", data=team_form(client.get(base(a) + "/admission").text, "Кедр.xlsx")
                | {"do": "all_docs"})
    page = client.get(base(b) + "/admission").text
    form = team_form(page, "Кедр.xlsx")
    assert all(form.get(f"p-{i}-d-{d}") for i in range(3) for d in ("id", "med", "book", "oms", "ins"))
    assert "в других соревнованиях фестиваля" in page and "Чемпионат города N" in page
    # снять галочку можно только там, где отметили: здесь снятие ничего не меняет
    client.post(base(b) + "/admission/team", data={k: v for k, v in form.items() if not k.startswith("p-0-d-")})
    assert team_form(client.get(base(b) + "/admission").text, "Кедр.xlsx").get("p-0-d-id")
    stored = b.admission()["teams"]["Кедр.xlsx"]["people"]
    assert not any(stored.get(k, {}).get("docs", {}).get("id") for k in stored)  # в своём соревновании не записаны


def test_delegation_doctor_fee_in_one_line_and_scans_folder(client, tmp_path, psr_card, opened):
    store = client.app.state.store
    f = store.create(psr_card)  # взнос 3000 ₽ за команду
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Кедр-2.xlsx", kedr2(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    page = client.get(base(f) + "/admission?by=delegation").text
    assert page.count('class="adm-deleg"') == 2 and "2 команды, 6 чел." in page
    key = next(x.split('"')[0] for x in page.split('name="key" value="')[1:] if x.startswith("красноярск|лебедев"))

    # заявка делегации с печатью врача — мед. допуск всем её участникам (обеих команд), у другой делегации — нет
    r = client.post(base(f) + "/admission/delegation", data={"key": key, "doctor": "1", "fee_mode": "one",
                                                             "fee_paid": "4000", "fee_method": "перевод"},
                    follow_redirects=False)
    assert "by=delegation" in r.headers["location"]
    page = client.get(base(f) + "/admission").text
    for file in ("Кедр.xlsx", "Кедр-2.xlsx"):
        assert all(team_form(page, file).get(f"p-{i}-d-med") for i in range(3))
    assert not team_form(page, "Сосна.xlsx").get("p-0-d-med")
    assert "платит делегация одной строкой" in page and "взносов собрано из" in page

    # взнос одной строкой: в ведомости — одна строка делегации, 4000 из 6000
    data = client.get(base(f) + "/admission/report.xlsx").content
    ws = load_workbook(io.BytesIO(data))["Ведомость взносов"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    deleg = next(r for r in rows if str(r[2]).startswith("Делегация:"))
    assert "Кедр" in deleg[2] and "Кедр-2" in deleg[2] and deleg[5] == 6 and deleg[6] == 6000 and deleg[7] == 4000
    assert deleg[9] == "частично" and not any(r[2] == "Кедр" for r in rows)

    # по командам — снова у каждой команды своя оплата
    client.post(base(f) + "/admission/delegation", data={"key": key, "fee_mode": "teams"})
    page = client.get(base(f) + "/admission").text
    assert "платит делегация одной строкой" not in page and 'name="fee_paid"' in page
    assert not all(team_form(page, "Кедр.xlsx").get(f"p-{i}-d-med") for i in range(3))  # врач делегации снят

    # папка сканов делегации: внутри — по людям; скан человека виден в окне проверки его команды
    client.post(base(f) + "/admission/delegation/folder", data={"key": key})
    folder = opened[-1]
    people = sorted(p.name for p in folder.iterdir())
    assert len(people) == 6 and "Зуева Мария Олеговна 05.06.1996" in people
    (folder / "Зуева Мария Олеговна 05.06.1996" / "паспорт.jpg").write_bytes(b"\xff\xd8\xff")
    check = client.get(base(f) + "/admission/check?file=" + "Кедр.xlsx").text
    assert "Зуева Мария Олеговна: паспорт" in check
    href = html.unescape(check.split('"' + base(f) + "/docs/person?")[1].split('"')[0])
    r = client.get(base(f) + "/docs/person?" + href)
    assert r.status_code == 200 and r.content == b"\xff\xd8\xff"


def test_own_columns_in_application_read_without_errors_and_kept(client, tmp_path, psr_card):
    path = make_application(tmp_path / "Кедр.xlsx", "Кедр", "Красноярск", "Лебедев Антон Игоревич", "89135550000", 2, [
        ["Кедр", "Красноярск", "Лебедев Антон Игоревич", "Лебедев Антон Игоревич", "02.02.1990", "I", "м", "М/Ж", 3],
        ["Кедр", "Красноярск", "Лебедев Антон Игоревич", "Зуева Мария Олеговна", "05.06.1996", "II", "ж", "М/Ж", 3]])
    wb = load_workbook(path)
    ws = wb.active
    ws.cell(6, 16, "Размер футболки")
    ws.cell(6, 17, "Питание")
    ws.cell(10, 16, "L")
    ws.cell(11, 16, "S")
    ws.cell(11, 17, "вегетарианское")
    wb.save(path)
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", path.read_bytes())
    result, _ = client.app.state.store.review(f, f.load())
    assert not [i for i in result.issues if "колонк" in i.text]
    zueva = next(e for e in result.teams[0].entries if e.name.full == "Зуева Мария Олеговна")
    assert zueva.extra == {"Размер футболки": "S", "Питание": "вегетарианское"}
    assert "Питание: вегетарианское" in client.get(base(f) + "/preapps/team?file=Кедр.xlsx").text

    # заявку исправили в программе — свои колонки остались у тех же людей
    rows = [{"fio": e.name.full, "birth": e.birth, "qual": e.qual.label, "sex": e.sex, "group": e.group,
             "cls": e.distance_class} for e in reversed(result.teams[0].entries)]
    f.save_preapp("Кедр.xlsx", {"team": "Кедр", "territory": "Красноярск", "representative": "Лебедев Антон Игоревич"},
                  rows)
    from st_secretary.importers.preapp_xlsx import read_preapplication

    raw = read_preapplication(f.preapp_dir / "Кедр.xlsx")
    by = {r.values["fio"]: r.extra for r in raw.rows}
    assert by["Зуева Мария Олеговна"] == {"Размер футболки": "S", "Питание": "вегетарианское"}
    assert by["Лебедев Антон Игоревич"] == {"Размер футболки": "L"}
