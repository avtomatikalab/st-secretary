"""Своя редакция разрядных норм в Excel (Правки, п. 59): копия выбранной редакции — секретарь правит и загружает.

Лист «Нормы»: название своей редакции, на чём основана, таблица строк ранга — «Ранг от», «Ранг (как в документе)»,
проценты I, II, III, Iю, IIю, IIIю («результат в процентах от результата победителя, не более»; пусто — разряд в
строке не присваивается; у IIIю — «условие», если III юношеский в редакции текстовым условием) и столбцы классов 1–6
(«да» — строка относится к классу, как незакрашенные ячейки в документе). Лист «Условия»: возраст, статус
соревнований, северная ходьба, баллы ранга, срок действия. Загрузка проверяет всё и пишет ошибки с адресом ячейки;
устройство результата — как у справочника (reference/data/norms/*.toml), его читает reference.edition_from_dict.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from fractions import Fraction
from itertools import pairwise

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from st_secretary.reference import OWN_PREFIX, edition_from_dict, norm_editions

INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # жёлтые — для правки, как в карточке
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
SHEET, COND_SHEET = "Нормы", "Условия"
QUALS = [("I", "I"), ("II", "II"), ("III", "III"), ("Y1", "Iю"), ("Y2", "IIю"), ("Y3", "IIIю")]
CLASSES = range(1, 7)
CONDITION = "условие"
YES = {"да", "x", "х", "+", "✓", "1", "yes"}
STATUS = {"any": "любой", "municipal": "не ниже муниципального", "regional": "не ниже субъекта РФ",
          "interregional": "не ниже межрегионального", "all_russian": "всероссийский"}
# лист «Условия»: подпись, где в справочнике, вид значения
CONDITIONS = [
    ("I разряд выполняется с, лет", ("min_age", "I"), "int"),
    ("II разряд — с, лет", ("min_age", "II"), "int"),
    ("III разряд — с, лет", ("min_age", "III"), "int"),
    ("Юношеские разряды — с, лет", ("min_age", "junior"), "int"),
    ("I разряд — статус соревнований", ("status", "I"), "status"),
    ("II разряд — статус соревнований", ("status", "II"), "status"),
    ("III разряд — статус соревнований", ("status", "III"), "status"),
    ("Юношеские — статус соревнований", ("status", "junior"), "status"),
    ("Северная ходьба: I разряд присваивается (да/нет)", ("nordic_walking", "I_allowed"), "bool"),
    ("Северная ходьба: I–III разряды — с, лет", ("nordic_walking", "min_age_I_II_III"), "int"),
    ("III юношеский — текст условия (если в таблице «условие»)", ("junior_III", "text"), "text"),
    ("Баллы ранга за МС", ("rank_points", "MS"), "num"),
    ("Баллы ранга за КМС", ("rank_points", "KMS"), "num"),
    ("Баллы ранга за I", ("rank_points", "I"), "num"),
    ("Баллы ранга за II", ("rank_points", "II"), "num"),
    ("Баллы ранга за III", ("rank_points", "III"), "num"),
    ("Баллы ранга за Iю", ("rank_points", "Y1"), "num"),
    ("Баллы ранга за IIю", ("rank_points", "Y2"), "num"),
    ("Баллы ранга за IIIю", ("rank_points", "Y3"), "num"),
    ("Ранг считается, если участников не меньше", ("rank_points", "min_participants"), "int"),
    ("Баллы ранга — за места с 1 по", ("rank_points", "max_place"), "int"),
    ("Действует с (дд.мм.гггг)", ("valid_from",), "date"),
    ("Действует по (дд.мм.гггг)", ("valid_to",), "date"),
]


def _num_text(v: Fraction) -> str:
    """Число для справочника: «12», «1.2» (как в TOML)."""
    return str(v.numerator) if v.denominator == 1 else format(float(v), "g")


def _get(d: dict, path: tuple[str, ...]):
    for k in path:
        if not isinstance(d, dict) or k not in d:
            return None
        d = d[k]
    return d


def write_norms(d: dict, title: str = "") -> bytes:
    """Редакция (словарь справочника) → книга Excel для своей копии."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A1"] = "Своя редакция разрядных норм"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"], ws["B2"] = "Название своей редакции", title or f"на основе {d.get('name') or d['edition']}"
    ws["B2"].fill = INPUT_FILL
    ws["A3"], ws["B3"] = "Годы действия (как у основы)", d["edition"]
    ws["B3"].fill = INPUT_FILL
    ws["A4"] = ("Жёлтое — правьте. Проценты — «результат в % от результата победителя, не более»; пусто — разряд в "
                "строке не присваивается; у IIIю — «условие», если III юношеский — текстом (лист «Условия»). Столбцы "
                "классов: «да» — строка относится к классу (в документе — незакрашенные ячейки); строки класса — "
                "подряд. Строки — по возрастанию «Ранг от». Сохраните и загрузите на странице «Разрядные нормы».")
    ws.merge_cells("A4:N4")
    ws["A4"].alignment = WRAP
    ws.row_dimensions[4].height = 64
    head = ["Ранг от", "Ранг (как в документе)"] + [q for _, q in QUALS] + [f"{c} кл" for c in CLASSES]
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=6, column=i, value=h)
        c.font, c.fill, c.alignment = BOLD, HEAD_FILL, WRAP
    spans = {int(k): (Fraction(str(v[0])), Fraction(str(v[1]))) for k, v in d["class_rows"].items()}
    for r, row in enumerate(d["row"], start=7):
        lo = Fraction(str(row["min"]))
        vals = [float(lo) if lo.denominator != 1 else int(lo), row["label"]]
        for key, _ in QUALS:
            vals.append(row[key] if key in row else (CONDITION if key == "Y3" and row.get("Y3_condition") else None))
        vals += ["да" if c in spans and spans[c][0] <= lo <= spans[c][1] else None for c in CLASSES]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v).fill = INPUT_FILL
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22
    for i in range(3, len(head) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 8
    ws.freeze_panes = "C7"

    wc = wb.create_sheet(COND_SHEET)
    wc["A1"], wc["B1"] = "Условие", "Значение"
    wc["A1"].font = wc["B1"].font = BOLD
    wc["C1"] = "Статус: " + "; ".join(STATUS.values())
    for r, (label, path, kind) in enumerate(CONDITIONS, start=2):
        v = _get(d, path)
        if path == ("nordic_walking", "min_age_I_II_III") and v is None:
            v = _get(d, ("nordic_walking", "min_age_II_III"))
        if kind == "status":
            v = STATUS.get(v, v)
        elif kind == "bool":
            v = "да" if v else "нет"
        elif kind == "num" and v is not None:
            f = Fraction(str(v))
            v = int(f) if f.denominator == 1 else float(f)
        elif kind == "date" and v:
            v = date.fromisoformat(str(v)).strftime("%d.%m.%Y")
        wc.cell(row=r, column=1, value=label)
        cell = wc.cell(row=r, column=2, value=v)
        cell.fill, cell.alignment = INPUT_FILL, WRAP
    wc.column_dimensions["A"].width = 58
    wc.column_dimensions["B"].width = 60
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def read_norms(data: bytes, base: dict | None = None) -> tuple[dict | None, list[str]]:
    """Книга Excel → словарь своей редакции и ошибки (понятным текстом, с адресом ячейки). base — основа: из неё
    берётся то, чего в книге нет (например, лист «Условия» удалили)."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:  # noqa: BLE001 — не Excel или повреждён
        return None, ["Файл не читается как книга Excel (.xlsx). Скачайте копию со страницы «Разрядные нормы», "
                      "поправьте её и загрузите снова."]
    if SHEET not in wb.sheetnames:
        return None, [f"В книге нет листа «{SHEET}» — загрузите копию, скачанную со страницы «Разрядные нормы»."]
    ws = wb[SHEET]
    errors: list[str] = []

    def at(r: int, c: int) -> str:
        return f"{ws.title}!{get_column_letter(c)}{r}"

    name = str(ws["B2"].value or "").strip()
    if not name:
        errors.append(f"{ws.title}!B2: впишите название своей редакции.")
    elif len(name) > 60:
        errors.append(f"{ws.title}!B2: название длиннее 60 знаков — сократите.")
    elif name.removeprefix(OWN_PREFIX) in norm_editions():
        errors.append(f"{ws.title}!B2: «{name}» — так называется редакция программы; выберите другое название.")
    years = str(ws["B3"].value or "").strip().replace("–", "-")
    if not re.fullmatch(r"\d{4}-\d{4}", years):
        errors.append(f"{ws.title}!B3: годы действия — как «2026-2029».")

    head_row = next((r for r in range(1, 20) if str(ws.cell(row=r, column=1).value or "").strip() == "Ранг от"), None)
    if head_row is None:
        return None, errors + [f"На листе «{SHEET}» нет заголовка таблицы «Ранг от» — загрузите скачанную копию."]
    rows: list[dict] = []
    marks: dict[int, list[int]] = {c: [] for c in CLASSES}
    r = head_row
    while True:
        r += 1
        lo_v, label = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        if lo_v in (None, "") and label in (None, ""):
            break
        try:
            lo = Fraction(str(lo_v).strip().replace(",", "."))
        except (ValueError, ZeroDivisionError):
            errors.append(f"{at(r, 1)}: «Ранг от» — число (например, 20 или 0,5).")
            continue
        row: dict = {"min": _num_text(lo), "label": str(label or "").strip()}
        if not row["label"]:
            errors.append(f"{at(r, 2)}: впишите ранг, как в документе (например, «20-24,9»).")
        for i, (key, q) in enumerate(QUALS, start=3):
            v = ws.cell(row=r, column=i).value
            if v in (None, ""):
                continue
            if key == "Y3" and str(v).strip().lower() == CONDITION:
                row["Y3_condition"] = True
                continue
            try:
                f = Fraction(str(v).strip().replace(",", "."))
            except (ValueError, ZeroDivisionError):
                f = None
            if f is None or f.denominator != 1 or not 1 <= f <= 1000:
                errors.append(f"{at(r, i)}: {q} — целый процент от 1 до 1000 или пусто"
                              + (" (или «условие»)" if key == "Y3" else "") + f", а здесь «{v}».")
                continue
            row[key] = int(f)
        got = [(q, row[k]) for k, q in QUALS if k in row]
        for (q1, a), (q2, b) in pairwise(got):
            if b < a:
                errors.append(f"{ws.title}, строка {r} («{row['label']}»): {q2} ({b} %) меньше, чем {q1} ({a} %) — "
                              "у младшего разряда процент не может быть строже.")
        for c in CLASSES:
            if str(ws.cell(row=r, column=2 + len(QUALS) + c).value or "").strip().lower() in YES:
                marks[c].append(len(rows))
        rows.append(row)
    if not rows:
        errors.append(f"На листе «{SHEET}» нет строк таблицы под заголовком «Ранг от».")
    mins = [Fraction(x["min"]) for x in rows]
    for a, b in pairwise(mins):
        if b <= a:
            errors.append(f"{ws.title}: «Ранг от» должен расти сверху вниз без повторов — {_num_text(b)} после "
                          f"{_num_text(a)}.")
            break
    class_rows = {}
    for c in CLASSES:
        idx = marks[c]
        if not idx:
            errors.append(f"{ws.title}, столбец «{c} кл»: отметьте «да» строки, которые относятся к {c} классу.")
        elif idx != list(range(idx[0], idx[-1] + 1)):
            errors.append(f"{ws.title}, столбец «{c} кл»: строки класса должны идти подряд, без пропусков.")
        else:
            class_rows[str(c)] = [rows[idx[0]]["min"], rows[idx[-1]]["min"]]

    d = dict(base or {})
    d.pop("source", None)
    d.pop("row", None)
    if COND_SHEET in wb.sheetnames:
        wc = wb[COND_SHEET]
        by_label = {str(wc.cell(row=i, column=1).value or "").strip(): (i, wc.cell(row=i, column=2).value)
                    for i in range(2, wc.max_row + 1)}
        rev_status = {v: k for k, v in STATUS.items()}
        for label, path, kind in CONDITIONS:
            if label not in by_label:
                continue
            i, v = by_label[label]
            where = f"{wc.title}!B{i}"
            text = str(v).strip() if v is not None else ""
            try:
                if kind == "int":
                    val = int(Fraction(text.replace(",", ".")))
                elif kind == "num":
                    val = _num_text(Fraction(text.replace(",", ".")))
                elif kind == "status":
                    val = rev_status.get(text, text)
                    if val not in STATUS:
                        raise ValueError
                elif kind == "bool":
                    if text.lower() not in ("да", "нет"):
                        raise ValueError
                    val = text.lower() == "да"
                elif kind == "date":
                    val = (v.date() if isinstance(v, datetime) else datetime.strptime(text, "%d.%m.%Y").date()
                           ).isoformat()
                else:
                    val = text
            except (ValueError, ZeroDivisionError, KeyError):
                hint = {"int": "целое число", "num": "число", "status": "одно из: " + "; ".join(STATUS.values()),
                        "bool": "«да» или «нет»", "date": "дата дд.мм.гггг"}.get(kind, "текст")
                errors.append(f"{where}: «{label}» — {hint}, а здесь «{text}».")
                continue
            if kind == "text" and not val:
                continue
            node = d
            for k in path[:-1]:
                node[k] = dict(node[k]) if isinstance(node.get(k), dict) else {}
                node = node[k]
            node[path[-1]] = val
    if any(r.get("Y3_condition") for r in rows) and not _get(d, ("junior_III", "text")):
        errors.append(f"В таблице у IIIю «{CONDITION}», а на листе «{COND_SHEET}» нет текста условия III юношеского.")
    if errors:
        return None, errors
    d |= {"edition": years, "title": f"Своя редакция «{name.removeprefix(OWN_PREFIX)}» (годы {years})",
          "name": OWN_PREFIX + name.removeprefix(OWN_PREFIX), "row": rows, "class_rows": class_rows}
    try:
        edition_from_dict(d, d["name"], own=True)
    except (KeyError, ValueError, TypeError) as e:
        return None, [f"Не получилось собрать редакцию: {e}. Сравните книгу с копией, скачанной со страницы."]
    return d, []
