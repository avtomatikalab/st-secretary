"""Комиссия по допуску: документы, решения, взносы, протокол по форме Правил, перезаявки."""

from dataclasses import replace
from datetime import datetime

from conftest import make_application
from openpyxl import load_workbook

from st_secretary.commission import (
    ADMITTED,
    PENDING,
    REJECTED,
    clean_own_docs,
    evaluate,
    person_key,
    protocol_row,
    reentry_record,
    required_docs,
)
from st_secretary.exporters.commission_xlsx import write_commission_report
from st_secretary.importers.preapp_xlsx import read_preapplication
from st_secretary.preapp import process

ALL_DOCS = {"id": True, "med": True, "book": True, "oms": True, "ins": True}


def apps(tmp_path):
    a = make_application(tmp_path / "Лесовики.xlsx", "Лесовики", "Красноярск", "Иванов Пётр Сергеевич",
                         "89131234567", 3, [
                             ["Лесовики", "Красноярск", "Иванов Пётр Сергеевич", "Иванов Пётр Сергеевич",
                              "17.10.1989", "II", "м", "М/Ж", 3],
                             ["Лесовики", "Красноярск", "Иванов Пётр Сергеевич", "Смирнова Анна Олеговна",
                              "12.03.2001", "КМС", "ж", "М/Ж", 3],
                             ["Лесовики", "Красноярск", "Иванов Пётр Сергеевич", "Кузьмин Олег Игоревич",
                              "28.02.1960", "б/р", "м", "М/Ж", 3],
                         ])
    b = make_application(tmp_path / "Сосна.xlsx", "Сосна", "Томск", "Орлов Павел Ильич", "89137654321", 3, [
        ["Сосна", "Томск", "Орлов Павел Ильич", "Орлов Павел Ильич", "11.05.1985", "I", "м", "М/Ж", 3],
        ["Сосна", "Томск", "Орлов Павел Ильич", "Белова Ирина Петровна", "03.07.1999", "III", "ж", "М/Ж", 3],
        ["Сосна", "Томск", "Орлов Павел Ильич", "Пестов Юрий Андреевич", "01.01.2009", "2ю", "м", "М/Ж", 3],
    ])
    return [read_preapplication(a), read_preapplication(b)]


def check(tmp_path, psr_card, data):
    result = process(apps(tmp_path), psr_card)
    return {t.file: t for t in evaluate(result, ["Лесовики.xlsx", "Сосна.xlsx"], psr_card, data)}


def all_docs(names):
    return {"team_docs": {"app": True, "doctor": True},
            "people": {person_key(n): {"docs": dict(ALL_DOCS)} for n in names}}


LES = ["Иванов Пётр Сергеевич", "Смирнова Анна Олеговна", "Кузьмин Олег Игоревич"]
SOS = ["Орлов Павел Ильич", "Белова Ирина Петровна", "Пестов Юрий Андреевич"]


def test_nothing_marked_everyone_waits(tmp_path, psr_card):
    teams = check(tmp_path, psr_card, {})
    les = teams["Лесовики.xlsx"]
    assert les.status == PENDING and all(p.status == PENDING for p in les.persons)
    assert "нет: паспорт, мед. допуск, книжка, омс, страховка" in les.persons[0].why
    assert any("нет: заявка, допуск врача" in p for p in les.problems)
    assert les.fee_due == 3000 and les.fee_status == "не оплачено"  # 3000 ₽ за команду по карточке


def test_all_documents_admit_team_and_decisions(tmp_path, psr_card):
    data = {"teams": {"Лесовики.xlsx": all_docs(LES) | {"fee_paid": 3000, "number": 7},
                      "Сосна.xlsx": all_docs(SOS)}}
    teams = check(tmp_path, psr_card, data)
    les, sos = teams["Лесовики.xlsx"], teams["Сосна.xlsx"]
    assert les.status == ADMITTED and les.fee_status == "оплачено" and les.number == 7
    # Пестову 16 лет — допуск только по решению ГСК: сам не допускается
    pestov = next(p for p in sos.persons if p.entry.name.full == "Пестов Юрий Андреевич")
    assert pestov.status == PENDING and "нужно решение комиссии" in pestov.why and sos.status == PENDING

    data["teams"]["Сосна.xlsx"]["people"][person_key("Пестов Юрий Андреевич")] |= {
        "decision": ADMITTED, "reason": "решение ГСК от 19.09"}
    sos = check(tmp_path, psr_card, data)["Сосна.xlsx"]
    assert sos.status == ADMITTED
    assert next(p for p in sos.persons if p.status == ADMITTED and "решением комиссии" in p.why)


def test_age_needs_commission_decision_even_if_checked_in_preapps(tmp_path, psr_card):
    """«Проверено» на предзаявках не заменяет решение ГСК о допуске младше, чем в Положении."""
    from dataclasses import replace

    from st_secretary.issues import CHECKED

    result = process(apps(tmp_path), psr_card)
    result.issues[:] = [replace(i, severity=CHECKED) if i.severity == "warning" else i for i in result.issues]
    sos = {t.file: t for t in evaluate(result, ["Лесовики.xlsx", "Сосна.xlsx"], psr_card,
                                       {"teams": {"Сосна.xlsx": all_docs(SOS)}})}["Сосна.xlsx"]
    pestov = next(p for p in sos.persons if p.entry.name.full == "Пестов Юрий Андреевич")
    assert pestov.status == PENDING and "нужно решение комиссии" in pestov.why


def test_rejected_member_breaks_team_composition(tmp_path, psr_card):
    data = {"teams": {"Лесовики.xlsx": all_docs(LES)}}
    data["teams"]["Лесовики.xlsx"]["people"][person_key("Кузьмин Олег Игоревич")] |= {
        "decision": REJECTED, "reason": "нет страховки"}
    les = check(tmp_path, psr_card, data)["Лесовики.xlsx"]
    assert les.status == PENDING
    assert any("осталось 2 чел., нужно 3 — нужна перезаявка" in p for p in les.problems)
    assert len(les.counted) == 2
    data["teams"]["Лесовики.xlsx"]["decision"] = REJECTED
    les = check(tmp_path, psr_card, data)["Лесовики.xlsx"]
    assert les.status == REJECTED and len(les.counted) == 3  # у недопущенной команды в протоколе — весь состав


def test_documents_list_follows_settings(tmp_path, psr_card):
    data = {"settings": {"docs": ["id", "med", "pd"], "team_docs": ["app"]}}
    pdocs, tdocs = required_docs(data)
    assert [d.key for d in pdocs] == ["id", "med", "pd"] and [d.key for d in tdocs] == ["app"]
    les = check(tmp_path, psr_card, data)["Лесовики.xlsx"]
    assert "нет: паспорт, мед. допуск, согласие" in les.persons[0].why


def test_protocol_row_counts_like_the_form(tmp_path, psr_card):
    les = check(tmp_path, psr_card, {"teams": {"Лесовики.xlsx": all_docs(LES)}})["Лесовики.xlsx"]
    row = protocol_row(les, psr_card)
    # МС КМС I II III 1ю 2ю 3ю б/р
    assert row["total"] == 3 and row["quals"] == [0, 1, 0, 1, 0, 0, 0, 0, 1]
    assert (row["men"], row["women"]) == (2, 1)
    # <18, 18–21, 22–25, 26–40, >40 — полных лет на 20.09.2025: 35, 24, 65
    assert row["ages"] == [0, 0, 1, 1, 1]
    assert row["decision"] == "допущена" and row["territory"] == "Красноярск"


def test_report_workbook(tmp_path, psr_card):
    teams = list(check(tmp_path, psr_card, {"teams": {"Лесовики.xlsx": all_docs(LES) | {"fee_paid": 1000}}}).values())
    out = write_commission_report(teams, psr_card, {}, tmp_path / "k.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Протокол комиссии", "Ведомость взносов", "Документы участников"]
    ws = wb["Протокол комиссии"]
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Протокол комиссии по допуску участников" in text and "Итого:" in text
    assert any(t.startswith("Председатель комиссии по допуску участников") for t in text)
    assert any("С.А. Секретарёва" in t or "А.И. Секретарёва" in t for t in text)  # подпись главного секретаря
    assert "Смирнова Анна Олеговна" not in text  # ФИО — только в листе «Документы участников»
    fees = [str(c.value) for row in wb["Ведомость взносов"].iter_rows() for c in row if c.value is not None]
    assert "частично" in fees and "не оплачено" in fees
    assert "Смирнова Анна Олеговна" in [c.value for c in wb["Документы участников"]["D"]]


def test_reentry_record(tmp_path, psr_card):
    les = process(apps(tmp_path), psr_card).teams[0]
    after = [{"fio": "Иванов Пётр Сергеевич", "group": "М/Ж", "cls": "3"},
             {"fio": "Смирнова Анна Олеговна", "group": "М/Ж", "cls": "3"},
             {"fio": "Новиков Илья Петрович", "group": "М/Ж", "cls": "3"}]
    rec = reentry_record(les.entries, after, datetime(2025, 9, 20, 9, 30), datetime(2025, 9, 20, 10, 0), [])
    assert "выбыл(а): Кузьмин Олег Игоревич" in rec["text"] and "включён(а): Новиков Илья Петрович" in rec["text"]
    assert rec["late"] and not rec["repeat"]  # за 30 минут до старта — позже, чем за час
    rec = reentry_record(les.entries, after, datetime(2025, 9, 20, 8, 0), datetime(2025, 9, 20, 10, 0), [rec])
    assert not rec["late"] and rec["repeat"]  # вовремя, но повторная — по Правилам не принимается


def test_doctor_admission_sets_med_except_reentered_and_unticked(tmp_path, psr_card):
    """Правки.md, п. 23: у команды «Допуск врача в заявке» → «Мед. допуск» всем участникам сам; кроме добавленных
    перезаявкой и тех, у кого галочку сняли (врач не допустил); сняли «Допуск врача» — уходят только поставленные сами."""
    no_med = {k: True for k in ALL_DOCS if k != "med"}
    tm = {"team_docs": {"app": True, "doctor": True},
          "people": {person_key(n): {"docs": dict(no_med)} for n in LES}}
    data = {"teams": {"Лесовики.xlsx": tm}}
    les = check(tmp_path, psr_card, data)["Лесовики.xlsx"]
    assert les.status == ADMITTED and all(p.docs["med"] and p.auto_docs == {"med"} for p in les.persons)

    tm["people"][person_key(LES[2])]["med_off"] = True  # врач не допустил Кузьмина — галочку сняли
    kuz = check(tmp_path, psr_card, data)["Лесовики.xlsx"].persons[2]
    assert not kuz.docs["med"] and kuz.status == PENDING and "мед. допуск" in kuz.why

    tm["reentries"] = [{"text": "включён(а): Смирнова Анна Олеговна", "added": [person_key(LES[1])]}]
    smirnova = check(tmp_path, psr_card, data)["Лесовики.xlsx"].persons[1]
    assert not smirnova.docs["med"] and not smirnova.auto_docs  # добавлена перезаявкой — мед. допуск отдельно
    tm["reentries"] = [{"text": "выбыл(а): Носов Глеб Андреевич; включён(а): Смирнова Анна Олеговна"}]
    assert not check(tmp_path, psr_card, data)["Лесовики.xlsx"].persons[1].docs["med"]  # запись прежней версии

    tm["people"][person_key(LES[0])]["docs"]["med"] = True  # у Иванова — отметка секретаря (справка)
    tm["team_docs"]["doctor"] = False
    les = check(tmp_path, psr_card, data)["Лесовики.xlsx"]
    assert [p.docs["med"] for p in les.persons] == [True, False, False]
    assert not any(p.auto_docs for p in les.persons)


def test_unreviewed_application_needs_commission_decision_with_basis(tmp_path, psr_card):
    """Правки.md, п. 24: заявку секретарь не отметил «Проверено» — сама команда не допускается; допустить можно
    решением комиссии с обязательным основанием, в протоколе — «допущена решением комиссии: …»; участника без
    документов — тоже только с основанием. Всё это — в списке «допущены без проверки секретаря»."""
    from st_secretary.commission import without_check

    result = process(apps(tmp_path), psr_card)
    files = ["Лесовики.xlsx", "Сосна.xlsx"]
    data = {"teams": {"Лесовики.xlsx": all_docs(LES)}}

    def run(reviewed):
        return {t.file: t for t in evaluate(result, files, psr_card, data, None, reviewed)}["Лесовики.xlsx"]

    les = run({"Лесовики.xlsx": False})
    assert les.status == PENDING and les.unreviewed
    assert les.problems[0].startswith("заявку секретарь не проверил") and les.problem_targets[0] == "preapp:Лесовики.xlsx"
    assert run({"Лесовики.xlsx": True}).status == ADMITTED
    assert run(None).status == ADMITTED  # без сведений о проверке (как раньше)

    data["teams"]["Лесовики.xlsx"]["decision"] = ADMITTED
    les = run({"Лесовики.xlsx": False})
    assert les.status == PENDING and any("укажите основание" in x for x in les.problems)
    data["teams"]["Лесовики.xlsx"]["note"] = "заявка проверена на комиссии, решение ГСК"
    les = run({"Лесовики.xlsx": False})
    assert les.status == ADMITTED and les.by_decision
    row = protocol_row(les, psr_card)
    assert row["decision"] == "допущена решением комиссии: заявка проверена на комиссии, решение ГСК"
    assert "заявку секретарь не проверил" in row["remarks"]
    assert without_check([les]) == ["«Лесовики» — заявку секретарь не проверил (основание: заявка проверена на "
                                    "комиссии, решение ГСК)"]

    pm = data["teams"]["Лесовики.xlsx"]["people"][person_key(LES[0])]
    pm["docs"]["ins"] = False
    pm["decision"] = ADMITTED  # без страховки, основания нет
    p = run({"Лесовики.xlsx": True}).persons[0]
    assert p.status == PENDING and "укажите основание" in p.why
    pm["reason"] = "справка будет до старта"
    les = run({"Лесовики.xlsx": True})
    assert les.persons[0].status == ADMITTED
    assert without_check([les])[-1].startswith("Иванов Пётр Сергеевич («Лесовики») — допущен решением комиссии")


def test_person_documents_marked_once_count_in_all_his_teams(tmp_path, psr_card):
    """Правки.md, п. 20: один человек (ФИО + дата рождения) в двух командах — документы отмечают один раз, в любой;
    тёзка с другой датой рождения — другой человек."""
    kedr = make_application(tmp_path / "Кедр.xlsx", "Кедр", "Красноярск", "Иванов Пётр Сергеевич", "89130000000", 3, [
        ["Кедр", "Красноярск", "Иванов Пётр Сергеевич", "Иванов Пётр Сергеевич", "17.10.1989", "II", "м", "М/Ж", 3],
        ["Кедр", "Красноярск", "Иванов Пётр Сергеевич", "Кузьмин Олег Игоревич", "01.01.1990", "б/р", "м", "М/Ж", 3],
        ["Кедр", "Красноярск", "Иванов Пётр Сергеевич", "Белова Ирина Петровна", "03.07.1999", "III", "ж", "М/Ж", 3],
    ])
    result = process(apps(tmp_path) + [read_preapplication(kedr)], psr_card)
    files = ["Лесовики.xlsx", "Сосна.xlsx", "Кедр.xlsx"]
    data = {"teams": {"Лесовики.xlsx": all_docs(LES)}}  # у «Лесовиков» всё отмечено
    teams = {t.file: t for t in evaluate(result, files, psr_card, data)}
    ivanov, kuzmin, belova = teams["Кедр.xlsx"].persons
    assert all(ivanov.docs.values()) and set(ivanov.shared_docs) == set(ALL_DOCS)
    assert ivanov.shared_docs["med"] == "Лесовики" and not ivanov.missing
    assert not any(kuzmin.docs.values())  # Кузьмин Олег Игоревич 01.01.1990 — не тот, что в «Лесовиках» (1960)
    assert not any(belova.docs.values())  # Белова Ирина Петровна есть в «Сосне», но там ничего не отмечено


def test_own_documents_for_minors_with_or_pair(tmp_path, psr_card):
    """Правки, п. 30 (ИБ Кубка г. Красноярска, п. 9.2): свои документы по Положению — расписка родителей (только
    несовершеннолетним) ИЛИ приказ о полномочиях представителя у команды (если в ней есть несовершеннолетний);
    журнал инструктажа — у команды; ОМС не требуется."""
    comp = replace(psr_card, zachety=[replace(psr_card.zachety[0], age_from=14, age_from_by_gsk=None)])
    own = clean_own_docs([{"title": "Расписка родителей несовершеннолетнего", "scope": "person", "when": "minor"},
                          {"title": "Приказ о полномочиях представителя", "scope": "team", "when": "team_minor"},
                          {"title": "Журнал инструктажа по технике безопасности", "scope": "team"}])
    rasp, prikaz, journal = (d["key"] for d in own)
    own[0]["alt"] = prikaz  # «или» — в обе стороны
    docs = {k: v for k, v in ALL_DOCS.items() if k != "oms"}
    data = {"settings": {"docs": ["id", "med", "book", "ins", rasp], "team_docs": ["app", "doctor", prikaz, journal],
                         "own_docs": own},
            "teams": {f: {"team_docs": {"app": True, "doctor": True, journal: True},
                          "people": {person_key(n): {"docs": dict(docs)} for n in names}}
                      for f, names in (("Лесовики.xlsx", LES), ("Сосна.xlsx", SOS))}}
    pdocs, tdocs = required_docs(data)
    assert [d.short for d in pdocs][-1] == "Расписка родителей несовершеннолетнего" and "oms" not in [d.key for d in pdocs]
    assert [d.short for d in tdocs][-2:] == ["Приказ о полномочиях…", "Журнал инструктажа по…"]

    teams = check(tmp_path, comp, data)
    les, sos = teams["Лесовики.xlsx"], teams["Сосна.xlsx"]
    # у «Лесовиков» нет несовершеннолетних — ни расписка, ни приказ не нужны
    assert les.status == ADMITTED and prikaz in les.team_not_needed and all(rasp in p.not_needed for p in les.persons)
    pestov = next(p for p in sos.persons if p.entry.name.full == "Пестов Юрий Андреевич")  # 16 лет
    orlov = next(p for p in sos.persons if p.entry.name.full == "Орлов Павел Ильич")
    # без расписки и без приказа — ждут оба: участник (нет расписки) и команда (нет приказа)
    assert [d.key for d in pestov.missing] == [rasp] and pestov.status == PENDING and rasp in orlov.not_needed
    assert [d.key for d in sos.missing_team_docs] == [prikaz] and sos.status == PENDING

    # приказ у команды — несовершеннолетний без расписки допущен
    data["teams"]["Сосна.xlsx"]["team_docs"][prikaz] = True
    sos = check(tmp_path, comp, data)["Сосна.xlsx"]
    pestov = next(p for p in sos.persons if p.entry.name.full == "Пестов Юрий Андреевич")
    assert pestov.status == ADMITTED and not pestov.missing and "Приказ" in pestov.covered[rasp]
    assert sos.status == ADMITTED

    # приказа нет, но расписки у всех несовершеннолетних — приказ не нужен
    data["teams"]["Сосна.xlsx"]["team_docs"][prikaz] = False
    data["teams"]["Сосна.xlsx"]["people"][person_key("Пестов Юрий Андреевич")]["docs"][rasp] = True
    sos = check(tmp_path, comp, data)["Сосна.xlsx"]
    assert sos.status == ADMITTED and "Расписка" in sos.team_covered[prikaz]

    # журнал инструктажа нужен всем командам
    data["teams"]["Лесовики.xlsx"]["team_docs"][journal] = False
    les = check(tmp_path, comp, data)["Лесовики.xlsx"]
    assert [d.key for d in les.missing_team_docs] == [journal] and les.status == PENDING
