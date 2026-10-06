"""Страницы карточки соревнования: подписи, примеры, подсказки, ссылки (Правки, п. 56–65). Данные — выдуманные."""

import re
from dataclasses import replace

import pytest

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_web import base

from st_secretary.competition import Official, territory_suggestions
from st_secretary.importers.card_xlsx import load_card, write_card
from st_secretary.issues import ERROR
from st_secretary.norms import PercentMethod, percent_of_winner
from st_secretary.web.app import create_app
from st_secretary.web.forms import card_to_form, form_from_data, form_to_card


@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(tmp_path / "данные", opener=lambda p: None, docs_dir=tmp_path / "документы",
                                 board_host="127.0.0.1"))


def test_one_class_label_is_plain(client, psr_card):
    """П. 56: галочка «один класс» подписана понятно — и в редактировании, и в просмотре карточки."""
    f = client.app.state.store.create(replace(psr_card, one_class=True))
    edit = client.get(base(f) + "/card/edit").text
    assert "Один человек — только в одном классе" in edit and "нельзя сразу во 2 и 3 кл." in edit
    assert "так бывает в Положении или ИБ" in edit
    view = client.get(base(f) + "/card").text
    assert "Один человек — только в одном классе" in view and "проверка заявок покажет ошибку" in view


def test_new_competition_from_example_is_full(client):
    """П. 57: «Новое соревнование» по примеру (учебный режим) — сразу с ГСК и зачётами учебного соревнования."""
    ex = client.get("/example/card").json()["fields"]
    assert ex["learn_full"] == "1" and 'name="learn_full"' in client.get("/new").text
    main = {k: v for k, v in ex.items() if not k.startswith(("g-", "z-"))}
    r = client.post("/new", data=main, follow_redirects=False)
    assert r.status_code == 303
    comp = client.app.state.store.all()[0].load()
    assert [z.key for z in comp.zachety] == ["М/Ж_3", "М/Ж_2"] and comp.official("Главный судья").fio
    assert not [i for i in comp.check() if i.severity == ERROR]
    r = client.post("/new", data=main | {"learn_full": ""}, follow_redirects=False)  # без галочки — как раньше
    assert r.status_code == 303 and len(client.app.state.store.all()) == 2


def test_section_examples_gsk_and_zachet(client, psr_card):
    """П. 57: в разделах «ГСК» и «Зачёты» — свои «Заполнить примером»; ГСК меняет только g-поля (спросит, если судьи
    уже вписаны), зачёт — добавляется к существующим: первый учебный, которого ещё нет (группа_класс)."""
    f = client.app.state.store.create(psr_card)
    edit = client.get(base(f) + "/card/edit").text
    assert 'data-example="/example/card?part=gsk" data-example-confirm="g-fio"' in edit
    assert 'data-example="/example/card?part=zachet"' in edit
    gsk = client.get("/example/card?part=gsk").json()
    assert gsk["clear"] == ["g"] and gsk["fields"] and all(k.startswith("g-") for k in gsk["fields"])
    z = client.get("/example/card?part=zachet").json()
    assert z["fields"] == {} and z["append"]["box"] == "z-rows" and z["append"]["key"] == ["group", "distance_class"]
    keys = [(r["group"], r["distance_class"]) for r in z["append"]["rows"]]
    assert len(set(keys)) == len(keys) >= 6 and all("zid" not in r for r in z["append"]["rows"])
    # любой учебный зачёт ложится в карточку без ошибок (как его вставит app.js — в новую строку формы)
    rows = card_to_form(psr_card)["zachety"]
    for r in z["append"]["rows"]:
        if (r["group"], r["distance_class"]) in {(x["group"], x["distance_class"]) for x in rows}:
            continue
        data = {}
        for i, row in enumerate(rows + [r | {"zid": ""}]):
            data |= {f"z-{i}-{k}": v for k, v in row.items()}
        form = card_to_form(psr_card)
        for i, o in enumerate(form["officials"]):
            data |= {f"g-{i}-{k}": v for k, v in o.items() if k != "own"}
        comp, errors = form_to_card(form_from_data(data | form["main"]))
        assert not errors, (r, errors)
        assert len(comp.zachety) == len(rows) + 1 and not [i for i in comp.check() if i.severity == ERROR], r


def test_territory_suggestions_merge_spellings():
    """П. 58: «г. Красноярск» и «Красноярск» — одна территория; сначала — как в этом соревновании, потом частые."""
    got = territory_suggestions(["Энск", "", "г. Энск"], ["Красноярск", "г. Красноярск", "г. Красноярск", "г Ачинск",
                                                         "энск", "Город Ачинск", "Дивногорск"])
    assert got == ["Энск", "г. Красноярск", "г Ачинск", "Дивногорск"]


def test_territory_list_in_gsk(client, psr_card):
    """П. 58: у «Территории» в ГСК — список с вводом (datalist): территории этого и прошлых соревнований на компьютере."""
    store = client.app.state.store
    store.create(replace(psr_card, title="Прошлые соревнования", host_territory="Дивногорск",
                         officials=[Official("Главный судья", "Петров Пётр Петрович", "СС1К", "г. Ачинск")]))
    f = store.create(replace(psr_card, host_territory="Красноярск",
                             officials=[Official("Главный судья", "Сидоров Иван Иванович", "СС1К", "г. Красноярск")]))
    page = client.get(base(f) + "/card/edit").text
    assert 'list="dl-terr"' in page
    terr = re.search(r'<datalist id="dl-terr">(.*?)</datalist>', page).group(1)
    assert re.findall(r'value="([^"]+)"', terr) == ["Красноярск", "г. Ачинск", "Дивногорск"]


def test_percent_method_in_plain_words(client, psr_card, tmp_path):
    """П. 60: способ «% от победителя» — подпись словами, пример серым, формула — в подсказке «?»; расчёт тот же;
    карточка Excel с прежней подписью (формулой) читается."""
    comp = replace(psr_card, percent_method=PercentMethod.POINTS_RELATIVE_TO_WINNER)
    f = client.app.state.store.create(comp)
    view = client.get(base(f) + "/card").text
    assert "Баллы: отставание от победителя в процентах (как в ПСР)" in view
    assert "Победитель 200, команда 250 → 125 %" in view and 'class="tip" title="Как считается:' in view
    assert "(1 + (результат − победитель) / |победитель|) × 100" in view  # формула — в подсказке
    edit = client.get(base(f) + "/card/edit").text
    assert "data-notes=" in edit and "data-note>Победитель 200, команда 250 → 125 %" in edit
    assert percent_of_winner(250, 200, PercentMethod.POINTS_RELATIVE_TO_WINNER) == 125
    # старая карточка Excel: в ячейке — прежняя подпись-формула
    path = write_card(tmp_path / "Карточка.xlsx", comp)
    wb = load_workbook(path)
    cells = [c for row in wb["Карточка"].iter_rows() for c in row
             if c.value == "Баллы: отставание от победителя в процентах (как в ПСР)"]
    assert len(cells) == 1
    cells[0].value = "Баллы: (1 + (результат − победитель) / |победитель|) × 100"
    wb.save(path)
    assert load_card(path).percent_method is PercentMethod.POINTS_RELATIVE_TO_WINNER


def test_zachet_in_card_view_links_to_its_edit(client, psr_card):
    """П. 64: в просмотре карточки название зачёта — ссылка в редактирование сразу к этому зачёту (как замечания)."""
    from html import unescape

    f = client.app.state.store.create(psr_card)
    view = client.get(base(f) + "/card").text
    links = [unescape(h) for h in re.findall(r'<a class="z-edit" href="([^"]+)"', view)]
    assert len(links) == len(psr_card.zachety) >= 1
    for i, h in enumerate(links):
        assert h == f"{base(f)}/card/edit?focus=z-{i}-group#zachety"
        page = client.get(h)
        assert page.status_code == 200 and f'name="z-{i}-group"' in page.text
        field = re.search(rf'<input[^>]*name="z-{i}-group"[^>]*>', page.text).group(0)
        assert f'value="{psr_card.zachety[i].group}"' in field  # тот самый зачёт
