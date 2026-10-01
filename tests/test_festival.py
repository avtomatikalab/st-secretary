"""Фестиваль (Правки, п. 19; решение 039): общая ГСК с заменой в соревновании, судейская бригада, режимы договоров
и табеля, взноса и стартовых номеров. Все данные — выдуманные."""

import io
import zipfile
from dataclasses import replace
from urllib.parse import quote, unquote

import pytest

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient
from test_web import FormFields, base, kedr, sosna, team_form  # заготовки страниц и заявок

from st_secretary.competition import Official
from st_secretary.web.app import create_app


@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(tmp_path / "данные", opener=lambda p: None, docs_dir=tmp_path / "документы",
                                 board_host="127.0.0.1"))


def make_festival(client, *folders, title="Осенний выезд") -> str:
    r = client.post("/festival/new", data={"title": title, "member": [f.id for f in folders]}, follow_redirects=False)
    assert "done=festival_made" in r.headers["location"]
    return unquote(r.headers["location"].split("/festival/")[1].split("?")[0])


def fest_url(fid: str) -> str:
    return "/festival/" + quote(fid, safe="")


def gsk(folder) -> dict[str, str]:
    return {o.role: o.fio for o in folder.load().officials}


def test_festival_gsk_common_with_own_replacement_in_one_competition(client, tmp_path, psr_card):
    store = client.app.state.store
    a = store.create(psr_card)
    b = store.create(replace(psr_card, title="Кубок города N по спортивному туризму", officials=[
        Official("Главный судья", "Другов Пётр Сергеевич", "СС1К", "г. Ачинск"),
        Official("Главный секретарь", "Секретарёва Анна Ивановна", "СС2К", "г. Красноярск")]))
    fid = make_festival(client, a, b)
    fest = store.festival(fid)
    # ГСК фестиваля — из первого соревнования; у второго свой главный судья стал «своей» заменой — ничего не потеряно
    assert [o["fio"] for o in fest["officials"]] == ["Судьин Иван Петрович", "Секретарёва Анна Ивановна"]
    assert fest["own_gsk"] == {b.id: ["Главный судья"]}
    assert gsk(b)["Главный судья"] == "Другов Пётр Сергеевич"

    # ГСК фестиваля вводится один раз — попадает в карточки обоих соревнований, кроме своей замены
    page = client.get(fest_url(fid)).text
    assert "Главная судейская коллегия фестиваля" in page and "Своя ГСК в этом соревновании: главный судья" in page
    form = FormFields(page.replace('action="/festival/' + quote(fid, safe="") + '/gsk"', 'id="g"'), "g").fields
    form["g-1-fio"] = "Новикова Ольга Петровна"
    form["g-3-fio"] = "Безопасов Олег Игоревич"  # заместитель по безопасности
    r = client.post(fest_url(fid) + "/gsk", data=form, follow_redirects=False)
    assert "done=festival_gsk" in r.headers["location"]
    assert gsk(a) == {"Главный судья": "Судьин Иван Петрович", "Главный секретарь": "Новикова Ольга Петровна",
                      "Заместитель главного судьи по безопасности": "Безопасов Олег Игоревич"}
    assert gsk(b)["Главный судья"] == "Другов Пётр Сергеевич" and gsk(b)["Главный секретарь"] == "Новикова Ольга Петровна"

    # в карточке второго соревнования сняли «Своя» у главного судьи — снова общий
    page = client.get(base(b) + "/card/edit").text
    assert "Соревнование входит в фестиваль" in page and 'name="g-0-own" value="1" checked' in page
    form = FormFields(page, "cardform").fields
    form.pop("g-0-own")
    assert client.post(base(b) + "/card/edit", data=form, follow_redirects=False).status_code == 303
    assert gsk(b)["Главный судья"] == "Судьин Иван Петрович" and store.festival(fid)["own_gsk"][b.id] == []

    # в первом — свой секретарь (вписали и отметили «Своя»): ГСК фестиваля его больше не меняет
    form = FormFields(client.get(base(a) + "/card/edit").text, "cardform").fields
    form.update({"g-1-fio": "Своева Дарья Андреевна", "g-1-own": "1"})
    client.post(base(a) + "/card/edit", data=form)
    form = FormFields(client.get(fest_url(fid)).text.replace(
        'action="/festival/' + quote(fid, safe="") + '/gsk"', 'id="g"'), "g").fields
    form["g-1-fio"] = "Другая Секретарша Ивановна"
    client.post(fest_url(fid) + "/gsk", data=form)
    assert gsk(a)["Главный секретарь"] == "Своева Дарья Андреевна"
    assert gsk(b)["Главный секретарь"] == "Другая Секретарша Ивановна"

    # копия фестиваля хранит ГСК и замены; после восстановления — у новых папок
    data = client.get(fest_url(fid) + "/backup.zip").content
    rec = zipfile.ZipFile(io.BytesIO(data)).read("Фестиваль.json").decode("utf-8")
    assert "Своева" not in rec and "Безопасов Олег Игоревич" in rec  # своя замена — в карточке, не в фестивале
    r = client.post("/restore", files={"backup": ("ф.zip", data, "application/zip")}, follow_redirects=False)
    new = store.festival(unquote(r.headers["location"].split("/festival/")[1].split("?")[0]))
    assert len(new["members"]) == 2 and set(new["members"]).isdisjoint({a.id, b.id})
    assert new["own_gsk"].get(new["members"][0]) == ["Главный секретарь"]

    # разъединили — в карточках остаётся полная ГСК
    client.post(fest_url(fid) + "/edit", data={"do": "split"})
    assert store.festival(fid) is None and gsk(a)["Главный секретарь"] == "Своева Дарья Андреевна"
    assert gsk(b)["Заместитель главного судьи по безопасности"] == "Безопасов Олег Игоревич"


def test_festival_brigade_shared_and_one_contract_for_whole_festival(client, tmp_path, psr_card):
    from datetime import date

    from openpyxl import load_workbook

    store = client.app.state.store
    a = store.create(psr_card)
    b = store.create(replace(psr_card, title="Кубок города N по спортивному туризму", date_from=date(2025, 9, 22),
                             date_to=date(2025, 9, 23)))
    fid = make_festival(client, a, b)

    # бригада общая: коменданта добавили у одного соревнования — он и у другого; дни — свои у каждого
    client.post(base(a) + "/contracts/add", data={"fio": "Комендантов Семён Ильич", "role": "Комендант",
                                                  "category": "б/к"})
    assert "Комендантов Семён Ильич" in client.get(base(b) + "/contracts").text
    assert "бригада (добавленные ниже) — общая" in client.get(base(b) + "/contracts").text
    from st_secretary import staff as sf

    assert all(d <= date(2025, 9, 22) for p in sf.people(a.load(), a.contracts()) for d in p.days)

    # один договор и табель на фестиваль: период — от первого до последнего дня, ГСК и бригада — в одном табеле
    r = client.post(fest_url(fid) + "/modes", data={"contracts": "festival", "fee": "each", "numbers": "festival"},
                    follow_redirects=False)
    assert "done=festival_modes" in r.headers["location"]
    page = client.get(base(b) + "/contracts").text
    assert "Один договор и табель на весь фестиваль" in page and "Комендантов Семён Ильич" in page
    form = FormFields(page.replace('action="' + base(b) + '/contracts/days"', 'id="d"'), "d").fields
    assert form, "форма табеля не найдена"
    key = next(k for k, v in form.items() if k.endswith("-key") and v.startswith("комендантов"))
    i = key.split("-")[1]
    marks = {k: v for k, v in form.items() if not k.startswith(f"p-{i}-d-")}
    marks |= {f"p-{i}-d-20250919": "on", f"p-{i}-d-20250923": "on"}
    client.post(base(b) + "/contracts/days", data=marks)
    from st_secretary import festival as fv

    whole = fv.joint_comp(store.festival(fid), [a.load(), b.load()])
    team = {p.fio: p for p in sf.people(whole, a.contracts())}  # те же отметки — у другого соревнования
    assert team["Комендантов Семён Ильич"].days == [date(2025, 9, 19), date(2025, 9, 23)]
    assert team["Судьин Иван Петрович"].days == [date(2025, 9, d) for d in (20, 21, 22, 23)]
    xlsx = client.get(base(a) + "/contracts/file/tabel").content
    ws = load_workbook(io.BytesIO(xlsx)).active
    assert "Фестиваля «Осенний выезд»" in ws["A3"].value and "Кубок города N" in ws["A3"].value
    assert ws["A4"].value.startswith("20–23 сентября 2025 г.")  # даты всего фестиваля

    # разъединили — бригада фестиваля остаётся в договорах каждого соревнования
    client.post(fest_url(fid) + "/edit", data={"do": "split"})
    assert any(x["fio"] == "Комендантов Семён Ильич" for x in b.contracts().get("extra", []))
    assert any(x["fio"] == "Комендантов Семён Ильич" for x in a.contracts().get("extra", []))


def numbers(folder) -> dict[str, int | None]:
    return {k: v.get("number") for k, v in folder.admission().get("teams", {}).items()}


def test_festival_start_numbers_common_or_own_in_each_competition(client, tmp_path, psr_card):
    store = client.app.state.store
    a = store.create(psr_card)
    b = store.create(replace(psr_card, title="Кубок города N по спортивному туризму"))
    a.add_preapp("Сосна.xlsx", sosna(tmp_path))
    a.add_preapp("Кедр.xlsx", kedr(tmp_path))
    b.add_preapp("Кедр команда.xlsx", kedr(tmp_path))  # та же команда (название и территория), другой файл
    fid = make_festival(client, a, b)

    # общие (по умолчанию): нумерация сразу во всех соревнованиях, у «Кедра» — один номер
    r = client.post(base(b) + "/admission/numbers", data={"mode": "missing"}, follow_redirects=False)
    assert "adm_numbers_fest" in r.headers["location"]
    assert numbers(a) == {"Кедр.xlsx": 1, "Сосна.xlsx": 2} and numbers(b) == {"Кедр команда.xlsx": 1}
    assert "Общие на фестиваль" in client.get(base(a) + "/admission").text

    # номер поправили в одном соревновании — он и в другом; совпал с чужим — предупреждение по фестивалю
    page = client.get(base(b) + "/admission").text
    form = team_form(page, "Кедр команда.xlsx") | {"number": "7"}
    r = client.post(base(b) + "/admission/team", data=form, headers={"x-autosave": "1"})
    assert numbers(a)["Кедр.xlsx"] == 7
    form = team_form(client.get(base(a) + "/admission").text, "Сосна.xlsx") | {"number": "7"}
    r = client.post(base(a) + "/admission/team", data=form, headers={"x-autosave": "1"})
    assert "уже у команды «Кедр»" in r.json()["team"]
    assert "номер 7 — у разных команд" in client.get(base(a) + "/admission").text

    # свои номера в каждом соревновании: правка в одном другое не трогает
    client.post(fest_url(fid) + "/modes", data={"contracts": "each", "fee": "each", "numbers": "each"})
    form = team_form(client.get(base(b) + "/admission").text, "Кедр команда.xlsx") | {"number": "3"}
    client.post(base(b) + "/admission/team", data=form, headers={"x-autosave": "1"})
    assert numbers(a)["Кедр.xlsx"] == 7 and numbers(b)["Кедр команда.xlsx"] == 3
    client.post(base(b) + "/admission/numbers", data={"mode": "all"})
    assert numbers(b) == {"Кедр команда.xlsx": 1} and numbers(a) == {"Кедр.xlsx": 7, "Сосна.xlsx": 7}


def test_festival_one_fee_for_whole_festival_or_by_competition(client, tmp_path, psr_card):
    from openpyxl import load_workbook

    store = client.app.state.store
    a = store.create(psr_card)
    b = store.create(replace(psr_card, title="Кубок города N по спортивному туризму"))
    a.add_preapp("Кедр.xlsx", kedr(tmp_path))
    a.add_preapp("Сосна.xlsx", sosna(tmp_path))
    b.add_preapp("Кедр.xlsx", kedr(tmp_path))
    fid = make_festival(client, a, b)

    # по каждому соревнованию (по умолчанию) — как в карточке: 3000 ₽ за команду в каждом
    assert "взносов собрано из 6 000 ₽" in client.get(base(a) + "/admission").text
    assert "взносов собрано из 3 000 ₽" in client.get(base(b) + "/admission").text

    # один взнос за фестиваль: 1500 ₽ за участника, человек считается один раз на весь фестиваль
    client.post(fest_url(fid) + "/modes", data={"contracts": "each", "fee": "festival", "numbers": "festival",
                                                "fee_amount": "1500", "fee_per": "участника"})
    page = client.get(fest_url(fid)).text
    assert "Взнос за фестиваль" in page and "Кубок города N" in page
    form = FormFields(page.replace('action="/festival/' + quote(fid, safe="") + '/fees"', 'id="fees"'), "fees").fields
    kedr_i = next(k.split("-")[1] for k, v in form.items() if k.endswith("-key") and v.startswith("кедр"))
    assert ">3</td><td>4500</td>" in page.replace("\n", "")  # Кедр: 3 человека × 1500, хоть и в двух соревнованиях
    form |= {f"f-{kedr_i}-paid": "4500", f"f-{kedr_i}-method": "перевод"}
    r = client.post(fest_url(fid) + "/fees", data=form, follow_redirects=False)
    assert "festival_fees" in r.headers["location"] and "оплачено" in client.get(fest_url(fid)).text
    adm = client.get(base(a) + "/admission").text
    assert "Взнос — один за фестиваль" in adm and 'name="fee_paid"' not in adm
    ws = load_workbook(io.BytesIO(client.get(fest_url(fid) + "/fees.xlsx").content)).active
    assert ws["A2"].value == "Ведомость стартовых взносов" and any(c.value == 4500 for c in ws["I"])
    sheet = load_workbook(io.BytesIO(client.get(base(a) + "/admission/report.xlsx").content))["Ведомость взносов"]
    assert any("один за фестиваль" in str(c.value) for row in sheet.iter_rows() for c in row)


def test_festival_break_between_starts_of_one_person_in_two_competitions(client, tmp_path, psr_card):
    """Правки, п. 25.3: перерыв — общий для соревнований фестиваля; жеребьёвка второго соревнования его соблюдает."""
    from urllib.parse import urlencode

    store = client.app.state.store
    a = store.create(psr_card)
    b = store.create(replace(psr_card, title="Кубок города N по спортивному туризму"))
    a.add_preapp("Кедр.xlsx", kedr(tmp_path))
    a.add_preapp("Сосна.xlsx", sosna(tmp_path))
    b.add_preapp("Кедр.xlsx", kedr(tmp_path))  # те же люди
    b.add_preapp("Сосна.xlsx", sosna(tmp_path))
    fid = make_festival(client, a, b)
    q = "?" + urlencode({"z": "М/Ж_3"})
    times = {"day": "2025-09-20", "first": "10:00", "interval": "10", "break": "30"}
    client.post(base(a) + "/start/draw" + q, data=times)
    assert store.festival(fid)["start_break"] == 30 and "start_break" not in a.run_data()
    client.post(base(a) + "/start/draw" + q, data={"method": "number", "action": "draw"})
    first_a = a.run_data()["zachety"]["М/Ж_3"]["draw"]["order"][0]
    # во втором соревновании та же раскладка: кто стартовал первым в первом — не раньше чем через 30 мин
    r = client.post(base(b) + "/start/draw" + q, data={k: v for k, v in times.items() if k != "break"},
                    follow_redirects=False)
    assert "start_times_fit" in r.headers["location"]
    dr = b.run_data()["zachety"]["М/Ж_3"]["draw"]
    assert dr["times"].get(first_a) and dr["times"][first_a] >= "10:30" and dr["fit"]
    page = client.get(base(b) + "/start" + q).text
    assert "Перерыв участника соблюдён" in page and "всех соревнований фестиваля" in page
