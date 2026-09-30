"""Учёт судейской практики: кто, где и в какой должности судил — по всем соревнованиям в программе.

Для присвоения и подтверждения квалификационных категорий (Квалификационные требования к спортивным судьям,
приказ Минспорта № 1101) учитываются соревнования, должность, вхождение в ГСК, статус соревнований и оценка
судейства. Программа собирает это из карточек соревнований (лист «ГСК»), оценок на странице итогов и
судей, добавленных на странице договоров, и считает баллы по таблице приказа (reference/data/judge_points_1101.json):
на присвоение следующей категории и на подтверждение текущей — за период из таблицы (2 года, 1 год, 4 года),
только соревнования с оценкой «хорошо» и «отлично» (примечание 3); без оценки — отдельно, «ждут оценки».
Период подтверждения отсчитывается от сегодняшнего дня, а не от даты приказа о категории — её программа не знает.
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from st_secretary.competition import GSK_ROLES, LEVEL_LABELS, Competition
from st_secretary.reference import Level, judge_points
from st_secretary.results import person_key
from st_secretary.textclean import alpha_key

# Должности судей (а не комендантской бригады) среди добавленных на странице договоров.
_JUDGING = re.compile(r"суд|секретар|инспектор|начальник дистанции|хронометр|стартер|стартёр", re.IGNORECASE)
_GSK = {r.lower() for r in GSK_ROLES}


@dataclass
class Record:
    fio: str
    title: str
    date_from: date
    date_to: date
    kind: str
    level: str
    role: str
    category: str
    in_gsk: bool
    grade: str = ""
    competition: str = ""  # папка соревнования
    status: str = ""  # статус по таблице приказа № 1101: ЧР, КР, …, ЧМ, ПМ, ДМ (пусто — не спортивное соревнование)


@dataclass
class Judge:
    key: str
    fio: str
    records: list[Record] = field(default_factory=list)

    @property
    def category(self) -> str:
        """Категория из последнего по дате соревнования."""
        dated = [r for r in self.records if r.category]
        return max(dated, key=lambda r: r.date_from).category if dated else ""

    @property
    def gsk(self) -> int:
        return sum(r.in_gsk for r in self.records)

    @property
    def grades(self) -> Counter:
        return Counter(r.grade for r in self.records if r.grade)

    @property
    def last(self) -> date | None:
        return max((r.date_to for r in self.records), default=None)


def records_of(comp: Competition, folder: str, grades: dict[str, str], brigade: list[dict]) -> list[Record]:
    """Судьи одного соревнования: из карточки (лист «ГСК») и добавленные на странице договоров."""
    out, seen = [], set()
    status = f"{comp.kind}, {LEVEL_LABELS.get(comp.level, '').lower()}"

    def add(fio, role, category):
        k = person_key(fio)
        if not fio.strip() or (k, role) in seen:
            return
        seen.add((k, role))
        out.append(Record(" ".join(fio.split()), comp.title, comp.date_from, comp.date_to, comp.kind, status,
                          role, category if category != "б/к" else "", role.strip().lower() in _GSK,
                          grades.get(k, ""), folder, status_code(comp.kind, comp.level)))

    for o in comp.officials:
        add(o.fio, o.role, o.category)
    for x in brigade:
        if _JUDGING.search(str(x.get("role", ""))):
            add(str(x.get("fio", "")), str(x.get("role", "")), str(x.get("category", "")))
    return out


def status_code(kind: str, level) -> str:
    """Статус соревнования по таблице приказа № 1101 (колонки F–R)."""
    k = kind.lower()
    if k.startswith("физкультурное"):
        return ""  # засчитываются только официальные спортивные соревнования
    if level is Level.ALL_RUSSIAN:
        return {"чемпионат": "ЧР", "кубок": "КР", "первенство": "ПР"}.get(k, "ВС")
    if level is Level.INTERREGIONAL:
        return {"чемпионат": "ЧФО", "первенство": "ПФО"}.get(k, "")
    if level is Level.REGIONAL:
        return {"чемпионат": "ЧС", "кубок": "КС", "первенство": "ПС"}.get(k, "ДС")
    return {"чемпионат": "ЧМ", "первенство": "ПМ"}.get(k, "ДМ")


CATEGORY_CODE = {"ССВК": "ВК", "СС1К": "1К", "СС2К": "2К", "СС3К": "3К", "ЮС": "ЮС"}
NEXT = {"": "2К", "ЮС": "2К", "3К": "2К", "2К": "1К", "1К": "ВК"}
ALL_RUSSIAN = {"ЧР", "КР", "ПР", "ВС"}
SUBJECT = {"ЧС", "КС", "ПС", "ДС"}
ORDER = ["ЧР", "КР", "ПР", "ВС", "ЧФО", "ПФО", "ЧС", "КС", "ПС", "ДС", "ЧМ", "ПМ", "ДМ"]  # от высшего статуса
COUNTED = ("хорошо", "отлично")  # примечание 3: только с оценкой «хорошо» и «отлично»
_ROLE_ALIASES = {"судья-секретарь": "судья секретарь", "секретарь": "судья секретарь",
                 "секретарь соревнований": "судья секретарь", "зам. главного судьи": "заместитель главного судьи",
                 "зам. главного секретаря": "заместитель главного секретаря"}


def role_key(role: str) -> str:
    r = " ".join(role.lower().replace("ё", "е").replace("—", "-").replace("–", "-").split()).replace(" - ", "-")
    return _ROLE_ALIASES.get(r, r)


def at_least(status: str, lowest: str) -> bool:
    """Статус не ниже заданного (по порядку колонок таблицы)."""
    return status in ORDER and ORDER.index(status) <= ORDER.index(lowest)


@dataclass
class Progress:
    """Выполнение требований на присвоение или подтверждение категории."""
    mode: str  # assign — присвоение, confirm — подтверждение
    target: str  # ВК, 1К, 2К, 3К
    years: int
    points: int  # засчитано (оценка «хорошо», «отлично»)
    pending: int  # без оценки — засчитаются после оценки «хорошо» или «отлично»
    need: int | None
    lines: list[tuple[Record, int | None, bool]]  # запись, баллы (None — не по таблице), засчитана
    missing: list[str]
    conditions: list[str]

    @property
    def done(self) -> bool:
        return not self.missing


def progress(j: Judge, today: date) -> list[Progress]:
    """На присвоение следующей категории и на подтверждение текущей (если категория известна)."""
    table = judge_points()
    cur = CATEGORY_CODE.get(j.category, "")
    out = []
    for mode, target in (("assign", NEXT.get(cur)), ("confirm", "3К" if cur == "ЮС" else cur or None)):
        spec = table[mode].get(target or "")
        if not spec:
            continue
        start = date(today.year - spec["years"], today.month, min(today.day, 28))
        lines, counted = [], []
        for r in j.records:
            if not (start <= r.date_from <= today):
                continue
            by_role = {role_key(k): v for k, v in spec["points"].items()}
            pts = by_role.get(role_key(r.role), {}).get(r.status)
            ok = r.grade in COUNTED
            lines.append((r, pts, ok))
            if ok:
                counted.append(r)
        points = sum(p for _, p, ok in lines if ok and p)
        pending = sum(p for r, p, ok in lines if not ok and not r.grade and p)
        missing = []
        need = spec.get("sum")
        if need and points < need:
            missing.append(f"баллов {points} из {need}")
        checks = [("competitions", len(counted), "соревнований"),
                  ("gsk", sum(1 for r in counted if r.in_gsk and at_least(r.status, spec.get("gsk_min_status", "ДМ"))),
                   "в составе ГСК" + (f" (не ниже «{spec['gsk_min_status']}»)" if spec.get("gsk_min_status") else "")),
                  ("high", sum(1 for r in counted if at_least(r.status, spec.get("high_min_status", "ДМ"))),
                   f"статуса не ниже «{spec.get('high_min_status', '')}»"),
                  ("all_russian", sum(1 for r in counted if r.status in ALL_RUSSIAN), "всероссийских"),
                  ("subject", sum(1 for r in counted if r.status in SUBJECT), "субъекта РФ")]
        for key, have, what in checks:
            if spec.get(key) and have < spec[key]:
                missing.append(f"{what} {have} из {spec[key]}")
        out.append(Progress(mode, target, spec["years"], points, pending, need, lines, missing, spec["conditions"]))
    return out


def judges(records: list[Record]) -> list[Judge]:
    by: dict[str, Judge] = {}
    for r in records:
        k = person_key(r.fio)
        by.setdefault(k, Judge(k, r.fio)).records.append(r)
    for j in by.values():
        j.records.sort(key=lambda r: r.date_from, reverse=True)
    return sorted(by.values(), key=lambda j: alpha_key(j.fio))


THIN = Side(style="thin", color="7F7F7F")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="DCE6F1")


def write_practice(people: list[Judge], path: str | Path, today: date | None = None) -> Path:
    """Excel: «Сводка» — строка на судью (с баллами по приказу № 1101, если задан today); «Практика» — строка на
    судью и соревнование."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    heads = ["ФИО", "Категория (последняя)", "Соревнований", "В составе ГСК", "Последнее", "Отлично", "Хорошо",
             "Удовлетворительно", "Неудовлетворительно"]
    if today:
        heads += ["Присвоение: категория", "баллы", "нужно", "не выполнено", "Подтверждение: категория", "баллы",
                  "нужно", "не выполнено"]
    for c, h in enumerate(heads, start=1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True), HEAD, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, j in enumerate(people, start=2):
        g = j.grades
        vals = [j.fio, j.category, len(j.records), j.gsk, f"{j.last:%d.%m.%Y}" if j.last else "", g["отлично"],
                g["хорошо"], g["удовлетворительно"], g["неудовлетворительно"]]
        if today:
            pp = {p.mode: p for p in progress(j, today)}
            for mode in ("assign", "confirm"):
                p = pp.get(mode)
                vals += [p.target, p.points, p.need or "", "; ".join(p.missing)] if p else ["", "", "", ""]
        for c, v in enumerate(vals, start=1):
            ws.cell(r, c, v).border = BOX
    for c, w in enumerate([34, 14, 13, 12, 12, 9, 9, 12, 14, 12, 8, 8, 40, 12, 8, 8, 40], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    ws2 = wb.create_sheet("Практика")
    heads2 = ["ФИО", "Соревнования", "Даты", "Статус", "Должность", "В составе ГСК", "Категория", "Оценка"]
    for c, h in enumerate(heads2, start=1):
        cell = ws2.cell(1, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True), HEAD, BOX
    r = 2
    for j in people:
        for x in j.records:
            dates = f"{x.date_from:%d.%m.%Y}" + (f"–{x.date_to:%d.%m.%Y}" if x.date_to != x.date_from else "")
            vals = [j.fio, x.title, dates, x.level, x.role, "да" if x.in_gsk else "", x.category, x.grade]
            for c, v in enumerate(vals, start=1):
                cell = ws2.cell(r, c, v)
                cell.border, cell.alignment = BOX, Alignment(wrap_text=True, vertical="top")
            r += 1
    for c, w in enumerate([30, 44, 22, 34, 28, 10, 11, 14], start=1):
        ws2.column_dimensions[get_column_letter(c)].width = w
    ws2.freeze_panes = "B2"
    path = Path(path)
    wb.save(path)
    return path
