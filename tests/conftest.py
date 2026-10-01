"""Общие заготовки тестов: карточка соревнования и генератор предзаявок по форме бланка."""

from datetime import date

import pytest
from openpyxl import Workbook

from st_secretary.competition import Competition, Official, Zachet
from st_secretary.norms import PercentMethod
from st_secretary.qualification import Qual
from st_secretary.reference import Level


@pytest.fixture
def psr_card() -> Competition:
    """ПСР, 3 класс, команды по 3 человека (не менее 1 мужчины и 1 женщины), с 22 лет, с 16 — по решению ГСК."""
    return Competition(
        title="Чемпионат города N по спортивному туризму",
        kind="Чемпионат",
        level=Level.MUNICIPAL,
        date_from=date(2025, 9, 20),
        date_to=date(2025, 9, 21),
        place="окрестности г. N",
        host_territory="Красноярск",
        organizers=["Федерация спортивного туризма"],
        norms_edition="2022-2025",
        percent_method=PercentMethod.POINTS_RELATIVE_TO_WINNER,
        preapp_deadline=date(2025, 9, 17),
        officials=[Official("Главный судья", "Судьин Иван Петрович", "СС1К", "г. Красноярск"),
                   Official("Главный секретарь", "Секретарёва Анна Ивановна", "СС2К", "г. Красноярск")],
        zachety=[Zachet("М/Ж", 3, "0840161811Я", age_from=22, age_from_by_gsk=16, min_qual=Qual.BR,
                        team_size=3, min_men=1, min_women=1, fee=3000)],
    )


HEAD = ["№ п/п", "Команда", "Территория", "Представитель", "Фамилия, имя", "Дата рождения", "Разряд по СТ",
        "Пол", "Группа", "Класс дистанции", "Номер чипа (указать, если чип свой)", "Участие в личной дистанции",
        "Участие в дистанции связок", "Номер связки (если больше одной связки)", "Группа"]


def make_class_sheets_application(path, sheets):
    """Бланк делегации спелео (Правки, п. 29): по листу на класс — «2 КЛАСС», «3 КЛАСС», шапка как у командного
    бланка, строка «ОБРАЗЕЦ» с примером. sheets — {название листа: строки участников (колонки B–O)}."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        ws["A1"] = "Бланк электронной заявки для участия в соревнованиях"
        ws["A2"] = "ВНИМАНИЕ!!! Перед заполнением прочитайте примечания"
        for col, v in zip("BCDEF", ["Команда", "Территория", "Фамилия, имя, отчество представителя",
                                    "Контактный телефон, адрес эл. почты", "Кол-во участников в делегации"]):
            ws[f"{col}3"] = v
        for col, v in zip("BCDEF", ["Сталактит", "Энск", "Пещерин Олег Петрович", "89000000000", 4]):
            ws[f"{col}4"] = v
        for i, h in enumerate(HEAD, start=1):
            ws.cell(6, i, h)
        ws["A7"] = "ОБРАЗЕЦ"
        ws.append([0, "Жигули", "Волжский район", "Петров Иван", "Иванов Петр", "19.08.1996", "III", "м", "ЮН", 2,
                   None, 1, "см", 1, 1])
        for n in range(1, 19):  # пустые строки бланка с номерами
            ws.cell(8 + n, 1, n)
        for n, r in enumerate(rows, start=1):
            for c, v in enumerate(r, start=2):
                ws.cell(8 + n, c, v)
    wb.save(path)
    return path


def make_application(path, team, territory, rep, contacts, count, rows):
    """Предзаявка в раскладке БЛАНК_ПРЕД_ЗАЯВКИ_ST. rows — списки значений колонок B–O (без №)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявка"
    ws["A1"] = "Бланк электронной заявки для участия в соревнованиях"
    for col, v in zip("BCDEF", ["Команда", "Территория", "Фамилия, имя, отчество представителя",
                                "Контактный телефон, адрес эл. почты", "Кол-во участников в делегации"]):
        ws[f"{col}3"] = v
    for col, v in zip("BCDEF", [team, territory, rep, contacts, count]):
        ws[f"{col}4"] = v
    for i, h in enumerate(HEAD, start=1):
        ws.cell(6, i, h)
    ws["A7"] = "ОБРАЗЕЦ"
    ws.append([0, "Турклуб", "Красноярск", "Сидоров Николай Николаевич", "Иванов Иван", "19.08.1981", "1", "м",
               "м/ж", 2, None, 1, "м", 1, 1])
    for n, r in enumerate(rows, start=1):
        ws.cell(9 + n, 1, n)
        for c, v in enumerate(r, start=2):
            ws.cell(9 + n, c, v)
    wb.save(path)
    return path
