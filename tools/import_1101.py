"""Таблица баллов практики судейства (Квалификационные требования к спортивным судьям, приказ Минспорта № 1101)
из Excel-версии с сайта ФСТР → src/st_secretary/reference/data/judge_points_1101.json.

    uv run python tools/import_1101.py "…\\Квалиф_требования_судьи_СТ_1101.xlsx"

Листы «1 Присвоение (практика)» и «2 Подтверждение (практика)»: по категориям — должности и баллы за судейство
соревнований каждого статуса (колонки F–R), период, условия. Объединённые ячейки разворачиваются (значение —
во все ячейки диапазона), поэтому у «Главного секретаря» те же баллы, что у «Главного судьи» в одной строке
таблицы. Числовые условия (сумма баллов, число соревнований, ГСК, статус) записаны ниже рядом с текстом из таблицы.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "src" / "st_secretary" / "reference" / "data" / "judge_points_1101.json"
STATUSES = ["ЧР", "КР", "ПР", "ВС", "ЧФО", "ПФО", "ЧС", "КС", "ПС", "ДС", "ЧМ", "ПМ", "ДМ"]  # колонки F–R
TITLES = {
    "ЧР": "Чемпионат России", "КР": "Кубок России", "ПР": "Первенство России",
    "ВС": "Другие официальные всероссийские спортивные соревнования",
    "ЧФО": "Чемпионат федерального округа, двух и более федеральных округов, чемпионаты г. Москвы и г. Санкт-Петербурга",
    "ПФО": "Первенство федерального округа, двух и более федеральных округов, первенства г. Москвы и г. Санкт-Петербурга",
    "ЧС": "Чемпионат субъекта Российской Федерации", "КС": "Кубок субъекта Российской Федерации",
    "ПС": "Первенство субъекта Российской Федерации", "ДС": "Другие официальные спортивные соревнования субъекта РФ",
    "ЧМ": "Чемпионат муниципального образования", "ПМ": "Первенство муниципального образования",
    "ДМ": "Другие официальные соревнования муниципального образования",
}
CATEGORY = {"всероссийская": "ВК", "первая": "1К", "вторая": "2К", "третья, юс": "3К"}
SHEETS = {"assign": "1 Присвоение (практика)", "confirm": "2 Подтверждение (практика)"}
# Числовые условия — по тексту условий в таблице (текст сохраняется рядом): сумма баллов, минимум соревнований,
# в составе ГСК, «не ниже» какого статуса и сколько таких, минимум всероссийских.
RULES = {
    "assign": {
        "ВК": {"sum": 100, "competitions": 6, "gsk": 3, "gsk_min_status": "ПС", "all_russian": 2, "subject": 2},
        "1К": {"sum": 60, "competitions": 6, "gsk": 2, "high": 2, "high_min_status": "ПФО"},
        "2К": {"sum": 25, "competitions": 3},
    },
    "confirm": {
        "ВК": {"sum": 120, "all_russian_each_year": 1},
        "1К": {"sum": 55, "competitions": 4, "gsk": 1, "high": 1, "high_min_status": "ПФО"},
        "2К": {"sum": 40, "competitions_each_year": 2, "gsk": 1},
        "3К": {"competitions": 1},
    },
}


def grid(ws) -> dict:
    g = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=18):
        for c in row:
            g[(c.row, c.column)] = c.value
    for mr in ws.merged_cells.ranges:
        v = ws.cell(mr.min_row, mr.min_col).value
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, min(mr.max_col, 18) + 1):
                g[(r, c)] = v
    return g


def text(v) -> str:
    return " ".join(str(v or "").split())


def read(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {"source": "Квалификационные требования к спортивным судьям по виду спорта «спортивный туризм» "
                     "(приказ Минспорта России № 1101) — Excel-версия с сайта ФСТР (tssr.ru)",
           "statuses": STATUSES, "titles": TITLES}
    for mode, sheet in SHEETS.items():
        ws = wb[sheet]
        g = grid(ws)
        cats: dict = {}
        for r in range(1, ws.max_row + 1):
            cat = CATEGORY.get(text(g.get((r, 1))).lower())
            if not cat:
                continue
            c = cats.setdefault(cat, {"age": text(g.get((r, 2))), "years": int(float(g.get((r, 3)) or 0)),
                                      "points": {}, "conditions": [], **RULES[mode].get(cat, {})})
            role = text(g.get((r, 5))) or text(g.get((r, 4)))
            first = g.get((r, 6))
            if isinstance(first, str) and len(first) > 20:  # строка условий, а не баллы
                t = text(first)
                if t not in c["conditions"]:
                    c["conditions"].append(t)
                continue
            pts = {s: int(g[(r, 6 + i)]) for i, s in enumerate(STATUSES)
                   if isinstance(g.get((r, 6 + i)), (int, float))}
            if role and pts:
                c["points"][role] = pts
        out[mode] = cats
    return out


def main():
    src = Path(sys.argv[1])
    data = read(src)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    for mode in SHEETS:
        for cat, c in data[mode].items():
            print(mode, cat, c["years"], "лет;", len(c["points"]), "должностей;", len(c["conditions"]), "условий")
    print("→", OUT)


if __name__ == "__main__":
    main()
