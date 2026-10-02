"""Интерфейс в браузере: страницы открываются, карточка сохраняется без потерь, заявки проверяются."""

import io
import os
import re
from dataclasses import replace
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import quote, urlencode

import pytest

pytest.importorskip("fastapi")
from conftest import make_application
from fastapi.testclient import TestClient

from st_secretary.importers.card_xlsx import load_card, write_card
from st_secretary.web.app import create_app
from st_secretary.web.review import issue_key

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MAIN = {"title": "Чемпионат города N по спортивному туризму", "kind": "Чемпионат", "level": "MUNICIPAL",
        "date_from": "2025-09-20", "date_to": "2025-09-21", "place": "окрестности г. N",
        "host_territory": "Красноярск", "organizers": "Федерация спортивного туризма", "calendar_number": "",
        "norms_edition": "2022-2025", "percent_method": "POINTS_RELATIVE_TO_WINNER", "preapp_deadline": "2025-09-17"}


class FormFields(HTMLParser):
    """Поля формы так, как их отправит браузер (содержимое <template> — заготовки строк — не отправляется)."""

    def __init__(self, html: str, form_id: str):
        super().__init__()
        self.form_id, self.inside, self.tpl = form_id, False, 0
        self.fields: dict[str, str] = {}
        self._select = self._textarea = None
        self._first = None
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "form" and a.get("id") == self.form_id:
            self.inside = True
        if tag == "template":
            self.tpl += 1
        if not self.inside or self.tpl:
            return
        if tag == "input" and a.get("type") in ("checkbox", "radio"):
            if a.get("name") and "checked" in a:  # браузер отправляет только отмеченные галочки и переключатели
                self.fields[a["name"]] = a.get("value") or "on"
        elif tag == "input" and a.get("name") and a.get("type") not in ("file", "submit"):
            self.fields[a["name"]] = a.get("value") or ""
        elif tag == "select":
            self._select, self._first = a.get("name"), None
        elif tag == "option" and self._select:
            if self._first is None:
                self._first = a.get("value", "")
                self.fields[self._select] = self._first
            if "selected" in a:
                self.fields[self._select] = a.get("value", "")
        elif tag == "textarea":
            self._textarea = a.get("name")
            self.fields[self._textarea] = ""

    def handle_endtag(self, tag):
        if tag == "template":
            self.tpl -= 1
        elif tag == "form":
            self.inside = False
        elif tag == "select":
            self._select = None
        elif tag == "textarea":
            self._textarea = None

    def handle_data(self, data):
        if self._textarea and self.inside and not self.tpl:
            self.fields[self._textarea] += data


@pytest.fixture
def opened():
    return []


@pytest.fixture
def client(tmp_path, opened):
    return TestClient(create_app(tmp_path / "данные", opener=opened.append, docs_dir=tmp_path / "документы",
                                 board_host="127.0.0.1"))  # табло в тестах — не в сеть (без окна брандмауэра)


def base(folder) -> str:
    return "/c/" + quote(folder.id, safe="")


def test_home_and_health(client):
    r = client.get("/")
    assert r.status_code == 200 and "Пока нет ни одного соревнования" in r.text
    assert client.get("/health").json()["app"] == "st-secretary"
    assert client.get("/static/style.css").status_code == 200
    assert 'action="/shutdown"' not in r.text  # запущено не из окна программы — кнопки «Выключить» нет
    assert client.post("/shutdown").status_code == 409


def test_shutdown_button_stops_server_after_answering(tmp_path):
    stopped = []
    client = TestClient(create_app(tmp_path / "данные", opener=lambda p: None, shutdown=lambda: stopped.append(1)))
    assert 'action="/shutdown"' in client.get("/").text
    r = client.post("/shutdown")
    assert r.status_code == 200 and "СТ-Секретарь выключен" in r.text and 'action="/shutdown"' not in r.text
    assert stopped == [1]  # сервер останавливается после того, как страница отдана


def test_new_competition_then_fill_card(client, tmp_path, psr_card):
    r = client.post("/new", data=MAIN, follow_redirects=False)
    assert r.status_code == 303
    card_url = r.headers["location"].split("?")[0]
    page = client.get(r.headers["location"])
    assert "Соревнование создано" in page.text and "Главный секретарь" in page.text
    folder = tmp_path / "данные" / "2025-09-20 Чемпионат города N по спортивному туризму"
    assert (folder / "Карточка_соревнования.xlsx").is_file() and (folder / "Предзаявки").is_dir()

    data = FormFields(page.text, "cardform").fields
    roles = {v: k.split("-")[1] for k, v in data.items() if k.endswith("-role")}
    js, gs = roles["Главный судья"], roles["Главный секретарь"]
    data |= {f"g-{js}-fio": "Судьин Иван Петрович", f"g-{js}-category": "СС1К", f"g-{js}-territory": "г. Красноярск",
             f"g-{gs}-fio": "Секретарёва Анна Ивановна", f"g-{gs}-category": "СС2К", f"g-{gs}-territory": "г. Красноярск",
             "z-0-group": "М/Ж", "z-0-distance_class": "3", "z-0-discipline_code": "0840161811Я", "z-0-age_from": "22",
             "z-0-age_from_by_gsk": "16", "z-0-team_size": "3", "z-0-min_men": "1", "z-0-min_women": "1",
             "z-0-fee": "3000"}
    assert card_url.endswith("/card/edit")  # новое соревнование — сразу в форму карточки
    r = client.post(card_url, data=data, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].split("?")[0].endswith("/card")
    assert load_card(folder / "Карточка_соревнования.xlsx") == psr_card
    view = client.get(r.headers["location"]).text  # после сохранения — карточка для просмотра
    assert "Карточка заполнена, замечаний нет" in view and "Карточка сохранена" in view
    assert 'id="cardform"' not in view and "Редактировать карточку" in view
    assert "И.П. Судьин, СС1К, г. Красноярск" in view and "М/Ж_3" in view and "с 16 лет — по решению ГСК" in view


def test_card_page_roundtrip_loses_nothing(client, psr_card):
    """Открыли карточку и нажали «Сохранить», ничего не меняя, — файл тот же по содержанию."""
    psr_card.officials.append(replace(psr_card.officials[0], role="Судья-информатор", fio="Инфо Олег Олегович",
                                      category="судья без категории"))  # категория не из списка — не теряется
    f = client.app.state.store.create(psr_card)
    url = base(f) + "/card/edit"
    page = client.get(url).text
    assert "(не из списка)" in page
    r = client.post(url, data=FormFields(page, "cardform").fields, follow_redirects=False)
    assert r.status_code == 303
    assert load_card(f.card_path) == psr_card


def test_errors_shown_next_to_fields_and_input_kept(client):
    r = client.post("/new", data=MAIN | {"title": "", "date_to": "2025-09-19"})
    assert r.status_code == 422
    assert "обязательное поле" in r.text and "дата окончания раньше даты начала" in r.text
    assert 'value="окрестности г. N"' in r.text  # введённое не пропало
    assert client.app.state.store.all() == []


def test_card_changed_in_excel_is_not_overwritten(client, psr_card):
    f = client.app.state.store.create(psr_card)
    url = base(f) + "/card/edit"
    data = FormFields(client.get(url).text, "cardform").fields | {"place": "моё место"}
    write_card(f.card_path, replace(psr_card, place="место из Excel"))  # кто-то поправил в Excel
    os.utime(f.card_path, ns=(10**18, 10**18))
    r = client.post(url, data=data)
    assert r.status_code == 409 and "изменили в Excel" in r.text
    assert load_card(f.card_path).place == "место из Excel"
    r = client.post(url, data=data | {"force": "1"}, follow_redirects=False)
    assert r.status_code == 303 and load_card(f.card_path).place == "моё место"


def test_zachet_field_errors(client, psr_card):
    f = client.app.state.store.create(psr_card)
    url = base(f) + "/card/edit"
    data = FormFields(client.get(url).text, "cardform").fields | {"z-0-distance_class": "", "z-0-fee": "три тысячи"}
    r = client.post(url, data=data)
    assert r.status_code == 422 and "нужно целое число" in r.text
    assert load_card(f.card_path) == psr_card


def lesoviki(tmp_path):
    return make_application(tmp_path / "Лесовики.xlsx", "Лесовики", "Красноярск", "Иванов Пётр Сергеевич",
                            "89131234567, les@mail.ru", 3, [
                                ["Лесовики", "Красноярск", "Иванов Пётр Сергеевич", "Иванов Пётр Сергеевич",
                                 "17.10.1989", "II", "м", "М/Ж", 3],
                                ["Лесовики", "Красноярск", "Иванов Пётр Сергеевич", "Смирнова Анна Олеговна",
                                 "12.03.2001", "КМС", "ж", "М/Ж", 3],
                                ["Лесовики", "Красноярск", "Иванов Пётр Сергеевич", "Кузьмин Олег Игоревич",
                                 "29.02.1995", "б/р", "м", "М/Ж", 3],
                            ]).read_bytes()


def test_preapps_upload_check_summary_remove(client, tmp_path, psr_card, opened):
    f = client.app.state.store.create(psr_card)
    b = base(f)
    data = lesoviki(tmp_path)
    r = client.post(b + "/preapps/upload", files=[("files", ("Лесовики.xlsx", data, XLSX)),
                                                  ("files", ("заметки.txt", b"x", "text/plain"))])
    assert "добавлено заявок: 1" in r.text and "пропущено файлов не Excel: 1" in r.text
    assert "такой даты не существует" in r.text  # 29.02.1995 — ошибка видна
    assert "+7 913 123-45-67" in r.text  # телефон представителя — рядом с замечаниями
    people = r.text.split('class="team-people"')[1].split("</div>")[0]  # в списке команд — все участники
    assert all(n in people for n in ("Иванов Пётр Сергеевич", "Смирнова Анна Олеговна", "Кузьмин Олег Игоревич"))
    assert 'sev-error"' in people  # у кого ошибка — выделен
    assert re.search(r'<b>1</b><span>ошибок', r.text)

    r = client.post(b + "/preapps/summary", follow_redirects=False)
    assert r.status_code == 303 and opened == [f.summary_path] and f.summary_path.is_file()
    r = client.get(b + "/preapps/summary.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"

    r = client.post(b + "/preapps/upload", files=[("files", ("Лесовики.xlsx", data, XLSX))])
    assert "заменено исправленными: 1" in r.text
    assert len(list((f.preapp_dir / "Прежние версии").iterdir())) == 1  # прежний вариант не затёрт

    r = client.post(b + "/preapps/remove", data={"name": "Лесовики.xlsx"})
    assert "убрана из обработки" in r.text
    assert not (f.preapp_dir / "Лесовики.xlsx").exists()
    assert (f.preapp_dir / "Убранные" / "Лесовики.xlsx").is_file()  # не удалена, а перенесена


def test_issue_explains_why_and_links_to_form(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Лесовики.xlsx", lesoviki(tmp_path))
    r = client.get(base(f) + "/preapps")
    assert "Почему:" in r.text and "1995 год не високосный" in r.text and "Что сделать:" in r.text
    assert "строка 12 в файле" in r.text
    assert "/preapps/edit?file=%D0%9B%D0%B5%D1%81%D0%BE%D0%B2%D0%B8%D0%BA%D0%B8.xlsx#p-12" in r.text
    assert 'class="team-link" href="' + base(f) + "/preapps/team?file=" in r.text  # название команды — ссылка


def preapp_url(f, name="Лесовики.xlsx", page="edit"):
    return base(f) + f"/preapps/{page}?file=" + quote(name)


def test_edit_application_in_form(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Лесовики.xlsx", lesoviki(tmp_path))
    page = client.get(preapp_url(f))
    assert page.status_code == 200 and "Редактирование заявки «Лесовики»" in page.text
    data = FormFields(page.text, "preappform").fields
    assert data["h-team"] == "Лесовики" and data["p-2-birth"] == "29.02.1995" and data["p-2-zachet"] == "М/Ж_3"
    assert data["p-0-birth"] == "17.10.1989" and data["p-0-qual"] == "II" and data["p-1-sex"] == "ж"
    assert 'id="p-12"' in page.text and "mk-error" in page.text  # строка с ошибкой отмечена

    r = client.post(base(f) + "/preapps/save", data=data | {"p-2-birth": "28.02.1995"}, follow_redirects=False)
    assert r.status_code == 303 and "/preapps/team?" in r.headers["location"] and "done=psaved" in r.headers["location"]
    page = client.get(r.headers["location"])  # после сохранения — карточка команды
    assert "Карточка команды" in page.text and "Прежние версии" in page.text
    assert "Программа замечаний не нашла" in page.text and "проверять нечего" not in page.text
    # замечаний не осталось — «Проверено» ставится сразу, и видно, что поставлен при сохранении
    assert "сразу отмечена «Проверено»" in page.text and "Поставлен сам" in page.text
    review = client.app.state.store.review(f, psr_card)[1]["Лесовики.xlsx"]
    assert review.status == "done" and review.by == "save"
    [old] = (f.preapp_dir / "Прежние версии").iterdir()  # файл команды не пропал
    assert old.name.startswith("Лесовики (") and [p.name for p in f.preapp_files()] == ["Лесовики.xlsx"]
    result = client.app.state.store.preapps(f, psr_card)
    assert result.count("error") == 0 and len(result.entries) == 3
    assert [e.name.full for e in result.entries] == ["Иванов Пётр Сергеевич", "Смирнова Анна Олеговна",
                                                    "Кузьмин Олег Игоревич"]

    # удалить участника, добавить нового
    data = FormFields(client.get(preapp_url(f)).text, "preappform").fields
    data = {k: v for k, v in data.items() if not k.startswith("p-1-")}
    data |= {"p-7-fio": "Орлова Мария Ивановна", "p-7-birth": "01.02.2000", "p-7-qual": "III", "p-7-sex": "ж",
             "p-7-zachet": "М/Ж_3", "p-7-team_dist": "1"}
    client.post(base(f) + "/preapps/save", data=data)
    names = [e.name.full for e in client.app.state.store.preapps(f, psr_card).entries]
    assert names == ["Иванов Пётр Сергеевич", "Кузьмин Олег Игоревич", "Орлова Мария Ивановна"]


def test_unchanged_form_save_keeps_application(client, tmp_path, psr_card):
    """Открыли заявку и сохранили, ничего не меняя, — участники, даты, разряды, зачёты те же."""
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Лесовики.xlsx", lesoviki(tmp_path))
    store = client.app.state.store

    def snapshot():
        r = store.preapps(f, psr_card)
        return [(e.name.full, e.birth, e.qual, e.sex, e.zachet, e.team_dist) for e in r.entries], \
            [(t.team, t.territory, t.representative, t.phone, t.email) for t in r.teams], r.count("error")

    before = snapshot()
    data = FormFields(client.get(preapp_url(f)).text, "preappform").fields
    assert client.post(base(f) + "/preapps/save", data=data, follow_redirects=False).status_code == 303
    assert snapshot() == before


def test_new_application_and_form_errors(client, psr_card):
    f = client.app.state.store.create(psr_card)
    page = client.get(base(f) + "/preapps/new")
    assert page.status_code == 200 and "Новая заявка" in page.text
    data = FormFields(page.text, "preappform").fields
    r = client.post(base(f) + "/preapps/save", data=data)
    assert r.status_code == 422 and "Впишите название команды" in r.text
    data |= {"h-team": "Ураган", "h-territory": "Красноярск", "p-0-birth": "01.01.1990"}
    r = client.post(base(f) + "/preapps/save", data=data)
    assert r.status_code == 422 and "Впишите ФИО или удалите строку" in r.text and 'value="Ураган"' in r.text
    data |= {"p-0-fio": "Петров Иван Ильич", "p-0-sex": "м", "p-0-zachet": "М/Ж_3"}
    r = client.post(base(f) + "/preapps/save", data=data, follow_redirects=False)
    assert r.status_code == 303 and "done=pcreated" in r.headers["location"]
    assert [p.name for p in f.preapp_files()] == ["Ураган.xlsx"]
    page = client.get(r.headers["location"])
    assert "Заявка сохранена в файл «Ураган.xlsx»" in page.text and "Карточка команды" in page.text


def test_application_changed_meanwhile_is_not_overwritten(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Лесовики.xlsx", lesoviki(tmp_path))
    data = FormFields(client.get(preapp_url(f)).text, "preappform").fields | {"h-team": "Лесовики-2"}
    os.utime(f.preapp_dir / "Лесовики.xlsx", ns=(10**18, 10**18))  # команда прислала новую версию
    r = client.post(base(f) + "/preapps/save", data=data)
    assert r.status_code == 409 and "Файл заявки изменился" in r.text
    r = client.post(base(f) + "/preapps/save", data=data | {"force": "1"}, follow_redirects=False)
    assert r.status_code == 303
    assert client.app.state.store.preapps(f, psr_card).teams[0].team == "Лесовики-2"


def sosna(tmp_path, contacts="89137654321"):
    """Заявка без ошибок, но с «проверить»: участнику 16 лет — допуск только по решению ГСК."""
    return make_application(tmp_path / "Сосна.xlsx", "Сосна", "Красноярск", "Орлов Павел Ильич",
                            contacts, 3, [
                                ["Сосна", "Красноярск", "Орлов Павел Ильич", "Орлов Павел Ильич",
                                 "11.05.1985", "II", "м", "М/Ж", 3],
                                ["Сосна", "Красноярск", "Орлов Павел Ильич", "Белова Ирина Петровна",
                                 "03.07.1999", "КМС", "ж", "М/Ж", 3],
                                ["Сосна", "Красноярск", "Орлов Павел Ильич", "Пестов Юрий Андреевич",
                                 "01.01.2009", "б/р", "м", "М/Ж", 3],
                            ]).read_bytes()


def test_back_to_list_lands_on_the_team_and_next_team_link(client, tmp_path, psr_card):
    from st_secretary.web.app import team_anchor

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Лесовики.xlsx", lesoviki(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    anchor = team_anchor("Лесовики.xlsx")
    assert f'id="{anchor}"' in client.get(base(f) + "/preapps").text  # у команды в списке — постоянная метка
    card = client.get(preapp_url(f, page="team")).text
    assert f'href="{base(f)}/preapps#{anchor}"' in card  # «← Все заявки» ведёт прямо к этой команде
    assert "команда 1 из 2" in card and "Следующая команда: Сосна" in card
    assert preapp_url(f, "Сосна.xlsx", "team").replace(" ", "%20") in card
    last = client.get(preapp_url(f, "Сосна.xlsx", "team")).text
    assert "Это последняя команда в списке" in last and "Предыдущая:" in last
    # «Проверено» в списке возвращает к той же команде
    result, _ = client.app.state.store.review(f, psr_card)
    [warn] = [i for i in result.issues if i.source == "Сосна.xlsx" and i.severity == "warning"]
    r = client.post(base(f) + "/preapps/check", follow_redirects=False,
                    data={"file": "Сосна.xlsx", "key": issue_key(warn), "back": f"{base(f)}/preapps#{team_anchor('Сосна.xlsx')}"})
    assert r.headers["location"].endswith("#" + team_anchor("Сосна.xlsx"))


def test_team_card_is_read_only_with_edit_button(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Лесовики.xlsx", lesoviki(tmp_path))
    page = client.get(preapp_url(f, page="team"))
    assert page.status_code == 200 and "Карточка команды" in page.text and "Лесовики" in page.text
    assert 'name="p-0-fio"' not in page.text and "Кузьмин Олег Игоревич" in page.text  # только просмотр
    assert "Редактировать заявку" in page.text and preapp_url(f).replace(" ", "%20") in page.text
    assert "status-opt-fix\" aria-pressed=\"true\"" in page.text  # есть ошибки → «Исправить» сам
    assert "1995 год не високосный" in page.text
    people = client.get(base(f) + "/preapps").text.split('id="panel-people"')[1].split("</section>")[0]
    # ФИО в списке участников — без ссылок; ссылки — только значки «Ошибка» / «Проверить» (п. 22 правок)
    assert "Кузьмин Олег Игоревич" in people
    assert people.count("<a ") == people.count('<a class="badge ')


def test_preapps_page_links_open(client, tmp_path, psr_card):
    """Все ссылки страницы предзаявок (обе вкладки) открываются: значки «Ошибка» / «Проверить» на вкладке
    «Участники» ведут в карточку своей команды (п. 49 правок)."""
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Лесовики.xlsx", lesoviki(tmp_path))  # ошибка
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))  # «проверить»
    page = client.get(base(f) + "/preapps").text
    people = page.split('id="panel-people"')[1].split("</section>")[0]
    badges = [unescape(h) for h in re.findall(r'<a class="badge [^"]*" href="([^"]+)"', people)]
    assert len(badges) == 2
    assert any("file=" + quote("Лесовики.xlsx") + "&only=error#issues" in h for h in badges)
    assert any("file=" + quote("Сосна.xlsx") + "&only=warning#issues" in h for h in badges)
    links = {unescape(h).split("#")[0] for h in re.findall(r'href="([^"]+)"', page)}
    links = {h for h in links if h.startswith("/") and not h.startswith("/static/")}
    assert len(links) > 10
    for h in sorted(links):
        r = client.get(h)
        assert r.status_code == 200, h


def test_check_marks_and_statuses(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    f.add_preapp("Лесовики.xlsx", lesoviki(tmp_path))
    store = client.app.state.store
    card = preapp_url(f, "Сосна.xlsx", "team")

    result, reviews = store.review(f, psr_card)
    [warn] = [i for i in result.issues if i.source == "Сосна.xlsx" and i.severity == "warning"]
    assert reviews["Сосна.xlsx"].status == "check" and reviews["Лесовики.xlsx"].status == "fix"
    lst = client.get(base(f) + "/preapps").text
    assert re.search(r'data-team-filter="fix"[^>]*>Ошибки <b>1</b>', lst)
    assert re.search(r'data-team-filter="check"[^>]*>Проверить <b>1</b>', lst)

    # «Проверено» у замечания: оно больше не в «Проверить»
    r = client.post(base(f) + "/preapps/check", data={"file": "Сосна.xlsx", "key": issue_key(warn), "on": "1",
                                                     "back": base(f) + "/preapps#t-2"}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].endswith("done=checked#t-2")
    result, _ = store.review(f, psr_card)
    assert result.count("warning") == 0 and result.count("checked") == 1
    assert "Проверено замечаний: 1" in client.get(base(f) + "/preapps").text
    client.post(base(f) + "/preapps/check", data={"file": "Сосна.xlsx", "key": issue_key(warn), "on": "0"})
    assert store.review(f, psr_card)[0].count("warning") == 1

    # статус «Проверено» отмечает и оставшиеся «проверить»; с ошибками — нельзя
    r = client.post(base(f) + "/preapps/status", data={"file": "Сосна.xlsx", "status": "done"}, follow_redirects=False)
    assert r.status_code == 303 and "/preapps/team?" in r.headers["location"]
    result, reviews = store.review(f, psr_card)
    assert reviews["Сосна.xlsx"].status == "done" and reviews["Сосна.xlsx"].manual and result.count("warning") == 0
    assert "Выставлен вручную" in client.get(card).text
    assert client.post(base(f) + "/preapps/status", data={"file": "Лесовики.xlsx", "status": "done"}).status_code == 409
    client.post(base(f) + "/preapps/status", data={"file": "Лесовики.xlsx", "status": "check"})
    assert store.review(f, psr_card)[1]["Лесовики.xlsx"].status == "check"  # вручную можно и «Проверить»

    # команда прислала новый файл — «Проверено» снимается, а отметка у того же замечания остаётся
    f.add_preapp("Сосна.xlsx", sosna(tmp_path, contacts="89130000000, sosna@mail.ru"))
    result, reviews = store.review(f, psr_card)
    assert reviews["Сосна.xlsx"].status == "check" and not reviews["Сосна.xlsx"].manual
    assert "снят: заявку с тех пор изменили" in reviews["Сосна.xlsx"].reset
    assert result.count("checked") == 1  # то же замечание — отметка «проверено» сохранилась

    # сохранили в форме: неотмеченных «проверить» нет (отмеченное не мешает) — «Проверено» сразу
    data = FormFields(client.get(preapp_url(f, "Сосна.xlsx")).text, "preappform").fields
    r = client.post(base(f) + "/preapps/save", data=data, follow_redirects=False)
    assert "done=psaved-done" in r.headers["location"]
    review = store.review(f, psr_card)[1]["Сосна.xlsx"]
    assert review.status == "done" and review.by == "save"
    client.post(base(f) + "/preapps/check", data={"file": "Сосна.xlsx", "key": issue_key(warn), "on": "0"})
    review = store.review(f, psr_card)[1]["Сосна.xlsx"]  # появилось неотмеченное «проверить» — статус снят
    assert review.status == "check" and "появились замечания «Проверить»" in review.reset
    client.post(base(f) + "/preapps/save", data=data)  # сохранили с неотмеченным «проверить» — сам не ставится
    assert store.review(f, psr_card)[1]["Сосна.xlsx"].status == "check"

    # статус — в сводке Excel
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get(base(f) + "/preapps/summary.xlsx").content))
    teams = {row[1]: row[12] for row in wb["Команды"].iter_rows(min_row=2, values_only=True) if row[1]}
    assert teams["Лесовики"] == "Проверить" and teams["Сосна"] == "Проверить"


def test_preapps_wait_for_card_without_errors(client, tmp_path, psr_card):
    psr_card.zachety = []
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Лесовики.xlsx", lesoviki(tmp_path))
    r = client.get(base(f) + "/preapps")
    assert "Сначала исправьте карточку" in r.text and "Лесовики.xlsx" in r.text
    assert client.post(base(f) + "/preapps/summary").status_code == 409
    assert client.get(preapp_url(f)).status_code == 409


def test_overview_steps_and_errors(client, psr_card, opened):
    f = client.app.state.store.create(psr_card)
    r = client.get(base(f))
    assert "Дальше по порядку" in r.text and "Комиссия по допуску" in r.text and "И.П. Судьин, СС1К" in r.text
    assert f'href="{base(f)}/start"' in r.text and f'href="{base(f)}/admission"' in r.text  # шаг — своя страница
    assert client.get("/c/нет такого").status_code == 404
    assert "Страница не найдена" in client.get("/c/..%2F..%2Fsecret").text
    client.post(base(f) + "/open/folder")
    assert opened == [f.path]


def test_import_card_from_excel(client, tmp_path, psr_card):
    card = write_card(tmp_path / "card.xlsx", psr_card)
    r = client.post("/import", files={"card": ("Карточка.xlsx", card.read_bytes(), XLSX)})
    assert "Карточка загружена из Excel" in r.text and psr_card.title in r.text
    [f] = client.app.state.store.all()
    assert load_card(f.card_path) == psr_card
    r = client.post("/import", files={"card": ("заявка.xlsx", b"not an excel file", XLSX)})
    assert r.status_code == 422 and "не похож на карточку" in r.text


# ------------------------------------------------------------------ комиссия по допуску


def kedr(tmp_path):
    """Заявка без замечаний: три взрослых участника."""
    return make_application(tmp_path / "Кедр.xlsx", "Кедр", "Красноярск", "Лебедев Антон Игоревич",
                            "89135550000", 3, [
                                ["Кедр", "Красноярск", "Лебедев Антон Игоревич", "Лебедев Антон Игоревич",
                                 "02.02.1990", "I", "м", "М/Ж", 3],
                                ["Кедр", "Красноярск", "Лебедев Антон Игоревич", "Зуева Мария Олеговна",
                                 "05.06.1996", "II", "ж", "М/Ж", 3],
                                ["Кедр", "Красноярск", "Лебедев Антон Игоревич", "Носов Глеб Андреевич",
                                 "09.09.1993", "б/р", "м", "М/Ж", 3],
                            ]).read_bytes()


def team_form(page_html, file):
    """Поля формы команды на странице комиссии (как их отправит браузер)."""
    from st_secretary.web.app import team_anchor

    start = page_html.index(f'id="{team_anchor(file)}"')
    block = page_html[start:page_html.index("</article>", start)]
    return FormFields('<form id="t">' + block.split(">", 1)[1], "t").fields


def test_admission_documents_decisions_fees_numbers(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    url = base(f) + "/admission"
    page = client.get(url).text
    assert "Комиссия по допуску" in page and "Зуева Мария Олеговна" in page and "Паспорт" in page
    assert "0 из 2" in page and "Ожидают <b>2</b>" in page  # пока ничего не отмечено

    # «Отметить все документы» — со страницы приходит без перезагрузки (X-Autosave), ответ — блок команды и итоги
    data = team_form(page, "Кедр.xlsx") | {"do": "all_docs", "fee_paid": "3000", "fee_method": "наличные"}
    j = client.post(url + "/team", data=data, headers={"X-Autosave": "1"}).json()
    assert j["status"] == "pending" and "заявку секретарь не проверил" in j["team"]  # п. 24 правок: сама — нет
    for name in ("Кедр.xlsx", "Сосна.xlsx"):
        client.post(base(f) + "/preapps/status", data={"file": name, "status": "done"})
    r = client.post(url + "/team", data=data, headers={"X-Autosave": "1"})
    j = r.json()
    assert j["status"] == "admitted" and "Допущена" in j["team"] and "оплачено" in j["team"]
    assert "1 из 2" in j["tiles"] and "3 000 ₽" in j["tiles"]

    # Сосна: Пестову 16 лет — ждёт решения комиссии; допустить решением комиссии с основанием
    data = team_form(client.get(url).text, "Сосна.xlsx") | {"do": "all_docs"}
    j = client.post(url + "/team", data=data, headers={"X-Autosave": "1"}).json()
    assert j["status"] == "pending" and "нужно решение комиссии" in j["team"]
    data = team_form(client.get(url).text, "Сосна.xlsx")
    pestov = next(k.split("-")[1] for k, v in data.items() if k.endswith("-key") and v == "пестов юрий андреевич")
    data |= {f"p-{pestov}-decision": "admitted", f"p-{pestov}-reason": "решение ГСК"}
    r = client.post(url + "/team", data=data, follow_redirects=False)  # без скриптов — обычная форма
    assert r.status_code == 303 and "#t-" in r.headers["location"]
    page = client.get(url).text
    assert "2 из 2" in page and "допущен решением комиссии: решение ГСК" in page

    # номера командам по порядку списка
    client.post(url + "/numbers", data={"mode": "missing"})
    adm = f.admission()["teams"]
    assert sorted(t["number"] for t in adm.values()) == [1, 2]
    data = team_form(client.get(url).text, "Сосна.xlsx") | {"number": str(adm["Кедр.xlsx"]["number"])}
    j = client.post(url + "/team", data=data, headers={"X-Autosave": "1"}).json()
    assert "уже у команды «Кедр»" in j["team"]  # одинаковый номер — предупреждение

    # протокол и ведомость
    r = client.get(url + "/report.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Протокол комиссии", "Ведомость взносов", "Документы участников"]
    text = [str(c.value) for row in wb["Протокол комиссии"].iter_rows() for c in row if c.value is not None]
    assert text.count("допущена") == 2  # обе команды — в графе «Решения по замечаниям»
    r = client.post(url + "/report", follow_redirects=False)
    assert r.status_code == 303 and f.commission_report_path.is_file()


def test_admission_settings_and_reentry(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    url = base(f) + "/admission"
    client.post(url + "/settings", data={"doc-id": "1", "doc-med": "1", "doc-pd": "1", "tdoc-app": "1",
                                         "start_at": "2000-01-01T10:00"})  # старт давно прошёл — перезаявка поздняя
    page = client.get(url).text
    assert "Согласие" in page and "Книжка" not in page.split("<thead>")[1].split("</thead>")[0]

    edit = client.get(base(f) + "/preapps/edit?file=" + quote("Кедр.xlsx") + "&reentry=1").text
    assert "Перезаявка «Кедр»" in edit and "меньше часа" in edit
    data = FormFields(edit, "preappform").fields
    assert data["reentry"] == "1"
    fio_key = next(k for k, v in data.items() if v == "Носов Глеб Андреевич")
    data[fio_key] = "Носова Галина Андреевна"
    data[fio_key.replace("fio", "sex")] = "ж"
    r = client.post(base(f) + "/preapps/save", data=data, follow_redirects=False)
    assert r.status_code == 303 and "/admission?done=reentry_late" in r.headers["location"]
    [rec] = f.admission()["teams"]["Кедр.xlsx"]["reentries"]
    assert "выбыл(а): Носов Глеб Андреевич" in rec["text"] and "включён(а): Носова Галина Андреевна" in rec["text"]
    assert rec["added"] == ["носова галина андреевна"]  # в заявке с печатью врача её не было (п. 23 правок)
    assert rec["late"] and "позже, чем за час до старта" in client.get(url).text
    edit = client.get(base(f) + "/preapps/edit?file=" + quote("Кедр.xlsx") + "&reentry=1").text
    assert "повторные не принимаются" in edit


# ------------------------------------------------------------------ проверка снаряжения и документы команд

PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde"
       b"\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0\x00\x00\x03\x01\x01\x00\xc9\xfe\x92\xef\x00\x00\x00\x00IEND\xaeB`\x82")


def test_equipment_page_check_and_report(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    url = base(f) + "/equipment"
    assert "Перечень снаряжения не задан" in client.get(url).text
    client.post(url + "/settings", data={"do": "template"})
    page = client.get(url).text
    assert "Не проверено" in page and "Начать проверку: всё есть" in page
    assert "снаряжение не проверено" in client.get(base(f) + "/admission").text  # без проверки не допустить

    j = client.post(url + "/team", data={"file": "Кедр.xlsx", "do": "all"}, headers={"X-Autosave": "1"}).json()
    assert "0 б. — допущена" in j["team"] and j["by"]["ok"] == 1
    data = team_form(client.get(url).text, "Кедр.xlsx")
    lost = next(k for k, v in data.items() if k.startswith("p-0-") and v == "on")  # первая галочка первого участника
    data = {k: v for k, v in data.items() if k != lost}  # у участника одного предмета нет
    j = client.post(url + "/team", data=data, headers={"X-Autosave": "1"}).json()
    assert "1 б. — допущена" in j["team"] and "Нет:" in j["team"]

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get(base(f) + "/admission/report.xlsx").content))
    assert "Проверка снаряжения" in wb.sheetnames
    text = [str(c.value) for row in wb["Проверка снаряжения"].iter_rows() for c in row if c.value is not None]
    assert "Акт проверки снаряжения" in text and "допущена" in text


def test_team_documents_stored_locally_and_checked_in_one_window(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    card = base(f) + "/preapps/team?file=" + quote("Кедр.xlsx")
    assert "Документы команды (0)" in client.get(card).text
    r = client.post(base(f) + "/docs/upload", data={"file": "Кедр.xlsx", "back": card},
                    files=[("files", ("паспорт Лебедев.png", PNG, "image/png")),
                           ("files", ("полис.pdf", b"%PDF-1.4 test", "application/pdf")),
                           ("files", ("вирус.exe", b"MZ", "application/octet-stream"))], follow_redirects=False)
    assert r.status_code == 303 and "done=docs_skipped" in r.headers["location"]
    docs_dir = tmp_path / "документы" / f.id / "Кедр"
    assert sorted(p.name for p in docs_dir.iterdir()) == ["паспорт Лебедев.png", "полис.pdf"]
    assert not any((f.path).rglob("*.png"))  # в папке соревнования (её копируют, держат в облаке) сканов нет

    page = client.get(card).text
    assert "Документы команды (2)" in page and "Проверить в одном окне" in page
    view = client.get(base(f) + "/docs/view?" + urlencode({"file": "Кедр.xlsx", "name": "паспорт Лебедев.png"}))
    assert view.content == PNG and view.headers["cache-control"] == "no-store"
    assert "inline" in view.headers["content-disposition"]
    assert client.get(base(f) + "/docs/view?" + urlencode({"file": "Кедр.xlsx", "name": "../../x.png"})).status_code == 404

    check = client.get(base(f) + "/admission/check?file=" + quote("Кедр.xlsx")).text
    assert 'data-doc-kind="image"' in check and 'data-doc-kind="pdf"' in check
    assert "Отметить все документы" in check and "Лебедев Антон Игоревич" in check
    assert "проверить в одном окне" not in check  # уже в нём
    data = team_form(check, "Кедр.xlsx") | {"do": "all_docs"}
    client.post(base(f) + "/preapps/status", data={"file": "Кедр.xlsx", "status": "done"})
    j = client.post(base(f) + "/admission/team", data=data, headers={"X-Autosave": "1"}).json()
    assert j["status"] == "admitted" and "проверить в одном окне" not in j["team"]

    client.post(base(f) + "/docs/remove", data={"file": "Кедр.xlsx", "name": "полис.pdf"})
    assert (docs_dir / "Убранные" / "полис.pdf").is_file() and not (docs_dir / "полис.pdf").exists()


# ------------------------------------------------------------------ награждение и документы по итогам


def test_awards_manual_results_and_documents(client, tmp_path, psr_card, opened):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    url = base(f) + "/awards"
    page = client.get(url).text
    assert "Результатов пока нет" in page and "Ввести или поправить места вручную" in page
    r = client.post(url + "/import", data={"zachet": "М/Ж_3"},
                    files={"protocol": ("протокол.xls", b"not an excel file", "application/vnd.ms-excel")},
                    follow_redirects=False)
    assert "done=res_bad" in r.headers["location"]  # не протокол — объяснение, а не ошибка

    data = {"zachet": "М/Ж_3", "r-0-team": "Кедр", "r-0-place": "1", "r-0-result": "-283", "r-0-norm": "II",
            "r-1-team": "Сосна", "r-1-place": "2", "r-1-result": "-213", "r-1-norm": ""}
    r = client.post(url + "/manual", data=data, follow_redirects=False)
    assert "done=res_saved" in r.headers["location"]
    page = client.get(url).text
    assert "Места введены вручную" in page and "Лебедев Антон Игоревич (I)" in page

    from docx import Document
    d = client.get(url + "/file/diplomas")
    assert d.status_code == 200 and d.content[:2] == b"PK"
    texts = [p.text for p in Document(io.BytesIO(d.content)).paragraphs if p.text]
    assert texts.count("Награждаются") == 6 and "за II место" in texts
    r = client.post(url + "/doc/stickers", follow_redirects=False)
    assert r.status_code == 303 and opened[-1] == f.out_dir / "Наклейки на медали.xlsx" and opened[-1].is_file()
    assert client.get(url + "/file/nothing").status_code == 404


def test_judge_grades_go_to_certificates(client, psr_card):
    f = client.app.state.store.create(psr_card)
    url = base(f) + "/awards"
    page = client.get(url).text
    assert "Судьин Иван Петрович" in page and "Сохранить оценки" in page
    r = client.post(url + "/grades", data={"g-0": "отлично", "g-1": "хорошо"}, follow_redirects=False)
    assert "done=grades_saved" in r.headers["location"]
    from docx import Document
    texts = [p.text for p in Document(io.BytesIO(client.get(url + "/file/judging").content)).paragraphs]
    assert "Оценка судейства: «отлично»." in texts and "Оценка судейства: «хорошо»." in texts
    sk = client.get(url + "/file/sk")
    assert sk.status_code == 200 and sk.content[:2] == b"PK"


def test_report_texts_and_extracts_download(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    url = base(f) + "/awards"
    client.post(url + "/manual", data={"zachet": "М/Ж_3", "r-0-team": "Кедр", "r-0-place": "1", "r-0-norm": "III"})
    r = client.post(url + "/report", data={"protests": "подан один протест, отклонён", "base": "Всё хорошо."},
                    follow_redirects=False)
    assert "done=report_saved" in r.headers["location"]
    assert "подан один протест, отклонён" in client.get(url).text
    from docx import Document
    texts = [p.text for p in Document(io.BytesIO(client.get(url + "/file/report").content)).paragraphs]
    assert "Подан один протест, отклонён." in texts and "Всё хорошо." in texts
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get(url + "/file/extracts").content))
    rows = [[c.value for c in r] for r in wb.active.iter_rows() if r[5].value == "III"]
    assert len(rows) == 3 and rows[0][2] == "02.02.1990"  # дата рождения — из заявки


def test_contracts_tabel_personal_data_and_documents(client, tmp_path, psr_card, opened):
    from docx import Document
    from openpyxl import load_workbook
    from test_contracts import CUSTOMER, PERSON

    f = client.app.state.store.create(psr_card)
    url = base(f) + "/contracts"
    page = client.get(url).text
    assert "Судьин Иван Петрович" in page and "Секретарёва Анна Ивановна" in page and "ставка не задана" not in page
    assert page.count('type="checkbox" name="p-0-d-') == 4  # 19–22 сентября: день до и день после
    r = client.post(url + "/add", data={"fio": "Работяга  Семён Ильич", "role": "Рабочий комендантской бригады",
                                        "category": "б/к"}, follow_redirects=False)
    assert "done=ct_added" in r.headers["location"]
    assert "done=ct_exists" in client.post(url + "/add", data={"fio": "Судьин Иван Петрович", "role": "Судья"},
                                           follow_redirects=False).headers["location"]

    r = client.post(url + "/settings", data={"do": "sample"}, follow_redirects=False)
    assert "done=ct_settings" in r.headers["location"]
    data = f.contracts()
    assert data["rates"] == {"главный судья|1": 850, "рабочий комендантской бригады|б/к": 460}  # чего нет в образце — пусто
    client.post(url + "/settings", data={"from": "2025-09-19", "to": "2025-09-22", "accrual": "30",
                                         "rate-главный судья|1": "850", "rate-главный секретарь|2": "700",
                                         "rate-рабочий комендантской бригады|б/к": "460"})
    days = {"p-0-key": "судьин иван петрович", "p-1-key": "секретарева анна ивановна", "p-2-key": "работяга семен ильич",
            "p-0-d-20250919": "on", "p-0-d-20250920": "on", "p-0-d-20250921": "on", "p-0-d-20250922": "on",
            "p-1-d-20250920": "on", "p-1-d-20250921": "on", "p-2-d-20250921": "on", "p-2-unpaid": "on"}
    r = client.post(url + "/days", data=days, headers={"X-Autosave": "1"})
    j = r.json()
    assert "<b>3 400</b>" in j["team"] and "4 800,00" in j["team"]  # 4×850 + 2×700; рабочий без оплаты
    assert "6 240,00" in j["team"] and "6 240,00 ₽" in j["tiles"]

    person = url + "/person?" + urlencode({"key": "судьин иван петрович"})
    p = client.get(person)
    assert p.headers["cache-control"] == "no-store" and "в качестве главного судьи 1 категории" in p.text
    r = client.post(url + "/person", data={"key": "судьин иван петрович", **dict(PERSON, inn="123456789012")},
                    follow_redirects=False)
    assert "done=ct_person" in r.headers["location"]
    assert "ИНН не сходится" in client.get(person).text
    client.post(url + "/person", data={"key": "судьин иван петрович", **PERSON})
    assert client.app.state.store.personal()["судьин иван петрович"]["inn"] == PERSON["inn"]
    assert client.app.state.store.personal_path.parent == tmp_path / "документы"  # не в папке соревнования
    client.post(url + "/customer", data=CUSTOMER)

    r = client.post(url + "/doc/person", data={"key": "судьин иван петрович"}, follow_redirects=False)
    assert "done=ct_doc" in r.headers["location"]
    doc_path = opened[-1]
    assert doc_path.parent == tmp_path / "документы" / f.id / "Договоры и табель"
    assert doc_path.name == "Судьин И.П. — главный судья.docx"
    text = "\n".join(p.text for p in Document(str(doc_path)).paragraphs)
    assert "в лице директора Начальникова Петра Сергеевича" in text and "Срок оказания услуг: 19–22 сентября" in text
    assert "3400 (три тысячи четыреста) рублей 00 копеек" in text

    ws = load_workbook(io.BytesIO(client.get(url + "/file/tabel").content)).active
    names = [ws.cell(r, 2).value for r in range(8, 11)]
    assert names == ["Судьин Иван Петрович", "Секретарёва Анна Ивановна", "Итого начислено"]
    all_docs = client.get(url + "/file/all")
    assert all_docs.status_code == 200 and all_docs.headers["cache-control"] == "no-store"
    assert "\n".join(p.text for p in Document(io.BytesIO(all_docs.content)).paragraphs).count("ДОГОВОР № ______") == 2

    r = client.post(url + "/template", follow_redirects=False)
    assert "done=ct_template" in r.headers["location"] and f.contract_template.is_file()
    page = client.get(url).text
    assert "Свой шаблон <b>Шаблон договора.docx</b> — для:" in page and "Всем — встроенный" not in page

    # у главного судьи — своя форма договора (у заказчика так и есть: с отчётом главного судьи)
    r = client.post(url + "/template", data={"role": "Главный судья"}, follow_redirects=False)
    own = f.path / "Шаблон договора — главный судья.docx"
    assert own.is_file() and opened[-1] == own
    doc = Document(str(own))
    doc.add_paragraph("Особое условие для {{Должность}}: отчёт в течение 5 дней.")
    doc.save(str(own))
    all_text = "\n".join(p.text for p in Document(io.BytesIO(client.get(url + "/file/all").content)).paragraphs)
    assert all_text.count("Особое условие") == 1 and "Особое условие для Главный судья" in all_text
    page = re.sub(r"\s+", " ", client.get(url).text)
    assert "Свой шаблон <b>Шаблон договора — главный судья.docx</b> — для: главный судья." in page
    client.post(url + "/remove", data={"key": "работяга семен ильич"})
    assert "Работяга" not in client.get(url).text


def test_verify_page_program_documents_and_upload(client, psr_card, opened):
    from docx import Document

    f = client.app.state.store.create(psr_card)
    url = base(f) + "/verify"
    assert "Пока нечего сверять" in client.get(url).text
    client.post(base(f) + "/awards/doc/judging")  # программа сохранила справки и табель
    client.post(base(f) + "/contracts/doc/tabel")
    page = client.get(url).text
    assert r"Документы по итогам\Справки о судействе.docx" in page and r"Договоры и табель\Табель-наряд.xlsx" in page
    assert page.count("расхождений нет") == 2

    bad = Document()
    bad.add_paragraph("Отчёт главного судьи Чемпионата города N по спортивному туризму, «04» октября 2024 г., "
                      "код ВРВС 0840271811Я")
    buf = io.BytesIO()
    bad.save(buf)
    docx_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    r = client.post(url, files=[("files", ("старый отчёт.docx", buf.getvalue(), docx_type)),
                                ("files", ("скан.pdf", b"%PDF-1.4 broken", "application/pdf"))])
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store"
    text = re.sub(r"\s+", " ", r.text)
    assert "старый отчёт.docx" in text and "дата ««04» октября 2024» — не 2025 год" in text
    assert "код ВРВС 0840271811Я — это «дистанция - спелео - группа»" in text
    assert "Не прочитаны" in text and "скан.pdf" in text
    assert "Сейчас добавлены: старый отчёт.docx, скан.pdf" in text
    assert not any(p.name.startswith("старый") for p in f.path.rglob("*"))  # присланные файлы не сохраняются


def test_results_stages_points_and_live_table(client, tmp_path, psr_card):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    url = base(f) + "/results"
    page = client.get(url).text
    assert "Этапы дистанции М/Ж_3" in page and "Баллы по этапам" not in page  # сначала — этапы
    r = client.post(url + "/stages?z=М/Ж_3", data={
        "st-0-tour": "Тур 1", "st-0-name": "Узлы", "st-0-max": "100",
        "st-1-tour": "Тур 1", "st-1-name": "Переправа", "st-1-max": "",
        "st-2-tour": "Бонус", "st-2-name": "Ориентирование", "st-3-name": "",
        "km": "", "modes": "", "kv_hours": "", "tie": "same"}, follow_redirects=False)
    assert "done=run_stages" in r.headers["location"]
    stages = f.run_data()["zachety"]["М/Ж_3"]["stages"]
    assert [(s["id"], s["tour"], s["name"]) for s in stages] == [("s1", "Тур 1", "Узлы"), ("s2", "Тур 1", "Переправа"),
                                                                  ("s3", "Бонус", "Ориентирование")]
    page = client.get(url).text
    assert "Баллы по этапам" in page and 'name="p-0-s1"' in page and page.count("data-team=") == 2
    files = re.findall(r'name="p-(\d)-file" value="([^"]+)"', page)
    data = {f"p-{i}-file": name for i, name in files}
    by = {name: i for i, name in files}
    data.update({f"p-{by['Кедр.xlsx']}-s1": "20", f"p-{by['Кедр.xlsx']}-s2": "-5", f"p-{by['Кедр.xlsx']}-s3": "-100",
                 f"p-{by['Сосна.xlsx']}-s1": "10,5", f"p-{by['Сосна.xlsx']}-s2": "абв",
                 f"p-{by['Кедр.xlsx']}-status": "finished", f"p-{by['Сосна.xlsx']}-status": "finished"})
    j = client.post(url + "/points?z=М/Ж_3", data=data, headers={"X-Autosave": "1"}).json()
    assert j["cells"]["Кедр.xlsx"] == {"total": "-85", "place": "1", "bad": []}
    assert j["cells"]["Сосна.xlsx"] == {"total": "10,5", "place": "2", "bad": ["s2"]}
    assert "«Сосна», Тур 1 · Переправа: «абв» — не число" in j["results"]
    assert "Лебедев Антон Игоревич (I)" in j["results"]

    data[f"p-{by['Сосна.xlsx']}-status"] = "removed"
    j = client.post(url + "/points?z=М/Ж_3", data=data, headers={"X-Autosave": "1"}).json()
    assert j["cells"]["Сосна.xlsx"]["place"] == "снята" and "снята" in j["results"]
    page = client.get(url).text
    assert 'value="абв"' in page and 'class="is-invalid"' in page  # введённое не пропадает, ошибка подсвечена
    assert client.get(url + "?z=нет").status_code == 404


def test_preliminary_protests_official_to_awards(client, tmp_path, psr_card, opened):
    from datetime import datetime

    from openpyxl import load_workbook

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    url = base(f) + "/results"
    q = "?z=М/Ж_3"
    client.post(url + "/stages" + q, data={"st-0-tour": "Тур 1", "st-0-name": "Узлы", "tie": "same"})
    page = client.get(url).text
    by = {name: i for i, name in re.findall(r'name="p-(\d)-file" value="([^"]+)"', page)}
    pts = {f"p-{i}-file": n for n, i in by.items()}
    pts.update({f"p-{by['Кедр.xlsx']}-s1": "20", f"p-{by['Сосна.xlsx']}-s1": "35"})
    client.post(url + "/points" + q, data=pts)

    clock = {"now": datetime(2025, 9, 21, 15, 0)}
    client.app.state.clock = lambda: clock["now"]
    assert "done=run_not_published" in client.post(url + "/approve" + q, follow_redirects=False).headers["location"]
    r = client.post(url + "/publish" + q, follow_redirects=False)
    assert "done=run_published" in r.headers["location"] and "until=16%3A00" in r.headers["location"]
    prelim = opened[-1]
    assert prelim.parent == f.protocols_dir and prelim.name == "Предварительный протокол М-Ж_3 21.09 15-00.xlsx"
    cells = [c for row in load_workbook(prelim)["Протокол"].iter_rows(values_only=True) for c in row if c]
    assert "ПРЕДВАРИТЕЛЬНЫЙ ПРОТОКОЛ РЕЗУЛЬТАТОВ" in cells and "Кедр" in cells and "20" in cells
    assert any(str(c).startswith("Опубликован 21.09.2025 в 15:00. Протесты по результатам принимаются в течение 1 часа — "
                                 "до 16:00") for c in cells)

    clock["now"] = datetime(2025, 9, 21, 15, 30)
    r = client.post(url + "/protest" + q, data={"at": "2025-09-21T15:25", "team": "Сосна", "text": "Узлы: 35 → 25"},
                    follow_redirects=False)
    assert "done=run_protest" in r.headers["location"]
    assert "done=run_open_protests" in client.post(url + "/approve" + q, follow_redirects=False).headers["location"]
    client.post(url + "/protest/decide" + q, data={"id": "p1", "decision": "удовлетворён", "note": "видеозапись"})
    assert f.run_data()["zachety"]["М/Ж_3"]["protests"][0]["decision"] == "удовлетворён"

    pts[f"p-{by['Сосна.xlsx']}-s1"] = "25"  # по протесту
    client.post(url + "/points" + q, data=pts)
    assert "done=run_changed" in client.post(url + "/approve" + q, follow_redirects=False).headers["location"]
    clock["now"] = datetime(2025, 9, 21, 15, 40)
    client.post(url + "/publish" + q)
    page = client.get(url).text
    assert "принимаются до <b>16:40</b>" in page and "Опубликовать заново" in page

    clock["now"] = datetime(2025, 9, 21, 16, 45)
    r = client.post(url + "/approve" + q, follow_redirects=False)
    assert "done=run_official" in r.headers["location"] and opened[-1].name == "Протокол результатов М-Ж_3.xlsx"
    assert client.get(url + "/file/official" + q).status_code == 200
    awards = f.results_data()["zachety"]["М/Ж_3"]
    assert awards["source"] == "СТ-Секретарь, утверждён 21.09.2025 16:45"
    assert [(row["team"], row["place"], row["result"]) for row in awards["rows"]] == [("Кедр", 1, "20"), ("Сосна", 2, "25")]
    assert "Лебедев Антон Игоревич (I)" in client.get(base(f) + "/awards").text

    r = client.post(url + "/protest" + q, data={"at": "2025-09-21T16:50", "team": "Кедр", "text": "поздно"},
                    follow_redirects=False)
    assert "done=run_protest_late" in r.headers["location"]


def test_draw_start_protocol_publish_and_board(client, tmp_path, psr_card, opened):
    from datetime import datetime

    from openpyxl import load_workbook

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    url = base(f) + "/start"
    q = "?z=М/Ж_3"
    page = client.get(url).text
    assert "1. Жеребьёвка — зачёт М/Ж_3" in page and "Лебедев Антон Игоревич (I)" in page
    assert "Провести жеребьёвку" in page and "по стартовым номерам" in client.get(base(f) + "/results").text

    clock = {"now": datetime(2025, 9, 19, 20, 15)}
    client.app.state.clock = lambda: clock["now"]
    assert "done=start_bad_time" in client.post(url + "/draw" + q, data={"first": "25:00", "action": "save"},
                                                follow_redirects=False).headers["location"]
    client.post(url + "/draw" + q, data={"day": "2025-09-20", "first": "10:00", "interval": "5", "action": "save"})
    r = client.post(url + "/draw" + q, data={"method": "random", "groups": "2", "strong": "last", "action": "draw"},
                    follow_redirects=False)
    assert "done=start_drawn" in r.headers["location"]
    dr = f.run_data()["zachety"]["М/Ж_3"]["draw"]
    assert sorted(dr["order"]) == ["Кедр.xlsx", "Сосна.xlsx"] and 100000 <= dr["seed"] <= 999999
    assert dr["first"] == "10:00" and dr["interval"] == "5" and dr["done_method"] == "random"  # время не затёрто
    page = client.get(url).text
    assert "Провести жеребьёвку заново" in page and f"число жребия {dr['seed']}" in page
    assert "как в <a" in client.get(base(f) + "/results").text  # таблица результатов — в порядке старта

    first = dr["order"][0]
    other = dr["order"][1]
    r = client.post(url + "/order" + q, data={"file-0": first, "pos-0": "2", "time-0": "",
                                              "file-1": other, "pos-1": "1", "time-1": "10:30"},
                    follow_redirects=False)
    assert "done=start_saved" in r.headers["location"]
    dr = f.run_data()["zachety"]["М/Ж_3"]["draw"]
    assert dr["order"] == [other, first] and dr["times"] == {other: "10:30"} and dr["edited"] == "2025-09-19T20:15"

    clock["now"] = datetime(2025, 9, 20, 8, 30)
    r = client.post(url + "/publish" + q, follow_redirects=False)
    assert "done=start_published" in r.headers["location"] and "until=09%3A30" in r.headers["location"]
    path = opened[-1]
    assert path.parent == f.protocols_dir and path.name == "Стартовый протокол М-Ж_3.xlsx"
    cells = [c for row in load_workbook(path).active.iter_rows(values_only=True) for c in row if c]
    assert "СТАРТОВЫЙ ПРОТОКОЛ" in cells and "10:30" in cells and "10:05" in cells
    assert any("порядок изменён вручную 19.09.2025 в 20:15" in str(c) for c in cells)
    page = client.get(url).text
    assert "Опубликован 20.09 в 08:30" in page and "Принимаются до <b>09:30</b>" in page
    assert client.get(url + "/file" + q).status_code == 200

    client.post(url + "/order" + q, data={"file-0": other, "pos-0": "2", "file-1": first, "pos-1": "1"})
    assert "опубликуйте заново" in client.get(url).text  # после публикации порядок поменяли

    client.post(base(f) + "/board/toggle", data={"on": "1"})
    board = TestClient(client.app.state.board.app).get("/").text
    assert "Стартовый протокол" in board and "Кедр" in board and "10:30" not in board  # время Кедра теперь 10:00
    assert "после публикации менялся" in board


def test_help_opens_manual_next_to_program(client, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)  # переносная версия: инструкция лежит рядом с «СТ-Секретарь.bat»
    (tmp_path / "Инструкция секретаря.pdf").write_bytes(b"%PDF-1.4 test")
    r = client.get("/help")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content == b"%PDF-1.4 test"
    assert "inline" in r.headers["content-disposition"]
    assert 'href="/help"' in client.get("/").text


def test_backup_download_and_restore(client, tmp_path, psr_card, opened):
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    assert "Копий пока нет" in client.get(base(f)).text
    r = client.post(base(f) + "/backup", follow_redirects=False)
    assert "done=backup_made" in r.headers["location"]
    assert "Последняя копия" in client.get(base(f)).text
    z = client.get(base(f) + "/backup.zip")
    assert z.status_code == 200 and z.headers["content-type"] == "application/zip"
    r = client.post("/restore", files={"backup": ("копия.zip", z.content, "application/zip")}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].endswith("?done=restored")
    restored = client.app.state.store.all()
    assert len(restored) == 2 and all(len(x.preapp_files()) == 1 for x in restored)  # прежняя папка на месте
    assert "Соревнование восстановлено" in client.get(r.headers["location"]).text
    bad = client.post("/restore", files={"backup": ("x.zip", b"nope", "application/zip")}, follow_redirects=False)
    assert "done=restore_bad" in bad.headers["location"]
    assert "Восстановить из резервной копии" in client.get("/").text


def test_training_competition_from_home_page(client):
    assert "Создать учебное соревнование" in client.get("/").text
    r = client.post("/training", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].endswith("?done=training")
    page = client.get(r.headers["location"]).text
    assert "Учебное соревнование создано" in page and "Учебный чемпионат г. Энска" in page
    f = client.app.state.store.all()[0]
    assert len(f.preapp_files()) == 14
    assert "Ориентирование" in client.get(base(f) + "/results").text


def test_individual_discipline_each_athlete_gets_place(client, tmp_path, psr_card, opened):
    """Северная ходьба — личная: в результатах, жеребьёвке и протоколе каждый спортсмен со своим номером."""
    from openpyxl import load_workbook

    card = replace(psr_card, zachety=[replace(psr_card.zachety[0], group="Ж", distance_class=2,
                                              discipline_code="0840291811Л", team_size=None, min_men=0,
                                              min_women=0, age_from=18, age_from_by_gsk=None)])
    f = client.app.state.store.create(card)
    rows = [["Ива", "Красноярск", "Орлова Анна Петровна", fio, "01.02.1970", "б/р", "ж", "Ж", 2, None, 1]
            for fio in ("Орлова Анна Петровна", "Белова Ирина Сергеевна", "Котова Вера Ивановна")]
    f.add_preapp("Ива.xlsx", make_application(tmp_path / "Ива.xlsx", "Ива", "Красноярск", "Орлова Анна Петровна",
                                              "89130000000", 3, rows).read_bytes())
    client.post(base(f) + "/admission/numbers", data={"mode": "missing"})
    url = base(f) + "/results?z=" + quote("Ж_2")
    client.post(base(f) + "/results/stages?z=" + quote("Ж_2"),
                data={"st-0-name": "Точка 1", "kv": "120", "system": "penalty", "removal": "dsq"})
    page = client.get(url).text
    assert page.count("data-team=") == 3 and "Кр. карт." in page and "Штраф. время" in page
    keys = re.findall(r'name="p-(\d)-file" value="([^"]+)"', page)
    assert all("#" in k for _, k in keys)  # ключ — «файл#спортсмен»
    data = {f"p-{i}-file": k for i, k in keys}
    for n, (i, _) in enumerate(keys):
        data.update({f"p-{i}-start": "10:00:00", f"p-{i}-finish": f"10:5{n}:00", f"p-{i}-status": "finished"})
    data[f"p-{keys[2][0]}-red"] = "к"
    client.post(base(f) + "/results/points?z=" + quote("Ж_2"), data=data)
    page = client.get(url).text
    assert "<th>Участник</th>" in page and "1.1" in page and "аннулирован" in page
    client.post(base(f) + "/results/publish?z=" + quote("Ж_2"))
    ws = load_workbook(opened[-1])["Протокол"]
    cells = [c for row in ws.iter_rows(values_only=True) for c in row if c]
    assert "Участник" in cells and "Орлова Анна Петровна" in cells and "Ива" in cells
    start = client.get(base(f) + "/start?z=" + quote("Ж_2")).text
    assert "Белова Ирина Сергеевна" in start and "Ива" in start


def test_board_shows_only_enabled_competitions_and_only_results(client, tmp_path, psr_card):
    from datetime import datetime

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    url = base(f) + "/results"
    client.post(url + "/stages?z=М/Ж_3", data={"st-0-tour": "Тур 1", "st-0-name": "Узлы"})
    client.post(url + "/points?z=М/Ж_3", data={"p-0-file": "Кедр.xlsx", "p-0-s1": "20"})
    board = TestClient(client.app.state.board.app)
    assert "Табло пока не включено ни для одного соревнования" in board.get("/").text
    assert board.get("/c/" + quote(f.id, safe="")).status_code == 404  # табло не включено — результатов не видно
    assert "Кедр" in client.get(base(f) + "/board/preview").text  # секретарю — предпросмотр

    assert "done=board_on" in client.post(base(f) + "/board/toggle", data={"on": "1"},
                                          follow_redirects=False).headers["location"]
    page = board.get("/").text  # одно соревнование — сразу его результаты
    assert "Кедр" in page and "Текущие результаты на" in page and 'http-equiv="refresh" content="30"' in page
    assert "Лебедев" not in page  # составы на табло не нужны — только команды, баллы и места
    for private in ("/c/" + quote(f.id, safe="") + "/preapps", "/c/" + quote(f.id, safe="") + "/contracts", "/shutdown"):
        assert board.get(private).status_code in (404, 405)  # табло — только чтение результатов

    client.app.state.clock = lambda: datetime(2025, 9, 21, 15, 0)
    client.post(url + "/publish?z=М/Ж_3")
    assert "Предварительные результаты — опубликованы в 15:00, протесты принимаются до 16:00" in board.get("/").text
    admin = client.get(base(f) + "/board").text
    assert "Раздать табло по Wi-Fi" in admin and "включено" in admin


def test_board_server_starts_and_stops(tmp_path):
    import json
    import urllib.request

    app = create_app(tmp_path / "данные", opener=lambda p: None, board_host="127.0.0.1")
    srv = app.state.board
    assert srv.start() and srv.running
    with urllib.request.urlopen(f"http://127.0.0.1:{srv.port}/health", timeout=5) as r:
        assert json.load(r) == {"app": "st-secretary-board"}
    srv.stop()
    assert not srv.running


def test_board_qr_and_addresses():
    from st_secretary.web.board import lan_addresses, qr_svg

    svg = qr_svg("http://192.168.0.2:8780/")
    assert svg.startswith("<svg") and "</svg>" in svg
    assert all(not ip.startswith("127.") for ip in lan_addresses())


def test_judge_phone_link_sync_and_conflicts(client, tmp_path, psr_card):
    import json

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    q = "?z=М/Ж_3"
    client.post(base(f) + "/results/stages" + q, data={"st-0-tour": "Тур 1", "st-0-name": "Узлы", "st-0-kv": "15",
                                                        "st-1-tour": "Тур 1", "st-1-name": "Бивак"})
    page = client.get(base(f) + "/judges" + q).text
    assert "Ссылки этапов — зачёт М/Ж_3" in page and page.count("Выдать ссылку") == 2
    r = client.post(base(f) + "/judges/link" + q, data={"stage": "s1"}, follow_redirects=False)
    assert "done=judge_issued" in r.headers["location"]
    token = next(iter(f.run_data()["judge_links"]))

    phone = TestClient(client.app.state.board.app)
    assert "Ссылка не действует" in phone.get("/j/нетакой").text
    p = phone.get(f"/j/{token}")
    assert p.status_code == 200 and p.headers["cache-control"] == "no-store" and "Тур 1 · Узлы" in p.text
    data = json.loads(re.search(r'<script type="application/json" id="data">(.*?)</script>', p.text, re.DOTALL).group(1))
    assert [t["team"] for t in data["teams"]] == ["Кедр", "Сосна"] and data["stage"]["kv"] == 15
    assert data["sync_url"] == f"/j/{token}/sync" and data["records"] == {}

    body = {"device": "т-abc", "records": [
        {"file": "Кедр.xlsx", "points": "20", "arrive": "10:05:00", "leave": "10:17:00", "updated": 1000},
        {"file": "Сосна.xlsx", "points": "", "arrive": "10:20:00", "removed": True, "reason": "опасно", "updated": 1000},
        {"file": "../../Карточка_соревнования.xlsx", "points": "1", "updated": 1000}]}
    j = phone.post(f"/j/{token}/sync", json=body).json()
    assert j["ok"] and j["saved"] == ["Кедр.xlsx", "Сосна.xlsx"]  # чужой файл — не записан
    assert f.run_data()["zachety"]["М/Ж_3"]["teams"]["Кедр.xlsx"]["points"]["s1"] == "20"
    res = client.get(base(f) + "/results" + q).text
    assert 'value="20" data-stage="s1" class=" pts-phone"' in res
    assert "«Сосна», Тур 1 · Узлы: судья этапа отметил снятие с этапа: опасно" in res
    jp = client.get(base(f) + "/judges" + q).text
    assert "2 из 2" in jp and f"/judges/phone/{token}" in jp  # судья прислал обе команды; открыть страницу судьи здесь
    local = client.get(base(f) + f"/judges/phone/{token}").text
    assert f"/judges/phone/{token}/sync" in local  # на ноутбуке — отправка в саму программу, без Wi-Fi

    # секретарь поправил по протесту — телефон таблицу не перезаписывает, расхождение видно
    by = {name: i for i, name in re.findall(r'name="p-(\d)-file" value="([^"]+)"', res)}
    client.post(base(f) + "/results/points" + q, data={f"p-{i}-file": n for n, i in by.items()}
                | {f"p-{by['Кедр.xlsx']}-s1": "25"})
    body["records"] = [{"file": "Кедр.xlsx", "points": "18", "updated": 2000}]
    phone.post(f"/j/{token}/sync", json=body)
    res = client.get(base(f) + "/results" + q).text
    assert 'value="25"' in res and "«Кедр», Тур 1 · Узлы: судья этапа прислал 18" in res and "в таблице 25" in res

    again = phone.get(f"/j/{token}").text  # на телефон приходят уже принятые записи (например, с другого телефона)
    recs = json.loads(re.search(r'id="data">(.*?)</script>', again, re.DOTALL).group(1))["records"]
    assert recs["Кедр.xlsx"]["points"] == "18" and recs["Кедр.xlsx"]["file"] == "Кедр.xlsx"
    client.post(base(f) + "/judges/link" + q, data={"stage": "s1", "do": "revoke"})
    r = phone.post(f"/j/{token}/sync", json=body)
    assert r.status_code == 404 and "ссылка больше не действует" in r.json()["error"]
    assert phone.post(f"/j/{token}/sync", content=b"not json").status_code == 400
    assert "Раздача по Wi-Fi не включена" in client.get(base(f) + "/judges/print" + q).text


def test_speleo_zachet_times_protocol_and_board(client, tmp_path, psr_card):
    from dataclasses import replace

    from openpyxl import load_workbook

    from st_secretary.competition import Zachet

    comp = replace(psr_card, zachety=[Zachet("М/Ж", 3, "0840271811Я", age_from=16, team_size=3)])
    f = client.app.state.store.create(comp)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    url, q = base(f) + "/results", "?z=М/Ж_3"
    client.post(url + "/stages" + q, data={"st-0-name": "Колодец", "st-1-name": "Шкуродёр", "spp": "", "expected": "25",
                                            "kv": "60", "removed_order": "after"})
    page = client.get(url).text
    assert "Время и штрафные баллы" in page and "1 балл = <b>15 с</b>" in page and 'name="p-0-start"' in page
    by = {name: i for i, name in re.findall(r'name="p-(\d)-file" value="([^"]+)"', page)}
    k, s = by["Кедр.xlsx"], by["Сосна.xlsx"]
    data = {f"p-{k}-file": "Кедр.xlsx", f"p-{s}-file": "Сосна.xlsx",
            f"p-{k}-start": "10:00:00", f"p-{k}-finish": "10:20:30", f"p-{k}-cutoffs": "2:00", f"p-{k}-s1": "0,3",
            f"p-{k}-chip": "2001234", f"p-{s}-start": "10:05:00", f"p-{s}-finish": "10:15:00", f"p-{s}-s2": "с"}
    j = client.post(url + "/points" + q, data=data, headers={"X-Autosave": "1"}).json()
    assert j["cells"]["Кедр.xlsx"] == {"total": "18:34,5", "place": "1", "bad": []}  # 18:30 + 0,3 × 15 с
    assert j["cells"]["Сосна.xlsx"]["place"] == "2"  # быстрее, но со снятием — после прошедших полностью
    assert "Время на дистанции" in j["results"] and "18:30" in j["results"]
    assert f.run_data()["zachety"]["М/Ж_3"]["teams"]["Кедр.xlsx"]["chip"] == "2001234"

    data[f"p-{s}-finish"] = "10:6"
    j = client.post(url + "/points" + q, data=data, headers={"X-Autosave": "1"}).json()
    assert j["cells"]["Сосна.xlsx"]["bad"] == ["finish"] and "«Сосна»: финиш «10:6» — не время" in j["results"]

    data[f"p-{s}-finish"] = "10:15:00"
    client.post(url + "/points" + q, data=data)
    client.post(url + "/publish" + q)
    ws = load_workbook(next(f.protocols_dir.glob("Предварительный*.xlsx")))["Протокол"]
    cells = [c for row in ws.iter_rows(values_only=True) for c in row if c not in (None, "")]
    # протокол спелео — как у СЕКРЕТАРЬ_ST (Правки, п. 33): время с часами, «Снятий с этапов» — если были
    assert "Время прохождения дистанции" in cells and "0:18:34,5" in cells and "Снятий с этапов" in cells
    client.post(base(f) + "/board/toggle", data={"on": "1"})
    assert "18:34,5" in TestClient(client.app.state.board.app).get("/").text

    # своё число секунд за балл из Условий (Правки, п. 47): 0,3 × 60 с = 18 с
    stages = {"st-0-id": "s1", "st-0-name": "Колодец", "st-1-id": "s2", "st-1-name": "Шкуродёр", "expected": "25",
              "kv": "60", "removed_order": "after"}
    r = client.post(url + "/stages" + q, data={**stages, "spp": "own", "spp_own": "60"})
    assert "Этапы дистанции сохранены" in r.text and "1 балл = <b>60 с</b>" in r.text
    assert '<option value="own" selected>' in r.text and 'name="spp_own" type="text" inputmode="numeric" value="60"' in r.text
    j = client.post(url + "/points" + q, data=data, headers={"X-Autosave": "1"}).json()
    assert j["cells"]["Кедр.xlsx"]["total"] == "18:48"  # 18:30 + 0,3 × 60 с
    r = client.post(url + "/stages" + q, data={**stages, "spp": "own", "spp_own": "минута"})
    assert "впишите в «своё» целое число секунд" in r.text and "1 балл = <b>60 с</b>" in r.text  # прежнее — на месте


def test_si_reader_upload_fills_times(client, tmp_path, psr_card):
    from dataclasses import replace

    from test_si_reader import csv, line

    from st_secretary.competition import Zachet

    comp = replace(psr_card, zachety=[Zachet("М/Ж", 3, "0840271811Я", team_size=3)])
    f = client.app.state.store.create(comp)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    url, q = base(f) + "/results", "?z=М/Ж_3"
    client.post(url + "/stages" + q, data={"st-0-name": "Колодец", "cutoff_pairs": "31-32", "kv": ""})
    client.post(url + "/points" + q, data={"p-0-file": "Кедр.xlsx", "p-0-chip": "2001234"})
    data = csv(line(1, "2001234", "10:00:00", "10:30:00", [("31", "10:05:00"), ("32", "10:08:00")]),
               line(2, "2007777", "10:00:00", "10:40:00"))
    r = client.post(url + "/si" + q, files={"csv": ("si_reader.csv", data, "text/csv")}, follow_redirects=False)
    loc = r.headers["location"]
    assert "done=si_done" in loc and "unknown=2007777" in loc
    d = f.run_data()["zachety"]["М/Ж_3"]["teams"]["Кедр.xlsx"]
    assert (d["start"], d["finish"], d["cutoffs"]) == ("10:00:00", "10:30:00", "3:00")
    page = client.get(url + q).text
    assert "27:00" in page  # 30 мин − 3 мин отсечки
    bad = client.post(url + "/si" + q, files={"csv": ("x.csv", b"a;b\n1;2\n", "text/csv")}, follow_redirects=False)
    assert "done=si_bad" in bad.headers["location"]


def test_practice_page_and_excel(client, psr_card, opened):
    f = client.app.state.store.create(psr_card)
    client.post(base(f) + "/awards/grades", data={"g-0": "отлично"})
    client.post(base(f) + "/contracts/add", data={"fio": "Этапов Семён Игоревич", "role": "Судья этапа", "category": "СС3К"})
    assert "Судейская практика" in client.get("/").text
    page = client.get("/practice").text
    assert "Судьин Иван Петрович" in page and "Этапов Семён Игоревич" in page and "оценка «отлично»" in page
    assert "<th>Присвоение</th>" in page and "ВК</b>:" in page  # у СС1К — на присвоение ВК, по таблице приказа № 1101
    assert "2К</b>:" in page  # у СС3К — на присвоение 2К
    assert client.get("/practice.xlsx").content[:2] == b"PK"
    r = client.post("/practice/open", follow_redirects=False)
    assert r.status_code == 303 and opened[-1].name == "Судейская практика.xlsx"


def test_start_order_drag_markup(client):
    """Порядок старта перетаскивается (Правки.md, п. 2): «ручка» у каждой строки, первый старт и интервал —
    для пересчёта времени в браузере; сохраняют прежние поля pos-i / file-i."""
    client.post("/training")
    f = client.app.state.store.all()[0]
    b = base(f) + "/start?" + urlencode({"z": "М/Ж_3"})
    client.post(base(f) + "/start/draw?" + urlencode({"z": "М/Ж_3"}),
                data={"action": "save", "day": "2026-10-03", "first": "10:00", "interval": "3"})
    html = client.get(b).text
    assert 'data-reorder data-first="10:00" data-interval="3"' in html
    assert html.count("data-drag") == html.count('name="pos-') > 0 and "data-time" in html
    assert "стрелками" in html and "Перетащите команду" in html


def test_own_documents_in_settings_remembered_on_computer(client, tmp_path, psr_card):
    """Правки, п. 30: свой документ — в «Какие документы проверять» (название, у кого, кому нужен, «или»); набор
    запоминается на компьютере — новое соревнование начинается с него; стёрли название — забыт."""
    store = client.app.state.store
    f = store.create(psr_card)
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    page = client.get(base(f) + "/admission").text
    assert "Свои документы — по Положению" in page
    form = {"doc-id": "on", "doc-med": "on", "doc-book": "on", "doc-ins": "on", "tdoc-app": "on", "tdoc-doctor": "on",
            "od-0-key": "", "od-0-on": "on", "od-0-title": "Расписка  родителей несовершеннолетнего",
            "od-0-short": "", "od-0-scope": "person", "od-0-when": "minor", "od-0-alt": "",
            "od-1-key": "", "od-1-on": "on", "od-1-title": "Журнал инструктажа по ТБ", "od-1-short": "Журнал ТБ",
            "od-1-scope": "team", "od-1-when": "all", "od-1-alt": "",
            "od-2-key": "", "od-2-title": "", "start_at": ""}
    r = client.post(base(f) + "/admission/settings", data=form, follow_redirects=False)
    assert "adm_settings" in r.headers["location"]
    s = f.admission()["settings"]
    rasp, journal = (d["key"] for d in s["own_docs"])
    assert "oms" not in s["docs"] and rasp in s["docs"] and journal in s["team_docs"]
    assert s["own_docs"][0]["title"] == "Расписка родителей несовершеннолетнего"
    page = client.get(base(f) + "/admission").text
    assert 'title="Расписка родителей несовершеннолетнего">Расписка родителей несовершеннолетнего</th>' in page
    assert "Журнал ТБ — " not in page and "Журнал инструктажа по ТБ" in page
    assert "не нужен: только несовершеннолетним" in page  # взрослым расписка не нужна — «—» в клетке

    # новое соревнование на этом компьютере — с тем же набором
    g = store.create(replace(psr_card, title="Кубок города N по спортивному туризму"))
    store.apply_doc_set(g)
    assert g.admission()["settings"]["docs"] == s["docs"] and len(g.admission()["settings"]["own_docs"]) == 2
    # стёрли название — документ забыт и здесь, и на компьютере
    form |= {"od-0-key": rasp, "od-0-title": "", "od-1-key": journal}
    client.post(base(f) + "/admission/settings", data=form)
    assert [d["key"] for d in f.admission()["settings"]["own_docs"]] == [journal]
    assert [d["key"] for d in store.doc_set()["docs"]] == [journal]


def test_draw_without_touching_time_fields_sets_start_times(client):
    """Правки, п. 40: в полях времени — настоящие значения по умолчанию (10:00, 5 мин), а не серые подсказки; кнопка
    жеребьёвки отправляет и время — после жеребьёвки у всех команд есть время старта."""
    client.post("/training")
    f = client.app.state.store.all()[0]
    q = "?" + urlencode({"z": "М/Ж_3"})
    page = client.get(base(f) + "/start" + q).text
    action = f'action="{base(f)}/start/draw{q}"'.replace("&", "&amp;")
    assert page.count(action) == 1  # одна форма на жеребьёвку и время старта
    form = FormFields(page.replace(action, 'id="d"'), "d").fields
    assert form["interval"] == "5" and re.fullmatch(r"\d\d:\d\d", form["first"])
    r = client.post(base(f) + "/start/draw" + q, data={**form, "action": "draw"}, follow_redirects=False)
    flash = client.get(r.headers["location"]).text
    assert f"Время старта — с {form['first']} через 5 мин" in flash
    import json

    sched = client.get(base(f) + "/schedule").text
    lanes = json.loads(sched.split('id="schedule-data">')[1].split("</script>")[0])["lanes"]
    times = [r["t"] for r in next(x for x in lanes if x["zkey"] == "М/Ж_3")["rows"]]
    assert len(times) > 1 and None not in times and times[1] - times[0] == 300
    page = client.get(base(f) + "/start" + q).text
    assert "Время старта ещё не сохранено" not in page and "первый старт" in page


def test_old_draw_request_without_time_takes_defaults(client, tmp_path, psr_card):
    """Жеребьёвка без полей времени (старая форма) — время по умолчанию: «Начало соревнований» или 10:00, 5 мин."""
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    q = "?" + urlencode({"z": "М/Ж_3"})
    page = client.get(base(f) + "/start" + q).text
    assert 'name="first" value="10:00" placeholder="чч:мм"' in page and 'name="interval" value="5"' in page
    r = client.post(base(f) + "/start/draw" + q, data={"method": "random", "action": "draw"}, follow_redirects=False)
    assert "first=10%3A00" in r.headers["location"] and "interval=5" in r.headers["location"]
    dr = f.run_data()["zachety"]["М/Ж_3"]["draw"]
    assert (dr["first"], dr["interval"]) == ("10:00", "5")
    client.post(base(f) + "/start/draw" + q, data={"action": "save", "first": "", "interval": "5"})
    page = client.get(base(f) + "/start" + q).text
    assert 'name="first" value="" placeholder="чч:мм"' in page  # стёрли время — без времени, только очерёдность


def test_rename_stage_keeps_points_judge_log_and_link(client, tmp_path, psr_card):
    """Правки.md, п. 5 (переделать): название этапа можно менять — баллы, журнал судьи и ссылка привязаны к id этапа;
    на телефоне судьи — новое название. В «Этапах дистанции» название — поле, которое растёт по тексту."""
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    q = "?z=М/Ж_3"
    stages = {"st-0-tour": "Тур 1", "st-0-name": "Узлы", "st-0-nv": "5", "st-0-kv": "10", "st-0-tsh": "20",
              "st-0-vsh": "10", "st-1-tour": "Тур 1", "st-1-name": "Бивак"}
    client.post(base(f) + "/results/stages" + q, data=stages)
    client.post(base(f) + "/judges/link" + q, data={"stage": "s1"})
    token = next(iter(f.run_data()["judge_links"]))
    phone = TestClient(client.app.state.board.app)
    phone.post(f"/j/{token}/sync", json={"device": "т-1", "records": [
        {"file": "Кедр.xlsx", "points": "10", "arrive": "10:00:00", "leave": "10:07:12", "updated": 1}]})
    z = f.run_data()["zachety"]["М/Ж_3"]
    assert z["teams"]["Кедр.xlsx"]["points"]["s1"] == "15"

    ids = {"st-0-id": "s1", "st-1-id": "s2"}
    client.post(base(f) + "/results/stages" + q, data={**stages, **ids, "st-0-name": "Транспортировка пострадавшего"})
    z = f.run_data()["zachety"]["М/Ж_3"]
    assert [s["name"] for s in z["stages"]] == ["Транспортировка пострадавшего", "Бивак"]
    assert z["teams"]["Кедр.xlsx"]["points"]["s1"] == "15" and "Кедр.xlsx" in z["judge"]["s1"]
    assert js_token(f) == token and "Тур 1 · Транспортировка пострадавшего" in phone.get(f"/j/{token}").text
    page = client.get(base(f) + "/results" + q).text
    assert '<textarea name="st-0-name" rows="1" class="st-name" data-autosize' in page
    assert "res-stages res-stages-psr" in page and 'class="c-name"' in page


def test_extra_result_part_column_protocol_and_judge_phone(client, tmp_path, psr_card):
    """Правки, п. 32: составляющая результата задаётся у зачёта по времени, вносится в таблице (своя колонка) или с
    телефона судьи («Прибыла» — «Убыла», как время этапа) и прибавляется к результату; в протоколе — своя колонка."""
    import json

    from openpyxl import load_workbook

    speleo = replace(psr_card, zachety=[replace(psr_card.zachety[0], discipline_code="0840271811Я", team_size=3)])
    f = client.app.state.store.create(speleo)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    q = "?" + urlencode({"z": "М/Ж_3"})
    client.post(base(f) + "/results/stages" + q, data={"st-0-name": "Колодец", "add-0-id": "", "add-0-name": "Топосъёмка",
                                                        "add-0-kind": "time", "add-1-id": "", "add-1-name": ""})
    assert f.run_data()["zachety"]["М/Ж_3"]["adds"] == [{"id": "a1", "name": "Топосъёмка", "kind": "time",
                                                          "in_percent": False}]
    page = client.get(base(f) + "/results" + q).text
    assert '<th rowspan="2">Топосъёмка</th>' in page and 'name="p-0-add-a1"' in page
    # % от победителя — по умолчанию без составляющей, как СЕКРЕТАРЬ_ST (Правки, п. 48)
    assert "% от результата победителя посчитан без составляющей «Топосъёмка»." in page

    # с телефона судьи: ссылка на «Топосъёмку», время между «Прибыла» и «Убыла»
    client.post(base(f) + "/judges/link" + q, data={"stage": "add-a1"})
    token = js_token(f)
    phone = TestClient(client.app.state.board.app)
    p = phone.get(f"/j/{token}").text
    data = json.loads(re.search(r'<script type="application/json" id="data">(.*?)</script>', p, re.DOTALL).group(1))
    assert "Топосъёмка" in p and data["stage"]["add"] == "time" and data["penalties"] is None
    phone.post(f"/j/{token}/sync", json={"device": "т-1", "records": [
        {"file": "Кедр.xlsx", "arrive": "11:00:00", "leave": "11:10:39", "points": "0", "updated": 1}]})
    team_ = f.run_data()["zachety"]["М/Ж_3"]["teams"]["Кедр.xlsx"]
    assert team_["add-a1"] == "10:39" and "s1" not in team_.get("points", {})
    client.post(base(f) + "/results/points" + q, data={"p-0-file": "Кедр.xlsx", "p-0-start": "10:00:00",
                                                        "p-0-finish": "10:18:05", "p-0-add-a1": "10:39"})
    r = client.post(base(f) + "/results/publish" + q, follow_redirects=False)
    assert r.status_code == 303
    proto = next((f.path / "Протоколы").glob("Предварительный протокол*.xlsx"))
    cells = [str(c.value) for row in load_workbook(proto)["Протокол"].iter_rows() for c in row if c.value is not None]
    assert "Топосъёмка" in cells and "0:10:39" in cells and "0:28:44" in cells  # как у СЕКРЕТАРЬ_ST (п. 33)
    assert "% от результата победителя посчитан без составляющей «Топосъёмка»." in cells  # под таблицей (п. 48)

    # «с ней» — от полного результата: настройка сохраняется и видна на странице
    page = client.post(base(f) + "/results/stages" + q, data={"st-0-id": "s1", "st-0-name": "Колодец", "add-0-id": "a1",
                                                               "add-0-name": "Топосъёмка", "add-0-kind": "time",
                                                               "add-0-pct": "in"}).text
    assert f.run_data()["zachety"]["М/Ж_3"]["adds"][0]["in_percent"] is True
    assert '<option value="in" selected>' in page and "посчитан от полного результата (с составляющей" in page


def test_named_application_word_from_preapp(client, tmp_path, psr_card):
    """Правки, п. 36: «Именная заявка (Word)» у команды — шапка из карточки, участники из предзаявки (ФИО полностью,
    дата рождения, разряд), пустые поля для врача и подписей; свой бланк соревнования — с метками {…}."""
    from docx import Document

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    page = client.get(base(f) + "/preapps/team?file=Кедр.xlsx").text
    assert "/preapps/named.docx?file=" in page and "Именная заявка (Word)" in page
    r = client.get(base(f) + "/preapps/named.docx?file=Кедр.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "В Главную судейскую коллегию Чемпионата города N по спортивному туризму" in text and "ЗАЯВКА" in text
    assert "команду «Кедр» (Красноярск)" in text and "Лебедев Антон Игоревич" in text  # представитель
    assert "«С правилами техники безопасности ознакомлен»" in text and "Приложения: 1. Документы о возрасте" in text
    rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert rows[0][1] == "Фамилия, имя, отчество участника" and len(rows) == 4
    assert rows[1][:4] == ["1", "Лебедев Антон Игоревич", "02.02.1990", "I"] and rows[1][4:6] == ["", ""]
    assert rows[1][6] == "М/Ж_3"

    # свой бланк соревнования: метки в тексте и строка участника с {ФИО}
    tpl = Document()
    tpl.add_paragraph("Заявка команды {команда} на {соревнование}")
    tpl.add_paragraph("В ГСК {соревнования}")
    t = tpl.add_table(rows=2, cols=3)
    for c, v in zip(t.rows[0].cells, ["№", "ФИО", "Год"], strict=True):
        c.text = v
    for c, v in zip(t.rows[1].cells, ["{№}", "{ФИО}", "{дата рождения}"], strict=True):
        c.text = v
    buf = io.BytesIO()
    tpl.save(buf)
    r = client.post(base(f) + "/preapps/named-template", files={"template": ("бланк.docx", buf.getvalue())},
                    follow_redirects=False)
    assert "named_saved" in r.headers["location"]
    assert "Сейчас — <b>свой бланк</b>" in client.get(base(f) + "/forms").text
    doc = Document(io.BytesIO(client.get(base(f) + "/preapps/named.docx?file=Кедр.xlsx").content))
    assert doc.paragraphs[0].text == f"Заявка команды Кедр на {psr_card.title}"
    assert doc.paragraphs[1].text == "В ГСК Чемпионата города N по спортивному туризму"
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows][1:] == [
        ["1", "Лебедев Антон Игоревич", "02.02.1990"], ["2", "Зуева Мария Олеговна", "05.06.1996"],
        ["3", "Носов Глеб Андреевич", "09.09.1993"]]
    client.post(base(f) + "/preapps/named-template", data={"do": "delete"})
    assert "ЗАЯВКА" in "\n".join(p.text for p in Document(io.BytesIO(
        client.get(base(f) + "/preapps/named.docx?file=Кедр.xlsx").content)).paragraphs)


def js_token(f) -> str:
    return next(iter(f.run_data()["judge_links"]))


def test_check_badges_are_clickable(client, tmp_path, psr_card):
    """Правки.md, п. 22: «Проверить» / «Ошибка» — не просто значки, а ссылки: плитки сводки → список с фильтром,
    у команды → карточка заявки с фильтром замечаний, у замечания → «Исправить» или к «Проверено»."""
    f = client.app.state.store.create(psr_card)
    b = base(f)
    client.post(b + "/preapps/upload", files=[("files", ("Лесовики.xlsx", lesoviki(tmp_path), XLSX))])
    page = client.get(b + "/preapps").text
    assert 'class="tile tile-warning tile-link" href="?show=' in page and 'data-show="' in page
    assert re.search(r'class="badge badge-(warning|error) badge-link" href="[^"]*preapps/team\?file=[^"]*&only=', page)
    card = client.get(b + "/preapps/team?file=" + quote("Лесовики.xlsx")).text
    assert "badge-link" in card and ("data-to-check" in card or "Исправить" in card)
    js = (Path(__file__).parents[1] / "src" / "st_secretary" / "web" / "static" / "app.js").read_text(encoding="utf-8")
    assert "function onlyFilter" in js and 'qs.get("show")' in js and "[data-to-check]" in js


def test_wait_cut_setting_on_results_page_and_phone(client, tmp_path, psr_card):
    """Правки.md, п. 3: «Ожидание очереди на этапе» — рядом с неполным интервалом; смена пересчитывает итог
    программы; телефон судьи знает настройку."""
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    q = "?z=М/Ж_3"
    stages = {"st-0-tour": "Тур 1", "st-0-name": "Узлы", "st-0-nv": "5", "st-0-kv": "10", "st-0-tsh": "20",
              "st-0-vsh": "10"}
    client.post(base(f) + "/results/stages" + q, data=stages)
    page = client.get(base(f) + "/results" + q).text
    assert 'name="wait_cut"' in page and page.index('name="vsh_round"') < page.index('name="wait_cut"')
    client.post(base(f) + "/judges/link" + q, data={"stage": "s1"})
    token = js_token(f)
    phone = TestClient(client.app.state.board.app)
    phone.post(f"/j/{token}/sync", json={"device": "т-1", "records": [
        {"file": "Кедр.xlsx", "points": "10", "arrive": "10:00:00", "leave": "10:07:12", "cutoff": "1:00",
         "updated": 1}]})
    assert f.run_data()["zachety"]["М/Ж_3"]["teams"]["Кедр.xlsx"]["points"]["s1"] == "13"
    assert '"wait_cut": true' in phone.get(f"/j/{token}").text
    client.post(base(f) + "/results/stages" + q, data={**stages, "st-0-id": "s1", "wait_cut": "no"})
    z = f.run_data()["zachety"]["М/Ж_3"]
    assert z["wait_cut"] == "no" and z["teams"]["Кедр.xlsx"]["points"]["s1"] == "15"
    html = phone.get(f"/j/{token}").text
    assert '"wait_cut": false' in html and "ожидание не вычитается" in html


def test_removed_mark_in_table_protocol_and_unmark(client, tmp_path, psr_card, opened):
    """Правки.md, п. 5 (ответ 2): снята → в клетке МШ и отметка «снята ✕» (причина в подсказке), в протоколе
    результатов и на листе «По этапам» — «снята»; секретарь снимает отметку — итог этапа заново."""
    from openpyxl import load_workbook

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    q = "?z=М/Ж_3"
    client.post(base(f) + "/results/stages" + q, data={
        "st-0-tour": "Тур 1", "st-0-name": "Узлы", "st-0-nv": "5", "st-0-kv": "10", "st-0-tsh": "20", "st-0-vsh": "10"})
    client.post(base(f) + "/judges/link" + q, data={"stage": "s1"})
    phone = TestClient(client.app.state.board.app)
    phone.post(f"/j/{js_token(f)}/sync", json={"device": "т-1", "records": [
        {"file": "Кедр.xlsx", "points": "10", "arrive": "10:00:00", "leave": "10:07:12", "removed": True,
         "reason": "опасная страховка", "updated": 1}]})
    page = client.get(base(f) + "/results" + q).text
    assert 'class="pts-mark"' in page and "снята ✕" in page and "опасная страховка" in page
    assert "pts-removed" in page and "Узлы — снята с этапа" in page  # и в таблице результатов под командой
    client.post(base(f) + "/results/publish" + q)
    wb = load_workbook(opened[-1])
    cells = [c for row in wb["Протокол"].iter_rows(values_only=True) for c in row if c]
    assert "Снятия с этапов" in cells and "Узлы" in cells
    stage_cells = [c for row in wb["По этапам"].iter_rows(values_only=True) for c in row if c]
    assert "30 снята" in stage_cells

    r = client.post(base(f) + "/results/unmark" + q + "&file=" + quote("Кедр.xlsx") + "&sid=s1", follow_redirects=False)
    assert r.status_code == 303 and "focus=c-t-" in r.headers["location"]
    z = f.run_data()["zachety"]["М/Ж_3"]
    assert z["teams"]["Кедр.xlsx"]["points"]["s1"] == "15" and z["teams"]["Кедр.xlsx"]["unremoved"] == ["s1"]
    page = client.get(base(f) + "/results" + q).text
    assert "↺ снята" in page and "pts-removed" not in page


def test_admission_doctor_mark_sets_med_for_team(client, tmp_path, psr_card):
    """Правки.md, п. 23: отметили у команды «Допуск врача» — мед. допуск у всех отмечен сам (не отметкой
    секретаря); сняли галочку у одного — больше сама не ставится; сняли «Допуск врача» — уходят только свои."""
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    url = base(f) + "/admission"
    data = team_form(client.get(url).text, "Кедр.xlsx") | {"td-doctor": "on"}
    j = client.post(url + "/team", data=data, headers={"X-Autosave": "1"}).json()
    assert "по допуску врача в заявке" in j["team"] and "adm-auto" in j["team"]
    people = f.admission()["teams"]["Кедр.xlsx"]["people"]
    assert not any(p["docs"].get("med") for p in people.values())  # не отметка секретаря — стоит сама

    data = team_form(client.get(url).text, "Кедр.xlsx")
    meds = sorted(k for k in data if k.endswith("-d-med"))
    assert len(meds) == 3  # на странице — отмечены у всех
    off_key = data[meds[0].replace("-d-med", "-key")]
    data.pop(meds[0])  # врач не допустил первого
    client.post(url + "/team", data=data, headers={"X-Autosave": "1"})
    assert f.admission()["teams"]["Кедр.xlsx"]["people"][off_key]["med_off"]
    data = team_form(client.get(url).text, "Кедр.xlsx")
    assert meds[0] not in data and meds[1] in data

    data.pop("td-doctor")  # сняли «Допуск врача»
    client.post(url + "/team", data=data, headers={"X-Autosave": "1"})
    data = team_form(client.get(url).text, "Кедр.xlsx")
    assert not any(k.endswith("-d-med") for k in data)


def test_author_in_footer_quietly(client):
    """Правки.md, п. 11: «СТ-Секретарь <версия> · автор Udnikov Denis» — мелко в подвале страниц программы; на печати и
    на страницах для судей и табло — нет."""
    import tomllib

    from st_secretary import __version__

    home = client.get("/").text
    assert f'<p class="page-foot">СТ-Секретарь {__version__} бета · автор Udnikov Denis' in home
    css = (Path(__file__).parents[1] / "src" / "st_secretary" / "web" / "static" / "style.css").read_text(encoding="utf-8")
    assert "@media print { .page-foot { display: none; } }" in css
    tpl = Path(__file__).parents[1] / "src" / "st_secretary" / "web" / "templates"
    assert all("Udnikov" not in (tpl / n).read_text(encoding="utf-8") for n in ("judge.html", "board.html",
                                                                               "judges_print.html"))
    meta = tomllib.loads((Path(__file__).parents[1] / "pyproject.toml").read_text(encoding="utf-8"))
    assert meta["project"]["authors"] == [{"name": "Udnikov Denis"}]


def test_judge_names_self_and_secretary_sees_who_and_phone(client, tmp_path, psr_card):
    """Правки.md, п. 12: на телефоне — список судей соревнования (номер подставляется, если известен); на странице
    «Телефоны судей этапов» — кто судит, номер ссылкой tel:, журнал этапа; другой номер — «Принять номер»."""
    import json

    from st_secretary import staff as sf

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    q = "?z=М/Ж_3"
    client.post(base(f) + "/results/stages" + q, data={"st-0-tour": "Тур 1", "st-0-name": "Узлы"})
    client.post(base(f) + "/judges/link" + q, data={"stage": "s1"})
    token = js_token(f)
    store = client.app.state.store
    chief = sf.people(psr_card, f.contracts())[0]
    store.save_personal(chief.key, {"phone": "+7 913 000-00-01"})
    phone = TestClient(client.app.state.board.app)
    html = phone.get(f"/j/{token}").text
    data = json.loads(re.search(r'<script type="application/json" id="data">(.*?)</script>', html, re.DOTALL).group(1))
    assert "Кто судит этап?" in html
    assert {"fio": chief.fio, "role": "Главный судья", "phone": "+7 913 000-00-01"} in data["judges"]

    phone.post(f"/j/{token}/sync", json={"device": "т-1", "judge": {"fio": chief.fio, "phone": "8 913 000 00 01"},
                                         "records": [{"file": "Кедр.xlsx", "points": "5", "updated": 1}]})
    page = client.get(base(f) + "/judges" + q).text
    assert chief.fio in page and 'href="tel:89130000001"' in page and "судья указал другой номер" not in page
    assert "Журнал этапа: кто что прислал (1)" in page
    phone.post(f"/j/{token}/sync", json={"device": "т-2", "judge": {"fio": chief.fio, "phone": "+7 999 111-22-33"},
                                         "records": [{"file": "Кедр.xlsx", "points": "6", "updated": 2}]})
    page = client.get(base(f) + "/judges" + q).text
    assert "судья указал другой номер (в личных данных +7 913 000-00-01)" in page
    client.post(base(f) + "/judges/accept-phone" + q, data={"key": chief.key, "phone": "+7 999 111-22-33"})
    assert store.personal()[chief.key]["phone"] == "+7 999 111-22-33"
    assert client.get(base(f) + "/judges" + q).text.count("судья указал другой номер") == 1  # у прежней записи
    results = client.get(base(f) + "/results" + q).text
    assert f'title="С телефона судьи этапа: {chief.fio}, +7 999 111-22-33"' in results
    assert chief.fio not in TestClient(client.app.state.board.app).get("/").text  # на табло — нет


def test_judge_phone_contacts_heads_and_stage_judges(client, tmp_path, psr_card):
    """Правки.md, п. 14: на телефоне судьи — «Связь»: ГСК с номерами и судьи всех этапов (кто присылал с этапа);
    этап без судьи — «судья не указан»; список обновляется с каждой отправкой."""
    import json

    from st_secretary import staff as sf

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    q = "?z=М/Ж_3"
    client.post(base(f) + "/results/stages" + q, data={"st-0-tour": "Тур 1", "st-0-name": "Узлы",
                                                       "st-1-tour": "Тур 1", "st-1-name": "Бивак"})
    client.post(base(f) + "/judges/link" + q, data={"stage": "*"})
    links = {v["stage"]: t for t, v in f.run_data()["judge_links"].items()}
    chief = sf.people(psr_card, f.contracts())[0]
    client.app.state.store.save_personal(chief.key, {"phone": "+7 913 000-00-01"})
    phone = TestClient(client.app.state.board.app)
    html = phone.get(f"/j/{links['s1']}").text
    data = json.loads(re.search(r'<script type="application/json" id="data">(.*?)</script>', html, re.DOTALL).group(1))
    assert "Связь: главный судья, судьи этапов" in html and data["stage_id"] == "s1" and data["zachet"] == "М/Ж_3"
    assert data["contacts"]["heads"] == [{"role": "Главный судья", "fio": chief.fio, "phone": "+7 913 000-00-01"}]
    assert [(s["title"], s["fio"]) for s in data["contacts"]["stages"]] == [("Тур 1 · Узлы", ""), ("Тур 1 · Бивак", "")]
    r = phone.post(f"/j/{links['s2']}/sync", json={"device": "т-2", "judge": {"fio": "Петров Пётр", "phone": "+7 900 1"},
                                                   "records": [{"file": "Кедр.xlsx", "points": "3", "updated": 1}]})
    stages = r.json()["contacts"]["stages"]
    assert (stages[1]["fio"], stages[1]["phone"]) == ("Петров Пётр", "+7 900 1") and stages[0]["fio"] == ""


def test_penalty_table_choose_print_and_phone(client, tmp_path, psr_card):
    """Правки.md, п. 15: у зачёта выбирается таблица штрафов (в ПСР по умолчанию нет — можно взять пешеходную или
    загрузить свою из Excel); её можно распечатать и скачать; у судьи на телефоне — «Таблица штрафов» с поиском."""
    import json

    from openpyxl import Workbook

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    q = "?z=М/Ж_3"
    client.post(base(f) + "/results/stages" + q, data={"st-0-tour": "Тур 1", "st-0-name": "Узлы"})
    client.post(base(f) + "/judges/link" + q, data={"stage": "s1"})
    page = client.get(base(f) + "/judges" + q).text
    assert "3. Таблица штрафов" in page and "в Правилах нет (ПСР, горные" in page and "Таблица не выбрана" in page
    assert client.get(base(f) + "/penalties" + q).status_code == 404

    client.post(base(f) + "/judges/penalties" + q, data={"table": "pedestrian"})
    page = client.get(base(f) + "/judges" + q).text
    assert "Пешеходные дистанции — система оценки нарушений</b> — 26 строк" in page
    printed = client.get(base(f) + "/penalties" + q).text
    assert "Не заблокирована защёлка карабина" in printed and "Бесштрафовая" in printed
    x = client.get(base(f) + "/penalties.xlsx" + q)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    phone = TestClient(client.app.state.board.app)
    html = phone.get(f"/j/{js_token(f)}").text
    data = json.loads(re.search(r'<script type="application/json" id="data">(.*?)</script>', html, re.DOTALL).group(1))
    assert 'id="pen-open"' in html and data["penalties"]["systems"] and data["penalties"]["rows"][0]["code"] == "1"

    wb = Workbook()
    wb.active.append(["№", "Нарушение", "Баллы", "Разъяснение"])
    wb.active.append(["1", "Нет каски на этапе", "5", ""])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(base(f) + "/judges/penalties" + q, data={"do": "upload"},
                    files={"file": ("Условия штрафы.xlsx", buf.getvalue(), XLSX)}, follow_redirects=False)
    assert "penalty_loaded" in r.headers["location"]
    page = client.get(base(f) + "/judges" + q).text
    assert "своя: Условия штрафы.xlsx" in page and "1 строк" in page
    r = client.post(base(f) + "/judges/penalties" + q, data={"do": "upload"},
                    files={"file": ("x.xlsx", b"nope", XLSX)}, follow_redirects=False)
    assert "penalty_bad" in r.headers["location"]


def test_judge_penalty_items_journal_mapping_and_protocol(client, tmp_path, psr_card, opened):
    """Правки.md, п. 16: в журнале этапа — пункты таблицы с расшифровкой; «без пункта» сопоставляется; в протоколе
    (по желанию) — номера пунктов у баллов; на телефон приходит словарь слов судей."""
    import json

    from openpyxl import load_workbook

    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    q = "?z=М/Ж_3"
    client.post(base(f) + "/results/stages" + q, data={"st-0-tour": "Тур 1", "st-0-name": "Узлы"})
    client.post(base(f) + "/judges/link" + q, data={"stage": "s1"})
    client.post(base(f) + "/judges/penalties" + q, data={"table": "pedestrian", "protocol_codes": "1",
                                                         "jargon": "полез не туда = опоры\nбез знака"})
    token = js_token(f)
    phone = TestClient(client.app.state.board.app)
    html = phone.get(f"/j/{token}").text
    data = json.loads(re.search(r'<script type="application/json" id="data">(.*?)</script>', html, re.DOTALL).group(1))
    assert ["полез не туда", "опоры"] in data["penalties"]["jargon"] and ["замуфт", "защелк"] in data["penalties"]["jargon"]
    pens = [{"id": "a", "code": "1", "title": "Не заблокирована защёлка карабина", "v": 1, "pts": "1 балл", "n": 2},
            {"id": "c", "code": "", "title": "полез не туда", "v": None, "n": 1}]
    phone.post(f"/j/{token}/sync", json={"device": "т-1", "records": [
        {"file": "Кедр.xlsx", "points": "2", "pens": pens, "updated": 1}]})
    page = client.get(base(f) + "/judges" + q).text
    assert 'id="log-s1"' in page and "п. 1 ×2 = 2, без пункта: «полез не туда» → 2" in page and "Сопоставить" in page
    r = client.post(base(f) + "/judges/pen-code" + q, data={"sid": "s1", "file": "Кедр.xlsx", "id": "c", "code": "99"},
                    follow_redirects=False)
    assert "pen_code_bad" in r.headers["location"]
    client.post(base(f) + "/judges/pen-code" + q, data={"sid": "s1", "file": "Кедр.xlsx", "id": "c", "code": "10.1"})
    page = client.get(base(f) + "/judges" + q).text
    assert "пункт сопоставлен на ноутбуке" in page and "→ 12" in page
    results = client.get(base(f) + "/results" + q).text
    assert "п. 1 ×2 = 2, п. 10.1 = 10 → 12" in results
    client.post(base(f) + "/results/publish" + q)
    cells = [c for row in load_workbook(opened[-1])["По этапам"].iter_rows(values_only=True) for c in row if c]
    assert "2 (п. 1×2, 10.1)" in cells


def test_unofficial_card_remembers_own_values_on_this_computer(client, psr_card):
    """Правки.md, п. 17: свои группы, названия зачётов и дисциплины неофициальных соревнований запоминаются на этом
    компьютере (рядом с личными данными, не в папке соревнования) и подсказываются в карточке; лишнее — убрать."""
    from dataclasses import replace

    from st_secretary.competition import Zachet

    store = client.app.state.store
    f = store.create(psr_card)
    form = FormFields(client.get(base(f) + "/card/edit").text, "cardform").fields
    form |= {"unofficial": "1", "z-1-group": "СЕМЬИ", "z-1-distance_class": "0", "z-1-name": "Семейные команды",
             "z-1-discipline_text": "Полоса препятствий", "z-1-result": "time", "z-1-unit": "team"}
    r = client.post(base(f) + "/card/edit", data=form, follow_redirects=False)
    assert r.status_code == 303 and "done=saved" in r.headers["location"]
    own = store.own_values()
    assert own == {"groups": ["СЕМЬИ"], "names": ["Семейные команды"],
                   "disciplines": [{"name": "Полоса препятствий", "result": "time", "unit": "team"}]}
    assert store.own_values_path.parent == store.docs_root and not (f.path / store.own_values_path.name).exists()
    page = client.get(base(f) + "/card/edit").text
    assert '<option value="СЕМЬИ" label="своя группа (неофициальные)">' in page and 'data-result="time"' in page
    assert f.load().zachety[1].key == "Семейные команды"
    client.post(base(f) + "/card/forget", data={"kind": "disciplines", "value": "Полоса препятствий"})
    assert store.own_values()["disciplines"] == [] and store.own_values()["groups"] == ["СЕМЬИ"]
    store.remember_own(replace(psr_card, zachety=[Zachet("НОВЫЕ", 1, "0840161811Я")]))  # официальные — не запоминаем
    assert "НОВЫЕ" not in store.own_values()["groups"]


def test_admission_by_delegation_view(client, tmp_path, psr_card):
    """Правки.md, п. 20: комиссия «по делегациям» — команды одной территории и представителя подряд, их взнос."""
    f = client.app.state.store.create(psr_card)
    f.add_preapp("Кедр.xlsx", kedr(tmp_path))
    f.add_preapp("Сосна.xlsx", sosna(tmp_path))
    page = client.get(base(f) + "/admission?by=delegation").text
    assert "<b>по делегациям</b>" in page and page.count('class="adm-deleg"') == 2
    assert "Делегация: Красноярск" in page and "Представитель: Лебедев Антон Игоревич" in page and "1 команда" in page
    assert "по делегациям</a>" in client.get(base(f) + "/admission").text


def test_festival_group_summary_switch_backup_and_restore(client, tmp_path, psr_card):
    """Правки.md, п. 19 (неспорная часть, решение 039): фестиваль — группа соревнований на главной, переключение,
    сводка «кто заявлен в нескольких соревнованиях», копия одним архивом и восстановление; данные соревнований
    не трогаются."""
    from dataclasses import replace
    from urllib.parse import unquote

    store = client.app.state.store
    a = store.create(psr_card)
    b = store.create(replace(psr_card, title="Кубок города по спортивному туризму"))
    a.add_preapp("Кедр.xlsx", kedr(tmp_path))
    b.add_preapp("Кедр.xlsx", kedr(tmp_path))  # те же люди заявлены и там, и там
    before = sorted(p.name for p in a.path.rglob("*"))
    r = client.post("/festival/new", data={"title": "Осенний выезд", "member": [a.id, b.id]}, follow_redirects=False)
    assert r.status_code == 303 and "done=festival_made" in r.headers["location"]
    fid = unquote(r.headers["location"].split("/festival/")[1].split("?")[0])
    assert store.festival(fid)["members"] == [a.id, b.id] and sorted(p.name for p in a.path.rglob("*")) == before
    home = client.get("/").text
    assert 'class="fest-group"' in home and "Осенний выезд" in home
    page = client.get(f"/festival/{fid}").text
    assert "Кто заявлен в нескольких соревнованиях: 3" in page and "Лебедев Антон Игоревич" in page
    assert 'class="side-fest"' in client.get(base(a)).text and b.id in client.get(base(a)).text

    data = client.get(f"/festival/{fid}/backup.zip").content
    assert data[:2] == b"PK"
    r = client.post("/restore", files={"backup": ("фестиваль.zip", data, "application/zip")}, follow_redirects=False)
    assert "done=restored" in r.headers["location"] and "/festival/" in r.headers["location"]
    fests = store.festivals()
    assert len(fests) == 2 and all(len(x["members"]) == 2 for x in fests)
    client.post(f"/festival/{fid}/edit", data={"do": "split"})
    assert store.festival(fid) is None and store.get(a.id) and store.get(b.id)
    r = client.post("/festival/new", data={"title": "Один", "member": [a.id]}, follow_redirects=False)
    assert "festival_few" in r.headers["location"]
