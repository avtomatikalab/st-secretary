"""Чтение и запись предварительных заявок по бланку БЛАНК_ПРЕД_ЗАЯВКИ_ST (xlsx или xls).

Бланк: лист «Заявка»; строка с заголовками «Команда | Территория | Представитель | Контакты |
Кол-во участников» и строка значений под ней; таблица участников с заголовком «№ п/п» — колонки
A–O те же, что на листе «Заявка» книги СЕКРЕТАРЬ_ST. Строка «ОБРАЗЕЦ» и пример под ней пропускаются.

Заявка, исправленная в программе, записывается в ту же раскладку (write_preapplication) — её можно
открыть в Excel, и программа читает её так же, как файлы команд.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from st_secretary.textclean import clean_spaces

# Колонки таблицы участников (как в бланке и в СЕКРЕТАРЬ_ST): ключ → (номер колонки, ожидаемое начало заголовка)
COLUMNS = {
    "num": (0, "№"),
    "team": (1, "Команда"),
    "territory": (2, "Территория"),
    "representative": (3, "Представитель"),
    "fio": (4, "Фамилия"),
    "birth": (5, "Дата"),
    "qual": (6, "Разряд"),
    "sex": (7, "Пол"),
    "group": (8, "Группа"),
    "cls": (9, "Класс"),
    "chip": (10, "Номер чипа"),
    "personal": (11, "Участие в личной"),
    "pair": (12, "Участие в дистанции связок"),
    "pair_num": (13, "Номер связки"),
    "team_dist": (14, ""),  # «Участие в дистанции-группа» или просто «Группа»
}

BLANK_HOW = ("Программа ищет лист «Заявка» (или первый лист), строку заголовков «Команда | Территория | …» "
             "и таблицу участников с заголовком «№ п/п … Фамилия, имя» — как в бланке предзаявки.")


@dataclass
class RawRow:
    row: int  # номер строки в файле (как видит человек)
    values: dict[str, object]
    extra: dict[str, object] = field(default_factory=dict)  # свои колонки справа от бланка: заголовок → значение


@dataclass
class RawApplication:
    path: Path
    team: str = ""
    territory: str = ""
    representative: str = ""
    contacts: str = ""
    declared_count: object = None
    rows: list[RawRow] = field(default_factory=list)
    problems: list[tuple[str, str]] = field(default_factory=list)  # (что не так, почему) — файл не прочитан как бланк
    notes: list[tuple[str, str]] = field(default_factory=list)  # (что не так, почему) — прочитан, но не всё
    sheets: list[str] = field(default_factory=list)  # листы с таблицей участников, по порядку
    form: str = ""  # прочитан по своей форме (Правки, п. 37) — её название; пусто — стандартный бланк
    team_in_row: bool = False  # по форме команда — в каждой строке (заявка делегации)
    similar_form: str = ""  # не прочитан, но шапка похожа на эту свою форму — «Изменить форму по этому файлу» (п. 41)


SHEET_ROWS = 1000  # строка участника на втором листе — 1000 + номер строки, на третьем — 2000 + … (Правки, п. 29)


def row_place(row: int) -> str:
    """«строка 12» или «лист 2, строка 12» — где в файле строка участника."""
    return f"лист {row // SHEET_ROWS + 1}, строка {row % SHEET_ROWS}" if row >= SHEET_ROWS else f"строка {row}"


def _grids(path: Path) -> list[tuple[str, list[list[object]]]]:
    """Все листы книги: [(название, строки значений)] — заявка может быть на нескольких листах (по классам)."""
    if path.suffix.lower() == ".xls":
        import xlrd  # необязательная зависимость

        book = xlrd.open_workbook(str(path))
        out = []
        for sh in book.sheets():
            grid = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                    row.append(v)
                grid.append(row)
            out.append((sh.name, grid))
        return out
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    wb.close()
    return out


_CLASS_IN_NAME = re.compile(r"(\d)\s*[- ]?\s*(?:кл|класс)", re.IGNORECASE)


def read_preapplication(path: str | Path, forms: list[dict] | None = None) -> RawApplication:
    """Заявка из всех листов книги, где есть таблица участников по бланку (по листу на класс — тоже). forms —
    свои формы заявок (Правки, п. 37): файл с шапкой одной из них читается по ней."""
    path = Path(path)
    app = RawApplication(path)
    try:
        grids = _grids(path)
    except Exception as e:  # noqa: BLE001 — любой нечитаемый файл должен попасть в отчёт, а не уронить обработку
        app.problems.append((f"файл не открывается: {e}",
                             "Файл повреждён, защищён паролем или это не таблица Excel."))
        return app
    found = None
    if forms:
        from st_secretary import forms as fm

        found = fm.match(forms, grids)
        if found and found[1]:
            return fm.read_with_form(path, grids, found[0])
    tables = [(name, grid, head) for name, grid in grids if (head := _table_head(grid)) is not None]
    if not tables:
        if found:  # похоже на свою форму, но шапка изменилась
            app.similar_form = found[0].get("name", "")
            app.problems.append((f"шапка таблицы похожа на форму «{found[0].get('name', '')}», но не совпадает с ней",
                                 "В файле переименовали, добавили или убрали колонки. Откройте «Свои формы заявок» "
                                 "и добавьте форму по этому файлу (или поправьте файл)."))
            return app
        app.problems.append(("не найдена таблица участников (строка заголовков «№ п/п … Фамилия, имя»)",
                             f"Похоже, заявка заполнена не по бланку. {BLANK_HOW} Если у соревнования своя форма "
                             "заявки — покажите её программе один раз: «Свои формы заявок» на странице заявок."))
        return app
    many = len(grids) > 1
    for name, grid, head in tables:
        for c, expected in COLUMNS.values():
            got = clean_spaces(grid[head][c] if c < len(grid[head]) else None)
            if expected and not got.startswith(expected):
                app.problems.append((("лист «" + name + "», " if many else "") + f"колонка {chr(65 + c)}: ожидался "
                                     f"заголовок «{expected}…», а в файле «{got}» — бланк изменён",
                                     "Колонки бланка сдвинуты или переименованы. Если читать такой файл, данные "
                                     "попадут не в те графы — поэтому программа его не обрабатывает."))
    if app.problems:
        return app
    for name, grid in grids:  # лист с данными, но без таблицы по бланку — сказать, а не молчать
        if not any(n == name for n, _, _ in tables) and \
                sum(1 for row in grid if sum(1 for v in row if clean_spaces(v)) >= 3) >= 2:
            app.notes.append((f"лист «{name}» не прочитан: на нём нет таблицы участников по бланку",
                              "Если на этом листе участники — перенесите их в таблицу бланка (на любой лист с шапкой "
                              "«№ п/п … Фамилия, имя») или впишите в программе. Если это справочный лист — ничего "
                              "делать не нужно."))
    for idx, (name, grid, head) in enumerate(tables):
        app.sheets.append(name)
        _read_sheet(app, name, grid, head, idx * SHEET_ROWS)
    return app


def _table_head(grid: list[list[object]]) -> int | None:
    def cell(r, c):
        return grid[r][c] if r < len(grid) and c < len(grid[r]) else None

    return next((r for r in range(min(len(grid), 25)) if clean_spaces(cell(r, 0)).startswith("№")
                 and clean_spaces(cell(r, 4)).startswith("Фамилия")), None)


def _read_sheet(app: RawApplication, name: str, grid: list[list[object]], head: int, offset: int) -> None:
    def cell(r, c):
        return grid[r][c] if r < len(grid) and c < len(grid[r]) else None

    head_team = next((r for r in range(min(len(grid), 15)) if clean_spaces(cell(r, 1)) == "Команда"
                      and clean_spaces(cell(r, 2)).startswith("Территория")), None)
    if head_team is not None and head_team + 1 < head and not (app.team or app.territory or app.representative):
        r = head_team + 1  # шапка команды — с первого листа, где она заполнена
        app.team, app.territory, app.representative = (clean_spaces(cell(r, c)) for c in (1, 2, 3))
        app.contacts = clean_spaces(cell(r, 4))
        app.declared_count = cell(r, 5)
    sheet_class = m.group(1) if (m := _CLASS_IN_NAME.search(name)) else ""
    # свои колонки справа от бланка (чип, размер футболки, питание…): читаются как есть и ничего не ломают
    width = max((len(row) for row in grid[head:]), default=0)
    own = {c: clean_spaces(cell(head, c)) for c in range(len(COLUMNS), width) if clean_spaces(cell(head, c))}
    for r in range(head + 1, len(grid)):
        first = clean_spaces(cell(r, 0))
        if first.upper() == "ОБРАЗЕЦ" or first in ("0", "0.0"):
            continue  # строка «ОБРАЗЕЦ» и пример под ней
        values = {key: cell(r, c) for key, (c, _) in COLUMNS.items()}
        if not clean_spaces(values["fio"]):
            continue  # пустые строки бланка (№ проставлен заранее)
        if not clean_spaces(values["cls"]) and sheet_class:
            values["cls"] = sheet_class  # класс не вписан — из названия листа («3 КЛАСС»)
        extra = {h: cell(r, c) for c, h in own.items() if clean_spaces(cell(r, c))}
        app.rows.append(RawRow(offset + r + 1, values, extra))


# ------------------------------------------------------------------ запись

HEAD_LABELS = ["Команда", "Территория", "Фамилия, имя, отчество представителя",
               "Контактный телефон, адрес эл. почты", "Кол-во участников в делегации"]
TABLE_LABELS = ["№ п/п", "Команда", "Территория", "Представитель", "Фамилия, имя, отчество", "Дата рождения",
                "Разряд по СТ", "Пол", "Группа", "Класс дистанции", "Номер чипа (указать, если чип свой)",
                "Участие в личной дистанции", "Участие в дистанции связок", "Номер связки (если больше одной связки)",
                "Участие в дистанции-группа (номер группы)"]
TABLE_ROW = 6  # строка заголовков таблицы участников; участники — с 7-й
FIRST_ROW = TABLE_ROW + 1


def _value(v):
    """Числа — числами, даты — датами (как их ввёл бы человек в Excel), остальное — текстом."""
    if isinstance(v, (date, datetime, int)):
        return v
    s = clean_spaces(v)
    if not s:
        return None
    return int(s) if s.isdigit() and len(s) < 10 and not (len(s) > 1 and s[0] == "0") else s


def write_preapplication(path: str | Path, head: dict, rows: list[dict], note: str = "",
                         qual_labels: list[str] | None = None) -> Path:
    """Заявка в раскладке бланка. head — team, territory, representative, contacts, declared;
    rows — ключи COLUMNS без num/team/territory/representative (их программа проставляет из шапки) и «extra» —
    свои колонки заявки (заголовок → значение), они пишутся справа от бланка."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Заявка"
    ws["A1"] = "Предварительная заявка"
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(italic=True, color="5B6671")
    fill = PatternFill("solid", fgColor="DCE6F1")
    thin = Side(style="thin", color="A6A6A6")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    for c, label in enumerate(HEAD_LABELS, start=2):
        h = ws.cell(3, c, label)
        h.font, h.fill, h.border, h.alignment = Font(bold=True), fill, box, wrap
        ws.cell(4, c, _value(head.get(("team", "territory", "representative", "contacts", "declared")[c - 2])))
        ws.cell(4, c).border = box
    for c, label in enumerate(TABLE_LABELS, start=1):
        h = ws.cell(TABLE_ROW, c, label)
        h.font, h.fill, h.border, h.alignment = Font(bold=True), fill, box, wrap
    common = {"team": head.get("team", ""), "territory": head.get("territory", ""),
              "representative": head.get("representative", "")}
    own = list(dict.fromkeys(h for row in rows for h in (row.get("extra") or {})))
    for i, label in enumerate(own, start=len(TABLE_LABELS) + 1):
        h = ws.cell(TABLE_ROW, i, label)
        h.font, h.border, h.alignment = Font(bold=True), box, wrap
    for n, row in enumerate(rows, start=1):
        r = FIRST_ROW + n - 1
        values = {**row, **common, "num": n}
        for key, (c, _) in COLUMNS.items():
            cell = ws.cell(r, c + 1, _value(values.get(key)))
            cell.border = box
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "DD.MM.YYYY"
        for i, label in enumerate(own, start=len(TABLE_LABELS) + 1):
            cell = ws.cell(r, i, _value((row.get("extra") or {}).get(label)))
            cell.border = box
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "DD.MM.YYYY"
    last = max(FIRST_ROW + len(rows) + 20, 40)
    if qual_labels:
        dv = DataValidation(type="list", formula1='"' + ",".join(qual_labels) + '"', allow_blank=True)
        dv.add(f"G{FIRST_ROW}:G{last}")
        ws.add_data_validation(dv)
    sex = DataValidation(type="list", formula1='"м,ж"', allow_blank=True)
    sex.add(f"H{FIRST_ROW}:H{last}")
    ws.add_data_validation(sex)
    for c, w in zip("ABCDEFGHIJKLMNO", [6, 18, 16, 30, 34, 13, 9, 6, 10, 10, 14, 12, 12, 12, 14]):
        ws.column_dimensions[c].width = w
    ws.row_dimensions[TABLE_ROW].height = 45
    ws.freeze_panes = f"A{FIRST_ROW}"
    path = Path(path)
    wb.save(path)
    return path
