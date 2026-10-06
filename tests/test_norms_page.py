"""Страница «Разрядные нормы» и своя редакция (Правки, п. 59, решение 048). Данные — выдуманные."""

import io
import re
from dataclasses import replace
from fractions import Fraction
from urllib.parse import quote

import pytest

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_web import FormFields, base

from st_secretary.importers.norms_xlsx import read_norms, write_norms
from st_secretary.norms import achieved_norm
from st_secretary.qualification import Qual
from st_secretary.reference import (
    Level,
    all_norm_editions,
    edition_dict,
    edition_from_dict,
    norm_edition,
    norm_editions,
)
from st_secretary.web.app import create_app
from st_secretary.web.store import OWN_NORMS_COPY


@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(tmp_path / "данные", opener=lambda p: None, docs_dir=tmp_path / "документы",
                                 board_host="127.0.0.1"))


def test_norms_page_shows_the_table_the_program_reads(client):
    """Таблица на странице — та же, что в справочнике; документ ФСТР — копией в программе (без интернета)."""
    for e in norm_editions():
        n = norm_edition(e)
        page = client.get("/norms?e=" + e)
        assert page.status_code == 200 and n.title in page.text
        body = page.text.split("<tbody>")[1].split("</tbody>")[0]
        assert body.count("<tr>") == len(n.rows)
        assert [r.label for r in n.rows] == re.findall(r"<tr>\s*<td>([^<]+)</td>", body)
        assert n.source_doc and f'href="/norms/doc/{quote(n.source_doc)}"' in page.text
        doc = client.get("/norms/doc/" + quote(n.source_doc))
        assert doc.status_code == 200 and len(doc.content) > 100_000
    assert client.get("/norms/doc/rank_appl_2022_2025.pdf").headers["content-type"] == "application/pdf"
    assert 'href="https://www.tssr.ru/files/materials/2733/rank_appl_2022_2025.pdf"' in client.get(
        "/norms?e=2022-2025").text
    assert "Проект ФСТР" in client.get("/norms?e=2026-2029").text  # реквизитов приказа нет — так и сказано
    assert client.get("/norms/doc/..%2Fdata%2Fvrvs.toml").status_code == 404
    assert client.get("/norms/doc/vrvs.toml").status_code == 404
    assert 'href="/norms"' in client.get("/").text  # с главной


def test_card_links_to_its_norms(client, psr_card):
    """Просмотр карточки: редакция — ссылкой на её таблицу (с возвратом к соревнованию); в редактировании — тоже."""
    f = client.app.state.store.create(psr_card)
    view = client.get(base(f) + "/card").text
    assert re.search(r'href="/norms\?e=' + re.escape(quote(psr_card.norms_edition)) + r'&amp;c=', view)
    assert "сделать свою редакцию" in client.get(base(f) + "/card/edit").text
    page = client.get(f"/norms?e={psr_card.norms_edition}&c={quote(f.id)}").text
    assert f'href="{base(f)}/card"' in page  # назад — к карточке


def test_norms_copy_round_trip_is_exact():
    """Копия любой встроенной редакции в Excel читается обратно без потерь."""
    for e in norm_editions():
        d, errors = read_norms(write_norms(edition_dict(e), "проба"), edition_dict(e))
        assert errors == [] and d["name"] == "своя: проба"
        a = norm_edition(e)
        own = edition_from_dict(d, d["name"], own=True)
        assert (own.rows, own.class_rows, own.min_age, own.min_level, own.rank_points, own.junior_iii_text) == (
            a.rows, a.class_rows, a.min_age, a.min_level, a.rank_points, a.junior_iii_text)


def _edited(e: str, name: str, edit) -> bytes:
    wb = load_workbook(io.BytesIO(write_norms(edition_dict(e))))
    wb["Нормы"]["B2"] = name
    edit(wb["Нормы"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_own_edition_upload_choose_and_count(client, psr_card):
    """Своя редакция: правка копии → загрузка → выбор в карточке → разряд считается по ней; едет с соревнованием."""
    base_e = "2026-2029"
    n = norm_edition(base_e)
    row = next(r for r in n.rows if Qual.III in r.thresholds)
    r_idx = 7 + n.rows.index(row)

    def bump(ws):  # III и юношеские — на 50 % мягче в этой строке (младший разряд не строже старшего)
        for col in (5, 6, 7, 8):
            v = ws.cell(row=r_idx, column=col).value
            if isinstance(v, int):
                ws.cell(row=r_idx, column=col).value = v + 50

    r = client.post("/norms/upload", data={"e": base_e},
                    files={"file": ("копия.xlsx", _edited(base_e, "Кубок края", bump))}, follow_redirects=False)
    assert r.status_code == 303 and "done=norms_saved" in r.headers["location"]
    assert "своя: Кубок края" in all_norm_editions()
    own = norm_edition("своя: Кубок края")
    assert own.own and own.edition == base_e
    cls = next(c for c, (lo, hi) in own.class_rows.items() if lo <= row.min_rank <= hi)
    pct = Fraction(row.thresholds[Qual.III] + 30)
    assert achieved_norm(n, cls, row.min_rank, pct, Level.ALL_RUSSIAN).qual is None
    assert achieved_norm(own, cls, row.min_rank, pct, Level.ALL_RUSSIAN).qual == Qual.III

    store = client.app.state.store
    f = store.create(psr_card)
    assert "своя: Кубок края" in client.get(base(f) + "/card/edit").text  # в списке редакций
    data = FormFields(client.get(base(f) + "/card/edit").text, "cardform").fields | {"norms_edition": "своя: Кубок края"}
    r = client.post(base(f) + "/card/edit", data=data, follow_redirects=False)
    assert r.status_code == 303 and f.load().norms_edition == "своя: Кубок края"
    assert (f.path / OWN_NORMS_COPY).is_file()  # копия — в папке соревнования (едет с резервной копией)
    page = client.get(f"/norms?e={quote('своя: Кубок края')}").text
    assert "Своя редакция" in page and f.id in page

    # другой компьютер: тот же каталог данных, свои нормы — пусто; редакция ставится из папки соревнования
    other = TestClient(create_app(f.path.parent, opener=lambda p: None, docs_dir=f.path.parent.parent / "другой",
                                  board_host="127.0.0.1"))
    assert "своя: Кубок края" in all_norm_editions()
    assert other.get("/norms?e=" + quote("своя: Кубок края")).status_code == 200


def test_own_edition_errors_are_plain(client):
    """Ошибки в своей редакции — понятным текстом с адресом ячейки; ничего не сохраняется."""
    def spoil(ws):
        ws["C7"] = "много"  # I — не число
        ws["I9"] = None  # 1 класс: строка выпала из середины — строки класса не подряд

    r = client.post("/norms/upload", data={"e": "2026-2029"},
                    files={"file": ("копия.xlsx", _edited("2026-2029", "Ошибочная", spoil))})
    assert r.status_code == 422
    assert "Нормы!C7" in r.text and "целый процент" in r.text and "столбец «1 кл»" in r.text
    assert "своя: Ошибочная" not in all_norm_editions()
    r = client.post("/norms/upload", data={"e": "2026-2029"},
                    files={"file": ("копия.xlsx", _edited("2026-2029", "2026-2029", lambda ws: None))})
    assert r.status_code == 422 and "так называется редакция программы" in r.text
    r = client.post("/norms/upload", data={"e": "2026-2029"}, files={"file": ("x.xlsx", b"not excel")})
    assert r.status_code == 422 and "не читается" in r.text


def test_card_with_missing_own_edition_says_what_to_do(client, psr_card):
    """Карточка ссылается на свою редакцию, которой нет на этом компьютере, — ошибка с подсказкой, куда идти."""
    comp = replace(psr_card, norms_edition="своя: Чужая")
    msgs = [i.text for i in comp.check()]
    assert any("нет на этом компьютере" in m and "Разрядные нормы" in m for m in msgs)
