"""Учёт судейской практики по всем соревнованиям — на выдуманных судьях."""

from dataclasses import replace
from datetime import date

from openpyxl import load_workbook

from st_secretary import practice as pt
from st_secretary.competition import Official
from st_secretary.results import person_key


def test_records_and_summary(tmp_path, psr_card):
    c1 = psr_card  # 2025: Судьин — главный судья, СС1К
    c2 = replace(psr_card, title="Кубок города N по спортивному туризму", kind="Кубок", date_from=date(2026, 5, 9),
                 date_to=date(2026, 5, 10),
                 officials=[Official("Главный судья", "Секретарёва Анна Ивановна", "СС1К"),
                            Official("Начальник дистанции", "Судьин Иван Петрович", "ССВК")])
    brigade = [{"fio": "Этапов Семён Игоревич", "role": "Судья этапа", "category": "СС3К"},
               {"fio": "Работяга Семён Ильич", "role": "Рабочий комендантской бригады", "category": "б/к"}]
    recs = (pt.records_of(c1, "2025", {person_key("Судьин Иван Петрович"): "отлично"}, brigade)
            + pt.records_of(c2, "2026", {person_key("Судьин Иван Петрович"): "хорошо"}, []))
    people = {j.fio: j for j in pt.judges(recs)}
    assert set(people) == {"Судьин Иван Петрович", "Секретарёва Анна Ивановна", "Этапов Семён Игоревич"}  # без рабочих
    s = people["Судьин Иван Петрович"]
    assert len(s.records) == 2 and s.gsk == 2 and s.category == "ССВК" and s.last == date(2026, 5, 10)
    assert s.grades == {"отлично": 1, "хорошо": 1}
    assert s.records[0].level == "Кубок, муниципальный" and s.records[1].role == "Главный судья"
    e = people["Этапов Семён Игоревич"]
    assert (e.gsk, e.category, e.records[0].in_gsk) == (0, "СС3К", False)

    ws = load_workbook(pt.write_practice(pt.judges(recs), tmp_path / "p.xlsx"))
    rows = list(ws["Сводка"].iter_rows(min_row=2, values_only=True))
    assert ("Судьин Иван Петрович", "ССВК", 2, 2, "10.05.2026", 1, 1, 0, 0) in rows
    assert ws["Практика"].max_row == 6  # шапка + 5 записей: 3 в первом соревновании и 2 во втором
