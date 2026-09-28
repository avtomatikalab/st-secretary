"""Документы комиссии по допуску в Excel: протокол по форме Правил, ведомость взносов, документы участников.

Протокол — по образцу из приложений к разделу 3 Правил («Протокол комиссии по допуску участников»):
команды, субъект (муниципальное образование), разряды, пол, возраст, замечания и решения, подписи
председателя комиссии и главного секретаря. ФИО участников в протокол не попадают — только в лист
«Документы участников», который остаётся на ноутбуке секретариата.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from st_secretary.commission import (
    ADMITTED,
    AGE_COLUMNS,
    PERSON_LABEL,
    QUAL_COLUMNS,
    REJECTED,
    TeamCheck,
    protocol_row,
    required_docs,
)
from st_secretary.competition import Competition

THIN = Side(style="thin", color="7F7F7F")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)
HEAD = PatternFill("solid", fgColor="DCE6F1")
FILL = {ADMITTED: PatternFill("solid", fgColor="C6EFCE"), REJECTED: PatternFill("solid", fgColor="F8CBAD")}
PENDING_FILL = PatternFill("solid", fgColor="FFE699")


def _signature(comp: Competition, role: str) -> str:
    o = comp.official(role)
    return o.signature if o else " " * 30


def _title_block(ws, comp: Competition, heading: str, width: int) -> int:
    """Шапка: проводящие организации, название, заголовок, дата и место. Возвращает следующую строку."""
    r = 1
    for org in comp.organizers or [""]:
        ws.cell(r, 1, org.upper()).font = Font(bold=True, size=10)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        ws.cell(r, 1).alignment = CENTER
        r += 1
    ws.cell(r, 1, comp.title).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
    ws.cell(r, 1).alignment = CENTER
    ws.row_dimensions[r].height = 30
    r += 1
    ws.cell(r, 1, heading).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
    ws.cell(r, 1).alignment = CENTER
    r += 1
    ws.cell(r, 1, comp.dates_text)
    ws.cell(r, width, comp.place).alignment = Alignment(horizontal="right")
    return r + 2


def _box(ws, row: int, first: int, last: int, align=CENTER) -> None:
    for c in range(first, last + 1):
        ws.cell(row, c).border = BOX
        ws.cell(row, c).alignment = align


def _protocol(ws, teams: list[TeamCheck], comp: Competition) -> None:
    ws.title = "Протокол комиссии"
    nq, na = len(QUAL_COLUMNS), len(AGE_COLUMNS)
    c_total, c_q, c_sex = 5, 6, 6 + nq
    c_age = c_sex + 2
    c_rem = c_age + na
    width = c_rem + 1
    r = _title_block(ws, comp, "Протокол комиссии по допуску участников", width)
    # шапка таблицы в три строки, как в образце Правил
    h1, h2, h3 = r, r + 1, r + 2
    for c, text in ((1, "№ п/п"), (2, "Номер"), (3, "Команда"), (4, "Субъект РФ (муниципальное образование)"),
                    (c_rem, "Замечания"), (c_rem + 1, "Решения по замечаниям")):
        ws.cell(h1, c, text)
        ws.merge_cells(start_row=h1, start_column=c, end_row=h3, end_column=c)
    ws.cell(h1, c_total, "Сведения об участниках команды")
    ws.merge_cells(start_row=h1, start_column=c_total, end_row=h1, end_column=c_rem - 1)
    ws.cell(h2, c_total, "Всего")
    ws.merge_cells(start_row=h2, start_column=c_total, end_row=h3, end_column=c_total)
    for c, text, n in ((c_q, "Разряды (звания)", nq), (c_sex, "Пол", 2), (c_age, "Возраст", na)):
        ws.cell(h2, c, text)
        ws.merge_cells(start_row=h2, start_column=c, end_row=h2, end_column=c + n - 1)
    for i, (label, _) in enumerate(QUAL_COLUMNS):
        ws.cell(h3, c_q + i, label)
    ws.cell(h3, c_sex, "М")
    ws.cell(h3, c_sex + 1, "Ж")
    for i, (label, _, _) in enumerate(AGE_COLUMNS):
        ws.cell(h3, c_age + i, label)
    for row in (h1, h2, h3):
        _box(ws, row, 1, width)
        for c in range(1, width + 1):
            ws.cell(row, c).font = Font(bold=True, size=9)
            ws.cell(row, c).fill = HEAD
    r = h3 + 1
    totals = [0] * (c_rem - c_total)
    for n, t in enumerate(teams, start=1):
        row = protocol_row(t, comp)
        nums = [row["total"], *row["quals"], row["men"], row["women"], *row["ages"]]
        values = [n, row["number"], row["team"], row["territory"], *nums, row["remarks"], row["decision"]]
        for c, v in enumerate(values, start=1):
            ws.cell(r, c, v)
        _box(ws, r, 1, width)
        for c in (3, 4, c_rem, c_rem + 1):
            ws.cell(r, c).alignment = WRAP
        ws.cell(r, c_rem + 1).fill = FILL.get(t.status, PENDING_FILL)
        totals = [a + b for a, b in zip(totals, nums)]
        r += 1
    ws.cell(r, 3, "Итого:").font = Font(bold=True)
    for i, v in enumerate(totals):
        ws.cell(r, c_total + i, v).font = Font(bold=True)
    _box(ws, r, 1, width)
    r += 3
    ws.cell(r, 1, f"Председатель комиссии по допуску участников ________________ / "
                  f"{_signature(comp, 'Председатель комиссии по допуску')} /")
    ws.cell(r + 2, 1, f"Главный секретарь ________________ / {_signature(comp, 'Главный секретарь')} /")
    widths = {1: 5, 2: 7, 3: 24, 4: 18, c_rem: 38, c_rem + 1: 20}
    for c in range(1, width + 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 5.2)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{h1}:{h3}"


def _fees(ws, teams: list[TeamCheck], comp: Competition) -> None:
    head = ["№ п/п", "Номер", "Команда", "Территория", "Представитель", "Участников", "Взнос к оплате, ₽",
            "Оплачено, ₽", "Способ оплаты", "Отметка", "Подпись представителя"]
    r = _title_block(ws, comp, "Ведомость заявочных взносов", len(head))
    for c, text in enumerate(head, start=1):
        ws.cell(r, c, text).font = Font(bold=True, size=9)
        ws.cell(r, c).fill = HEAD
    _box(ws, r, 1, len(head))
    r += 1
    due = paid = 0
    for n, t in enumerate(teams, start=1):
        values = [n, t.number, t.title, t.team.territory if t.team else "", t.team.representative if t.team else "",
                  len(t.counted), t.fee_due or None, t.fee_paid or None, t.fee_method, t.fee_status, ""]
        for c, v in enumerate(values, start=1):
            ws.cell(r, c, v)
        _box(ws, r, 1, len(head), WRAP)
        ws.row_dimensions[r].height = 28
        due, paid = due + t.fee_due, paid + t.fee_paid
        r += 1
    ws.cell(r, 3, "Итого:").font = Font(bold=True)
    ws.cell(r, 7, due).font = Font(bold=True)
    ws.cell(r, 8, paid).font = Font(bold=True)
    _box(ws, r, 1, len(head))
    ws.cell(r + 3, 1, f"Главный секретарь ________________ / {_signature(comp, 'Главный секретарь')} /")
    for c, w in enumerate([5, 7, 24, 16, 30, 11, 12, 12, 13, 12, 18], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _people(ws, teams: list[TeamCheck], data: dict) -> None:
    """Для секретариата: у кого какие документы отмечены и почему не допущен (с ФИО — только на ноутбуке)."""
    pdocs, _ = required_docs(data)
    head = ["Номер", "Команда", "№", "ФИО", "Дата рождения", "Разряд", "Зачёт", *[d.short for d in pdocs],
            "Решение", "Замечание"]
    for c, text in enumerate(head, start=1):
        ws.cell(1, c, text).font = Font(bold=True, size=9)
        ws.cell(1, c).fill = HEAD
    _box(ws, 1, 1, len(head))
    r = 2
    for t in teams:
        for p in t.persons:
            e = p.entry
            birth = e.birth.strftime("%d.%m.%Y") if e.birth else (e.birth_year or "")
            values = [t.number, t.title, e.num_in_team, e.name.full, birth, e.qual.label if e.qual is not None else "",
                      e.zachet.key if e.zachet else "", *["да" if p.docs[d.key] else "—" for d in pdocs],
                      PERSON_LABEL[p.status], p.why]
            for c, v in enumerate(values, start=1):
                ws.cell(r, c, v)
            _box(ws, r, 1, len(head), WRAP)
            ws.cell(r, len(head) - 1).fill = FILL.get(p.status, PENDING_FILL)
            r += 1
    for c, w in enumerate([7, 22, 4, 32, 12, 8, 9, *[10] * len(pdocs), 11, 40], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{max(r - 1, 1)}"


def write_commission_report(teams: list[TeamCheck], comp: Competition, data: dict, path: str | Path) -> Path:
    wb = Workbook()
    _protocol(wb.active, teams, comp)
    _fees(wb.create_sheet("Ведомость взносов"), teams, comp)
    _people(wb.create_sheet("Документы участников"), teams, data)
    path = Path(path)
    wb.save(path)
    return path
