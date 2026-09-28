from datetime import date

import pytest
from openpyxl import load_workbook

from st_secretary.competition import Official, Zachet
from st_secretary.importers.card_xlsx import CardError, load_card, write_card
from st_secretary.issues import ERROR, WARNING


def test_template_has_dropdowns_and_sheets(tmp_path):
    p = write_card(tmp_path / "card.xlsx")
    wb = load_workbook(p)
    assert wb.sheetnames[:3] == ["Карточка", "ГСК", "Зачёты"]
    assert wb["Списки"].sheet_state == "hidden"
    assert len(wb["Карточка"].data_validations.dataValidation) >= 4
    with pytest.raises(CardError) as e:
        load_card(p)  # пустой шаблон — ошибки заполнения
    assert any("Наименование" in i.text for i in e.value.issues)


def test_roundtrip(tmp_path, psr_card):
    p = write_card(tmp_path / "card.xlsx", psr_card)
    assert load_card(p) == psr_card


def test_card_checks(psr_card):
    assert [i for i in psr_card.check() if i.severity in (ERROR, WARNING)] == []
    psr_card.date_to = date(2025, 9, 19)
    psr_card.officials = [Official("Главный судья", "Судьин Иван Петрович")]
    psr_card.zachety.append(Zachet("М/Ж", 3, "0840161811Я"))
    psr_card.percent_method = None
    texts = " | ".join(i.text for i in psr_card.check())
    assert "раньше даты начала" in texts
    assert "главный секретарь" in texts
    assert "указан дважды" in texts
    assert "методика" in texts
    assert "судейская категория" in texts


def test_norms_edition_year_mismatch(psr_card):
    psr_card.norms_edition = "2026-2029"
    assert any("редакция норм" in i.text for i in psr_card.check())


def test_signature_and_dates(psr_card):
    assert psr_card.official("Главный судья").signature == "И.П. Судьин, СС1К, г. Красноярск"
    assert psr_card.dates_text == "20–21 сентября 2025 г."
    assert psr_card.zachety[0].key == "М/Ж_3"
    assert psr_card.find_zachet("м/ж ", 3) is psr_card.zachety[0]
