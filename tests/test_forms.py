"""Свои формы предзаявок (Правки, п. 37): секретарь один раз показывает программе, где что в его форме. Все ФИО —
выдуманные."""

import json
from dataclasses import replace
from urllib.parse import parse_qs, urlsplit

import pytest

pytest.importorskip("fastapi")
from conftest import make_class_sheets_application
from fastapi.testclient import TestClient
from openpyxl import Workbook
from test_web import FormFields, base

from st_secretary import forms as fm
from st_secretary.competition import Zachet
from st_secretary.importers.preapp_xlsx import read_preapplication
from st_secretary.web.app import create_app

OWN_HEAD = ["№", "Фамилия", "Имя", "Отчество", "Год рождения", "Пол", "Спортивная квалификация", "Зачёт",
            "Размер футболки"]


@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(tmp_path / "данные", opener=lambda p: None, docs_dir=tmp_path / "документы",
                                 board_host="127.0.0.1"))


def own_form_file(path, team, people):
    """Заявка по «чужой» форме: команда — в шапке («Команда:» и значение справа), ФИО в трёх колонках, год
    рождения, зачёт одной колонкой, своя колонка в конце."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws["A1"] = "ЗАЯВКА на участие в Кубке города N"
    ws["A2"], ws["B2"] = "Команда:", team
    ws["A3"], ws["B3"] = "Представитель:", "Тренеров Пётр Ильич"
    ws.append([])
    ws.append(OWN_HEAD)
    for n, p in enumerate(people, start=1):
        ws.append([n, *p])
    ws.append(["", "ИТОГО: участников", len(people)])
    wb.save(path)
    return path


PEOPLE = [["Сводов", "Артём", "Игоревич", 1995, "МУЖ", "1 разряд", "М/Ж_3", "L"],
          ["Гротова", "Вера", "Олеговна", 1997, "ЖЕН", "кмс", "М/Ж_3", "S"],
          ["Колодцев", "Иван", "Петрович", 1990, "МУЖ", "б/р", "М/Ж_3", ""]]


def save_form_from_page(client, f, page_html, **changes):
    fields = FormFields(page_html.replace('class="form-edit"', 'id="fe"'), "fe").fields
    fields.update(changes, action="save")
    r = client.post(base(f) + "/forms/edit", data=fields, follow_redirects=False)
    assert r.status_code == 303, r.text.split("Форма не сохранена")[-1][:300]
    return r


def test_guess_columns_and_head():
    assert fm.guess_columns(OWN_HEAD) == ["num", "last", "first", "middle", "birth_year", "sex", "qual", "zachet",
                                          "own"]
    assert fm.guess_columns(["ФИО", "Д.р.", "Разряд", "Группа", "Класс", "Группа"])[:6] == \
        ["fio", "birth", "qual", "group", "cls", "team_dist"]
    grid = [["ЗАЯВКА"], ["Команда:", "Сталактит"], ["Представитель:", "Тренеров"], [], OWN_HEAD]
    assert fm.find_header(grid) == 4
    assert fm.guess_head(grid, 4) == {"team": "B2", "representative": "B3"}
    assert fm.cell_ref(3, 1) == "B4" and fm.parse_ref("aa12") == (11, 26)


def test_own_form_other_order_three_name_columns_year_and_own_column(client, tmp_path, psr_card):
    store = client.app.state.store
    f = store.create(psr_card)
    sample = own_form_file(tmp_path / "образец.xlsx", "Сталактит", PEOPLE)
    r = client.post(base(f) + "/forms/sample", files={"sample": ("образец.xlsx", sample.read_bytes())},
                    follow_redirects=False)
    edit = r.headers["location"]
    page = client.get(edit).text
    assert "Новая форма заявки" in page and "Так программа прочитает этот файл" in page
    assert "Сводов Артём Игоревич" in page and "Размер футболки: L" in page  # предпросмотр: ФИО из трёх колонок
    r = save_form_from_page(client, f, page, name="Кубок города N — команды")
    assert "form_saved" in r.headers["location"]
    form = store.forms()[0]
    assert form["team_in"] == "head" and form["head"]["team"] == "B2" and form["columns"][-1] == "own"
    token = parse_qs(urlsplit(edit).query)["sample"][0]
    assert not list(store.samples_dir.glob(token + "*"))  # образец удалён после сохранения

    # заявка команды по этой форме читается сама: ФИО, год, пол, разряд, зачёт, своя колонка
    other = own_form_file(tmp_path / "Сталагмит.xlsx", "Сталагмит", [
        ["Натёкова", "Лиза", "Андреевна", 1996, "ЖЕН", "II", "М/Ж_3", "M"],
        ["Пещерин", "Олег", "Петрович", 1994, "МУЖ", "III", "М/Ж_3", "XL"]])
    client.post(base(f) + "/preapps/upload", files={"files": ("Сталагмит.xlsx", other.read_bytes())})
    result, _ = store.review(f, f.load())
    t = result.teams[0]
    assert t.team == "Сталагмит" and [e.name.full for e in t.entries] == ["Натёкова Лиза Андреевна",
                                                                         "Пещерин Олег Петрович"]
    e = t.entries[0]
    assert e.birth_year == 1996 and e.sex == "ж" and e.qual.label == "II" and e.zachet.key == "М/Ж_3"
    assert e.extra == {"Размер футболки": "M"}
    assert not [i for i in result.issues if "таблица участников" in i.text]

    # шапку в файле поменяли — программа не гадает, а говорит, что форма похожа, но не та
    changed = own_form_file(tmp_path / "Свод.xlsx", "Свод", PEOPLE[:1])
    from openpyxl import load_workbook

    wb = load_workbook(changed)
    wb.active["I5"] = "Питание"
    wb.save(changed)
    app = read_preapplication(changed, store.forms())
    assert app.problems and "похожа на форму «Кубок города N — команды»" in app.problems[0][0]


def test_delegation_form_by_speleo_blank_splits_into_teams(client, tmp_path, psr_card):
    """Бланк делегации спелео (образец «Blank_predzayavki.xls»): листы по классам, команда в каждой строке."""
    store = client.app.state.store
    card = replace(psr_card, zachety=[Zachet("ЮН", 2, "0840271811Я"),
                                      Zachet("ЮН", 3, "0840271811Я")])
    f = store.create(card)
    rows = {"2 КЛАСС": [["Сталактит", "Энск", "Пещерин Олег Петрович", "Сводов Артём", "12.03.2009", "II", "м", "ЮН",
                         None, None, None, None, None, 1],
                        ["Сталагмит", "Энск", "Пещерин Олег Петрович", "Гротова Вера", "05.07.2010", "III", "ж", "ЮН",
                         None, None, None, None, None, 1]],
            "3 КЛАСС": [["Сталактит", "Энск", "Пещерин Олег Петрович", "Колодцев Иван", "21.01.2006", "I", "м", "ЮН",
                         None, None, None, None, None, 1],
                        ["Сталагмит", "Энск", "Пещерин Олег Петрович", "Натёкова Лиза", "30.09.2007", "II", "ж", "ЮН",
                         None, None, None, None, None, 1]]}
    sample = make_class_sheets_application(tmp_path / "Blank.xlsx", rows)
    r = client.post(base(f) + "/forms/sample", files={"sample": ("Blank.xlsx", sample.read_bytes())},
                    follow_redirects=False)
    page = client.get(r.headers["location"]).text
    assert 'value="row" checked' in page  # команда — в каждой строке (заявка делегации)
    assert "разделится на <b>2</b>" in page
    assert 'name="sheets" value="2 КЛАСС" checked' in page and 'name="sheets" value="3 КЛАСС" checked' in page
    save_form_from_page(client, f, page, name="Спелео — делегация", sheets=["2 КЛАСС", "3 КЛАСС"])
    assert store.forms()[0]["sheets"] == []  # отмечены все листы с такой шапкой — читать все

    r = client.post(base(f) + "/preapps/upload", files={"files": ("Делегация Энск.xlsx", sample.read_bytes())},
                    follow_redirects=False)
    assert "split=" in r.headers["location"]
    names = sorted(p.name for p in f.preapp_files())
    assert names == ["Сталагмит.xlsx", "Сталактит.xlsx"]
    assert (f.preapp_dir / "Заявки делегаций (исходные)" / "Делегация Энск.xlsx").is_file()
    result, _ = store.review(f, f.load())
    by = {t.team: t for t in result.teams}
    assert [e.name.full for e in by["Сталактит"].entries] == ["Сводов Артём", "Колодцев Иван"]
    assert [e.zachet.key for e in by["Сталактит"].entries] == ["ЮН_2", "ЮН_3"]  # класс — из названия листа
    assert by["Сталагмит"].representative == "Пещерин Олег Петрович"
    assert not [i for i in result.issues if i.severity == "error"]


def test_form_file_export_import_and_delete(client, tmp_path, psr_card):
    store = client.app.state.store
    f = store.create(psr_card)
    form = {"name": "Форма А", "headers": OWN_HEAD, "signature": fm.signature(OWN_HEAD),
            "columns": fm.guess_columns(OWN_HEAD), "team_in": "head", "head": {"team": "B2"}, "sheets": [],
            "values": {"sex": {"муж": "м"}}}
    store.save_form(form)
    body = client.get(base(f) + "/forms/file?name=Форма А").content
    assert json.loads(body)["columns"] == form["columns"]
    client.post(base(f) + "/forms/delete", data={"name": "Форма А"})
    assert store.forms() == []
    r = client.post(base(f) + "/forms/import", files={"form": ("форма.json", body)}, follow_redirects=False)
    assert "form_imported" in r.headers["location"] and store.forms()[0]["values"] == {"sex": {"муж": "м"}}
    r = client.post(base(f) + "/forms/import", files={"form": ("x.json", b"{}")}, follow_redirects=False)
    assert "form_bad_file" in r.headers["location"]
    assert "Форма А" in client.get(base(f) + "/forms").text
    # редактирование сохранённой формы без образца
    page = client.get(base(f) + "/forms/edit?name=Форма А").text
    assert "Что в каждой колонке" in page and "Размер футболки" in page


def test_offer_saved_forms_on_sample_upload(client, tmp_path, psr_card):
    """Правки, п. 41: при загрузке образца — выбрать сохранённую форму для изменения; образец совпал с сохранённой —
    «Такая форма уже есть», похож — «Похожа на …»: изменить её или сохранить как новую; заявка, похожая на форму, —
    «Исправить» открывает форму по этому файлу."""
    from openpyxl import load_workbook

    store = client.app.state.store
    f = store.create(psr_card)
    store.save_form({"name": "Форма А", "headers": OWN_HEAD, "signature": fm.signature(OWN_HEAD),
                     "columns": fm.guess_columns(OWN_HEAD), "team_in": "head", "head": {"team": "B2"}, "sheets": [],
                     "values": {"sex": {"муж": "м", "жен": "ж"}}})
    page = client.get(base(f) + "/forms").text
    assert "Изменить сохранённую" in page and '<option value="Форма А">' in page

    # та же шапка, загружен как новая — «Такая форма уже есть», «Открыть её» — образец в её настройке
    same = own_form_file(tmp_path / "образец.xlsx", "Сталактит", PEOPLE).read_bytes()
    r = client.post(base(f) + "/forms/sample", files={"sample": ("образец.xlsx", same)}, follow_redirects=False)
    page = client.get(r.headers["location"]).text
    assert "Такая форма уже есть: «Форма А»" in page and "Открыть её" in page
    link = page.split("Открыть её")[0].rsplit('href="', 1)[1].split('"')[0].replace("&amp;", "&")
    assert parse_qs(urlsplit(link).query)["name"] == ["Форма А"]
    page = client.get(link).text
    assert "Форма «Форма А»" in page and "Такая форма уже есть" not in page

    # выбрали сохранённую при загрузке — сразу её настройка
    r = client.post(base(f) + "/forms/sample", data={"name": "Форма А"},
                    files={"sample": ("образец.xlsx", same)}, follow_redirects=False)
    assert parse_qs(urlsplit(r.headers["location"]).query)["name"] == ["Форма А"]

    # шапку поменяли (своя колонка «Питание» вместо «Размер футболки») — «Похожа на …», два пути
    changed = own_form_file(tmp_path / "Свод.xlsx", "Свод", PEOPLE[:2])
    wb = load_workbook(changed)
    wb.active["I5"] = "Питание"
    wb.save(changed)
    r = client.post(base(f) + "/forms/sample", files={"sample": ("Свод.xlsx", changed.read_bytes())},
                    follow_redirects=False)
    page = client.get(r.headers["location"]).text
    assert "Похожа на форму «Форма А»" in page and "сохраните ниже как новую форму" in page

    # на странице заявок: заявка не прочиталась, но похожа на форму — «Исправить» открывает форму по этому файлу
    client.post(base(f) + "/preapps/upload", files={"files": ("Свод.xlsx", changed.read_bytes())})
    result, _ = store.review(f, f.load())
    issue = next(i for i in result.issues if "похожа на форму «Форма А»" in i.text)
    assert issue.target == "form:Свод.xlsx" and "откроет форму «Форма А»" in issue.todo
    r = client.get(base(f) + "/forms/from-preapp?file=Свод.xlsx", follow_redirects=False)
    q = parse_qs(urlsplit(r.headers["location"]).query)
    assert q["name"] == ["Форма А"] and q["file"] == ["Свод.xlsx"]
    page = client.get(r.headers["location"]).text
    assert "Форма «Форма А»" in page and "Питание" in page
    save_form_from_page(client, f, page)  # колонки формы — к новой шапке: «Питание» — своя колонка
    forms = store.forms()
    assert [x["name"] for x in forms] == ["Форма А"] and forms[0]["columns"][-1] == "own"
    assert forms[0]["signature"][-1] == "питание" and forms[0]["values"]["sex"]["жен"] == "ж"
    result, _ = store.review(f, f.load())
    assert [e.name.full for t in result.teams for e in t.entries] == ["Сводов Артём Игоревич",
                                                                    "Гротова Вера Олеговна"]
