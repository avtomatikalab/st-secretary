"""Жеребьёвка и стартовый протокол (Правила, раздел 3, п. 8.4) — на выдуманных командах."""

from dataclasses import replace
from datetime import date
from fractions import Fraction

import pytest
from openpyxl import load_workbook

from st_secretary import psr_run as pr
from st_secretary import start_list as sl
from st_secretary import time_run as tr
from st_secretary.competition import Zachet
from st_secretary.exporters.start_protocol import draw_line, write_start_protocol
from st_secretary.psr_run import Member, TeamInput
from st_secretary.qualification import Qual
from st_secretary.reference import norm_edition

NAMES = ["Кедр", "Сосна", "Ель", "Пихта", "Берёза", "Осина", "Липа", "Клён"]


def team(name, n, qual=Qual.II, size=4, admitted=True):
    return TeamInput(f"{name}.xlsx", name, "г. N", str(n), [Member(f"{name} {i}", qual, qual.label) for i in range(size)],
                     admitted=admitted, representative=f"Представитель {name}")


def teams(k=8):
    return [team(n, i) for i, n in enumerate(NAMES[:k], start=1)]


def test_random_draw_is_a_repeatable_permutation():
    ts = teams()
    a = sl.draw(ts, "random", 123456)
    assert sorted(a) == sorted(t.file for t in ts)
    assert a == sl.draw(list(reversed(ts)), "random", 123456)  # тот же жребий — тот же порядок (проверяемо)
    assert a != sl.draw(ts, "random", 654321)
    assert sl.draw(ts, "number", 1) == [t.file for t in ts]


def test_group_draw_by_rank_weak_first():
    ts = teams(6)
    ranks = {t.file: Fraction(i) for i, t in enumerate(ts)}  # Кедр слабее всех, Осина сильнее
    order = sl.draw(ts, "rank", 42, ranks, groups=2)
    assert set(order[:3]) == {"Кедр.xlsx", "Сосна.xlsx", "Ель.xlsx"}  # слабая группа — первой
    assert set(order[3:]) == {"Пихта.xlsx", "Берёза.xlsx", "Осина.xlsx"}
    rev = sl.draw(ts, "rank", 42, ranks, groups=2, strong_last=False)
    assert set(rev[:3]) == {"Пихта.xlsx", "Берёза.xlsx", "Осина.xlsx"}
    assert sl.draw(ts, "rank", 42, ranks, groups=3)[:2] in (["Кедр.xlsx", "Сосна.xlsx"], ["Сосна.xlsx", "Кедр.xlsx"])


def test_team_rank_like_start_protocol():
    norms = norm_edition("2022-2025")
    members = [Member("", Qual.KMS, "КМС"), Member("", Qual.I, "I"), Member("", Qual.II, "II"), Member("", None, "")]
    want = (norms.rank_points[Qual.KMS] + norms.rank_points[Qual.I] + norms.rank_points[Qual.II]) / 4
    assert sl.team_rank(members, "group", norms) == want
    six = members + [Member("", Qual.I, "I")] * 2
    assert sl.team_rank(six, "group", norms) == (want * 4 + 2 * norms.rank_points[Qual.I]) / 6  # больше 4 — на число
    assert sl.team_rank(members, None, norms) is None


def test_times_manual_fixes_late_teams_and_checks():
    ts = teams(4)
    z = Zachet("М/Ж", 3, "0840161811Я")
    zdata = {"draw": {"order": ["Ель.xlsx", "Кедр.xlsx", "Сосна.xlsx", "Ушедшая.xlsx"], "at": "2025-09-19T20:15",
                      "done_method": "random", "seed": 123456, "first": "10:00", "interval": "5",
                      "day": "2025-09-20", "times": {"Сосна.xlsx": "11:30", "Кедр.xlsx": "полдень"}}}
    lst = sl.build(z, zdata, ts + [team("Не допущена", 9, admitted=False)])
    assert [(r.pos, r.inp.team, sl.hm_text(r.time), r.manual_time) for r in lst.rows] == [
        (1, "Ель", "10:00", False), (2, "Кедр", "10:05", False), (3, "Сосна", "11:30", True), (4, "Пихта", "10:15", False)]
    assert lst.first_start.isoformat() == "2025-09-20T10:00:00"
    texts = [i.text for i in lst.issues]
    assert any(t.startswith("не было при жеребьёвке — поставлены в конец: Пихта") for t in texts)
    assert any("сейчас не в зачёте" in t and "Ушедшая.xlsx" in t for t in texts)
    assert any("не по форме чч:мм — не учтено: Кедр" in t for t in texts)
    assert "Сосна" not in [r.inp.team for r in lst.rows if not r.drawn]
    fp = lst.fingerprint
    zdata["draw"]["published"] = {"at": "2025-09-20T09:30", "fp": fp}
    lst = sl.build(z, zdata, ts)
    assert not lst.changed
    assert any("меньше чем за час до старта (10:00)" in i.text for i in lst.issues)
    zdata["draw"]["times"] = {}
    assert sl.build(z, zdata, ts).changed  # время Сосны теперь по расчёту — протокол изменился
    assert "число жребия 123456" in draw_line(lst) and "19.09.2025 в 20:15" in draw_line(lst)
    zdata["draw"].update(first="", interval="")
    assert all(r.time is None for r in sl.build(z, zdata, ts).rows)  # без времени — только очерёдность
    assert sl.build(z, {"draw": {"first": "25:00"}}, ts).issues[0].text.startswith("время первого старта «25:00»")
    assert sl.build(z, {}, ts, default_day=date(2025, 9, 20)).start_day == date(2025, 9, 20)


def test_same_time_warning_only_with_interval():
    ts = teams(2)
    z = Zachet("М/Ж", 3, "0840161811Я")
    zdata = {"draw": {"first": "10:00", "interval": "5", "times": {"Сосна.xlsx": "10:00"}}}
    assert any("одинаковое время старта" in i.text for i in sl.build(z, zdata, ts).issues)
    zdata["draw"]["interval"] = "0"  # одновременный старт
    assert not sl.build(z, zdata, ts).issues


def test_results_follow_start_order(psr_card):
    z = psr_card.zachety[0]
    ts = teams(3)
    stages = [{"id": "a", "tour": "Тур 1", "name": "Узлы"}]
    zdata = {"stages": stages, "tie": "start", "draw": {"order": ["Ель.xlsx", "Сосна.xlsx", "Кедр.xlsx"]},
             "teams": {f"{n}.xlsx": {"points": {"a": "10"}} for n in ("Кедр", "Сосна", "Ель")}}
    run = pr.compute(psr_card, z, zdata, ts)
    assert [(r.inp.team, r.start_order, r.place) for r in run.rows] == [("Ель", 1, 1), ("Сосна", 2, 2), ("Кедр", 3, 3)]


def test_time_discipline_takes_start_from_protocol(psr_card):
    speleo = replace(psr_card, zachety=[Zachet("М/Ж", 2, "0840271811Я", team_size=4)])
    z = speleo.zachety[0]
    ts = teams(2)
    zdata = {"stages": [{"id": "s1", "tour": "", "name": "Колодец"}],
             "draw": {"order": ["Сосна.xlsx", "Кедр.xlsx"], "first": "10:00", "interval": "10"},
             "teams": {"Кедр.xlsx": {"finish": "10:30:00"},  # старт по протоколу — 10:10
                       "Сосна.xlsx": {"start": "10:02:00", "finish": "10:25:00"}}}  # вписан — он и считается
    rows = {r.inp.team: r for r in tr.compute(speleo, z, zdata, ts).rows}
    assert rows["Кедр"].planned_start and rows["Кедр"].distance_time == 20 * 60
    assert not rows["Сосна"].planned_start and rows["Сосна"].distance_time == 23 * 60


def test_start_protocol_excel(psr_card, tmp_path):
    z = psr_card.zachety[0]
    ts = teams(3)
    ts[0].members[0].chip = "8012345"
    zdata = {"draw": {"order": [t.file for t in ts], "at": "2025-09-19T20:15", "done_method": "manual",
                      "first": "10:00", "interval": "5", "day": "2025-09-20"}}
    lst = sl.build(z, zdata, ts, {t.file: Fraction(4, 3) for t in ts})
    from datetime import datetime

    path = write_start_protocol(psr_card, lst, tmp_path / "Стартовый.xlsx", datetime(2025, 9, 20, 8, 30))
    ws = load_workbook(path).active
    cells = [c for row in ws.iter_rows(values_only=True) for c in row if c not in (None, "")]
    assert "СТАРТОВЫЙ ПРОТОКОЛ" in cells and "Время старта" in cells and "Чип" in cells and "8012345" in cells
    assert "10:05" in cells and "1,33" in cells and "Представитель Сосна" in cells
    assert any(str(c).startswith("Жеребьёвка: на совещании ГСК") for c in cells)
    assert any("Протесты по допуску — в течение 1 часа" in str(c) for c in cells)
    assert any(str(c).startswith("Главный судья ") for c in cells)


@pytest.mark.parametrize("text,sec", [("10:00", 36000), ("9.30", 34200), ("", None), ("07:05:30", 25530)])
def test_parse_hm(text, sec):
    assert sl.parse_hm(text) == sec


def test_parse_hm_bad():
    for bad in ("24:00", "10:60", "десять"):
        with pytest.raises(ValueError):
            sl.parse_hm(bad)
    assert sl.hm_text(36000) == "10:00" and sl.hm_text(25530) == "07:05:30"


def test_strict_rank_weak_before_strong_equal_ranks_by_lot():
    """Правки.md, п. 25: «строго по рангу» — Ёлки-палки (33) всегда раньше Кедра (62); равные ранги — жребием."""
    names = ["Кедр", "Ёлки-палки", "Сосна", "Горный ветер", "Пихта", "Перевал"]
    ts = [team(n, i) for i, n in enumerate(names, start=1)]
    ranks = {"Кедр.xlsx": Fraction(62), "Ёлки-палки.xlsx": Fraction(33), "Сосна.xlsx": Fraction(431, 10),
             "Горный ветер.xlsx": Fraction(90), "Пихта.xlsx": Fraction(1), "Перевал.xlsx": Fraction(1)}
    seen = set()
    for seed in range(100000, 100040):
        order = sl.draw(ts, "strict", seed, ranks)
        assert order[2:] == ["Ёлки-палки.xlsx", "Сосна.xlsx", "Кедр.xlsx", "Горный ветер.xlsx"]
        seen.add(tuple(order[:2]))
    assert len(seen) == 2  # Пихта и Перевал (ранг 1) — то так, то наоборот: жребий, а не номер
    assert sl.draw(ts, "strict", 7, ranks, strong_last=False)[:4] == [
        "Горный ветер.xlsx", "Кедр.xlsx", "Сосна.xlsx", "Ёлки-палки.xlsx"]
    lst = sl.build(Zachet("М/Ж", 3, "0840161811Я"), {"draw": {"order": order, "at": "2026-10-01T01:03",
                                                           "done_method": "strict", "seed": 812988}}, ts, ranks)
    line = draw_line(lst)
    assert "строго по рангу состава — слабые раньше сильных" in line and "число жребия 812988" in line


def test_group_boundary_by_lot_and_group_heads(tmp_path, psr_card):
    """Правки.md, п. 25: равные ранги на границе групп делит жребий, а не номер; в «Порядке старта» и протоколе —
    «Группа 1 — ранг 1–5, стартуют первыми»; все без разрядов — «без ранга»."""
    ts = teams(4)  # Кедр, Сосна, Ель, Пихта
    ranks = {"Кедр.xlsx": Fraction(1), "Сосна.xlsx": Fraction(5), "Ель.xlsx": Fraction(5), "Пихта.xlsx": Fraction(9)}
    firsts = set()
    for seed in range(1, 40):
        order, groups = sl.draw_groups(ts, "rank", seed, ranks, 2)
        g1 = {f for f, g in groups.items() if g == 1}
        assert "Кедр.xlsx" in g1 and "Пихта.xlsx" not in g1 and set(order[:2]) == g1
        firsts |= g1 - {"Кедр.xlsx"}
    assert firsts == {"Сосна.xlsx", "Ель.xlsx"}  # на границе — то одна, то другая

    z = Zachet("М/Ж", 3, "0840161811Я")
    zdata = {"draw": {"order": order, "at": "2026-10-01T01:03", "done_method": "rank", "seed": 39, "groups_of": groups}}
    lst = sl.build(z, zdata, ts, ranks)
    heads = [r.group_head for r in lst.rows]
    assert heads[0] == "Группа 1 — ранг 1–5, стартуют первыми" and heads[1] == ""
    assert heads[2] == "Группа 2 — ранг 5–9, стартуют последними" and heads[3] == ""
    from datetime import datetime

    path = write_start_protocol(psr_card, lst, tmp_path / "Стартовый.xlsx", datetime(2026, 10, 3, 8, 0))
    cells = [c for row in load_workbook(path).active.iter_rows(values_only=True) for c in row if c]
    assert "Группа 1 — ранг 1–5, стартуют первыми. Внутри группы — жребий (Правила, раздел 3, п. 8.4)" in cells

    zero = {**ranks, "Кедр.xlsx": Fraction(0)}
    assert sl.build(z, zdata, ts, zero).rows[0].group_head.startswith("Группа 1 — ранг 0–5 (есть команды без ранга)")
    assert sl.rank_word(Fraction(0)) == "без ранга" and sl.rank_word(Fraction(431, 10)) == "43,1"


def test_person_in_two_zachety_starts_too_close():
    """Правки.md, п. 25.3: один человек (ФИО + дата рождения) в двух зачётах — следующий старт не раньше, чем старт +
    расчётное время прошлой дистанции + перерыв; тёзка с другой датой рождения — другой человек."""
    z3, z2 = Zachet("М/Ж", 3, "0840161811Я"), Zachet("М/Ж", 2, "0840271811Я")
    ivan = Member("Иванов Иван", Qual.II, "II", birth="1990-01-02")
    namesake = Member("Иванов Иван", Qual.II, "II", birth="1985-05-05")
    a = TeamInput("A.xlsx", "Кедр", "г. N", "1", [ivan, Member("Петров Пётр", Qual.II, "II")])
    b = TeamInput("B.xlsx#x", "Иванов Иван", "г. N", "1.1", [ivan])
    c = TeamInput("C.xlsx", "Сосна", "г. N", "2", [namesake])
    day = {"day": "2026-10-03"}
    l3 = sl.build(z3, {"draw": {"order": ["A.xlsx"], "first": "10:20", **day}}, [a])
    l2 = sl.build(z2, {"draw": {"order": ["B.xlsx#x", "C.xlsx"], "first": "10:30", "interval": "5", **day}}, [b, c])
    out = sl.person_conflicts([(l3, None), (l2, None)], 60)
    assert [i.text for _, i in out] == ["Иванов Иван: старт в М/Ж_3 в 10:20 и в М/Ж_2 в 10:30 — меньше перерыва 60 мин"]
    assert out[0][0] == {"М/Ж_3", "М/Ж_2"} and out[0][1].target == "start:order"
    assert sl.person_conflicts([(l3, None), (l2, None)], 5) == []
    out = sl.person_conflicts([(l3, 40 * 60), (l2, None)], 5)
    assert out[0][1].text.endswith("меньше, чем расчётное время М/Ж_3 (40 мин) + перерыв 5 мин")
    same = sl.build(z2, {"draw": {"order": ["B.xlsx#x"], "first": "10:20", **day}}, [b])
    assert sl.person_conflicts([(l3, None), (same, None)], 0)[0][1].text.endswith("— одновременно")
