"""Стартовый протокол зачёта — Excel, как «Стартовый_ГРУППА» из СЕКРЕТАРЬ_ST.

Шапка (проводящие организации, наименование, даты и место, дисциплина, класс, группа), строка о жеребьёвке,
таблица: № п/п, стартовый номер, команда, территория, представитель, состав с разрядами, ранг состава, чип
(если есть), время старта; пустые колонки «Факт. старт» и «Отметка» — для судей старта (старший судья старта
фиксирует время в стартовом протоколе, раздел 3, п. 8.4). Подписи главного судьи и главного секретаря.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from st_secretary.competition import Competition
from st_secretary.results import group_words
from st_secretary.start_list import METHODS, StartList, hm_text, rank_word

THIN = Side(style="thin", color="7F7F7F")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="center", wrap_text=True)
HEAD = PatternFill("solid", fgColor="DCE6F1")
GROUP = PatternFill("solid", fgColor="F2F2F2")


def rank_text(x) -> str:
    return "" if x is None else f"{float(x):.2f}".rstrip("0").rstrip(".").replace(".", ",")


def draw_line(sl: StartList) -> str:
    """«Жеребьёвка: общая (компьютерная), 19.09.2025 в 20:15, число жребия 123456»."""
    if not sl.drawn:
        return "Порядок старта — по стартовым номерам."
    what = {"random": "общая, компьютерная (случайные числа)",
            "rank": f"групповая по рангу состава, групп: {sl.settings['groups']}, внутри группы — компьютерная",
            "strict": ("строго по рангу состава — " + ("слабые раньше сильных" if sl.settings["strong"] == "last"
                                                        else "сильные раньше слабых") + ", равные ранги — компьютерная"),
            "manual": "на совещании ГСК с представителями команд",
            "number": "не проводилась — по стартовым номерам"}.get(sl.method, METHODS.get(sl.method, ""))
    tail = f", число жребия {sl.seed}" if sl.method in ("random", "rank", "strict") and sl.seed is not None else ""
    if sl.method in ("random", "rank", "strict") and sl.settings.get("spread"):
        tail += "; старты команд одной делегации максимально разнесены (перестановки внутри групп жребия)"
    edited = f"; порядок изменён вручную {sl.edited_at:%d.%m.%Y в %H:%M}" if sl.edited_at else ""
    return f"Жеребьёвка: {what}; {sl.drawn_at:%d.%m.%Y в %H:%M}{tail}{edited}."


def write_start_protocol(comp: Competition, sl: StartList, path: str | Path, at: datetime | None = None) -> Path:
    z = sl.zachet
    wb = Workbook()
    ws = wb.active
    ws.title = "Стартовый протокол"
    chips = any(any(m.chip for m in r.inp.members) for r in sl.rows)
    person = any(r.inp.club for r in sl.rows) and all(len(r.inp.members) == 1 for r in sl.rows)
    who = (["Участник", "Команда", "Территория", "Представитель", "Разряд", "Ранг"] if person
           else ["Команда", "Территория", "Представитель", "Состав (разряд)", "Ранг состава"])
    head = ["№ п/п", "№"] + who + (["Чип"] if chips else []) + ["Время старта", "Факт. старт", "Отметка"]
    widths = [6, 6] + ([28, 20, 16, 20, 8, 7] if person else [22, 16, 20, 46, 8]) + ([9] if chips else []) + [9, 9, 12]
    width = len(head)

    def line(r, text, bold=False, size=11):
        c = ws.cell(r, 1, text)
        c.font, c.alignment = Font(bold=bold, size=size), Alignment(horizontal="center", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)

    r = 1
    for org in comp.organizers:
        line(r, org.upper(), size=10)
        r += 1
    line(r, comp.title, True, 13)
    line(r + 1, f"{comp.dates_text}, {comp.place}")
    line(r + 2, "СТАРТОВЫЙ ПРОТОКОЛ", True, 13)
    line(r + 3, z.header_text)
    day = f"; старт {sl.start_day:%d.%m.%Y}" if sl.start_day else ""
    zname = f"Зачёт «{z.name}». " if z.name else ""
    line(r + 4, f"{zname}Группа: {group_words(z.group, z.rank_format, long=True)}{day}")
    r += 6
    head_row = r
    for c, h in enumerate(head, start=1):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill, cell.border, cell.alignment = Font(bold=True, size=10), HEAD, BOX, CENTER
    r += 1
    for row in sl.rows:
        if row.group_head:  # групповая жеребьёвка: разделитель группы
            ws.cell(r, 1, f"{row.group_head}. Внутри группы — жребий (Правила, раздел 3, п. 8.4)")
            ws.cell(r, 1).font, ws.cell(r, 1).fill = Font(bold=True, size=10), GROUP
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
            r += 1
        t = row.inp
        members = ", ".join(f"{m.fio} ({m.qual_label or 'б/р'})" for m in t.members)
        chip = ", ".join(dict.fromkeys(m.chip for m in t.members if m.chip))
        who_values = ([t.team, t.club, t.territory, t.representative,
                       ", ".join(m.qual_label or "б/р" for m in t.members), rank_word(row.rank)] if person
                      else [t.team, t.territory, t.representative, members, rank_word(row.rank)])
        values = [row.pos, t.number or ""] + who_values + ([chip] if chips else []) + [hm_text(row.time), "", ""]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.border = BOX
            cell.font = Font(size=10, bold=head[c - 1] in ("№ п/п", "Время старта"))
            cell.alignment = WRAP if head[c - 1] in ("Участник", "Команда", "Территория", "Представитель",
                                                     "Состав (разряд)") else CENTER
        r += 1
    r += 1
    notes = [draw_line(sl)]
    if at:
        notes.append(f"Опубликован {at:%d.%m.%Y в %H:%M}. Протесты по допуску — в течение 1 часа после публикации "
                     "(Правила, раздел 3, п. 8.17).")
    for n in notes:
        ws.cell(r, 1, n).font = Font(size=10, italic=True)
        r += 1
    r += 1
    for role in ("Главный судья", "Главный секретарь"):
        o = comp.official(role)
        ws.cell(r, 1, f"{role} ________________ / {o.signature if o else ' ' * 30} /").font = Font(size=11)
        r += 2
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{head_row}:{head_row}"  # шапка таблицы — на каждой странице
    path = Path(path)
    wb.save(path)
    return path
