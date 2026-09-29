"""Документы награждения: тексты дипломов, наклейки на медали, список награждаемых.

Дипломы печатаются на готовых бланках: в файле — только текст, по странице на диплом. У группы каждый
участник получает свой диплом, где первым стоит его имя (остальные — по кругу), имена — «Имя Фамилия»,
как в наградных 2024–2025 гг. Даты, название и дисциплина — из карточки соревнования, а не из прошлогодней
заготовки.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from st_secretary.competition import Competition
from st_secretary.results import ROMAN, ZachetResults, date_text, group_label, rotations, title_in

FONT = "Times New Roman"


def _para(doc, text: str, size: int, bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after: int = 6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name, run.font.size, run.bold = FONT, Pt(size), bold
    return p


def _new_doc(landscape: bool = False, top_cm: float = 2.0) -> Document:
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    if landscape:
        s.orientation, s.page_width, s.page_height = WD_ORIENT.LANDSCAPE, Cm(29.7), Cm(21)
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin, s.bottom_margin = Cm(top_cm), Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(12)
    return doc


def _zachet(comp: Competition, key: str):
    return next(z for z in comp.zachety if z.key == key)


def discipline_line(comp: Competition, z: ZachetResults) -> str:
    zz = _zachet(comp, z.key)
    return f"в спортивной дисциплине «{zz.discipline_name}» {zz.distance_class} класса, {group_label(z, zz.group)}"


def write_diplomas(results: list[ZachetResults], comp: Competition, path: str | Path, top_cm: float = 9.0) -> Path:
    """Тексты дипломов за I–III места — по странице на диплом (отступ сверху — под бланк)."""
    doc = _new_doc(top_cm=top_cm)
    first = True
    for z in results:
        for p in z.medalists:
            names = [m.first_last for m in p.members] or [p.team]
            for order in rotations(names):
                if not first:
                    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
                first = False
                _para(doc, "Награждается" if len(order) == 1 else "Награждаются", 20, True)
                _para(doc, ", ".join(order), 18, True)
                if len(p.members) > 1 or not p.members:
                    _para(doc, f"команда «{p.team}»", 16)
                _para(doc, f"за {ROMAN[p.place]} место", 20, True, space_after=12)
                _para(doc, f"на {title_in(comp)}", 14)
                _para(doc, discipline_line(comp, z), 14, space_after=24)
                _para(doc, f"{comp.place}, {date_text(comp.date_to)}", 12)
    if first:
        _para(doc, "Призёров пока нет — внесите результаты.", 14)
    path = Path(path)
    doc.save(path)
    return path


def medal_count(results: list[ZachetResults]) -> int:
    return sum(max(len(p.members), 1) for z in results for p in z.medalists)


def write_stickers(results: list[ZachetResults], comp: Competition, path: str | Path, spare: int = 2) -> Path:
    """Наклейки на медали: по одной на каждого призёра и немного запасных; две колонки — печать и вырезать."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Наклейки"
    text = f"{comp.title.upper()}\n{comp.place}\n{comp.dates_text}"
    dotted = Side(style="dotted", color="A6A6A6")
    n = medal_count(results) + spare
    for i in range(n):
        r, c = i // 2 + 1, i % 2 + 1
        cell = ws.cell(r, c, text)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.font = Font(name="Calibri", size=10, bold=True)
        cell.border = Border(left=dotted, right=dotted, top=dotted, bottom=dotted)
        ws.row_dimensions[r].height = 90
    ws.column_dimensions["A"].width = ws.column_dimensions["B"].width = 31
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_options.horizontalCentered = True
    path = Path(path)
    wb.save(path)
    return path


def write_awardees(results: list[ZachetResults], comp: Competition, path: str | Path) -> Path:
    """Список награждаемых для ведущего: по зачётам, в порядке вызова — III, II, I место."""
    doc = _new_doc()
    _para(doc, "СПИСОК НАГРАЖДАЕМЫХ", 16, True)
    _para(doc, comp.title, 13)
    _para(doc, f"{comp.dates_text}, {comp.place}", 12, space_after=14)
    for z in results:
        zz = _zachet(comp, z.key)
        _para(doc, f"«{zz.discipline_name}» {zz.distance_class} класса, {group_label(z, zz.group)}", 13, True,
              WD_ALIGN_PARAGRAPH.LEFT)
        for p in sorted(z.medalists, key=lambda p: -p.place):
            names = ", ".join(m.first_last for m in p.members)
            team = f"команда «{p.team}», {p.territory}" if p.territory else f"команда «{p.team}»"
            _para(doc, f"{ROMAN[p.place]} место — {team}: {names}", 12, align=WD_ALIGN_PARAGRAPH.LEFT)
        if not z.medalists:
            _para(doc, "призёров нет", 12, align=WD_ALIGN_PARAGRAPH.LEFT)
    path = Path(path)
    doc.save(path)
    return path
