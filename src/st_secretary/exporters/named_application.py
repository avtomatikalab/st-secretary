"""Именная заявка команды (Word) по предзаявке — Правки, п. 36.

Команда печатает её, идёт к врачу и приносит на комиссию по допуску: ФИО, даты рождения и разряды уже совпадают с
предзаявкой — сверять проще. Бланк — как именные заявки ЧК и ПК Красноярского края и Кубка г. Красноярска (Правила,
раздел 3, приложение «заявка»): «В Главную судейскую коллегию …», «ЗАЯВКА», таблица (№, ФИО, дата рождения,
квалификация, медицинский допуск — слово «допущен», подпись и печать врача, подпись участника, примечания), «всего
допущено / не допущено», врач, тренер-представитель, «С правилами техники безопасности ознакомлен», руководитель
командирующей организации, М.П., приложения. Пустые поля — для подписей и печатей.

Свой шаблон соревнования (.docx) — с метками в тексте: {соревнование}, {даты}, {место}, {команда}, {территория},
{представитель}, {телефон}, {всего}; строка таблицы с меткой {ФИО} повторяется на каждого участника (в ней же
{№}, {дата рождения}, {разряд}, {зачёт}).
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from st_secretary.competition import Competition
from st_secretary.preapp import TeamApplication

HEAD = ["№ п/п", "Фамилия, имя, отчество участника", "Дата рождения", "Спортивная квалификация",
        "Медицинский допуск: слово «допущен», подпись и печать врача напротив каждого участника", "Подпись участника",
        "Примечания"]
CONSENT = ("Ставя подпись, участник спортивных соревнований даёт своё согласие на обработку персональных данных (сбор, "
           "систематизацию, накопление, хранение, уточнение, использование, распространение, обезличивание), а также "
           "иных действий, необходимых для обработки персональных данных в рамках проведения официальных спортивных "
           "соревнований в соответствии с ФЗ № 152-ФЗ от 27.07.2006 г.")
APPENDIX = ["Документы о возрасте", "Документы о квалификации", "Медицинский допуск"]


def birth_text(e) -> str:
    return e.birth.strftime("%d.%m.%Y") if e.birth else str(e.birth_year or "")


def person_values(n: int, e) -> dict[str, str]:
    """Метки строки участника — для своего шаблона и для стандартного бланка."""
    return {"№": str(n), "ФИО": e.name.full, "дата рождения": birth_text(e),
            "разряд": e.qual.label if e.qual is not None else "", "зачёт": e.zachet.title if e.zachet else ""}


def team_values(comp: Competition, team: TeamApplication) -> dict[str, str]:
    return {"соревнование": comp.title, "даты": comp.dates_text, "место": comp.place, "команда": team.team,
            "территория": team.territory, "представитель": team.representative,
            "телефон": ", ".join(x for x in (team.phone, team.email) if x), "всего": str(len(team.entries))}


def _para(doc, text: str = "", bold: bool = False, size: float = 11, align=None, space: float = 2):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold, r.font.size = bold, Pt(size)
    p.paragraph_format.space_after = Pt(space)
    if align is not None:
        p.alignment = align
    return p


def _caption(doc, text: str) -> None:
    p = _para(doc, text, size=8, space=6)
    p.runs[0].italic = True


def write_standard(comp: Competition, team: TeamApplication, path: str | Path,
                   appendix: list[str] | None = None) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    _para(doc, f"В Главную судейскую коллегию {comp.title}, {comp.dates_text}, {comp.place}", size=11,
          align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(doc, "от " + "_" * 60, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _caption(doc, "(название командирующей организации, адрес, телефон, e-mail)")
    _para(doc, "ЗАЯВКА", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space=6)
    where = f" ({team.territory})" if team.territory else ""
    _para(doc, f"Просим допустить к участию в соревнованиях команду «{team.team}»{where} в следующем составе:",
          space=6)
    table = doc.add_table(rows=1, cols=len(HEAD))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, HEAD, strict=True):
        cell.text = ""
        r = cell.paragraphs[0].add_run(text)
        r.bold, r.font.size = True, Pt(9)
    for n, e in enumerate(team.entries, start=1):
        v = person_values(n, e)
        row = table.add_row().cells
        for cell, text in zip(row, [v["№"], v["ФИО"], v["дата рождения"], v["разряд"], "", "", v["зачёт"]],
                              strict=True):
            cell.text = ""
            cell.paragraphs[0].add_run(text).font.size = Pt(10)
    for col, w in zip(table.columns, (1.0, 5.0, 2.4, 2.2, 4.0, 2.2, 2.2), strict=True):
        for cell in col.cells:
            cell.width = Cm(w)
    _para(doc, "", space=4)
    _para(doc, f"Всего допущено к соревнованиям ______ человек (из {len(team.entries)}). Не допущено ______ человек, "
               "в том числе " + "_" * 40)
    _caption(doc, "ФИО не допущенных")
    _para(doc, "М.П.            Врач ____________________ / ______________________ /")
    _caption(doc, "Печать медицинского учреждения            подпись врача            расшифровка подписи врача")
    rep = team.representative or "_" * 40
    contacts = ", ".join(x for x in (team.phone, team.email) if x)
    _para(doc, f"Тренер-представитель команды: {rep}" + (f", {contacts}" if contacts else ""))
    _para(doc, CONSENT, size=8, space=6)
    _para(doc, "«С правилами техники безопасности ознакомлен» ____________________ / "
               f"{team.representative or '_' * 24} /")
    _caption(doc, "подпись представителя            расшифровка подписи")
    _para(doc, "Руководитель командирующей организации ____________________ / ______________________ /")
    _caption(doc, "подпись            расшифровка подписи")
    _para(doc, "М.П.", space=6)
    items = appendix or APPENDIX
    _para(doc, "Приложения: " + ", ".join(f"{i}. {x}" for i, x in enumerate(items, start=1)) + ".", size=9)
    path = Path(path)
    doc.save(path)
    return path


_MARK = re.compile(r"\{([^{}]+)\}")


def _fill(paragraph, values: dict[str, str]) -> None:
    """Метки {…} в абзаце — значениями; Word дробит текст на куски, поэтому собираем абзац в первый кусок."""
    text = "".join(r.text for r in paragraph.runs)
    if "{" not in text:
        return
    new = _MARK.sub(lambda m: values.get(m.group(1).strip().lower(), m.group(0)), text)
    if new != text and paragraph.runs:
        paragraph.runs[0].text = new
        for r in paragraph.runs[1:]:
            r.text = ""


def _fill_cells(cells, values: dict[str, str]) -> None:
    for cell in cells:
        for p in cell.paragraphs:
            _fill(p, values)


def write_from_template(template: str | Path, comp: Competition, team: TeamApplication, path: str | Path) -> Path:
    """Свой шаблон соревнования: метки заменяются, строка с {ФИО} повторяется на каждого участника."""
    doc = Document(str(template))
    common = {k.lower(): v for k, v in team_values(comp, team).items()}
    for p in doc.paragraphs:
        _fill(p, common)
    for table in doc.tables:
        rows = list(table.rows)
        proto = next((r for r in rows if any("{фио}" in c.text.lower().replace(" ", "") for c in r.cells)), None)
        for r in rows:
            if r is not proto:
                _fill_cells(r.cells, common)
        if proto is None:
            continue
        for n, e in enumerate(team.entries, start=1):
            tr = copy.deepcopy(proto._tr)
            proto._tr.addprevious(tr)
            new_row = next(r for r in table.rows if r._tr is tr)
            _fill_cells(new_row.cells, {**common, **{k.lower(): v for k, v in person_values(n, e).items()}})
        proto._tr.getparent().remove(proto._tr)
    for sec in doc.sections:
        for p in (*sec.header.paragraphs, *sec.footer.paragraphs):
            _fill(p, common)
    path = Path(path)
    doc.save(path)
    return path
