"""Инструкция секретаря (Word и PDF) со снимками экрана — на учебном соревновании с выдуманными данными.

    uv run --with websocket-client python tools/make_manual.py   → docs/Инструкция секретаря.docx (+ .pdf, если есть Word)

Скрипт создаёт учебное соревнование во временной папке, проходит его шаг за шагом (заявки, комиссия,
жеребьёвка, баллы, протоколы, награждение) и снимает страницы программы браузером Chrome или Edge без окна.
Снимки всегда совпадают с тем, что видит секретарь: инструкцию пересобирают после изменений интерфейса.
Реальных данных в инструкции нет.
"""

from __future__ import annotations

import base64
import json
import os
import random
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import urllib.request
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote, urlencode

import uvicorn
import websocket  # websocket-client
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from fastapi.testclient import TestClient

from st_secretary import __version__, training
from st_secretary import commission as cm
from st_secretary import judge_sync as js
from st_secretary.web.app import create_app

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Инструкция секретаря.docx"
BROWSERS = [
    Path(os.environ.get("PROGRAMFILES", r"C:\Program Files")) / "Google/Chrome/Application/chrome.exe",
    Path(os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)")) / "Google/Chrome/Application/chrome.exe",
    Path(os.environ.get("LOCALAPPDATA", "")) / "Google/Chrome/Application/chrome.exe",
    Path(os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)")) / "Microsoft/Edge/Application/msedge.exe",
]
GREEN = RGBColor(0x1B, 0x6B, 0x57)
GRAY = RGBColor(0x55, 0x5F, 0x69)


# ------------------------------------------------------------------ снимки экрана


class Shooter:
    """Программа на свободном порту и браузер без окна (Chrome или Edge), управляемый через DevTools: открывает
    страницу, находит нужный блок и снимает ровно его — без меню слева и верхней полосы (крупнее в документе)."""

    def __init__(self, app, out: Path):
        self.out = out
        self.shots: dict[str, Path] = {}
        s = socket.socket()
        s.bind(("127.0.0.1", 0))
        self.port = s.getsockname()[1]
        s.close()
        self.server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=self.port, log_level="warning"))
        threading.Thread(target=self.server.run, daemon=True).start()
        while not self.server.started:
            time.sleep(0.05)
        self.profile = tempfile.mkdtemp(prefix="st-manual-browser-")
        exe = next((b for b in BROWSERS if b.is_file()), None)
        if exe is None:
            raise SystemExit("Для снимков экрана нужен Google Chrome или Microsoft Edge")
        self.proc = subprocess.Popen([str(exe), "--headless=new", "--disable-gpu", "--hide-scrollbars",
                                      "--no-first-run", "--remote-debugging-port=0", "--remote-allow-origins=*",
                                      f"--user-data-dir={self.profile}", "about:blank"],
                                     stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        port_file = Path(self.profile) / "DevToolsActivePort"
        for _ in range(200):
            if port_file.is_file() and port_file.read_text().strip():
                break
            time.sleep(0.05)
        dt = port_file.read_text().split()[0]
        targets = json.load(urllib.request.urlopen(f"http://127.0.0.1:{dt}/json/list", timeout=10))
        page = next(t for t in targets if t.get("type") == "page")
        self.ws = websocket.create_connection(page["webSocketDebuggerUrl"], timeout=60, suppress_origin=True)
        self.n = 0
        self.call("Page.enable")

    def call(self, method: str, **params):
        self.n += 1
        self.ws.send(json.dumps({"id": self.n, "method": method, "params": params}))
        while True:
            msg = json.loads(self.ws.recv())
            if msg.get("id") == self.n:
                if "error" in msg:
                    raise RuntimeError(f"{method}: {msg['error']}")
                return msg.get("result", {})

    def open(self, path: str, width: int, height: int):
        self.call("Emulation.setDeviceMetricsOverride", width=width, height=height, deviceScaleFactor=1,
                  mobile=width < 600)
        self.call("Emulation.setEmulatedMedia", features=[{"name": "prefers-color-scheme", "value": "light"}])
        self.call("Page.navigate", url=f"http://127.0.0.1:{self.port}{path}")
        for _ in range(100):  # страница загрузилась и ничего не догружает
            state = self.call("Runtime.evaluate", expression="document.readyState", returnByValue=True)
            if state["result"].get("value") == "complete":
                break
            time.sleep(0.05)
        time.sleep(0.2)

    def shot(self, name: str, path: str, width: int = 1366, height: int = 900, at: str = "",
             crop: str = "content", prep: str = "", then: str = "") -> Path:
        """at — CSS-селектор блока, с которого начинается снимок (иначе — верх страницы); height — высота снимка.
        crop: «content» — без меню слева и верхней полосы, «wide» — середина (главная), «full» — как есть.
        prep — JS до перезагрузки страницы (например, запись в память телефона), then — JS после загрузки."""
        if prep:
            self.open(path, width, height)
            self.call("Runtime.evaluate", expression=prep)
        self.open(path, width, height)
        if then:
            self.call("Runtime.evaluate", expression=then)
            time.sleep(0.3)
        top = 60 if crop != "full" and width >= 600 else 0  # у телефона своей верхней полосы нет
        if at:
            js_ = f"(() => {{ const e = document.querySelector({json.dumps(at)}); if (!e) return null; " \
                  "const r = e.getBoundingClientRect(); return [r.top + window.scrollY, r.height]; })()"
            found = self.call("Runtime.evaluate", expression=js_, returnByValue=True)["result"].get("value")
            if found is None:
                raise SystemExit(f"На странице {path} нет блока {at}")
            top = max(0, int(found[0]) - 12)
            height = min(height, int(found[1]) + 24)  # блок короче — без пустого места под ним
        doc_h = self.call("Runtime.evaluate", expression="document.documentElement.scrollHeight",
                          returnByValue=True)["result"].get("value") or height
        height = max(100, min(height, int(doc_h) - top))  # страница короче — без пустого места внизу
        x, w = {"content": (300, width - 300), "wide": (150, width - 300)}.get(crop, (0, width))
        if width < 600:
            x, w = 0, width
        shot = self.call("Page.captureScreenshot", format="png", captureBeyondViewport=True,
                         clip={"x": x, "y": top, "width": w, "height": height, "scale": 1})
        png = self.out / f"{name}.png"
        png.write_bytes(base64.b64decode(shot["data"]))
        self.shots[name] = png
        print(f"  снимок {name}")
        return png

    def close(self):
        try:
            self.ws.close()
        finally:
            self.proc.terminate()
            self.proc.wait(timeout=10)
            self.server.should_exit = True
            shutil.rmtree(self.profile, ignore_errors=True)


def walk(work: Path) -> tuple[dict[str, Path], dict]:
    """Учебное соревнование от заявок до награждения; снимок на каждом шаге."""
    clock = {"now": datetime(2026, 9, 28, 19, 30)}
    app = create_app(work / "данные", opener=lambda p: None, docs_dir=work / "документы", board_host="127.0.0.1")
    app.state.clock = lambda: clock["now"]
    c = TestClient(app)
    c.post("/training")
    f = app.state.store.all()[0]
    base = "/c/" + quote(f.id, safe="")
    z3, z2 = "?" + urlencode({"z": "М/Ж_3"}), "?" + urlencode({"z": "М/Ж_2"})
    (work / "снимки").mkdir(exist_ok=True)
    sh = Shooter(app, work / "снимки")
    info = {"title": f.load().title}
    try:
        sh.shot("home", "/", height=790, crop="wide")
        sh.shot("overview", base, height=820, crop="full")
        sh.shot("card", base + "/card", height=900)
        sh.shot("preapps", base + "/preapps", height=1000)
        sh.shot("team_error", base + "/preapps/team?" + urlencode({"file": "Бурундуки.xlsx"}), height=900)
        sh.shot("team_edit", base + "/preapps/edit?" + urlencode({"file": "Бурундуки.xlsx"}), height=760)

        # ошибки в заявках исправлены (как сделал бы секретарь в форме заявки)
        comp = f.load()
        for head, rows in training.applications(comp):
            if head["team"] == "Бурундуки":
                rows[1]["birth"] = date(1998, 2, 13)
                f.save_preapp("Бурундуки.xlsx", head, rows, None)
        # заявки просмотрены — «Проверено» (заявки с ошибками так не отметить — они остаются на комиссии)
        for p in f.preapp_files():
            c.post(base + "/preapps/status", data={"file": p.name, "status": "done"})
        # комиссия: у всех документы на месте, номера по порядку, взносы
        data = f.admission()
        data["settings"] = {**cm.settings(data), "start_at": "2026-10-03T10:00"}
        f.save_admission(data)
        pdocs, tdocs = cm.required_docs(data)
        result, _ = app.state.store.review(f, comp)
        for n, t in enumerate(result.teams, start=1):
            tm = data.setdefault("teams", {}).setdefault(t.source, {})
            tm.update(number=n, fee_method="наличные", team_docs={d.key: True for d in tdocs})
            tm["fee_paid"] = 3000 if any(e.zachet and e.zachet.distance_class == 3 for e in t.entries) else 2000
            tm["people"] = {cm.person_key(e.name.full): {"docs": {d.key: True for d in pdocs}} for e in t.entries}
        data["teams"]["Лиственница.xlsx"]["fee_paid"] = 0  # взнос ещё не сдан
        f.save_admission(data)
        sh.shot("admission", base + "/admission", height=1000)

        clock["now"] = datetime(2026, 10, 2, 20, 15)
        c.post(base + "/start/draw" + z3, data={"day": "2026-10-03", "first": "10:00", "interval": "10",
                                                "action": "save"})
        c.post(base + "/start/draw" + z3, data={"method": "random", "groups": "2", "strong": "last",
                                                "action": "draw"})
        sh.shot("start_draw", base + "/start" + z3, height=900)
        sh.shot("start_order", base + "/start" + z3, at="#order", height=900)
        clock["now"] = datetime(2026, 10, 3, 8, 30)
        c.post(base + "/start/publish" + z3)
        c.post(base + "/start/draw" + z2, data={"day": "2026-10-04", "first": "09:00", "interval": "15",
                                                "action": "save"})
        c.post(base + "/start/draw" + z2, data={"method": "random", "action": "draw"})
        c.post(base + "/start/publish" + z2)
        sh.shot("start_publish", base + "/start" + z3, at="#publish", height=520)

        # баллы ПСР — как из протоколов судей этапов
        rng = random.Random(7)
        run = f.run_data()["zachety"]["М/Ж_3"]
        files = run["draw"]["order"]
        form = {}
        for i, file in enumerate(files):
            form[f"p-{i}-file"] = file
            form[f"p-{i}-status"] = "finished"
            for s in run["stages"]:
                form[f"p-{i}-{s['id']}"] = str(-rng.choice([0, 20, 40]) if s["tour"] == "Бонус"
                                               else rng.choice([0, 0, 5, 10, 20, 30, 45, 60, 90]))
        form[f"p-{len(files) - 1}-status"] = "removed"
        clock["now"] = datetime(2026, 10, 3, 15, 40)
        c.post(base + "/results/points" + z3, data=form)
        sh.shot("results_points", base + "/results" + z3, at="#points", height=820)
        sh.shot("results_table", base + "/results" + z3, at="#results-live", height=820)
        c.post(base + "/results/publish" + z3)
        clock["now"] = datetime(2026, 10, 3, 16, 5)
        c.post(base + "/results/protest" + z3, data={"at": "2026-10-03T16:05", "team": "Кедр",
                                                    "text": "Навесная переправа: 30 баллов — не согласны"})
        c.post(base + "/results/protest/decide" + z3, data={"id": "p1", "decision": "отклонён",
                                                           "note": "видеозапись судьи этапа"})
        clock["now"] = datetime(2026, 10, 3, 16, 50)
        sh.shot("results_protocol", base + "/results" + z3, at="#protocol", height=900)
        c.post(base + "/results/approve" + z3)

        # спелео: время старта и финиша, отсечки, баллы этапов
        run2 = f.run_data()["zachety"]["М/Ж_2"]
        form = {}
        for i, file in enumerate(run2["draw"]["order"]):
            start = 9 * 3600 + i * 15 * 60
            fin = start + rng.randint(38, 70) * 60 + rng.randint(0, 59)
            form.update({f"p-{i}-file": file, f"p-{i}-status": "finished",
                         f"p-{i}-start": f"{start // 3600:02d}:{start % 3600 // 60:02d}:00",
                         f"p-{i}-finish": f"{fin // 3600:02d}:{fin % 3600 // 60:02d}:{fin % 60:02d}",
                         f"p-{i}-cutoffs": f"{rng.randint(2, 6)}:{rng.choice(['00', '30'])}"})
            for s in run2["stages"]:
                form[f"p-{i}-{s['id']}"] = rng.choice(["", "", "0,3", "1", "2"])
        form["p-2-s2"] = "с"
        clock["now"] = datetime(2026, 10, 4, 12, 10)
        c.post(base + "/results/points" + z2, data=form)
        sh.shot("speleo_points", base + "/results" + z2, at="#points", height=760)

        c.post(base + "/judges/link" + z3, data={"stage": "*"})
        sh.shot("judges", base + "/judges" + z3, height=950)
        token = js.stage_token(f.run_data(), "М/Ж_3", "s3")
        judge = json.dumps({"fio": comp.officials[2].fio if len(comp.officials) > 2 else "Судья этапа",
                            "phone": "+7 900 000-00-00"}, ensure_ascii=False)  # выдуманный судья и номер
        me = f"localStorage.setItem('st-judge-me', {json.dumps(judge)})"
        sh.shot("judge_phone", base + f"/judges/phone/{token}", width=390, height=844, prep=me)
        c.post(base + "/judges/link" + z2, data={"stage": "*"})
        token2 = js.stage_token(f.run_data(), "М/Ж_2", "s1")
        sh.shot("judge_penalties", base + f"/judges/phone/{token2}", width=390, height=844, prep=me,
                then="document.getElementById('pen-open').click(); const q = document.getElementById('pen-q');"
                     " q.value = 'муфт'; q.dispatchEvent(new Event('input'));")

        c.post(base + "/board/toggle", data={"on": "1"})
        sh.shot("board", base + "/board/preview", width=390, height=844)
        sh.shot("board_admin", base + "/board", height=760)

        sh.shot("awards", base + "/awards", height=1000)
        c.post(base + "/contracts/settings", data={"do": "sample"})  # ставки по образцу
        sh.shot("contracts", base + "/contracts", height=900)
        sh.shot("verify", base + "/verify", height=900)
        c.post(base + "/backup")
        sh.shot("backup", base, at="#backup", height=500)
    finally:
        sh.close()
    return sh.shots, info


# ------------------------------------------------------------------ документ


def _png_size(path: Path) -> tuple[int, int]:
    head = path.read_bytes()[16:24]
    return int.from_bytes(head[:4], "big"), int.from_bytes(head[4:], "big")


def _shade(p, color: str):
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    ppr.append(shd)


class Manual:
    def __init__(self, shots: dict[str, Path]):
        self.doc = Document()
        self.shots = shots
        self.fig = 0
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(13)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.1
        for name, size in (("Heading 1", 20), ("Heading 2", 15)):
            h = self.doc.styles[name]
            h.font.name, h.font.size, h.font.bold, h.font.color.rgb = "Calibri", Pt(size), True, GREEN
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(14), Pt(6)
            h.paragraph_format.keep_with_next = True

    def h1(self, text: str, new_page: bool = True):
        if new_page and len(self.doc.paragraphs) > 3:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.doc.add_heading(text, level=1)

    def h2(self, text: str):
        self.doc.add_heading(text, level=2)

    def p(self, text: str, bold_lead: str = ""):
        par = self.doc.add_paragraph()
        if bold_lead:
            par.add_run(bold_lead).bold = True
        par.add_run(text)
        return par

    def steps(self, items: list[str]):
        for i, t in enumerate(items, start=1):
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent, par.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.9)
            r = par.add_run(f"{i}.\t")
            r.bold = True
            r.font.color.rgb = GREEN
            par.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            par.add_run(t)

    def bullets(self, items: list[str]):
        for t in items:
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent, par.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.5)
            par.add_run("•  " + t)

    def tip(self, text: str, title: str = "Совет. "):
        par = self.doc.add_paragraph()
        _shade(par, "E6F2EE")
        par.paragraph_format.left_indent = par.paragraph_format.right_indent = Cm(0.2)
        r = par.add_run(title)
        r.bold = True
        r.font.color.rgb = GREEN
        par.add_run(text)

    def warn(self, text: str):
        par = self.doc.add_paragraph()
        _shade(par, "FFF4D4")
        par.paragraph_format.left_indent = par.paragraph_format.right_indent = Cm(0.2)
        r = par.add_run("Важно. ")
        r.bold = True
        par.add_run(text)

    def img(self, name: str, caption: str, width_cm: float = 17, max_h_cm: float = 12.5):
        w, h = _png_size(self.shots[name])
        width_cm = min(width_cm, max_h_cm * w / h)  # высокий снимок — уже, чтобы не занимал всю страницу
        self.fig += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(self.shots[name]), width=Cm(width_cm))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Рис. {self.fig}. {caption}")
        r.italic, r.font.size, r.font.color.rgb = True, Pt(11), GRAY

    def phone_pair(self, left: str, right: str, caption: str):
        """Два снимка телефона рядом."""
        self.fig += 1
        t = self.doc.add_table(rows=1, cols=2)
        for cell, name in zip(t.rows[0].cells, (left, right)):
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(str(self.shots[name]), width=Cm(6.2))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Рис. {self.fig}. {caption}")
        r.italic, r.font.size, r.font.color.rgb = True, Pt(11), GRAY

    def save(self, path: Path):
        self.doc.save(path)


def write(shots: dict[str, Path], info: dict, path: Path) -> Path:
    m = Manual(shots)
    d = m.doc
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(120)
    r = t.add_run("СТ-Секретарь")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(34), GREEN
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Инструкция секретаря")
    r.font.size = Pt(22)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(24)
    r = t.add_run("Программа для секретариата соревнований по спортивному туризму:\n"
                  "заявки, допуск, жеребьёвка, протоколы, результаты, награждение, договоры")
    r.font.size, r.font.color.rgb = Pt(14), GRAY
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(160)
    r = t.add_run(f"Версия программы {__version__} · автор Udnikov Denis\n"
                  "Снимки экрана — на учебном соревновании с выдуманными данными.\n"
                  "Исходный код и новые версии: github.com/avtomatikalab/st-secretary")
    r.font.size, r.font.color.rgb = Pt(11), GRAY

    m.h1("Коротко о программе")
    m.p("СТ-Секретарь работает на ноутбуке секретариата и открывается в браузере, как сайт, — но интернет ему не "
        "нужен. Всё, что вы вводите, сохраняется сразу; кнопки «Сохранить документ» нет и не нужно.")
    m.p("Название, даты, место и судейская коллегия вводятся один раз — в карточке соревнования — и сами попадают "
        "во все протоколы, дипломы, справки и договоры. Поэтому в документах не бывает прошлогодних дат и "
        "разных написаний одной фамилии.")
    m.p("Работа идёт по шагам — они слева в меню, в том же порядке, что и на соревнованиях:")
    m.steps(["Карточка соревнования — название, даты, место, судьи, зачёты.",
             "Предварительные заявки — программа проверяет каждую заявку и показывает ошибки.",
             "Комиссия по допуску — документы участников, решения, номера, взносы, протокол.",
             "Жеребьёвка и стартовые протоколы.",
             "Протоколы этапов и результаты — баллы и время, места, разряды, протоколы, протесты.",
             "Телефоны судей этапов — судьи отмечают команды на своих телефонах.",
             "Награждение и документы по итогам — дипломы, наклейки, справки, выписки, отчёт.",
             "Договоры, акты, табель судей.",
             "Сверка документов — чтобы в бумагах не было расхождений.",
             "Табло — результаты на телефонах участников."])
    m.tip("сначала потренируйтесь на учебном соревновании (раздел 2) — на нём можно нажимать что угодно.")

    m.h1("1. Запуск и выключение")
    m.steps(["Распакуйте архив программы в любую папку, например в «Документы» (правой кнопкой по архиву → "
             "«Извлечь все…»).",
             "Дважды щёлкните файл «СТ-Секретарь.bat». Откроется чёрное окно — это сама программа — и через "
             "несколько секунд страница в браузере.",
             "Работайте в браузере. Чёрное окно не закрывайте, пока работаете: его можно свернуть.",
             "Закончили — нажмите «Выключить» вверху страницы или просто закройте чёрное окно. Всё уже сохранено."])
    m.img("home", "Главная страница: список соревнований")
    m.tip("сделайте ярлык на рабочем столе: правой кнопкой по «СТ-Секретарь.bat» → «Отправить» → «Рабочий стол "
          "(создать ярлык)».")
    m.p("Соревнования хранятся в папке «данные» рядом с программой — одна папка на соревнование. Сканы "
        "документов участников и паспортные данные судей хранятся отдельно, только на этом компьютере, в папке "
        "«СТ-Секретарь — документы участников» в вашей папке пользователя.")
    m.warn("если браузер не открылся сам, откройте его и введите адрес http://127.0.0.1:8765 — он написан и в "
           "чёрном окне.")
    m.h2("На Mac и в Linux")
    m.p("Программа работает так же. На Mac проще всего поставить её одной командой: откройте Терминал "
        "(Cmd+Пробел → «Терминал»), вставьте строку ниже и нажмите Enter. Программа встанет в папку «СТ-Секретарь» "
        "в вашей домашней папке, на Рабочем столе появится ярлык «СТ-Секретарь», и программа запустится. Этой же "
        "командой её потом можно обновить. В Linux — то же, ярлык будет в меню программ.")
    m.p("curl -fsSL https://raw.githubusercontent.com/avtomatikalab/st-secretary/main/tools/install.sh | sh")
    m.p("Вместо чёрного окна на Mac открывается окно Терминала — это сама программа, его тоже не закрывайте, пока "
        "работаете. Если скачали архив для Mac через браузер, файл запуска называется «СТ-Секретарь.command»; при "
        "первом запуске macOS скажет, что не может проверить разработчика: откройте «Системные настройки» → "
        "«Конфиденциальность и безопасность» и нажмите «Всё равно открыть» — это нужно один раз.")

    m.h1("2. Учебное соревнование")
    m.p("На главной странице есть кнопка «Создать учебное соревнование». Программа создаст «Учебный чемпионат "
        "г. Энска»: карточку с судейской коллегией, 14 команд с выдуманными участниками (ПСР и спелео) и этапы "
        "дистанций. В четырёх заявках ошибки оставлены специально — такие же, как в настоящих заявках: дата, "
        "которой не бывает, нет женщины в команде, участник младше допустимого возраста, один человек в двух "
        "командах.")
    m.p("Пройдите на нём все шаги: исправьте заявки, проведите комиссию и жеребьёвку, внесите баллы, опубликуйте "
        "протоколы, откройте дипломы. Все снимки в этой инструкции сделаны на учебном соревновании. Когда оно "
        "станет не нужно, удалите его папку в «данных».")
    m.img("overview", "Страница соревнования: меню шагов слева")

    m.h1("3. Карточка соревнования")
    m.steps(["На главной странице нажмите «Новое соревнование».",
             "Заполните наименование, вид и уровень соревнований, даты, место, проводящие организации.",
             "Внизу — судейская коллегия (главный судья, главный секретарь, заместители) с категориями и "
             "территорией, и зачёты: группа, класс дистанции, дисциплина, возраст, состав команды, взнос.",
             "Нажмите «Сохранить». Если что-то не так, ошибка будет видна прямо у поля — введённое не пропадёт."])
    m.img("card", "Карточка соревнования: реквизиты, судейская коллегия, зачёты")
    m.tip("карточку можно заполнить и в Excel: на главной — «Загрузить карточку». Если карточку поправили в "
          "Excel, пока была открыта страница, программа не затрёт эти правки.")
    m.h2("Неофициальные соревнования")
    m.p("Клубные, учебные соревнования, слёты — отметьте «Неофициальные соревнования». Тогда у зачёта можно задать "
        "своё название («Новички», «Семейные команды»), класс «без класса» (неклассифицированная дистанция) и свою "
        "дисциплину, если её нет в ВРВС: чем выражен результат (баллы, время или время + штрафные баллы) и состав "
        "(личный, связка, команда). Квалификационный ранг, процент от победителя и разряды у неофициальных не "
        "считаются, допуск по Правилам (класс, разряд) не проверяется — только то, что задано в зачёте. Свои "
        "группы, названия и дисциплины программа запоминает на этом компьютере и подсказывает в следующий раз.")
    m.tip("зачёт можно переименовать в любой момент — заявки, жеребьёвка, результаты и ссылки судей не потеряются.")

    m.h1("4. Предварительные заявки")
    m.steps(["Откройте шаг «Предварительные заявки».",
             "Перетащите на страницу файлы заявок команд (Excel) — сразу все. Можно и кнопкой «Выбрать файлы».",
             "Программа проверит каждую заявку: возраст, разряды, состав команды, дубли, написание ФИО и "
             "территорий. У каждой команды видно число ошибок и того, что нужно проверить."])
    m.img("preapps", "Список заявок: красным — ошибки, жёлтым — что проверить")
    m.p("Три вида замечаний:")
    m.bullets(["«Ошибка» — так участника допустить нельзя, нужно исправить (например, даты рождения нет в "
               "календаре).",
               "«Проверить» — программа сомневается; посмотрите и нажмите «Проверено», если всё верно.",
               "«Исправлено» — программа сама привела запись к правильному виду (разряд «кмс» → «КМС», "
               "ФИО прописными → как положено). Смотреть не обязательно."])
    m.p("У каждого замечания написано, почему оно появилось и что сделать. Щёлкните по названию команды — "
        "откроется её карточка.")
    m.img("team_error", "Карточка команды: заявка и замечания с объяснением")
    m.steps(["В карточке команды нажмите «Редактировать заявку».",
             "Исправьте поле, обведённое красным, — объяснение под строкой участника.",
             "Нажмите «Сохранить и проверить». Прежний файл заявки сохранится в папке «Прежние версии»."])
    m.img("team_edit", "Заявка в форме: поле с ошибкой обведено")
    m.tip("«Открыть сводку в Excel» — готовые строки для листа «Заявка» СЕКРЕТАРЬ_ST, если вы ещё ведёте его.")

    m.h1("5. Комиссия по допуску")
    m.p("Здесь — все допущенные по заявкам участники. Отмечайте галочками документы, которые показал каждый "
        "участник, и документы команды. Всё сохраняется сразу.")
    m.img("admission", "Комиссия по допуску: документы, решения, номера, взносы")
    m.bullets(["Команда допускается сама, когда её заявка отмечена «Проверено» (на странице заявок), документы "
               "отмечены и замечаний нет. Иначе — только решением комиссии, с подтверждением и основанием; такие "
               "команды и участники собраны в списке «Допущены без проверки секретаря».",
               "«Допуск врача в заявке» у команды сам отмечает «Мед. допуск» её участникам (кроме добавленных "
               "перезаявкой). Врач кого-то не допустил — снимите галочку у этого участника.",
               "Один человек в двух командах (те же ФИО и дата рождения) — документы отмечаются один раз.",
               "«Показать по делегациям» — команды одной территории и представителя подряд, с общим взносом.",
               "«Отметить все документы» — если у команды всё в порядке, одной кнопкой.",
               "Решение «допустить решением комиссии» или «не допускать» — с причиной.",
               "Номера командам — «Присвоить тем, у кого нет»; номер можно поправить вручную.",
               "Взнос — сумма и способ оплаты; ведомость взносов собирается сама.",
               "Перезаявка — из карточки команды, с временем подачи; программа предупредит, если позже, "
               "чем за час до старта.",
               "«Открыть в Excel» — протокол комиссии по форме Правил и ведомость взносов."])
    m.tip("сканы документов (паспорт, полис, справка) можно добавить в карточку команды и на комиссии смотреть "
          "«в одном окне»: документ слева, галочки справа. Сканы хранятся только на этом компьютере.")

    m.h1("6. Жеребьёвка и стартовые протоколы")
    m.p("Жеребьёвка проводится по каждому зачёту отдельно (Правила, раздел 3, п. 8.4).")
    m.steps(["Выберите способ: общая компьютерная жеребьёвка, групповая по рангу состава (группы видно в порядке "
             "старта; равные ранги на границе групп делит жребий), «строго по рангу» (слабые раньше сильных, равные "
             "ранги — жребием) или порядок, который вытянули представители на совещании.",
             "Нажмите «Провести жеребьёвку». Программа запишет время и число жребия — они будут в протоколе.",
             "Во «Времени старта» укажите день, время первого старта и интервал. Время каждой команды "
             "посчитается само."])
    m.img("start_draw", "Жеребьёвка: способ и время старта")
    m.p("В таблице «Порядок старта» порядок можно поменять — впишите номера и нажмите «Сохранить порядок». Время "
        "отдельной команды тоже можно поправить; вписанное время выделено жирным.")
    m.img("start_order", "Порядок старта: составы, ранг, время")
    m.steps(["Нажмите «Опубликовать стартовый протокол». Файл Excel сохранится в папку «Протоколы» и откроется — "
             "распечатайте и вывесите.",
             "Публикуйте не позднее чем за час до старта — программа напомнит, до какого времени.",
             "С момента публикации идёт час на протесты по допуску."])
    m.img("start_publish", "Публикация стартового протокола и сроки")
    m.tip("порядок старта сразу действует и в таблице результатов (при равенстве выше стартовавший раньше), и на "
          "телефонах судей, и на табло.")
    m.tip("если участник заявлен в нескольких зачётах, задайте «Перерыв между стартами участника» — программа "
          "предупредит, если его старты стоят слишком близко.")

    m.h1("7. Телефоны судей этапов")
    m.p("Судья этапа может отмечать команды на своём телефоне: прибытие и убытие, баллы этапа, снятие с "
        "причиной. Связь на дистанции не нужна — всё хранится в телефоне и приходит в таблицу, когда телефон "
        "снова в Wi-Fi ноутбука.")
    m.steps(["Включите на ноутбуке точку доступа Wi-Fi (или подключите ноутбук и телефоны к одной сети).",
             "На странице «Телефоны судей» нажмите «Раздать по Wi-Fi». Windows спросит разрешение в брандмауэре — "
             "разрешите для частной сети.",
             "«Ссылки всем этапам без ссылки», затем «Карточки с QR для печати» — у каждого этапа свой QR-код.",
             "Судья наводит камеру телефона на QR-код своего этапа на базе, пока телефон в Wi-Fi, и не закрывает "
             "страницу."])
    m.img("judges", "Ссылки этапов: кто судит, номер, что уже прислали; журнал этапа")
    m.bullets(["При первом открытии ссылки судья выбирает себя из списка судей соревнования и указывает телефон — "
               "один раз на этом телефоне. У каждой записи в журнале этапа видно, кто её поставил; номер — ссылкой, "
               "чтобы позвонить. Номер не такой, как в личных данных, — «Принять номер».",
               "«Связь» на телефоне судьи — главный судья, секретарь и судьи всех этапов с номерами.",
               "Если этап занят, у подошедшей команды — «ждёт очереди». Вычитать ли ожидание из времени — настройка "
               "зачёта на странице результатов.",
               "«Таблица штрафов» (раздел 3 на этой странице): для пешеходных и спелео — из Правил, для ПСР — своя из "
               "Excel. У судьи — поиск по словам и номеру пункта; «+ Штраф по таблице» у команды: баллы этапа "
               "складываются из пунктов сами, у главного судьи видна расшифровка. Штраф без пункта таблицы — "
               "сопоставьте пункт в журнале этапа."])
    m.phone_pair("judge_phone", "judge_penalties", "Слева — страница судьи этапа; справа — таблица штрафов с поиском")
    m.tip("если секретарь уже исправил клетку (например, после протеста), присланное с телефона её не затрёт — "
          "расхождение будет видно в результатах.")

    m.h1("8. Протоколы этапов и результаты")
    m.p("Этапы дистанции задаются один раз (или загружаются из рабочей книги СЕКРЕТАРЬ_ST вместе с баллами). "
        "Дальше таблица «команды × этапы» заполняется, как в Excel.")
    m.steps(["Вносите баллы этапа из протокола судьи: штраф — плюсом, премия — минусом.",
             "Enter — на команду ниже, Tab — на следующий этап. Всё сохраняется сразу.",
             "Команда снята или не финишировала — выберите статус справа: место ей не присуждается.",
             "Сумма, места, ранг и выполненные разряды пересчитываются сразу."])
    m.img("results_points", "Баллы по этапам: синим — пришедшее с телефонов судей")
    m.p("В ПСР у этапа можно задать НВ, КВ, ТШ и ВШ: судья вносит только технический штраф, временной штраф по "
        "времени на этапе программа добавит сама. Сняли команду с этапа или время сверх КВ — в клетке МШ и под ней "
        "«снята ✕» («сверх КВ ✕»); нажмите, чтобы снять отметку, — итог этапа посчитается заново.")
    m.img("results_table", "Результаты: места, процент от победителя, разряды")
    m.h2("Спелео и пешеходные дистанции")
    m.p("В этих зачётах результат — время: старт и финиш (чч:мм:сс), отсечки (мм:сс) и штрафные баллы этапов "
        "(1 балл — 15 или 30 секунд). Снятие с этапа — буква «с» в клетке этапа. Время старта и финиша можно "
        "загрузить из SPORTident Reader (файл si_reader.csv) — кнопкой над таблицей.")
    m.img("speleo_points", "Спелео: время, отсечки, баллы, снятие «с»")
    m.h2("Протоколы и протесты")
    m.steps(["«Опубликовать предварительный протокол» — файл Excel откроется, распечатайте и вывесите. С этого "
             "момента идёт час на протесты по результатам.",
             "Протест — запишите внизу страницы: время подачи, команда, суть. Решение — в строке протеста.",
             "После часа и решений по всем протестам — «Утвердить — официальный протокол». Места и разряды сами "
             "уйдут в награждение."])
    m.img("results_protocol", "Предварительный протокол, час на протесты, официальный протокол")
    m.warn("если после публикации поменять баллы, программа попросит опубликовать протокол заново — с новым часом "
           "на протесты.")

    m.h1("9. Табло")
    m.p("Участники и зрители видят стартовые протоколы и результаты на своих телефонах через Wi-Fi ноутбука, "
        "без интернета. Табло показывает только команды, баллы и места — заявки, документы и договоры с него "
        "не видны.")
    m.steps(["Шаг «Табло» → «Включить для этого соревнования».",
             "«Раздать табло по Wi-Fi» — появятся адрес и QR-код. Распечатайте QR-код или покажите его на экране.",
             "Страница табло обновляется сама."])
    m.img("board_admin", "Включение табло для соревнования и раздача по Wi-Fi")

    m.h1("10. Награждение и документы по итогам")
    m.p("Утверждённые результаты приходят сюда сами. Если результаты считались в СЕКРЕТАРЬ_ST, загрузите его "
        "итоговый протокол (.xls) или внесите места вручную.")
    m.img("awards", "Награждение: места по зачётам и документы одной кнопкой")
    m.p("Каждая кнопка «Открыть» сохраняет документ в папку «Документы по итогам» и открывает его в Word или "
        "Excel:")
    m.bullets(["дипломы (у каждого участника группы свой диплом, его имя первое);",
               "наклейки на медали;", "список награждаемых для ведущего (III → II → I место);",
               "справки о судействе, о составе и квалификации судейской коллегии, о количестве субъектов РФ;",
               "выписки из протокола на присвоение разрядов;", "отчёт главного судьи."])
    m.tip("перед печатью документ можно поправить в Word как обычно.")

    m.h1("11. Договоры, акты, табель")
    m.p("Табель по дням: судьи из карточки уже в нём, коменданта и рабочих добавьте на странице. Галочка — "
        "рабочий день; суммы по ставкам считаются сразу.")
    m.img("contracts", "Табель: дни работы и суммы")
    m.bullets(["«Табель-наряд» — Excel с формулами.",
               "Договоры и акты — всем одним файлом или каждому отдельно; сумма прописью, срок и должность — по "
               "табелю, поэтому договор, акт и табель не расходятся.",
               "Паспорт, ИНН, СНИЛС и счёт вводятся один раз в карточке человека и хранятся только на этом "
               "компьютере; ИНН, СНИЛС и счёт проверяются по контрольным цифрам.",
               "Форма договора у заказчика своя: «Шаблон договора» сохраняется в папку соревнования — перенесите "
               "в него текст заказчика, оставив поля в двойных фигурных скобках."])

    m.h1("12. Сверка документов")
    m.p("Перед сдачей документов откройте «Сверку». Программа прочитает документы (Word, Excel, PDF) и покажет, "
        "что расходится с карточкой и между документами: даты не того года, чужие коды дисциплин, опечатки в "
        "названии дисциплины, ФИО и категории судей, суммы цифрами и прописью, табель против договоров и актов.")
    m.img("verify", "Сверка документов")
    m.tip("перетащите на страницу и присланные файлы — протоколы из СЕКРЕТАРЬ_ST, договоры заказчика. Программа "
          "их только читает и никуда не сохраняет.")

    m.h1("13. Резервная копия")
    m.p("Программа сама делает копию каждого соревнования при запуске, если в нём что-то менялось, и хранит 10 "
        "последних. На странице соревнования можно сделать копию в любой момент или скачать её — например, на "
        "флешку. Паспортные данные и сканы документов в копию не входят: они хранятся только на этом компьютере.")
    m.img("backup", "Резервная копия на странице соревнования")
    m.p("Несколько соревнований за один выезд можно объединить в фестиваль (на главной — «Объединить соревнования "
        "в фестиваль»): они остаются самостоятельными, но видны одной группой, между ними — переход в один клик, "
        "есть сводка «кто заявлен в нескольких соревнованиях» и копия всего фестиваля одним архивом.")
    m.steps(["Перед отъездом с соревнований нажмите «Скачать копию (.zip)» и сохраните файл на флешку.",
             "Чтобы восстановить или перенести соревнование, на главной странице в блоке «Восстановить из резервной "
             "копии» выберите файл и нажмите «Восстановить». Соревнование появится отдельной папкой — то, что уже "
             "есть, не затрётся."])

    m.h1("14. Новая версия программы")
    m.p("Если ноутбук подключён к интернету, программа при запуске сама проверяет, нет ли новой версии, и сообщает "
        "о ней на главной странице. Обновляйте между соревнованиями, а не во время них.")
    m.steps(["На главной странице нажмите «Что нового и обновить», прочитайте, что изменилось.",
             "Нажмите «Обновить до …». Программа сделает резервную копию, скачает новую версию и сама перезапустится "
             "— страница откроется заново. Чёрное окно не закрывайте.",
             "Соревнования, резервные копии и документы участников остаются как были. Прежняя версия сохраняется в "
             "папке «program.old» рядом с «СТ-Секретарь.bat»."])
    m.tip("без интернета программа просто не покажет сообщение — работать это не мешает.")

    m.h1("15. Если что-то пошло не так")
    for q, a in [
        ("Закрыли чёрное окно.", "Ничего не потеряно — всё сохраняется сразу. Запустите «СТ-Секретарь.bat» снова."),
        ("Страница пишет, что файл открыт в Excel.", "Закройте этот файл в Excel и нажмите кнопку ещё раз."),
        ("Страница не открывается.", "Проверьте, открыто ли чёрное окно. Если открыто — введите в браузере адрес "
                                     "из чёрного окна (обычно http://127.0.0.1:8765)."),
        ("Windows спрашивает про брандмауэр.", "Это при раздаче табло или ссылок судей по Wi-Fi. Разрешите для "
                                               "частной сети."),
        ("Нужно перенести соревнование на другой ноутбук.", "На странице соревнования — «Скачать копию (.zip)», "
                                                           "сохраните на флешку; на другом ноутбуке на главной — "
                                                           "«Восстановить из резервной копии»."),
        ("Что-то испортили и хочется вернуть как было.", "Программа сама делает копию соревнования при каждом "
                                                         "запуске (хранятся 10 последних) — они в папке «Резервные "
                                                         "копии» рядом с «данными». Восстановите нужную: соревнование "
                                                         "появится отдельной папкой, текущее не затрётся."),
        ("После обновления программа не запускается.", "Закройте чёрное окно. В папке рядом с «СТ-Секретарь.bat» "
                                                       "удалите папку «program» и переименуйте «program.old» в "
                                                       "«program» — вернётся прежняя версия."),
        ("Непонятное сообщение об ошибке.", "Сфотографируйте экран вместе с чёрным окном и отправьте "
                                            "разработчикам — данные при этом не меняются."),
    ]:
        m.p(" " + a, bold_lead=q)
    m.save(path)
    return path


def to_pdf(docx: Path) -> Path | None:
    """PDF через Word (если установлен) — для чтения с телефона и для архива программы."""
    if sys.platform != "win32":
        return None
    pdf = docx.with_suffix(".pdf")
    ps = ("$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"try {{ $d = $w.Documents.Open('{docx}'); $d.SaveAs([ref]'{pdf}', [ref]17); $d.Close($false) }} "
          "finally { $w.Quit() }")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, check=False)
    if not pdf.is_file():
        print("PDF не сделан (нужен Microsoft Word):", r.stderr[-300:])
        return None
    return pdf


def main():
    work = Path(tempfile.mkdtemp(prefix="st-manual-"))
    try:
        shots, info = walk(work)
        path = write(shots, info, OUT)
        print(f"Готово: {path}")
        pdf = to_pdf(path)
        if pdf:
            print(f"PDF: {pdf}")
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
