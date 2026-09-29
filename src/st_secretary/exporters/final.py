"""Выписки из протоколов для присвоения разрядов (ЕВСК, п. 67.8.2) и отчёт главного судьи.

Выписка: наименование соревнований, дисциплина по ВРВС (код), дата и место, пол и возрастная группа,
распределение мест, ФИО полностью, дата рождения, разряд, результат, выполненный разряд, субъект РФ;
подпись главного судьи. В выписку попадают только те, кто выполнил норматив.

Отчёт — по образцу отчёта главного судьи Чемпионата г. Красноярска 2025 г.: состав участников,
результаты, рекорды, жалобы, материальная база, состав судейской коллегии с оценками.
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from st_secretary.competition import Competition
from st_secretary.exporters.awards import _new_doc, _para
from st_secretary.exporters.judges import _sign, _table
from st_secretary.qualification import Qual
from st_secretary.results import ROMAN, ZachetResults, group_label, person_key

THIN = Side(style="thin", color="7F7F7F")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
HEAD = PatternFill("solid", fgColor="DCE6F1")
MISSING = PatternFill("solid", fgColor="F8CBAD")

REPORT_DEFAULTS = {
    "not_arrived": "все заявленные команды прибыли на соревнования",
    "records": "дистанции не являются эталонными, рекорды не устанавливались",
    "protests": "протесты, жалобы в ГСК не подавались",
    "base": "",
    "notes": "",
}


def write_extracts(results: list[ZachetResults], comp: Competition, path: str | Path) -> Path:
    """Выписки на разряды — лист на зачёт; только выполнившие норматив."""
    wb = Workbook()
    wb.remove(wb.active)
    for z in results:
        zz = next(x for x in comp.zachety if x.key == z.key)
        ws = wb.create_sheet(z.key.replace("/", "-")[:31])
        lines = [
            ("ВЫПИСКА ИЗ ПРОТОКОЛА СОРЕВНОВАНИЙ", True, 13),
            (comp.title, True, 12),
            (f"Спортивная дисциплина: «{zz.discipline_name}», код ВРВС {zz.discipline_code}; {zz.distance_class} класс", False, 11),
            (f"Дата и место проведения: {comp.dates_text}, {comp.place}", False, 11),
            (f"Пол и возрастная группа: {group_label(z, zz.group)}"
             + (f"; квалификационный ранг соревнований: {z.rank}" if z.rank else ""), False, 11),
        ]
        for r, (text, bold, size) in enumerate(lines, start=1):
            ws.cell(r, 1, text).font = Font(bold=bold, size=size)
        head = ["Место", "Фамилия, имя, отчество", "Дата рождения", "Разряд, звание", "Результат",
                "Выполнен разряд", "Команда", "Субъект РФ"]
        r0 = len(lines) + 2
        for c, h in enumerate(head, start=1):
            cell = ws.cell(r0, c, h)
            cell.font, cell.fill, cell.border, cell.alignment = Font(bold=True, size=10), HEAD, BOX, WRAP
        r = r0 + 1
        for p in z.rows:
            if not p.norm:
                continue
            for m in p.members:
                values = [p.place_text, m.fio, m.birth_text, m.qual, p.result, p.norm, p.team, p.territory]
                for c, v in enumerate(values, start=1):
                    cell = ws.cell(r, c, v)
                    cell.border, cell.alignment = BOX, WRAP
                if not m.birth_text:  # без даты рождения выписку не примут
                    ws.cell(r, 3, "нет в заявке").fill = MISSING
                r += 1
        if r == r0 + 1:
            ws.cell(r, 1, "Нормативы в этом зачёте никто не выполнил.")
            r += 1
        o = comp.official("Главный судья")
        ws.cell(r + 2, 1, f"Главный судья ________________ / {o.signature if o else ' ' * 30} /")
        for c, w in enumerate([7, 34, 13, 10, 11, 11, 20, 20], start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    if not wb.sheetnames:
        wb.create_sheet("Выписки").cell(1, 1, "Результатов пока нет — внесите их на странице итогов.")
    path = Path(path)
    wb.save(path)
    return path


def extract_count(results: list[ZachetResults]) -> int:
    return sum(len(p.members) for z in results for p in z.rows if p.norm)


# ------------------------------------------------------------------ отчёт главного судьи


def _day(d) -> str:
    from st_secretary.results import MONTHS_GEN

    return f"«{d.day:02d}» {MONTHS_GEN[d.month - 1]} {d.year} г."


def write_report(comp: Competition, results: list[ZachetResults], people: list, teams: list[str],
                 grades: dict[str, str], texts: dict[str, str], path: str | Path) -> Path:
    """people — допущенные участники (preapp.Entry), teams — территории допущенных команд."""
    t = {**REPORT_DEFAULTS, **{k: v for k, v in texts.items() if v}}
    doc = _new_doc()
    _para(doc, "ОТЧЁТ", 16, True, space_after=0)
    _para(doc, "главного судьи соревнований", 13, space_after=0)
    _para(doc, f"о проведении: {comp.title}", 13, True, space_after=0)
    _para(doc, f"в период с {_day(comp.date_from)} по {_day(comp.date_to)}", 12, space_after=0)
    _para(doc, f"Место проведения: {comp.place}", 12, space_after=12)

    left = WD_ALIGN_PARAGRAPH.LEFT
    _para(doc, "СОСТАВ УЧАСТНИКОВ", 13, True, left)
    terr = list(dict.fromkeys(x for x in teams if x))
    _para(doc, f"1. В соревнованиях принимали участие команды: {', '.join(terr) or '—'} "
               f"(команд — {len(teams)}).", 12, align=left, space_after=2)
    men = sum(e.sex == "м" for e in people)
    _para(doc, f"2. Общее количество участников, допущенных до соревнований, — {len(people)}, "
               f"из них мужчин — {men}, женщин — {len(people) - men}.", 12, align=left, space_after=2)
    _para(doc, f"3. Не прибыли на соревнования: {t['not_arrived']}.", 12, align=left, space_after=2)
    ages = [e.age_in(comp.year) for e in people]
    young = sum(1 for a in ages if a is not None and a < 16)
    q = Counter(e.qual for e in people if e.qual is not None)
    juniors = q[Qual.Y1] + q[Qual.Y2] + q[Qual.Y3]
    quals = (f"МС — {q[Qual.MS]}, КМС — {q[Qual.KMS]}, I разряд — {q[Qual.I]}, II разряд — {q[Qual.II]}, "
             f"III разряд — {q[Qual.III]}, юношеские разряды — {juniors}, без разряда — {q[Qual.BR]}")
    _para(doc, f"4. Из общего числа участников по возрасту: до 16 лет — {young or 'нет'}, "
               f"16 лет и старше — {len(people) - young}; по спортивно-технической подготовке: {quals}.",
          12, align=left, space_after=10)

    _para(doc, "РЕЗУЛЬТАТЫ СОРЕВНОВАНИЙ", 13, True, left)
    for z in results:
        zz = next(x for x in comp.zachety if x.key == z.key)
        _para(doc, f"«{zz.discipline_name}» — {zz.distance_class} класс, {group_label(z, zz.group)}"
                   + (f"; ранг {z.rank}" if z.rank else ""), 12, True, left, space_after=2)
        for p in z.medalists:
            names = ", ".join(m.fio.rsplit(" ", 1)[0] if len(m.fio.split()) > 2 else m.fio for m in p.members)
            _para(doc, f"{ROMAN[p.place]} место — команда «{p.team}» ({p.territory}): {names}", 12, align=left,
                  space_after=2)
        norms = Counter(p.norm for p in z.rows if p.norm for _ in p.members)
        if norms:
            _para(doc, "Выполнили нормативы: " + ", ".join(f"{k} — {v} чел." for k, v in norms.items()), 12,
                  align=left, space_after=6)
    if not results:
        _para(doc, "Результаты не внесены.", 12, align=left)
    _para(doc, f"На соревнованиях установлены рекорды: {t['records']}.", 12, align=left, space_after=10)

    _para(doc, "ЖАЛОБЫ УЧАСТНИКОВ СОРЕВНОВАНИЙ", 13, True, left)
    _para(doc, t["protests"][:1].upper() + t["protests"][1:] + ".", 12, align=left, space_after=10)
    _para(doc, "МАТЕРИАЛЬНАЯ БАЗА", 13, True, left)
    _para(doc, t["base"] or "________________________________________________", 12, align=left, space_after=10)
    if t["notes"]:
        _para(doc, "ЗАМЕЧАНИЯ И ПРЕДЛОЖЕНИЯ", 13, True, left)
        _para(doc, t["notes"], 12, align=left, space_after=10)

    _para(doc, "СОСТАВ СУДЕЙСКОЙ КОЛЛЕГИИ", 13, True, left, space_after=2)
    _para(doc, f"с {_day(comp.date_from)} по {_day(comp.date_to)}", 12, align=left, space_after=4)
    judges = [o for o in comp.officials if o.fio]
    _table(doc, ["№", "Фамилия, имя, отчество", "Категория", "Территория", "В качестве кого судил", "Оценка"],
           [[str(i), o.fio, o.category, o.territory, o.role, grades.get(person_key(o.fio), "")]
            for i, o in enumerate(judges, start=1)], [0.8, 4.4, 1.8, 3.2, 3.8, 2.2])
    _para(doc, "", 12)
    _sign(doc, comp, ("Главный судья",))
    path = Path(path)
    doc.save(path)
    return path
