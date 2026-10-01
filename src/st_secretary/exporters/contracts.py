"""Табель-наряд, договоры и акты судейской (комендантской) бригады.

Табель — по образцу табеля-наряда Чемпионата г. Красноярска 2025 г.: отметки «р» по дням, число дней,
ставка, сумма, начисления. Число дней и суммы — формулами Excel: если поправить отметки в Excel,
всё пересчитается само (в 2025 г. при правке вручную у двух человек разошлись отметки и дни к оплате).

Договор и акт — по шаблону Word с полями {{ФИО}}, {{Сумма прописью}} и т. п. Встроенный шаблон —
обычный договор возмездного оказания услуг; у заказчика обычно своя форма — в неё вставляют те же поля
и кладут файл «Шаблон договора.docx» в папку соревнования. Список полей — FIELDS.
"""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from st_secretary.competition import Competition
from st_secretary.exporters.awards import _new_doc, _para
from st_secretary.exporters.judges import disciplines_text
from st_secretary.money import money, rubles_text
from st_secretary.names import genitive, initials, role_genitive
from st_secretary.results import title_of
from st_secretary.staff import CAT_WORDS, MARK, Person

THIN = Side(style="thin", color="7F7F7F")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="center", wrap_text=True)
HEAD = PatternFill("solid", fgColor="DCE6F1")
BLANK = "________________"

CUSTOMER_FIELDS = [
    ("name", "Заказчик — полное наименование", "Муниципальное автономное учреждение «…»"),
    ("short", "Заказчик кратко (для подписей)", "МАУ «…»"),
    ("head_post", "Должность руководителя", "директор"),
    ("head_fio", "ФИО руководителя", "Фамилия Имя Отчество"),
    ("basis", "Действует на основании", "Устава"),
    ("city", "Место заключения договора", "г. Красноярск"),
    ("requisites", "Адрес и банковские реквизиты заказчика (как в договоре)", "660000, г. …, ул. …\nИНН … КПП …\n…"),
    ("rates_basis", "Основание ставок", "норм расходов, утверждённых приказом …"),
    ("pay_term", "Срок оплаты после подписания акта", "45 рабочих дней"),
    ("funding", "Источник финансирования", "за счёт средств субсидии из бюджета г. …"),
]

# Поля шаблона договора: {{имя}} → что подставляется.
FIELDS = [
    ("ФИО", "исполнитель полностью"), ("ФИО кратко", "Иванов И.И. — для подписи"),
    ("Должность", "как в табеле: «Главный судья»"), ("Услуга", "«Главный судья 1 категории» — для акта"),
    ("В качестве", "«главного судьи 1 категории»"), ("Соревнование", "наименование из карточки"),
    ("Соревнования", "наименование в родительном падеже: «Чемпионата г. …»"),
    ("Дисциплина", "«дистанция - комбинированная»"), ("Даты соревнований", "«19–21 сентября 2025 г.»"),
    ("Место", "место проведения"), ("Срок", "дни работы человека: «18–22 сентября 2025 г.»"),
    ("Дней", "число отмеченных дней"), ("Ставка", "«850,00»"), ("Сумма", "«4 250,00»"),
    ("Сумма прописью", "«4250 (четыре тысячи двести пятьдесят) рублей 00 копеек»"), ("Год", "год соревнований"),
    ("Заказчик", ""), ("Заказчик кратко", ""), ("Должность руководителя", "«Директор»"),
    ("Руководитель", "ФИО руководителя"), ("Руководитель кратко", "«Иванов И.И.»"),
    ("В лице", "«директора Иванова Ивана Ивановича»"), ("Основание", "«Устава»"),
    ("Место заключения", ""), ("Реквизиты заказчика", ""), ("Основание ставок", ""), ("Срок оплаты", ""),
    ("Финансирование", ""),
    ("Дата рождения", ""), ("Паспорт", "серия и номер"), ("Кем выдан", ""), ("Дата выдачи", ""),
    ("Код подразделения", ""), ("Адрес", "адрес регистрации"), ("ИНН", ""), ("СНИЛС", ""), ("Счёт", ""),
    ("Банк", ""), ("БИК", ""), ("Телефон", ""),
]
_PH = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def _cap(s: str) -> str:
    return s[:1].upper() + s[1:]


def contract_values(comp: Competition, p: Person, customer: dict, pd: dict) -> dict[str, str]:
    """Значения полей договора и акта ({{…}} в шаблоне): заказчик, исполнитель, суммы и даты."""
    c = {k: str(customer.get(k, "") or "").strip() for k, _, _ in CUSTOMER_FIELDS}
    head_post = c["head_post"] or "директор"
    in_person = f"{role_genitive(head_post) or head_post} {genitive(c['head_fio'])}".strip() if c["head_fio"] else ""
    city = c["city"] or (comp.host_territory if comp.host_territory.startswith(("г.", "п.", "с.")) else
                         f"г. {comp.host_territory}" if comp.host_territory else "")
    service = f"{p.role} {CAT_WORDS.get(p.cat, '')}".strip()
    v = {
        "ФИО": p.fio, "ФИО кратко": initials(p.fio), "Должность": p.role, "Услуга": service,
        "В качестве": p.role_text, "Соревнование": comp.title, "Соревнования": title_of(comp),
        "Дисциплина": disciplines_text(comp), "Даты соревнований": comp.dates_text, "Место": comp.place,
        "Срок": p.span_text, "Дней": str(len(p.days)), "Ставка": money(p.rate or 0), "Сумма": money(p.amount),
        "Сумма прописью": rubles_text(p.amount), "Год": str(comp.year),
        "Заказчик": c["name"], "Заказчик кратко": c["short"] or c["name"],
        "Должность руководителя": _cap(head_post), "Руководитель": c["head_fio"],
        "Руководитель кратко": initials(c["head_fio"]) if c["head_fio"] else "", "В лице": in_person,
        "Основание": c["basis"] or "Устава", "Место заключения": city, "Реквизиты заказчика": c["requisites"],
        "Основание ставок": c["rates_basis"], "Срок оплаты": c["pay_term"] or "45 рабочих дней",
        "Финансирование": c["funding"],
    }
    for key, label in (("birth", "Дата рождения"), ("passport", "Паспорт"), ("issued_by", "Кем выдан"),
                       ("issued_on", "Дата выдачи"), ("dept_code", "Код подразделения"), ("address", "Адрес"),
                       ("inn", "ИНН"), ("snils", "СНИЛС"), ("account", "Счёт"), ("bank", "Банк"), ("bik", "БИК"),
                       ("phone", "Телефон")):
        v[label] = str(pd.get(key, "") or "").strip()
    return {k: (x or BLANK) for k, x in v.items()}


# ------------------------------------------------------------------ шаблон договора


def default_template() -> Document:
    """Встроенный шаблон: договор возмездного оказания услуг и акт к нему (с новой страницы)."""
    doc = _new_doc(top_cm=1.5)
    left, just = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY

    def p(text, size=11, bold=False, align=just, after=4):
        return _para(doc, text, size, bold, align, space_after=after)

    p("ДОГОВОР № ______", 13, True, WD_ALIGN_PARAGRAPH.CENTER, 0)
    p("возмездного оказания услуг", 11, False, WD_ALIGN_PARAGRAPH.CENTER, 8)
    p("{{Место заключения}}\t\t\t\t\t\t\t«___» ______________ {{Год}} г.", 11, align=left, after=8)
    p("{{Заказчик}}, в лице {{В лице}}, действующего на основании {{Основание}}, именуемое в дальнейшем «Заказчик», "
      "с одной стороны, и гражданин(-ка) Российской Федерации {{ФИО}}, действующий(-ая) от своего имени, именуемый(-ая) "
      "в дальнейшем «Исполнитель», с другой стороны, вместе именуемые «Стороны», заключили настоящий договор "
      "о нижеследующем.")
    p("1. Предмет договора", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p("1.1. Исполнитель обязуется по заданию Заказчика оказать услуги в качестве {{В качестве}} на спортивных "
      "соревнованиях: {{Соревнование}}, спортивная дисциплина {{Дисциплина}}, {{Даты соревнований}}, {{Место}}, "
      "а Заказчик обязуется оплатить эти услуги.")
    p("1.2. Срок оказания услуг: {{Срок}}")  # срок кончается на «г.» — вторая точка не нужна
    p("1.3. Услуги считаются оказанными после подписания Сторонами акта приёма-сдачи оказанных услуг.")
    p("2. Цена услуг и порядок расчётов", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p("2.1. Стоимость услуг определяется исходя из {{Основание ставок}} ({{Ставка}} руб. в день, дней — {{Дней}}) "
      "и составляет {{Сумма прописью}}, в том числе НДФЛ 13 %.")
    p("2.2. Оплата производится на счёт Исполнителя либо наличными через кассу Заказчика в течение {{Срок оплаты}} "
      "с момента подписания акта приёма-сдачи оказанных услуг.")
    p("2.3. Финансирование: {{Финансирование}}.")
    p("3. Права и обязанности сторон", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p("3.1. Исполнитель обязуется оказать услуги лично, надлежащего качества, в соответствии с Положением о "
      "спортивных судьях и Положением о проведении соревнований; соблюдать требования техники безопасности, "
      "пожарной безопасности и охраны окружающей среды; сдать оказанные услуги по акту в последний день их оказания.")
    p("3.2. Заказчик обязуется принять и оплатить оказанные услуги в порядке и на условиях настоящего договора.")
    p("3.3. Исполнитель выражает согласие на обработку своих персональных данных, указанных в настоящем договоре, "
      "в объёме, необходимом для его исполнения (Федеральный закон от 27.07.2006 № 152-ФЗ).")
    p("4. Ответственность сторон и заключительные положения", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p("4.1. За неисполнение либо ненадлежащее исполнение обязательств Стороны несут ответственность в соответствии "
      "с законодательством Российской Федерации. Споры разрешаются путём переговоров.")
    p("4.2. Договор вступает в силу с даты начала оказания услуг и действует до полного исполнения Сторонами своих "
      "обязательств. Договор составлен в двух экземплярах, имеющих равную юридическую силу.")
    p("5. Адреса и реквизиты сторон", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _sides(doc,
           "Заказчик:\n{{Заказчик кратко}}\n{{Реквизиты заказчика}}",
           "Исполнитель:\n{{ФИО}}\nДата рождения: {{Дата рождения}}\nПаспорт: {{Паспорт}}\nВыдан: {{Кем выдан}}\n"
           "Дата выдачи: {{Дата выдачи}}, код подразделения {{Код подразделения}}\nАдрес регистрации: {{Адрес}}\n"
           "ИНН: {{ИНН}}\nСНИЛС: {{СНИЛС}}\nСчёт: {{Счёт}}\nБанк: {{Банк}}, БИК {{БИК}}\nТелефон: {{Телефон}}")
    _sides(doc, "{{Должность руководителя}} {{Заказчик кратко}}\n\n______________ / {{Руководитель кратко}} /\nМ.П.",
           "\n\n______________ / {{ФИО кратко}} /")

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p("АКТ № ______ от «___» ______________ {{Год}} г.", 13, True, WD_ALIGN_PARAGRAPH.CENTER, 0)
    p("приёма-сдачи оказанных услуг", 11, False, WD_ALIGN_PARAGRAPH.CENTER, 0)
    p("к договору № ______ от «___» ______________ {{Год}} г.", 11, False, WD_ALIGN_PARAGRAPH.CENTER, 8)
    p("Заказчик: {{Заказчик}}", align=left)
    p("Исполнитель: {{ФИО}}", align=left, after=8)
    t = doc.add_table(rows=2, cols=6)
    t.style = "Table Grid"
    head = ["№", "Наименование услуг, соревнование, сроки оказания услуг", "Ед. изм.", "Кол-во", "Цена, руб.",
            "Сумма (включая НДФЛ), руб."]
    row = ["1", "{{Услуга}} — {{Соревнование}}, спортивная дисциплина {{Дисциплина}} ({{Срок}})", "дн.", "{{Дней}}",
           "{{Ставка}}", "{{Сумма}}"]
    for i, (h, v) in enumerate(zip(head, row)):
        t.rows[0].cells[i].text, t.rows[1].cells[i].text = h, v
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for r in t.rows:
        for i, w in enumerate([0.8, 8.4, 1.4, 1.4, 2.2, 3.0]):
            r.cells[i].width = Cm(w)
            for par in r.cells[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(10)
    p("", 6)
    p("Всего оказано услуг на сумму: {{Сумма прописью}}, в том числе НДФЛ 13 %.", align=left)
    p("Вышеперечисленные услуги выполнены полностью и в срок. Заказчик претензий по объёму, качеству и срокам "
      "оказания услуг не имеет.", after=12)
    _sides(doc, "Заказчик:\n{{Должность руководителя}} {{Заказчик кратко}}\n\n______________ / {{Руководитель кратко}} /"
                "\nМ.П.", "Исполнитель:\n\n\n______________ / {{ФИО кратко}} /")
    return doc


def _sides(doc, left: str, right: str) -> None:
    """Две колонки без рамок: заказчик слева, исполнитель справа."""
    t = doc.add_table(rows=1, cols=2)
    for cell, text in zip(t.rows[0].cells, (left, right)):
        cell.width = Cm(8.5)
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.font.size = Pt(10.5)


def _paragraphs(doc):
    """Все абзацы документа: основной текст, таблицы (и вложенные), колонтитулы."""
    def walk(container):
        yield from container.paragraphs
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            yield from walk(part)


def fill(doc, values: dict[str, str]) -> set[str]:
    """Подставить значения в поля {{…}}. Возвращает имена полей, которых программа не знает."""
    unknown: set[str] = set()

    def repl(m):
        name = m.group(1)
        if name in values:
            return values[name]
        unknown.add(name)
        return m.group(0)

    for par in _paragraphs(doc):
        if "{{" not in par.text:
            continue
        for run in par.runs:
            if "{{" in run.text and _PH.search(run.text):
                run.text = _PH.sub(repl, run.text)
        if _PH.search(par.text):  # Word разбил поле на несколько кусков — собрать абзац в первый кусок
            text = _PH.sub(repl, par.text)
            runs = par.runs
            runs[0].text = text
            for run in runs[1:]:
                run.text = ""
    return unknown


def template_fields(path: Path) -> set[str]:
    """Какие поля есть в шаблоне (для проверки своего шаблона)."""
    return {m.group(1) for par in _paragraphs(Document(str(path))) for m in _PH.finditer(par.text)}


def _load(template: Path | None):
    return Document(str(template)) if template else default_template()


def write_contract(template: Path | None, values: dict[str, str], path: str | Path) -> set[str]:
    doc = _load(template)
    unknown = fill(doc, values)
    doc.save(str(path))
    return unknown


def template_name(role: str = "") -> str:
    """«Шаблон договора.docx» — общий; «Шаблон договора — главный судья.docx» — для одной должности
    (у заказчика бывают разные формы: главному судье — с отчётом, рабочим — «техническое обслуживание»)."""
    return f"Шаблон договора — {' '.join(role.lower().split())}.docx" if role else "Шаблон договора.docx"


def find_template(folders: list[Path], role: str) -> Path | None:
    """Шаблон для должности: сначала в первой папке (соревнования) — для должности, затем общий; потом в следующей."""
    for d in folders:
        for name in (template_name(role), template_name()):
            if (d / name).is_file():
                return d / name
    return None


def write_contracts(items: list[tuple[Path | None, dict[str, str]]], path: str | Path) -> set[str]:
    """Все договоры с актами одним файлом — каждый с новой страницы (для печати). items — (шаблон, значения)."""
    unknown: set[str] = set()
    base = None
    for template, values in items:
        doc = _load(template)
        unknown |= fill(doc, values)
        if base is None:
            base = doc
            continue
        base.add_paragraph().add_run().add_break(WD_BREAK.PAGE)  # в конце документа, перед его настройками
        sect = base.element.body.find(qn("w:sectPr"))
        for el in list(doc.element.body):
            if el.tag == qn("w:sectPr"):
                continue
            sect.addprevious(deepcopy(el))
    if base is None:
        base = _new_doc()
        _para(base, "Нет людей с отмеченными днями и ставкой — договоров нет.", 12)
    base.save(str(path))
    return unknown


# ------------------------------------------------------------------ табель-наряд


def _put(ws, r, c, v, bold=False, size=11, align=None, border=False, fill=None):
    """Значение в клетку табеля с оформлением."""
    cell = ws.cell(r, c, v)
    cell.font = Font(bold=bold, size=size)
    if align:
        cell.alignment = align
    if border:
        cell.border = BOX
    if fill:
        cell.fill = fill
    return cell


def _tabel_head(ws, comp: Competition, customer: dict, c_ok: int, c_sum: int) -> None:
    """Шапка табеля: название и даты слева, «УТВЕРЖДАЮ» руководителя заказчика — справа, на последние колонки."""
    _put(ws, 1, 1, "ТАБЕЛЬ – НАРЯД", True, 14)
    _put(ws, 2, 1, "на оплату судейской (комендантской) бригады")
    _put(ws, 3, 1, f"{title_of(comp)} в спортивной дисциплине {disciplines_text(comp)}", True)
    _put(ws, 4, 1, f"{comp.dates_text}, {comp.place}")
    head_post = _cap(str(customer.get("head_post") or "директор"))
    head_fio = str(customer.get("head_fio") or "")
    approve = ["УТВЕРЖДАЮ:", f"{head_post} {customer.get('short') or customer.get('name') or ''}".strip(),
               f"______________ {initials(head_fio, surname_first=False) if head_fio else ''}".rstrip(),
               "«___» ______________ 20__ г."]
    for r, text in enumerate(approve, start=1):
        _put(ws, r, c_ok, text, r == 1, align=Alignment(horizontal="left", vertical="center"))
        ws.merge_cells(start_row=r, start_column=c_ok, end_row=r, end_column=c_sum)


def _tabel_columns(ws, r0: int, days: list, c_days: int, c_count: int, c_rate: int, c_sum: int) -> None:
    """Две строки заголовков: человек, «Дни работы» с датами, дни, ставка, сумма."""
    nd = len(days)
    heads = ["№", "Фамилия, имя, отчество", "Должность", "Год рожд.", "Категория", "№ удостоверения"]
    for c, h in enumerate(heads, start=1):
        _put(ws, r0, c, h, True, 10, CENTER, True, HEAD)
        ws.merge_cells(start_row=r0, start_column=c, end_row=r0 + 1, end_column=c)
    if nd:
        _put(ws, r0, c_days, "Дни работы", True, 10, CENTER, True, HEAD)
        ws.merge_cells(start_row=r0, start_column=c_days, end_row=r0, end_column=c_days + nd - 1)
    for i, d in enumerate(days):
        _put(ws, r0 + 1, c_days + i, f"{d:%d.%m}", True, 10, CENTER, True, HEAD)
    for c, h in ((c_count, "Кол-во дней"), (c_rate, "Оплата за день, руб."), (c_sum, "Общая сумма, руб.")):
        _put(ws, r0, c, h, True, 10, CENTER, True, HEAD)
        ws.merge_cells(start_row=r0, start_column=c, end_row=r0 + 1, end_column=c)
    for r in (r0, r0 + 1):  # рамки у объединённых ячеек
        for c in range(1, c_sum + 1):
            ws.cell(r, c).border = BOX


def _tabel_people(ws, r: int, paid: list[Person], days: list, personal: dict[str, dict], c_days: int, c_count: int,
                  c_rate: int, c_sum: int) -> int:
    """Строки бригады: человек, отметки дней, формулы «дней × ставка»; возвращает следующую строку."""
    nd = len(days)
    for n, p in enumerate(paid, start=1):
        pd = personal.get(p.key, {})
        birth = str(pd.get("birth", "") or "")
        values = [n, p.fio, p.role, birth[-4:] if len(birth) >= 4 else "", p.cat, pd.get("judge_id", "")]
        for c, v in enumerate(values, start=1):
            _put(ws, r, c, v, size=10, align=CENTER if c in (1, 4, 5) else WRAP, border=True)
        marked = set(p.days)
        for i, d in enumerate(days):
            _put(ws, r, c_days + i, MARK if d in marked else None, size=10, align=CENTER, border=True)
        a, b = get_column_letter(c_days), get_column_letter(c_days + nd - 1)
        _put(ws, r, c_count, f'=COUNTIF({a}{r}:{b}{r},"{MARK}")' if nd else 0, size=10, align=CENTER, border=True)
        _put(ws, r, c_rate, p.rate or 0, size=10, align=CENTER, border=True).number_format = "#,##0"
        _put(ws, r, c_sum, f"={get_column_letter(c_count)}{r}*{get_column_letter(c_rate)}{r}", size=10, align=CENTER,
             border=True).number_format = "#,##0"
        r += 1
    return r


def write_tabel(comp: Competition, team: list[Person], days: list, accrual: float, customer: dict,
                personal: dict[str, dict], path: str | Path) -> Path:
    """Табель-наряд бригады (Excel): шапка с «УТВЕРЖДАЮ», люди по дням, суммы формулами, начисления, подписи."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Табель-наряд"
    paid = [p for p in team if p.paid]
    nd = len(days)
    c_days = 7  # первая колонка дней (G)
    c_count, c_rate, c_sum = c_days + nd, c_days + nd + 1, c_days + nd + 2
    last = get_column_letter(c_sum)
    c_ok = max(c_days, c_sum - 3)  # «Утверждаю» — справа, на четыре последние колонки (иначе обрезается при печати)
    _tabel_head(ws, comp, customer, c_ok, c_sum)
    r0 = 6
    _tabel_columns(ws, r0, days, c_days, c_count, c_rate, c_sum)
    first = r0 + 2
    r = _tabel_people(ws, first, paid, days, personal, c_days, c_count, c_rate, c_sum)
    s = get_column_letter(c_sum)
    total = f"SUM({s}{first}:{s}{r - 1})" if paid else "0"
    rows = [("Итого начислено", f"={total}"),
            (f"Начисления на оплату {accrual:g} %", f"=ROUND({s}{r}*{accrual:g}/100,2)"),
            ("Итого", f"={s}{r}+{s}{r + 1}")]
    for label, formula in rows:
        _put(ws, r, 2, label, True)
        _put(ws, r, c_sum, formula, True, align=CENTER, border=True).number_format = "#,##0.00"
        r += 1
    r += 1
    for role in ("Главный судья", "Главный секретарь"):
        o = comp.official(role)
        who = initials(o.fio, surname_first=False) if o and o.fio else " " * 20
        _put(ws, r, 2, f"{role} ______________ / {who} /")
        r += 2
    for c, w in enumerate([5, 34, 22, 8, 10, 13], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    for i in range(nd):
        ws.column_dimensions[get_column_letter(c_days + i)].width = 7
    for c, w in ((c_count, 9), (c_rate, 11), (c_sum, 13)):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{last}{r}"
    wb.calculation.fullCalcOnLoad = True  # формулы считаются при открытии в Excel
    path = Path(path)
    wb.save(path)
    return path
