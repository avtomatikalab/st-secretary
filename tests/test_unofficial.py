"""Неофициальные соревнования: свои зачёты и дисциплины, «без класса», без ранга и разрядов (Правки.md, п. 17–18,
решение 038)."""

from dataclasses import replace
from datetime import datetime
from fractions import Fraction

from conftest import make_application
from openpyxl import load_workbook

from st_secretary import psr_run as pr
from st_secretary import time_run as tr
from st_secretary.competition import Zachet
from st_secretary.exporters.results_protocol import write_protocol
from st_secretary.importers.card_xlsx import load_card, write_card
from st_secretary.importers.preapp_xlsx import read_preapplication
from st_secretary.issues import ERROR
from st_secretary.preapp import process
from st_secretary.psr_run import Member, TeamInput
from st_secretary.qualification import Qual
from st_secretary.web import forms

OBSTACLE = Zachet("М/Ж", 0, "", team_size=3, name="Новички", zid="Новички", discipline_text="Полоса препятствий",
                  result="time", unit="team")
FAMILY = Zachet("М/Ж", 2, "0840161811Я", name="Семейные команды", zid="Семейные команды")


def unofficial(card):
    return replace(card, unofficial=True, zachety=[OBSTACLE, FAMILY, *card.zachety])


def test_card_roundtrip_through_excel(tmp_path, psr_card):
    comp = unofficial(psr_card)
    back = load_card(write_card(tmp_path / "Карточка.xlsx", comp))
    assert back == comp and back.unofficial
    assert not [i for i in back.check() if i.severity == ERROR]
    z = back.zachety[0]
    assert (z.key, z.title, z.is_custom, z.rank_format, z.class_text) == (
        "Новички", "Новички", True, "group", "неклассифицированная")
    assert z.discipline_name == "Полоса препятствий" and z.header_text == "Дисциплина «Полоса препятствий»; неклассифицированная"
    assert back.zachety[2].key == "М/Ж_3" and back.zachety[2].title == "М/Ж_3"  # официальный зачёт — как раньше


def test_official_card_has_no_own_discipline_and_no_class(psr_card):
    errs = [i.text for i in replace(unofficial(psr_card), unofficial=False).check() if i.severity == ERROR]
    assert any("своя дисциплина — только у неофициальных" in e for e in errs)
    assert any("класс дистанции должен быть от 1 до 6" in e for e in errs)
    bad = replace(OBSTACLE, result="")
    errs = [i.text for i in replace(psr_card, unofficial=True, zachety=[bad]).check() if i.severity == ERROR]
    assert any("нужны название, вид результата и состав" in e for e in errs)


def test_rename_keeps_zachet_code(psr_card):
    """Переименовали зачёт (название, группу) — код тот же: заявки, жеребьёвка, результаты, ссылки на месте."""
    form = forms.card_to_form(unofficial(psr_card))
    form["zachety"][1] |= {"name": "Семьи"}
    form["zachety"][2] |= {"group": "МЖ"}  # официальный зачёт М/Ж_3: поправили группу
    comp, err = forms.form_to_card(form)
    assert not err
    fam, psr = comp.zachety[1], comp.zachety[2]
    assert (fam.key, fam.title) == ("Семейные команды", "Семьи")
    assert (psr.key, psr.title) == ("М/Ж_3", "МЖ_3")
    form["main"]["unofficial"] = ""  # сняли галочку — своя дисциплина и «без класса» не пройдут
    _, err = forms.form_to_card(form)
    assert "z-0-distance_class" in err and "z-0-discipline_code" in err


def team(name, n, qual=Qual.II):
    return TeamInput(f"{name}.xlsx", name, "г. N", str(n), [Member(f"{name} {i}", qual, qual.label) for i in range(3)],
                     admitted=True, representative="")


def test_unofficial_results_without_rank_percent_and_norms(tmp_path, psr_card):
    comp = unofficial(psr_card)
    z = comp.zachety[2]  # ПСР, 3 класс
    zdata = {"stages": [{"id": "s1", "tour": "Тур 1", "name": "Узлы", "max": "30"}],
             "teams": {"Кедр.xlsx": {"points": {"s1": "5"}}, "Сосна.xlsx": {"points": {"s1": "3"}}}}
    run = pr.compute(comp, z, zdata, [team("Кедр", 1), team("Сосна", 2)])
    assert run.rank is None and not run.norms_ok and [r.place for r in run.rows] == [1, 2]
    assert all(r.percent is None and not r.norm for r in run.rows)
    assert any("неофициальные соревнования" in i.text for i in run.issues)
    path = write_protocol(comp, run, "official", datetime(2026, 10, 3, 18, 0), tmp_path / "p.xlsx")
    wb = load_workbook(path)
    cells = [c for row in wb["Протокол"].iter_rows(values_only=True) for c in row if c]
    assert "Результат" in cells and "% от победителя" not in cells and "Выполнен разряд" not in cells
    assert not any("Квалификационный ранг" in str(c) for c in cells)
    assert str(wb["По этапам"].cell(2, 1).value).startswith("Результаты на ")


def test_own_time_discipline_counts_like_pedestrian_without_points(psr_card):
    comp = unofficial(psr_card)
    assert tr.is_time_discipline(OBSTACLE) and tr.profile(OBSTACLE) == "pedestrian"
    assert tr.settings(OBSTACLE, {})["system"] == "nopenalty"
    zdata = {"stages": [], "teams": {"Кедр.xlsx": {"start": "10:00:00", "finish": "10:12:00"},
                                     "Сосна.xlsx": {"start": "10:05:00", "finish": "10:15:30"}}}
    run = tr.compute(comp, OBSTACLE, zdata, [team("Кедр", 1), team("Сосна", 2)])
    assert [(r.inp.team, r.place, r.total) for r in run.rows] == [("Сосна", 1, Fraction(630)), ("Кедр", 2, Fraction(720))]
    assert run.rank is None


def test_application_names_own_zachet(tmp_path, psr_card):
    """В заявке у своего зачёта в «Группе» — его название; проверок по Правилам (класс, разряд) нет."""
    comp = unofficial(psr_card)
    app = make_application(tmp_path / "Ёжики.xlsx", "Ёжики", "Красноярск", "Иванов Пётр Сергеевич", "89131234567", 3, [
        ["Ёжики", "Красноярск", "Иванов Пётр Сергеевич", "Иванов Пётр Сергеевич", "17.10.1989", "б/р", "м", "Новички", ""],
        ["Ёжики", "Красноярск", "Иванов Пётр Сергеевич", "Смирнова Анна Олеговна", "12.03.2014", "б/р", "ж", "Новички", ""],
        ["Ёжики", "Красноярск", "Иванов Пётр Сергеевич", "Кузьмин Олег Игоревич", "28.02.2012", "б/р", "м", "новички", ""],
    ])
    result = process([read_preapplication(app)], comp)
    t = result.teams[0]
    assert [e.zachet.key for e in t.entries] == ["Новички"] * 3 and {e.group for e in t.entries} == {"Новички"}
    assert not [i for i in result.issues if i.text.startswith("Правила:")]
    assert comp.zachet_in_file(OBSTACLE) == ("Новички", "") and comp.zachet_in_file(comp.zachety[2]) == ("М/Ж", "3")


def test_own_personal_discipline_splits_team_into_athletes(psr_card):
    """Правки, п. 38: своя дисциплина с составом «личный» — формат ВРВС «individual»: каждый спортсмен — свой номер
    «3.1» и своё место (раньше «person» не узнавался, и место получала вся команда)."""
    from types import SimpleNamespace as NS

    from st_secretary import commission as cm
    from st_secretary.rank import INDIVIDUAL, PAIR
    from st_secretary.units import zachet_units

    run_z = Zachet("М/Ж", 0, "", name="Кросс", zid="Кросс", discipline_text="Кросс", result="time", unit="person")
    assert run_z.rank_format == INDIVIDUAL and pr.unit_kind(run_z.rank_format) == "person"
    assert replace(run_z, unit="pair").rank_format == PAIR and OBSTACLE.rank_format == "group"
    people = [NS(entry=NS(name=NS(full=fio), qual=Qual.BR, chip="", zachet=run_z, personal="", pair="", pair_num="",
                          team_dist="", num_in_team=n), status=cm.PENDING)
              for n, fio in ((1, "Иванов Иван"), (2, "Петрова Анна"))]
    t = NS(file="Кедр.xlsx", team=NS(team="Кедр", territory="г. N", representative="Пред"), number=3,
           status=cm.PENDING, persons=people)
    units = zachet_units([t], run_z, run_z.rank_format)
    assert [(u.team, u.number, len(u.members)) for u in units] == [("Иванов Иван", "3.1", 1), ("Петрова Анна", "3.2", 1)]
    zdata = {"stages": [], "teams": {units[0].file: {"start": "10:00:00", "finish": "10:20:00"},
                                     units[1].file: {"start": "10:01:00", "finish": "10:15:00"}}}
    run = tr.compute(unofficial(psr_card), run_z, zdata, units)
    assert [(r.inp.team, r.place) for r in run.rows] == [("Петрова Анна", 1), ("Иванов Иван", 2)]
