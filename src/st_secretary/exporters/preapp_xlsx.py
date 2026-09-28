"""Сводка предзаявок в Excel: данные для СЕКРЕТАРЬ_ST, замечания, список команд."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from st_secretary.competition import Competition
from st_secretary.issues import ERROR, FIXED, INFO, SEVERITY_LABEL, SEVERITY_ORDER, WARNING
from st_secretary.preapp import PreappResult

FILL = {
    ERROR: PatternFill("solid", fgColor="F8CBAD"),
    WARNING: PatternFill("solid", fgColor="FFE699"),
    FIXED: PatternFill("solid", fgColor="C6EFCE"),
    INFO: PatternFill("solid", fgColor="DDEBF7"),
}
HEAD = PatternFill("solid", fgColor="DCE6F1")
TEAM_ROW = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# Заголовки листа «Заявка» книги СЕКРЕТАРЬ_ST (колонки A–O) + колонка для сверки.
SEKRETAR_HEADERS = [
    "№ п/п", "Делегация", "Территория", "Представитель", "Фамилия Имя", "Дата или год рождения",
    "Разряд по СТ", "Пол", "Группа", "Класс дистанции", "Номер чипа (указать, если чип свой)",
    "Участие в личной дистанции", "Участие в дистанции связок", "Номер связки (если больше одной связки)",
    "Участие в дистанции-группа (вписать номер группы)", "Зачёт (для сверки, не копировать)",
]


def _head(ws, headers, widths):
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font, c.fill, c.border, c.alignment = Font(bold=True), HEAD, BOX, WRAP
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{ws.max_row + 1}"  # строкой — чтобы не создавать пустую ячейку ниже шапки


def _cell_value(s: str):
    """Числа — числами (СЕКРЕТАРЬ_ST сравнивает номера групп и связок как числа), текст — текстом."""
    if not s:
        return None
    return int(s) if s.isdigit() else s


def write_preapp_report(result: PreappResult, comp: Competition, path: str | Path) -> Path:
    wb = Workbook()
    _summary(wb.active, result, comp)
    _sekretar_sheet(wb.create_sheet("Заявка для СЕКРЕТАРЬ"), result)
    _issues_sheet(wb.create_sheet("Проверка"), result)
    _teams_sheet(wb.create_sheet("Команды"), result)
    path = Path(path)
    wb.save(path)
    return path


def _summary(ws, r: PreappResult, comp: Competition):
    ws.title = "Итог"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    rows = [
        ("Соревнования", comp.title),
        ("Даты", comp.dates_text),
        ("Зачёты", ", ".join(z.key for z in comp.zachety)),
        ("Обработано", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Команд", len(r.teams)),
        ("Участников", len(r.entries)),
        ("Ошибок (исправить до комиссии)", r.count(ERROR)),
        ("Проверить (нужно решение)", r.count(WARNING)),
        ("Исправлено автоматически", r.count(FIXED)),
    ]
    ws.append(["Сводка предварительных заявок"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for k, v in rows:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    for sev, row in ((ERROR, 9), (WARNING, 10), (FIXED, 11)):
        ws.cell(row, 2).fill = FILL[sev]
    ws.append([])
    steps = [
        "Как перенести в СЕКРЕТАРЬ_ST:",
        "1. Откройте лист «Проверка» и разберите красные строки — это ошибки, их нужно исправить в заявке "
        "или уточнить у представителя. Жёлтые — проверьте. Зелёные система исправила сама — просмотрите. "
        "Для каждого замечания в колонках «Почему» и «Что сделать» написано, в чём дело и как поступить.",
        "2. На листе «Заявка для СЕКРЕТАРЬ» выделите строки ОДНОЙ команды, колонки A–O, и скопируйте.",
        "3. В книге СЕКРЕТАРЬ_ST на листе «Заявка» встаньте в ячейку A2 → Специальная вставка → Значения.",
        "4. Проверьте, что колонка «Зачет» (P) в СЕКРЕТАРЬ_ST заполнилась так же, как колонка P здесь, "
        "и нажмите «ИМПОРТ В БАЗУ». Повторите для следующей команды.",
        "Колонку A (№) не меняйте: СЕКРЕТАРЬ_ST берёт из неё номер участника в команде.",
    ]
    for s in steps:
        ws.append(["", s])
        ws.cell(ws.max_row, 2).alignment = WRAP
    ws.cell(ws.max_row - len(steps) + 1, 2).font = Font(bold=True)


def _sekretar_sheet(ws, r: PreappResult):
    _head(ws, SEKRETAR_HEADERS, [5, 22, 16, 30, 36, 13, 9, 6, 10, 9, 12, 10, 10, 10, 12, 14])
    worst_by_person = {}
    for i in r.issues:
        if i.person:
            key = (i.source, i.person)
            cur = worst_by_person.get(key)
            if cur is None or SEVERITY_ORDER[i.severity] < SEVERITY_ORDER[cur]:
                worst_by_person[key] = i.severity
    for t in r.teams:
        ws.append([f"{t.team} — {t.territory} ({len(t.entries)} чел., файл {t.source})"])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=len(SEKRETAR_HEADERS))
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 1).fill = TEAM_ROW
        for e in t.entries:
            birth = e.birth if e.birth else e.birth_year
            ws.append([
                e.num_in_team, e.team, e.territory, e.representative, e.name.full, birth,
                e.qual.label if e.qual is not None else "", e.sex or "", e.group, e.distance_class,
                _cell_value(e.chip), _cell_value(e.personal), _cell_value(e.pair), _cell_value(e.pair_num),
                _cell_value(e.team_dist),
                e.zachet.key if e.zachet else "",
            ])
            row = ws.max_row
            if e.birth:
                ws.cell(row, 6).number_format = "DD.MM.YYYY"
            sev = worst_by_person.get((e.source, e.name.full))
            for c in ws[row]:
                c.border = BOX
                if sev in (ERROR, WARNING):
                    c.fill = FILL[sev]


def _issues_sheet(ws, r: PreappResult):
    _head(ws, ["Уровень", "Команда", "Участник", "Поле", "Что", "Почему", "Что сделать", "Было", "Стало", "Файл",
               "Строка"], [12, 22, 32, 16, 50, 60, 60, 24, 24, 26, 8])
    for i in sorted(r.issues, key=lambda x: (SEVERITY_ORDER[x.severity], x.team, x.person)):
        ws.append([SEVERITY_LABEL[i.severity], i.team, i.person, i.field, i.text, i.why, i.todo, i.before, i.after,
                   i.source, i.row or None])
        for c in ws[ws.max_row]:
            c.border, c.alignment = BOX, WRAP
        ws.cell(ws.max_row, 1).fill = FILL[i.severity]
    ws.auto_filter.ref = f"A1:K{ws.max_row}"


def _teams_sheet(ws, r: PreappResult):
    _head(ws, ["№", "Команда", "Территория", "Представитель", "Телефон", "E-mail", "Участников", "Мужчин",
               "Женщин", "Ошибок", "Проверить", "Файл"], [5, 24, 16, 32, 18, 28, 11, 9, 9, 9, 10, 28])
    for n, t in enumerate(r.teams, start=1):
        errs = sum(1 for i in r.issues if i.source == t.source and i.severity == ERROR)
        warns = sum(1 for i in r.issues if i.source == t.source and i.severity == WARNING)
        ws.append([n, t.team, t.territory, t.representative, t.phone, t.email, len(t.entries),
                   sum(e.sex == "м" for e in t.entries), sum(e.sex == "ж" for e in t.entries), errs, warns, t.source])
        for c in ws[ws.max_row]:
            c.border = BOX
        if errs:
            ws.cell(ws.max_row, 10).fill = FILL[ERROR]
        if warns:
            ws.cell(ws.max_row, 11).fill = FILL[WARNING]
    ws.append([])
    ws.append(["", "Итого", "", "", "", "", len(r.entries), sum(e.sex == "м" for e in r.entries),
               sum(e.sex == "ж" for e in r.entries)])
    ws.cell(ws.max_row, 2).font = Font(bold=True)
