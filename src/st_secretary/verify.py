"""Сверка документов: с карточкой соревнования и между собой.

Документы секретариата часто делают из прошлогодних — и забывают поправить даты, коды, фамилии. Программа
читает готовые файлы (Word, Excel, PDF — свои и присланные: протоколы СЕКРЕТАРЬ_ST, договоры заказчика,
табель) и ищет то, что расходится:

- даты не того года и не тех дней (в папке 2025 г. отчёт, дипломы и наклейки были с датами 2024 г.);
- коды ВРВС не тех дисциплин (в протоколе ПСР-2024 стоял код спелео-группы);
- опечатки в названии дисциплины («комбинированые» во всех договорах 2025 г.) и в названии соревнований;
- ФИО и судейские категории не как в карточке и заявках; имя и отчество слитно;
- сумма цифрами и прописью не совпадают;
- табель: отмеченных дней не столько, сколько к оплате; дни × ставка ≠ сумма; итоги;
- договор и акт: срок в договоре и в акте разный, сумма договора ≠ сумме акта, дни и суммы не как в табеле.

Всё, что находится, — замечания с цитатой из документа; исправляет человек (в Word или Excel).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from st_secretary.competition import MONTHS_GEN, Competition
from st_secretary.issues import ERROR, WARNING, Issue
from st_secretary.money import number_words
from st_secretary.reference import discipline_by_code

READABLE = {".docx", ".xlsx", ".xlsm", ".xls", ".pdf"}
_MONTH = {m: i for i, m in enumerate(MONTHS_GEN, start=1)}
_MONTHS_RE = "|".join(MONTHS_GEN)


@dataclass
class Doc:
    name: str
    text: str = ""
    tables: list[list[list[str]]] = field(default_factory=list)  # таблица → строки → ячейки
    error: str = ""


# ------------------------------------------------------------------ чтение


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "strftime"):
        return v.strftime("%d.%m.%Y")
    return " ".join(str(v).split())


def read_doc(path: Path, name: str | None = None) -> Doc:
    """Текст и таблицы документа. Не прочитался — Doc с error (замечание, а не падение)."""
    name = name or path.name
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            from docx import Document

            d = Document(str(path))
            lines, tables = [], []
            for block in d.iter_inner_content():  # абзацы и таблицы по порядку: акт — сразу после своего договора
                if hasattr(block, "rows"):
                    rows = [[" ".join(c.text.split()) for c in r.cells] for r in block.rows]
                    tables.append(rows)
                    lines += [" | ".join(dict.fromkeys(r)) for r in rows]  # объединённые ячейки повторяются
                else:
                    lines.append(block.text)
            for s in d.sections:
                lines += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
            return Doc(name, "\n".join(lines), tables)
        if ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            wb = load_workbook(path, data_only=True, read_only=True)
            tables = [[[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
            wb.close()
            return Doc(name, _table_text(tables), tables)
        if ext == ".xls":
            import xlrd

            wb = xlrd.open_workbook(str(path))
            tables = [[[_cell_text(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
                      for sh in wb.sheets()]
            return Doc(name, _table_text(tables), tables)
        if ext == ".pdf":
            from pypdf import PdfReader

            text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
            if not text.strip():
                return Doc(name, error="в PDF нет текста — похоже, это скан; сверить можно только текстовый PDF")
            return Doc(name, text)
        if ext in (".doc", ".rtf", ".odt"):
            return Doc(name, error="старый формат — откройте в Word и сохраните как .docx")
        return Doc(name, error="такие файлы программа не читает — нужны Word (.docx), Excel или PDF")
    except ImportError as e:
        return Doc(name, error=f"не установлена библиотека для чтения ({e.name}) — выполните в папке программы: uv sync")
    except Exception as e:  # noqa: BLE001 — повреждённый или запароленный файл: сказать, а не упасть
        return Doc(name, error=f"файл не открывается: {e}")


def _table_text(tables) -> str:
    return "\n".join(" | ".join(c for c in row if c) for t in tables for row in t if any(row))


def _quote(text: str, start: int, end: int, pad: int = 40) -> str:
    """Цитата вокруг найденного — чтобы было видно, где это в документе."""
    a, b = max(0, start - pad), min(len(text), end + pad)
    s = " ".join(text[a:b].split())
    return ("…" if a else "") + s + ("…" if b < len(text) else "")


def _issue(sev: str, doc: Doc, text: str, why: str = "", todo: str = "", quote: str = "") -> Issue:
    return Issue(sev, text, source=doc.name, why=why, todo=todo, before=quote)


# ------------------------------------------------------------------ даты

_DATE_WORDS = re.compile(
    rf"(?<!\d)[«\"“]?(\d{{1,2}})[»\"”]?(?:\s*(?:[-–—]|по)\s*[«\"“]?(\d{{1,2}})[»\"”]?)?\s+({_MONTHS_RE})\s+(\d{{4}})",
    re.IGNORECASE)
_DATE_NUM = re.compile(r"(?<![\d.])(\d{1,2})(?:\s*[-–]\s*(\d{1,2}))?\.(\d{1,2})\.(\d{4})(?![\d])")


def _make(y: int, m: int, d: int) -> date | None:
    try:
        return date(y, m, d)
    except ValueError:
        return None


def find_ranges(text: str) -> list[tuple[date | None, date | None, int, int]]:
    """Даты словами: «19–21 сентября 2025», «04» октября 2024. (начало, конец, позиция в тексте)."""
    out = []
    for m in _DATE_WORDS.finditer(text):
        month, year = _MONTH[m.group(3).lower()], int(m.group(4))
        a = _make(year, month, int(m.group(1)))
        b = _make(year, month, int(m.group(2))) if m.group(2) else a
        out.append((a, b, m.start(), m.end()))
    return out


def check_dates(doc: Doc, comp: Competition) -> list[Issue]:
    """Даты не того года — почти наверняка документ из прошлогоднего; далеко от соревнований — проверить."""
    early = comp.date_from - timedelta(days=31)  # подготовка: приём заявок, договоры
    late = comp.date_to + timedelta(days=150)  # отчёт, выписки, акты — после соревнований (выписки — до 4 мес.)
    out, seen = [], set()
    for a, b, s, e in find_ranges(doc.text):
        found = " ".join(doc.text[s:e].split())
        if found in seen:
            continue
        seen.add(found)
        q = _quote(doc.text, s, e)
        if a is None or b is None:
            out.append(_issue(ERROR, doc, f"несуществующая дата «{found}»", quote=q))
        elif a.year != comp.year:
            out.append(_issue(ERROR, doc, f"дата «{found}» — не {comp.year} год",
                              why=f"Соревнования {comp.dates_text} Так бывает, когда документ сделан из прошлогоднего "
                                  "и дату забыли поправить.",
                              todo="Поправьте дату в документе.", quote=q))
        elif b < early or a > late:
            out.append(_issue(WARNING, doc, f"дата «{found}» далеко от дат соревнований ({comp.dates_text})",
                              todo="Проверьте, та ли это дата.", quote=q))
    for m in _DATE_NUM.finditer(doc.text):
        if int(m.group(4)) == comp.year - 1 and abs(int(m.group(3)) - comp.date_from.month) <= 1:
            found = m.group(0)
            if found not in seen:
                seen.add(found)
                out.append(_issue(WARNING, doc, f"дата «{found}» — прошлый год, а месяц как у соревнований",
                                  why="Похоже на дату прошлогодних соревнований.", todo="Проверьте дату.",
                                  quote=_quote(doc.text, m.start(), m.end())))
    return out


# ------------------------------------------------------------------ коды ВРВС, дисциплина, название

_CODE = re.compile(r"(?<!\d)0840\d{6}[ЯЛ]")


def _discipline_name(code: str) -> str | None:
    try:
        return discipline_by_code(code).name
    except KeyError:
        return None


def check_codes(doc: Doc, comp: Competition) -> list[Issue]:
    card = {z.discipline_code: z.discipline_name for z in comp.zachety}
    expected = "; ".join(f"{c} — «{n}»" for c, n in card.items())
    out, seen = [], set()
    for m in _CODE.finditer(doc.text):
        code = m.group(0)
        if code in card or code in seen:
            continue
        seen.add(code)
        name = _discipline_name(code)
        what = f"это «{name}»" if name else "такого кода нет во ВРВС"
        out.append(_issue(ERROR, doc, f"код ВРВС {code} — {what}, а в карточке: {expected}",
                          why="Код дисциплины в протоколах и выписках должен быть как во ВРВС. В шаблоне СЕКРЕТАРЬ_ST "
                              "по умолчанию стоят коды спелео — в протоколе ПСР-2024 так и осталось.",
                          todo="Поправьте код (в СЕКРЕТАРЬ_ST — на листе «Настройка»).",
                          quote=_quote(doc.text, m.start(), m.end())))
    return out


def _lev(a: str, b: str) -> int:
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def _stem(word: str) -> str:
    return word[:-2] if word.endswith(("ая", "ый", "ий", "ое", "ые", "ой")) else word[:-1]


def check_spelling(doc: Doc, comp: Competition) -> list[Issue]:
    """Опечатки в словах названия дисциплины: «комбинированые» вместо «комбинированные»."""
    stems = {}
    for z in comp.zachety:
        for w in re.findall(r"[а-яё]+", z.discipline_name.lower()):
            if len(w) >= 7:
                stems[_stem(w)] = w
    out, seen = [], set()
    for m in re.finditer(r"[А-ЯЁа-яё]+", doc.text):
        w = m.group(0).lower()
        if w in seen:
            continue
        for stem, full in stems.items():
            if (len(stem) - 1 <= len(w) <= len(full) + 3 and not w.startswith(stem) and w[:4] == stem[:4]
                    and _lev(w[:len(stem)], stem) == 1):
                seen.add(w)
                out.append(_issue(WARNING, doc, f"«{m.group(0)}» — похоже на опечатку: во ВРВС «{full}»",
                                  todo="Поправьте слово во всём документе (в Word — «Заменить»).",
                                  quote=_quote(doc.text, m.start(), m.end())))
                break
    return out


_KINDS = r"чемпионат\w*|первенств\w*|кубк?\w*|соревновани\w*|фестивал\w*|турнир\w*|спартакиад\w*"
_TITLE = re.compile(rf"\b({_KINDS})\s+((?:\S+\s+){{0,8}}?)по\s+спортивному\s+туризму", re.IGNORECASE)


def _norm_title(s: str) -> str:
    s = s.lower().replace("ё", "е").replace("«", "").replace("»", "").replace('"', "")
    s = re.sub(r"\bгорода\b", "г.", s)
    s = re.sub(r"\bг\.\s*", "г. ", s)
    return " ".join(s.split())


def check_title(doc: Doc, comp: Competition) -> list[Issue]:
    ref = _TITLE.search(comp.title)
    if not ref:
        return []
    kind, rest = ref.group(1).lower()[:3], _norm_title(ref.group(2))
    out, seen = [], set()
    for m in _TITLE.finditer(doc.text):
        k, r = m.group(1).lower()[:3], _norm_title(m.group(2))
        found = " ".join(m.group(0).split())
        if (k, r) == (kind, rest) or found in seen or not r:
            continue
        seen.add(found)
        out.append(_issue(WARNING, doc, f"название соревнований «{found}» не как в карточке («{comp.title}»)",
                          todo="Если это то же соревнование — поправьте название.",
                          quote=_quote(doc.text, m.start(), m.end())))
    return out


# ------------------------------------------------------------------ ФИО и категории


@dataclass(frozen=True)
class Known:
    fio: str
    category: str = ""  # как в карточке: СС1К… (пусто — не судья или не известно)
    source: str = ""  # «в карточке», «в заявке», «в табеле программы»


def _key(s: str) -> str:
    return " ".join(s.lower().replace("ё", "е").split())


_FULL = re.compile(r"\b([А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?)\s+([А-ЯЁ][а-яё]+)\s+([А-ЯЁ][а-яё]+(?:вич|вна|чна|ич|оглы|кызы))\b")
_INI_BEFORE = re.compile(r"\b([А-ЯЁ])\.\s?([А-ЯЁ])\.\s?([А-ЯЁ][а-яё]+)")
_INI_AFTER = re.compile(r"\b([А-ЯЁ][а-яё]+)\s+([А-ЯЁ])\.\s?([А-ЯЁ])\.")
_MERGED = re.compile(r"\b([А-ЯЁ][а-яё]+)\s+([А-ЯЁ][а-яё]+)([А-ЯЁ][а-яё]+(?:вич|вна|чна|ич))\b")
_CAT = re.compile(r"\b(ССВК|СС1К|СС2К|СС3К|ЮС)\b")


def check_people(doc: Doc, people: list[Known]) -> list[Issue]:
    by_surname: dict[str, list[Known]] = {}
    for p in people:
        parts = _key(p.fio).split()
        if len(parts) >= 2:
            by_surname.setdefault(parts[0], []).append(p)
    out, seen = [], set()

    def add(sev, text, s, e, todo="Поправьте в документе — или в карточке и заявке, если ошибка там."):
        if text not in seen:
            seen.add(text)
            out.append(_issue(sev, doc, text, todo=todo, quote=_quote(doc.text, s, e)))

    def category_after(p: Known, end: int, s: int):
        m = _CAT.search(doc.text, end, end + 30)
        if m and p.category and m.group(1) != p.category and "\n" not in doc.text[end:m.start()]:
            add(ERROR, f"{p.fio}: категория {m.group(1)}, а в карточке — {p.category}", s, m.end())

    for m in _FULL.finditer(doc.text):
        sur, name, pat = (x.lower().replace("ё", "е") for x in m.groups())
        cands = by_surname.get(sur)
        if not cands:
            continue
        hit = [p for p in cands if _key(p.fio).split()[1:3] == [name, pat]]
        if hit:
            category_after(hit[0], m.end(), m.start())
            continue
        known = " или ".join(f"«{p.fio}» ({p.source})" for p in cands)
        same_name = any(_key(p.fio).split()[1][:3] == name[:3] for p in cands)
        add(ERROR if same_name else WARNING, f"«{' '.join(m.groups())}» — в документе, а {known}", m.start(), m.end())
    for rx, before in ((_INI_BEFORE, True), (_INI_AFTER, False)):
        for m in rx.finditer(doc.text):
            i1, i2, sur = m.groups() if before else (m.group(2), m.group(3), m.group(1))
            cands = by_surname.get(sur.lower().replace("ё", "е"))
            if not cands:
                continue
            ok = [p for p in cands if [w[0] for w in _key(p.fio).split()[1:3]] == [i1.lower(), i2.lower()]]
            if ok:
                category_after(ok[0], m.end(), m.start())
            elif len(cands) == 1:
                add(ERROR, f"«{m.group(0)}» — инициалы не сходятся: {cands[0].source} «{cands[0].fio}»",
                    m.start(), m.end())
    for m in _MERGED.finditer(doc.text):
        add(WARNING, f"«{m.group(0)}» — имя и отчество слитно", m.start(), m.end(),
            todo="Поставьте пробел между именем и отчеством.")
    return out


# ------------------------------------------------------------------ суммы прописью

_MONEY = re.compile(r"(?<![\d,])(\d{1,3}(?:[  ]\d{3})+|\d+)\s*\(([а-яё ]+?)\)\s*руб", re.IGNORECASE)


def check_money_words(doc: Doc) -> list[Issue]:
    out = []
    for m in _MONEY.finditer(doc.text):
        n = int(re.sub(r"\D", "", m.group(1)))
        words = " ".join(m.group(2).lower().split())
        if words != number_words(n):
            out.append(_issue(ERROR, doc, f"сумма {n} прописью «{words}» — не та: должно быть «{number_words(n)}»",
                              todo="Поправьте сумму цифрами или прописью.", quote=_quote(doc.text, m.start(), m.end())))
    return out


# ------------------------------------------------------------------ табель


@dataclass
class TabelRow:
    fio: str
    marks: int  # отмечено «р»
    days: float | None  # в колонке «кол-во дней»
    rate: float | None
    total: float | None
    doc: str


def _num(s: str) -> float | None:
    s = (s or "").replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


_DAY_HEAD = re.compile(rf"^\d{{1,2}}(\s+({_MONTHS_RE})|\.\d{{1,2}})", re.IGNORECASE)


def _columns(head: list[list[str]]) -> tuple[dict[str, int], list[int]]:
    """Колонки табеля по двум строкам шапки: ФИО, дни, ставка, сумма и колонки дней («18 сентября», «18.09»)."""
    width = max(len(r) for r in head)
    col: dict[str, int] = {}
    for c in range(width):
        h = " ".join(r[c] for r in head if c < len(r)).lower()
        if "фамилия" in h or "ф.и.о" in h:
            col.setdefault("fio", c)
        elif "кол" in h and "дн" in h:
            col.setdefault("days", c)
        elif "оплата за" in h or "ставк" in h:
            col.setdefault("rate", c)
        elif "сумм" in h:
            col["sum"] = c
    days = [c for c in range(width) if any(c < len(r) and _DAY_HEAD.match(r[c].strip()) for r in head)]
    return col, days


def parse_tabels(doc: Doc) -> tuple[list[TabelRow], list[Issue]]:
    """Табель: строки людей (ФИО, отметки, дни, ставка, сумма) и замечания по самому табелю."""
    rows_out: list[TabelRow] = []
    issues: list[Issue] = []
    for t in doc.tables:
        for hi in range(min(15, len(t))):
            head = t[hi:hi + 2]
            joined = " ".join(x.lower() for r in head for x in r)
            if not (any("фамилия" in c.lower() or "ф.и.о" in c.lower() for c in t[hi]) and "сумм" in joined):
                continue
            col, day_cols = _columns(head)
            if "fio" not in col:
                continue

            def cell(r, k, col=col):
                return r[col[k]] if k in col and col[k] < len(r) else ""

            people, totals = [], []
            for r in t[hi + 2:]:
                if any(c.lower().startswith("итого") for c in r[:3]):
                    totals.append(r)
                    continue
                fio = cell(r, "fio")
                if totals:  # после итогов — подписи («Гл. судья ___ / Федяев В.А.»), а не люди табеля
                    continue
                if len(fio.split()) >= 2 and re.match(r"[А-ЯЁ]", fio):
                    marks = sum(1 for c in day_cols if c < len(r) and r[c].strip().lower() in ("р", "р."))
                    people.append(TabelRow(" ".join(fio.split()), marks, _num(cell(r, "days")), _num(cell(r, "rate")),
                                           _num(cell(r, "sum")), doc.name))
            if not people:
                continue
            rows_out += people
            for p in people:
                if day_cols and p.days is not None and p.marks != p.days:
                    issues.append(_issue(ERROR, doc, f"табель, {p.fio}: отмечено дней «р» — {p.marks}, а к оплате — "
                                                     f"{p.days:g}", todo="Поправьте отметки или число дней.",
                                         why="Дни к оплате должны совпадать с отметками — по ним же составляется акт."))
                if None not in (p.days, p.rate, p.total) and abs(p.days * p.rate - p.total) > 0.5:
                    issues.append(_issue(ERROR, doc, f"табель, {p.fio}: {p.days:g} дн. × {p.rate:g} = "
                                                     f"{p.days * p.rate:g}, а в сумме — {p.total:g}"))
            sums = [p.total for p in people if p.total is not None]
            if sums and totals and "sum" in col:
                first = _num(cell(totals[0], "sum"))
                if first is not None and abs(first - sum(sums)) > 0.5:
                    issues.append(_issue(ERROR, doc, f"табель: «Итого начислено» — {first:g}, а сумма по людям — "
                                                     f"{sum(sums):g}"))
            break
    return rows_out, issues


# ------------------------------------------------------------------ договоры и акты


@dataclass
class ContractInfo:
    fio: str
    amount: float | None
    term: tuple[date, date] | None
    term_text: str
    act_days: float | None
    act_price: float | None
    act_sum: float | None
    act_term: tuple[date, date] | None
    act_term_text: str
    doc: str


_FIO_EXEC = re.compile(r"Российской\s+Федерации\s+([А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+[А-ЯЁ][а-яё]+\s*[А-ЯЁ]?[а-яё]*)")
_AMOUNT = re.compile(r"составляет\s+(\d[\d  ]*)")
_ACT_ROW = re.compile(r"[Дд]н\.[\s|]+(\d+(?:[.,]\d+)?)[\s|]+([\d  ]+,\d{2})[\s|]+([\d  ]+,\d{2})")


def parse_contracts(doc: Doc) -> list[ContractInfo]:
    """Договоры (с актами) в документе: в одном файле может быть много — делим по «ДОГОВОР №»."""
    out = []
    for part in re.split(r"(?=ДОГОВОР\s*№)", doc.text):
        fio = _FIO_EXEC.search(part) if part.startswith("ДОГОВОР") else None
        if not fio:
            continue
        act_at = re.search(r"\bАКТ\b|\bАкт\s*№", part)
        contract, act = (part[:act_at.start()], part[act_at.start():]) if act_at else (part, "")
        term_m = re.search(r"Срок оказания услуги?:?\s*([^\n]+)", contract)
        term = find_ranges(term_m.group(1)) if term_m else []
        act_ranges = [r for r in find_ranges(act) if r[0]]
        amount = _AMOUNT.search(contract)
        row = _ACT_ROW.search(act)
        out.append(ContractInfo(
            fio=" ".join(fio.group(1).split()),
            amount=_num(amount.group(1)) if amount else None,
            term=(term[0][0], term[0][1]) if term and term[0][0] else None,
            term_text=" ".join(term_m.group(1).split()).rstrip(".") if term_m else "",
            act_days=_num(row.group(1)) if row else None, act_price=_num(row.group(2)) if row else None,
            act_sum=_num(row.group(3)) if row else None,
            act_term=(act_ranges[-1][0], act_ranges[-1][1]) if act_ranges else None,
            act_term_text=" ".join(act[act_ranges[-1][2]:act_ranges[-1][3]].split()) if act_ranges else "",
            doc=doc.name))
    return out


def check_contract(c: ContractInfo, doc: Doc) -> list[Issue]:
    out = []
    who = f"договор {c.fio}"
    if c.term and c.act_term and c.term != c.act_term:
        out.append(_issue(ERROR, doc, f"{who}: срок в договоре «{c.term_text}», а в акте — «{c.act_term_text}»",
                          why="Срок оказания услуг в договоре и период в акте должны совпадать.",
                          todo="Поправьте срок в договоре или в акте."))
    if None not in (c.act_days, c.act_price, c.act_sum) and abs(c.act_days * c.act_price - c.act_sum) > 0.5:
        out.append(_issue(ERROR, doc, f"{who}: в акте {c.act_days:g} дн. × {c.act_price:g} ≠ {c.act_sum:g}"))
    if c.amount is not None and c.act_sum is not None and abs(c.amount - c.act_sum) > 0.5:
        out.append(_issue(ERROR, doc, f"{who}: в договоре {c.amount:g} руб., а в акте — {c.act_sum:g}"))
    if c.act_term and c.act_days is not None and c.act_days > (c.act_term[1] - c.act_term[0]).days + 1:
        out.append(_issue(ERROR, doc, f"{who}: в акте {c.act_days:g} дн., а в периоде «{c.act_term_text}» столько "
                                      "дней нет"))
    return out


def _person(fio: str) -> str:
    return "".join(_key(fio).split())  # «АленаВикторовна» и «Алена Викторовна» — один человек


def cross_check(tabel: list[TabelRow], contracts: list[ContractInfo]) -> list[Issue]:
    """Договоры и акты против табеля: дни, ставка, сумма; кто есть в одном, но нет в другом."""
    out = []
    by = {_person(r.fio): r for r in tabel}
    for c in contracts:
        r = by.get(_person(c.fio))
        src = f"{c.doc} ↔ {r.doc}" if r else c.doc
        if r is None:
            if tabel:
                out.append(Issue(WARNING, f"{c.fio}: есть договор, а в табеле такого человека нет", source=src))
            continue
        tdays = r.days if r.days is not None else r.marks
        if c.act_days is not None and tdays and c.act_days != tdays:
            out.append(Issue(ERROR, f"{c.fio}: в акте {c.act_days:g} дн., а в табеле — {tdays:g}", source=src))
        if c.act_price is not None and r.rate is not None and abs(c.act_price - r.rate) > 0.5:
            out.append(Issue(ERROR, f"{c.fio}: ставка в акте {c.act_price:g}, в табеле — {r.rate:g}", source=src))
        if c.amount is not None and r.total is not None and abs(c.amount - r.total) > 0.5:
            out.append(Issue(ERROR, f"{c.fio}: сумма по договору {c.amount:g}, в табеле — {r.total:g}", source=src))
    if contracts:
        have = {_person(c.fio) for c in contracts}
        missing = [r for r in tabel if r.total and _person(r.fio) not in have]
        if missing:  # одним замечанием: часто проверяют не все договоры сразу
            names = ", ".join(f"{r.fio} ({r.total:g} руб.)" for r in missing)
            out.append(Issue(WARNING, f"есть в табеле, а договоров среди проверенных файлов нет — {len(missing)} чел.: "
                                      f"{names}", source=missing[0].doc,
                             todo="Если проверяли не все договоры — добавьте остальные файлы."))
    return out


# ------------------------------------------------------------------ всё вместе


@dataclass
class Report:
    files: list[tuple[Doc, list[Issue]]]
    cross: list[Issue]
    tabel_rows: int = 0
    contracts: int = 0

    @property
    def all_issues(self) -> list[Issue]:
        return [i for _, iss in self.files for i in iss] + self.cross


def verify(docs: list[Doc], comp: Competition, people: list[Known]) -> Report:
    files, tabel, contracts = [], [], []
    for d in docs:
        if d.error:
            files.append((d, []))
            continue
        issues = (check_dates(d, comp) + check_codes(d, comp) + check_spelling(d, comp)
                  + check_title(d, comp) + check_people(d, people) + check_money_words(d))
        rows, tabel_issues = parse_tabels(d)
        tabel += rows
        issues += tabel_issues
        for c in parse_contracts(d):
            contracts.append(c)
            issues += check_contract(c, d)
        files.append((d, issues))
    return Report(files, cross_check(tabel, contracts), len(tabel), len(contracts))
