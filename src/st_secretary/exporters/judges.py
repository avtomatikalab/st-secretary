"""Документы о судьях: справки о судействе, справка о составе и квалификации судейской коллегии,
справка о количестве субъектов РФ (формы с tmmoscow.ru — Комитет дистанций ФСТР).

Судьи — из карточки соревнования (лист «ГСК»), оценки судейства — со страницы итогов.
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

from st_secretary.competition import Competition
from st_secretary.exporters.awards import _new_doc, _para
from st_secretary.names import dative, guess_sex, he_she
from st_secretary.results import person_key, title_of
from st_secretary.textclean import plural

CATEGORY_WORDS = [("ССВК", "Всероссийская категория"), ("СС1К", "Первая категория"), ("СС2К", "Вторая категория"),
                  ("СС3К", "Третья категория"), ("ЮС", "Юный спортивный судья"), ("б/к", "Без категории")]


def issuer(comp: Competition) -> str:
    """Кто выдаёт справку: федерация из проводящих организаций (или первая из них)."""
    feds = [o for o in comp.organizers if "федерац" in o.lower()]
    return (feds or comp.organizers or [""])[0]


def disciplines_text(comp: Competition) -> str:
    names = list(dict.fromkeys(z.discipline_name for z in comp.zachety))
    return ", ".join(f"«{n}»" for n in names)


def _sign(doc, comp: Competition, roles=("Главный судья", "Главный секретарь")) -> None:
    for role in roles:
        o = comp.official(role)
        sig = o.signature if o else " " * 30
        _para(doc, f"{role} ________________ / {sig} /", 11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)


def write_judging_certificates(comp: Competition, grades: dict[str, str], path: str | Path) -> Path:
    """Справки о судействе — по две на листе (разрезать): «Дана … в том, что … участвовал(а) в судействе …»."""
    doc = _new_doc(top_cm=1.5)
    judges = [o for o in comp.officials if o.fio]
    for n, o in enumerate(judges):
        if n and n % 2 == 0:
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        elif n:
            _para(doc, "- " * 40, 8, space_after=18)  # линия отреза
        _para(doc, issuer(comp).upper(), 12, True, space_after=0)
        _para(doc, "КОЛЛЕГИЯ СУДЕЙ", 12, True, space_after=10)
        _para(doc, "СПРАВКА", 16, True, space_after=10)
        sex = guess_sex(o.fio)
        he, took = he_she(sex)
        text = (f"Дана {dative(o.fio, sex)} в том, что {he} {took} в судействе {title_of(comp)} — "
                f"дисциплина {disciplines_text(comp)}, {comp.dates_text}, {comp.place}, в должности "
                f"{o.role[:1].lower() + o.role[1:]}.")
        _para(doc, text, 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)
        grade = grades.get(person_key(o.fio), "")
        _para(doc, f"Оценка судейства: «{grade or '________________'}».", 12, align=WD_ALIGN_PARAGRAPH.LEFT,
              space_after=14)
        _para(doc, "М.П.", 11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
        _sign(doc, comp)
    if not judges:
        _para(doc, "В карточке соревнования не указаны судьи (лист «ГСК»).", 12)
    path = Path(path)
    doc.save(path)
    return path


def _people(n: int) -> str:
    return f"{n} {plural(n, 'человек', 'человека', 'человек')}"


def _same_region(territory: str, comp: Competition) -> bool:
    """Судья из региона организаторов: в его территории есть название территории организаторов
    («Красноярский край, г. Уяр» и «Красноярск» — один регион). Грубо — проверьте в документе."""
    stem = comp.host_territory[:6].lower()
    return not territory or (bool(stem) and stem in territory.lower())


def _table(doc, head: list[str], rows: list[list[str]], widths_cm: list[float]):
    from docx.shared import Cm

    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, h in enumerate(head):
        cell = t.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
            for p in row.cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(11)
    return t


def _head(doc, comp: Competition, title: str, subtitle: str) -> None:
    for org in comp.organizers:
        _para(doc, org.upper(), 11, True, space_after=0)
    _para(doc, "", 6)
    _para(doc, title, 16, True, space_after=0)
    _para(doc, subtitle, 13, space_after=12)
    for label, value in (("Наименование соревнований", comp.title), ("Дата проведения", comp.dates_text),
                         ("Место проведения", comp.place)):
        _para(doc, f"{label}: {value}", 12, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)


def write_sk_certificate(comp: Competition, path: str | Path) -> Path:
    """Справка о составе и квалификации судейской коллегии (ЕВСК, п. 67.9)."""
    doc = _new_doc()
    _head(doc, comp, "СПРАВКА", "о составе и квалификации судейской коллегии")
    judges = [o for o in comp.officials if o.fio]
    other = sum(not _same_region(o.territory, comp) for o in judges)
    _para(doc, f"Всего судей {_people(len(judges))}, в том числе из других регионов {_people(other)}.", 12,
          align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    cats = Counter(o.category for o in judges)
    for code, words in CATEGORY_WORDS:
        if cats.get(code):
            _para(doc, f"{words} — {_people(cats[code])}", 12, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    _para(doc, "", 6)
    _table(doc, ["№ п/п", "Должность", "Фамилия, имя, отчество", "Квалиф. категория", "Территория"],
           [[str(i), o.role, o.fio, o.category, o.territory] for i, o in enumerate(judges, start=1)],
           [1.2, 5.2, 5.2, 2.4, 3.4])
    _para(doc, "", 12)
    _sign(doc, comp, ("Главный судья",))
    path = Path(path)
    doc.save(path)
    return path


def write_subjects_certificate(comp: Competition, territories: list[str], path: str | Path) -> Path:
    """Справка о количестве субъектов РФ (ЕВСК, п. 67.14 — для всероссийских и межрегиональных).
    territories — территории допущенных команд; субъект РФ по территории программа не определяет."""
    doc = _new_doc()
    _head(doc, comp, "СПРАВКА",
          "о количестве субъектов Российской Федерации, принявших участие в соревнованиях по спортивному туризму")
    unique = list(dict.fromkeys(t for t in territories if t))
    _para(doc, f"Количество субъектов РФ, принявших участие в соревнованиях: {len(unique)}", 12,
          align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)
    _para(doc, "Перечень субъектов Российской Федерации:", 12, True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
    for i, t in enumerate(unique, start=1):
        _para(doc, f"{i}. {t}", 12, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    _para(doc, "", 12)
    _sign(doc, comp, ("Главный судья",))
    path = Path(path)
    doc.save(path)
    return path
